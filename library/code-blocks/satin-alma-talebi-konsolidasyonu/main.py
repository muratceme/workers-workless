"""
Satın Alma Talebi Konsolidasyonu — Workers / Workless kod bloğu
Satın Alma › Satın Alma Uzman Yardımcısı

Departmanlardan gelen satın alma taleplerini malzeme bazında birleştirir ve teklif istenecek listeyi çıkarır:
  - Yalnız onaylı talepler birleştirilir (onay sütunu yoksa tümü onaylı sayılır); bekleyenler ayrı listelenir.
  - Malzeme anahtarı: malzeme kodu; kod yoksa katlanmış malzeme adı. Ölçü birimi eş adları birleştirilir (ad/adet,
    kilogram/kg...); aynı malzeme farklı birimle istenmişse birleştirilmez, uyarı verilir.
  - İsteğe bağlı stok dosyası: net ihtiyaç = talep − (eldeki − emniyet stoğu, en az 0) − açık sipariş (en az 0).
  - İsteğe bağlı malzeme kartı: asgari sipariş miktarı ve sipariş katına yukarı yuvarlama; tahmini tutar (son alış
    fiyatı); tedarik süresi ile en erken istenen teslim karşılaştırması; tedarikçi listesi.
  - Teklif kuralı: tahmini tutar --uc-teklif-esik ve üzeriyse en az 3 teklif (şirket politikası; değiştirilebilir).
  - Kontroller: olası mükerrer talep (aynı departman + malzeme, --mukerrer-gun içinde), kodsuz olası aynı malzeme
    (ad benzerliği), geçmiş istenen teslim, termin riski, tedarikçi havuzu yetersiz, birim uyuşmazlığı.
Rapor: özet, konsolide liste, teklif istek listesi (tedarikçi × malzeme), bekleyen talepler, talep detayı, uyarılar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 25 talep satırı, 7 departman
    python main.py --talepler talepler.xlsx --malzemeler malzemeler.xlsx --stok stok.xlsx --bugun 09.10.2026
    python main.py --talepler talepler.xlsx --ayri-dosya            # her tedarikçi için ayrı teklif formu
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
BIRIM_ES = {"adet": "Adet", "ad": "Adet", "pcs": "Adet", "tane": "Adet", "kg": "kg", "kilogram": "kg", "kilo": "kg", "gr": "g", "g": "g", "gram": "g",
            "lt": "lt", "l": "lt", "litre": "lt", "m": "m", "mt": "m", "metre": "m", "m2": "m²", "metrekare": "m²", "paket": "Paket", "pk": "Paket",
            "kutu": "Kutu", "kt": "Kutu", "koli": "Koli", "rulo": "Rulo", "top": "Top", "takim": "Takım", "set": "Takım", "cift": "Çift", "ton": "ton"}
ONAYLI = {"onaylandi", "onayli", "onay", "evet", "uygun"}
REDDEDILEN = {"reddedildi", "red", "ret", "iptal", "hayir", "iptal edildi"}

TALEP_SUTUNLARI = {"no": ("talep no", "talep numarasi", "talep"), "tarih": ("talep tarihi", "tarih"),
                   "departman": ("departman", "talep eden birim", "bolum", "maliyet merkezi"), "talep_eden": ("talep eden", "isteyen", "kisi"),
                   "kod": ("malzeme kodu", "stok kodu", "kod"), "ad": ("malzeme adi", "malzeme", "urun", "malzeme hizmet"),
                   "miktar": ("miktar", "talep miktari"), "birim": ("olcu birimi", "birim"), "teslim": ("istenen teslim", "istenen teslim tarihi", "ihtiyac tarihi", "termin"),
                   "onay": ("onay durumu", "durum", "onay"), "not": ("not", "gerekce", "aciklama")}
MALZEME_SUTUNLARI = {"kod": ("malzeme kodu", "stok kodu", "kod"), "ad": ("malzeme adi", "malzeme"), "kategori": ("kategori", "malzeme grubu"),
                     "tedarikci": ("tedarikciler", "tedarikci", "onayli tedarikciler"), "fiyat": ("son alis fiyati", "birim fiyat", "fiyat"),
                     "doviz": ("doviz", "para birimi"), "birim": ("olcu birimi", "birim"), "moq": ("asgari siparis", "asgari siparis miktari", "minimum siparis"),
                     "kat": ("siparis kati", "ambalaj kati", "kat"), "sure": ("tedarik suresi", "tedarik suresi gun", "teslim suresi")}
STOK_SUTUNLARI = {"kod": ("malzeme kodu", "stok kodu", "kod"), "eldeki": ("eldeki stok", "stok", "mevcut"), "emniyet": ("emniyet stogu", "minimum stok"),
                  "acik": ("acik siparis", "acik siparisler", "yoldaki")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.0f}".replace(",", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def birim_norm(b) -> str:
    k = katla(b).replace(" ", "")
    return BIRIM_ES.get(k, metin(b) or "Adet")


def kelimeler(ad: str) -> set[str]:
    return {k for k in katla(ad).split() if len(k) > 1}


def benzer(a: str, b: str) -> bool:
    x, y = kelimeler(a), kelimeler(b)
    return bool(x and y) and len(x & y) / len(x | y) >= 0.6


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Talep:
    satir: int
    no: str
    tarih: date | None
    departman: str
    talep_eden: str
    kod: str
    ad: str
    miktar: Decimal
    birim: str
    teslim: date | None
    durum: str          # onayli | bekleyen | red
    not_: str = ""
    anahtar: str = ""
    notlar: list[str] = field(default_factory=list)


def talepleri_oku(yol: Path) -> tuple[list[Talep], list[dict], bool]:
    talepler, uy = [], []
    satirlar = kayitlar(yol, TALEP_SUTUNLARI, ("miktar",))
    onay_var = any("onay" in r for r in satirlar[:1])
    for r in satirlar:
        kod, ad = metin(r.get("kod")).upper(), metin(r.get("ad"))
        if not kod and not ad:
            continue
        m = para(r.get("miktar"))
        if m is None or m <= 0:
            uy.append({"onem": "Yüksek", "tur": "Geçersiz miktar", "kim": metin(r.get("no")) or f"satır {r['_satir']}",
                       "aciklama": f"{ad or kod}: miktar '{metin(r.get('miktar'))}' okunamadı veya sıfır; talep alınmadı"})
            continue
        o = katla(r.get("onay"))
        durum = "onayli" if not onay_var or o in ONAYLI else ("red" if o in REDDEDILEN else "bekleyen")
        talepler.append(Talep(r["_satir"], metin(r.get("no")) or f"S{r['_satir']}", tarih(r.get("tarih")), metin(r.get("departman")) or "(belirtilmemiş)",
                              metin(r.get("talep_eden")), kod, ad, m, birim_norm(r.get("birim")), tarih(r.get("teslim")), durum, metin(r.get("not"))))
    if not onay_var:
        uy.append({"onem": "Bilgi", "tur": "Onay sütunu yok", "kim": yol.name, "aciklama": "Onay durumu sütunu bulunamadı; tüm talepler onaylı sayıldı"})
    return talepler, uy, onay_var


def malzemeleri_oku(yol: Path | None) -> dict[str, dict]:
    if not yol:
        return {}
    kart = {}
    for r in kayitlar(yol, MALZEME_SUTUNLARI, ("kod",)):
        kod = metin(r.get("kod")).upper()
        if not kod:
            continue
        ted = [t.strip() for t in re.split(r"[;/,|]", metin(r.get("tedarikci"))) if t.strip()]
        moq, kat, sure = para(r.get("moq")), para(r.get("kat")), para(r.get("sure"))
        kart[kod] = {"ad": metin(r.get("ad")), "kategori": metin(r.get("kategori")) or "Diğer", "tedarikciler": ted, "fiyat": para(r.get("fiyat")),
                     "doviz": (metin(r.get("doviz")) or "TRY").upper(), "birim": birim_norm(r.get("birim")) if metin(r.get("birim")) else "",
                     "moq": moq if moq and moq > 0 else None, "kat": kat if kat and kat > 0 else None, "sure": int(sure) if sure is not None else None}
    return kart


def stoku_oku(yol: Path | None) -> dict[str, dict]:
    if not yol:
        return {}
    return {metin(r.get("kod")).upper(): {"eldeki": para(r.get("eldeki")) or SIFIR, "emniyet": para(r.get("emniyet")) or SIFIR, "acik": para(r.get("acik")) or SIFIR}
            for r in kayitlar(yol, STOK_SUTUNLARI, ("kod", "eldeki")) if metin(r.get("kod"))}


# ----------------------------------------------------------------------------
# Konsolidasyon
# ----------------------------------------------------------------------------

def yukari_yuvarla(miktar: Decimal, moq: Decimal | None, kat: Decimal | None) -> Decimal:
    if miktar <= 0:
        return SIFIR
    if moq and miktar < moq:
        miktar = moq
    if kat:
        miktar = Decimal(math.ceil(miktar / kat)) * kat
    return miktar


def konsolide_et(talepler: list[Talep], kart: dict, stok: dict, bugun: date, mukerrer_gun: int = 7, uc_teklif_esik: Decimal = Decimal(100000),
                 teklif_gun: int = 7) -> dict:
    uy = []
    for t in talepler:
        t.anahtar = t.kod or "AD:" + katla(t.ad)
        if t.kod and t.kod in kart and not t.ad:
            t.ad = kart[t.kod]["ad"]
    onayli = [t for t in talepler if t.durum == "onayli"]

    # Mükerrer talep: aynı departman + malzeme, farklı talep satırı, tarih farkı ≤ mukerrer_gun
    gruplar = defaultdict(list)
    for t in onayli:
        gruplar[(katla(t.departman), t.anahtar)].append(t)
    for lst in gruplar.values():
        lst.sort(key=lambda t: (t.tarih or date.min, t.satir))
        for a, b in zip(lst, lst[1:]):
            if a.tarih and b.tarih and (b.tarih - a.tarih).days <= mukerrer_gun:
                b.notlar.append(f"olası mükerrer: {a.no}")
                uy.append({"onem": "Orta", "tur": "Olası mükerrer talep", "kim": f"{b.departman} · {b.ad or b.kod}",
                           "aciklama": f"{a.no} ({a.tarih:%d.%m.%Y}, {a.miktar} {a.birim}) ve {b.no} ({b.tarih:%d.%m.%Y}, {b.miktar} {b.birim}) "
                                       f"{(b.tarih - a.tarih).days} gün arayla; departmana teyit ettirin (listede ikisi de toplandı)"})

    # Kodsuz olası aynı malzeme
    kodsuz = {t.anahtar: t for t in onayli if not t.kod}
    adlar = {t.anahtar: t.ad for t in onayli} | {k: v["ad"] for k, v in kart.items() if v["ad"]}
    bildirilen = set()
    for ak, t in kodsuz.items():
        for bk, bad in adlar.items():
            if bk != ak and (ak, bk) not in bildirilen and (bk, ak) not in bildirilen and benzer(t.ad, bad):
                bildirilen.add((ak, bk))
                uy.append({"onem": "Orta", "tur": "Olası aynı malzeme", "kim": t.ad,
                           "aciklama": f"Kodsuz talep '{t.ad}' ile '{bad}'{' (' + bk + ')' if not bk.startswith('AD:') else ''} benzer; aynı malzemeyse "
                                       "koda bağlayın (ayrı satır olarak bırakıldı)"})

    satirlar = {}
    for t in onayli:
        anahtar = (t.anahtar, t.birim)
        s = satirlar.setdefault(anahtar, {"anahtar": t.anahtar, "kod": t.kod, "ad": t.ad, "birim": t.birim, "talep": SIFIR, "departmanlar": defaultdict(Decimal),
                                          "talepler": [], "teslim": None})
        s["talep"] += t.miktar
        s["departmanlar"][t.departman] += t.miktar
        s["talepler"].append(t.no)
        if t.teslim and (s["teslim"] is None or t.teslim < s["teslim"]):
            s["teslim"] = t.teslim
    birimleri = defaultdict(set)
    for a, b in satirlar:
        birimleri[a].add(b)
    for a, bs in birimleri.items():
        if len(bs) > 1:
            ad = next(s["ad"] for s in satirlar.values() if s["anahtar"] == a)
            uy.append({"onem": "Orta", "tur": "Birim uyuşmazlığı", "kim": ad,
                       "aciklama": f"Aynı malzeme farklı birimlerle istenmiş ({', '.join(sorted(bs))}); birimler çevrilmeden toplanmadı"})

    liste = []
    for (anahtar, birim), s in sorted(satirlar.items(), key=lambda i: (i[1]["ad"].lower(), i[0][1])):
        k = kart.get(s["kod"]) if s["kod"] else None
        s["kategori"] = k["kategori"] if k else "(kart yok)"
        s["tedarikciler"] = k["tedarikciler"] if k else []
        s["notlar"] = []
        if s["kod"] and not k and kart:
            s["notlar"].append("malzeme kartı yok")
        st = stok.get(s["kod"]) if s["kod"] else None
        kart_birim = (k or {}).get("birim") or birim
        s["kullanilabilir"] = s["acik"] = SIFIR
        if k and k["ad"]:
            s["ad"] = k["ad"]
        if kart_birim != birim:
            s["notlar"].append(f"kart birimi {kart_birim}, talep birimi {birim}: stok, asgari sipariş ve fiyat uygulanmadı")
        elif st:
            s["kullanilabilir"] = max(SIFIR, st["eldeki"] - st["emniyet"])
            s["acik"] = st["acik"]
        s["net"] = max(SIFIR, s["talep"] - s["kullanilabilir"] - s["acik"])
        moq, kat = ((k or {}).get("moq"), (k or {}).get("kat")) if kart_birim == birim else (None, None)
        s["siparis"] = yukari_yuvarla(s["net"], moq, kat)
        if s["siparis"] > s["net"] > 0:
            s["notlar"].append(f"asgari sipariş / kat nedeniyle {s['net']:g} → {s['siparis']:g}")
        s["fiyat"] = (k or {}).get("fiyat") if kart_birim == birim else None
        s["doviz"] = (k or {}).get("doviz", "TRY")
        s["tutar"] = s["siparis"] * s["fiyat"] if s["fiyat"] is not None else None
        s["sure"] = (k or {}).get("sure")
        s["gerekli_teklif"] = 3 if s["tutar"] is not None and s["doviz"] == "TRY" and s["tutar"] >= uc_teklif_esik else 1
        ad = s["ad"] or s["kod"]
        if s["siparis"] == 0:
            s["notlar"].append("stok ve açık siparişten karşılanıyor")
        else:
            if s["teslim"] and s["teslim"] < bugun:
                uy.append({"onem": "Yüksek", "tur": "İstenen teslim geçmiş", "kim": ad,
                           "aciklama": f"En erken istenen teslim {s['teslim']:%d.%m.%Y} rapor tarihinden önce; departmanla yeni termin belirleyin"})
            elif s["teslim"] and s["sure"] is not None and bugun + timedelta(days=teklif_gun + s["sure"]) > s["teslim"]:
                uy.append({"onem": "Yüksek", "tur": "Termin riski", "kim": ad,
                           "aciklama": f"Teklif süresi ({teklif_gun} gün) + tedarik süresi ({s['sure']} gün) → en erken "
                                       f"{bugun + timedelta(days=teklif_gun + s['sure']):%d.%m.%Y}; istenen {s['teslim']:%d.%m.%Y}"})
            if s["gerekli_teklif"] > len(s["tedarikciler"]):
                uy.append({"onem": "Orta" if s["gerekli_teklif"] == 3 else "Bilgi", "tur": "Tedarikçi havuzu yetersiz", "kim": ad,
                           "aciklama": f"Tahmini tutar {tl(s['tutar']) + ' ' + s['doviz'] if s['tutar'] is not None else '—'}; en az {s['gerekli_teklif']} "
                                       f"teklif gerekir, kayıtlı tedarikçi {len(s['tedarikciler'])}"})
            if s["fiyat"] is None:
                uy.append({"onem": "Bilgi", "tur": "Fiyat yok", "kim": ad, "aciklama": "Son alış fiyatı bilinmiyor; tahmini tutar hesaplanamadı"})
        liste.append(s)

    # Teklif istek listesi: tedarikçi × malzeme
    son_tarih = bugun + timedelta(days=teklif_gun)
    teklif = []
    for s in liste:
        if s["siparis"] <= 0:
            continue
        for ted in (s["tedarikciler"] or ["(Belirlenecek)"]):
            teklif.append({"tedarikci": ted, "kategori": s["kategori"], "kod": s["kod"], "ad": s["ad"], "miktar": s["siparis"], "birim": s["birim"],
                           "teslim": s["teslim"], "son_tarih": son_tarih})
    teklif.sort(key=lambda x: (x["tedarikci"] == "(Belirlenecek)", katla(x["tedarikci"]), katla(x["kategori"]), katla(x["ad"])))
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"liste": liste, "teklif": teklif, "uyarilar": uy, "bugun": bugun, "son_tarih": son_tarih, "uc_teklif_esik": uc_teklif_esik,
            "bekleyen": [t for t in talepler if t.durum == "bekleyen"], "red": [t for t in talepler if t.durum == "red"], "talepler": talepler}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
SARI = PatternFill("solid", fgColor="FFF4CE")
UST = Alignment(vertical="top", wrap_text=True)
PF, MF = "#,##0.00", "#,##0.##"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _f(x):
    return None if x is None else float(x)


def teklif_formu(ws, satirlar: list[dict], tedarikci: str, son_tarih: date) -> None:
    ws.append([f"Teklif İsteği — {tedarikci}"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=13)
    ws.append([f"Teklif son tarihi: {son_tarih:%d.%m.%Y}. Sarı alanları doldurup geri gönderiniz. Fiyatlar KDV hariç ve para birimi belirtilerek yazılmalıdır."])
    ws.append([])
    _baslik(ws, ["Kategori", "Malzeme Kodu", "Malzeme", "Miktar", "Birim", "İstenen Teslim", "Birim Fiyat", "Para Birimi", "Teslim Süresi (gün)",
                 "Teklif Geçerlilik", "Açıklama"], (16, 14, 36, 10, 8, 13, 12, 10, 12, 13, 26))
    ws.freeze_panes = "A5"
    for x in satirlar:
        ws.append([x["kategori"], x["kod"], x["ad"], float(x["miktar"]), x["birim"], x["teslim"], None, None, None, None, None])
        ws.cell(ws.max_row, 4).number_format = MF
        ws.cell(ws.max_row, 6).number_format = "DD.MM.YYYY"
        for j in range(7, 12):
            ws.cell(ws.max_row, j).fill = SARI


def rapor_yaz(cikti: Path, s: dict, ayri_dosya: bool = False) -> list[Path]:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer"], (44, 20))
    tutar_try = sum((x["tutar"] for x in s["liste"] if x["tutar"] is not None and x["doviz"] == "TRY"), SIFIR)
    for a, v in [("Rapor tarihi", s["bugun"].strftime("%d.%m.%Y")), ("Talep satırı (toplam)", len(s["talepler"])),
                 ("Onaylı talep satırı", sum(t.durum == "onayli" for t in s["talepler"])), ("Onay bekleyen", len(s["bekleyen"])),
                 ("Reddedilen / iptal (dahil edilmedi)", len(s["red"])), ("Konsolide malzeme satırı", len(s["liste"])),
                 ("Sipariş gereken malzeme", sum(x["siparis"] > 0 for x in s["liste"])),
                 ("Stoktan karşılanan malzeme", sum(x["siparis"] == 0 for x in s["liste"])),
                 ("Tahmini tutar (TL, fiyatı bilinenler)", float(tutar_try)),
                 ("Teklif istenecek tedarikçi", len({x["tedarikci"] for x in s["teklif"]} - {"(Belirlenecek)"})),
                 ("Tedarikçisi belirlenecek malzeme", len({x["ad"] for x in s["teklif"] if x["tedarikci"] == "(Belirlenecek)"})),
                 ("Teklif son tarihi", s["son_tarih"].strftime("%d.%m.%Y")), ("3 teklif eşiği (TL)", float(s["uc_teklif_esik"]))]:
        oz.append([a, v])
        if isinstance(v, float):
            oz.cell(oz.max_row, 2).number_format = PF

    ko = wb.create_sheet("Konsolide Liste")
    _baslik(ko, ["Kategori", "Malzeme Kodu", "Malzeme", "Birim", "Talep Toplamı", "Kullanılabilir Stok", "Açık Sipariş", "Net İhtiyaç", "Sipariş Miktarı",
                 "En Erken Teslim", "Son Alış Fiyatı", "Döviz", "Tahmini Tutar", "Gerekli Teklif", "Tedarikçiler", "Departmanlar", "Talep No", "Not"],
            (14, 13, 32, 7, 10, 10, 9, 10, 10, 12, 11, 6, 13, 8, 30, 36, 22, 36))
    for x in sorted(s["liste"], key=lambda x: (katla(x["kategori"]), katla(x["ad"]))):
        ko.append([x["kategori"], x["kod"], x["ad"], x["birim"], float(x["talep"]), float(x["kullanilabilir"]), float(x["acik"]), float(x["net"]),
                   float(x["siparis"]), x["teslim"], _f(x["fiyat"]), x["doviz"], _f(x["tutar"]), x["gerekli_teklif"], ", ".join(x["tedarikciler"]),
                   "; ".join(f"{d} {m:g}" for d, m in sorted(x["departmanlar"].items())), ", ".join(x["talepler"]), "; ".join(x["notlar"])])
        for j in range(5, 10):
            ko.cell(ko.max_row, j).number_format = MF
        ko.cell(ko.max_row, 10).number_format = "DD.MM.YYYY"
        ko.cell(ko.max_row, 11).number_format = ko.cell(ko.max_row, 13).number_format = PF
        if x["siparis"] == 0:
            ko.cell(ko.max_row, 9).fill = PatternFill("solid", fgColor="E3F4E1")
    ko.auto_filter.ref = f"A1:R{ko.max_row}"

    ti = wb.create_sheet("Teklif İstek Listesi")
    _baslik(ti, ["Tedarikçi", "Kategori", "Malzeme Kodu", "Malzeme", "Miktar", "Birim", "İstenen Teslim", "Teklif Son Tarihi", "Gönderildi mi?"],
            (28, 14, 13, 34, 10, 7, 12, 12, 12))
    for x in s["teklif"]:
        ti.append([x["tedarikci"], x["kategori"], x["kod"], x["ad"], float(x["miktar"]), x["birim"], x["teslim"], x["son_tarih"], ""])
        ti.cell(ti.max_row, 5).number_format = MF
        ti.cell(ti.max_row, 7).number_format = ti.cell(ti.max_row, 8).number_format = "DD.MM.YYYY"
        ti.cell(ti.max_row, 9).fill = SARI
        if x["tedarikci"] == "(Belirlenecek)":
            ti.cell(ti.max_row, 1).fill = PatternFill("solid", fgColor="FDE2E1")
    ti.auto_filter.ref = f"A1:I{ti.max_row}"

    be = wb.create_sheet("Bekleyen Talepler")
    _baslik(be, ["Talep No", "Tarih", "Departman", "Talep Eden", "Malzeme Kodu", "Malzeme", "Miktar", "Birim", "İstenen Teslim", "Durum"],
            (12, 11, 16, 18, 13, 32, 9, 7, 12, 12))
    for t in s["bekleyen"] + s["red"]:
        be.append([t.no, t.tarih, t.departman, t.talep_eden, t.kod, t.ad, float(t.miktar), t.birim, t.teslim, "Onay bekliyor" if t.durum == "bekleyen" else "Reddedildi / iptal"])
        be.cell(be.max_row, 2).number_format = be.cell(be.max_row, 9).number_format = "DD.MM.YYYY"

    td = wb.create_sheet("Talep Detayı")
    _baslik(td, ["Talep No", "Tarih", "Departman", "Talep Eden", "Malzeme Kodu", "Malzeme", "Miktar", "Birim", "İstenen Teslim", "Durum", "Gerekçe / Not",
                 "Kontrol"], (12, 11, 16, 18, 13, 32, 9, 7, 12, 12, 30, 30))
    for t in sorted(s["talepler"], key=lambda t: (t.tarih or date.min, t.satir)):
        td.append([t.no, t.tarih, t.departman, t.talep_eden, t.kod, t.ad, float(t.miktar), t.birim, t.teslim,
                   {"onayli": "Onaylı", "bekleyen": "Onay bekliyor", "red": "Reddedildi / iptal"}[t.durum], t.not_, "; ".join(t.notlar)])
        td.cell(td.max_row, 2).number_format = td.cell(td.max_row, 9).number_format = "DD.MM.YYYY"
    td.auto_filter.ref = f"A1:L{td.max_row}"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Malzeme / Kim", "Açıklama"], (9, 24, 30, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

    dosyalar = []
    if ayri_dosya:
        gruplar = defaultdict(list)
        for x in s["teklif"]:
            if x["tedarikci"] != "(Belirlenecek)":
                gruplar[x["tedarikci"]].append(x)
        klasor = cikti.parent / "teklif_formlari"
        klasor.mkdir(parents=True, exist_ok=True)
        for ted, lst in sorted(gruplar.items()):
            w = Workbook()
            teklif_formu(w.active, lst, ted, s["son_tarih"])
            w.active.title = "Teklif"
            yol = klasor / (re.sub(r"[^A-Za-z0-9]+", "_", katla(ted)).strip("_") + ".xlsx")
            w.save(yol)
            dosyalar.append(yol)
    return dosyalar


def calistir(talep_yolu: Path, cikti: Path, bugun: date, malzeme_yolu: Path | None = None, stok_yolu: Path | None = None, *, mukerrer_gun: int = 7,
             uc_teklif_esik: Decimal = Decimal(100000), teklif_gun: int = 7, ayri_dosya: bool = False) -> dict:
    talepler, uy, _ = talepleri_oku(talep_yolu)
    if not talepler:
        raise ValueError(f"{talep_yolu.name}: talep bulunamadı")
    s = konsolide_et(talepler, malzemeleri_oku(malzeme_yolu), stoku_oku(stok_yolu), bugun, mukerrer_gun, uc_teklif_esik, teklif_gun)
    s["uyarilar"] = uy + s["uyarilar"]
    s["dosyalar"] = rapor_yaz(cikti, s, ayri_dosya)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Departman satın alma taleplerini malzeme bazında birleştirir, teklif istenecek listeyi çıkarır.")
    p.add_argument("--talepler", type=Path, default=ORNEK / "talepler.csv",
                   help="Talep No, Tarih, Departman, Talep Eden, Malzeme Kodu, Malzeme Adı, Miktar, Birim, İstenen Teslim, Onay Durumu, Gerekçe")
    p.add_argument("--malzemeler", type=Path, help="İsteğe bağlı malzeme kartı: Malzeme Kodu, Malzeme Adı, Kategori, Tedarikçiler (; ile), Son Alış Fiyatı, "
                                                   "Döviz, Birim, Asgari Sipariş, Sipariş Katı, Tedarik Süresi (gün)")
    p.add_argument("--stok", type=Path, help="İsteğe bağlı: Malzeme Kodu, Eldeki Stok, Emniyet Stoğu, Açık Sipariş")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (örnek veride 09.10.2026)")
    p.add_argument("--mukerrer-gun", type=int, default=7, help="Aynı departman + malzeme bu kadar gün içinde tekrar istenirse uyarı (varsayılan 7)")
    p.add_argument("--uc-teklif-esik", type=float, default=100000, help="Bu tutar (TL) ve üzerinde en az 3 teklif (şirket politikası; varsayılan 100.000)")
    p.add_argument("--teklif-gun", type=int, default=7, help="Teklif son tarihi = rapor tarihi + bu kadar gün (varsayılan 7)")
    p.add_argument("--ayri-dosya", action="store_true", help="Her tedarikçi için teklif_formlari/ altında ayrı teklif formu yaz")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "satin_alma_konsolidasyonu.xlsx")
    a = p.parse_args(argv)
    ornek = a.talepler == ORNEK / "talepler.csv"
    malz = a.malzemeler or (ORNEK / "malzemeler.csv" if ornek else None)
    stok = a.stok or (ORNEK / "stok.csv" if ornek else None)
    for y in (a.talepler, malz, stok):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.talepler, a.cikti, bugun, malz, stok, mukerrer_gun=a.mukerrer_gun, uc_teklif_esik=Decimal(str(a.uc_teklif_esik)),
                     teklif_gun=a.teklif_gun, ayri_dosya=a.ayri_dosya)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    sip = [x for x in s["liste"] if x["siparis"] > 0]
    print(f"[OK] {sum(t.durum == 'onayli' for t in s['talepler'])} onaylı talep satırı → {len(s['liste'])} malzeme; sipariş gereken {len(sip)}, "
          f"stoktan karşılanan {len(s['liste']) - len(sip)}; onay bekleyen {len(s['bekleyen'])}")
    print(f"[OK] Teklif istenecek: {len({x['tedarikci'] for x in s['teklif']} - {'(Belirlenecek)'})} tedarikçi · son tarih {s['son_tarih']:%d.%m.%Y}")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    if s["dosyalar"]:
        print(f"[OK] {len(s['dosyalar'])} teklif formu: {s['dosyalar'][0].parent.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
