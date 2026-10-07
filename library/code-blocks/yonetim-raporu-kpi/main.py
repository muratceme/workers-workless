"""
Yönetim Raporu (KPI) — Workers / Workless kod bloğu
Finans › Bütçe ve Raporlama Uzmanı

Aylık yönetim raporunun tablo ve grafiklerini satış verisinden (ve isteğe bağlı mizandan) otomatik hazırlar:
  Satış KPI'ları (rapor ayı, önceki ay, geçen yıl aynı ay, yılbaşından bugüne ve geçen yıl YTD):
    net satış ve büyüme, satış adedi, fatura (sipariş) sayısı, ortalama fatura tutarı, aktif ve yeni müşteri,
    brüt kâr ve marjı (maliyet verilmişse), müşteri yoğunlaşması (ilk 5 müşterinin payı), kanal/kategori
    kırılımı, en büyük müşteri ve ürünler, 13 aylık trend.
  Finans KPI'ları (mizan verilirse, Tekdüzen Hesap Planı): net satış, brüt kâr, esas faaliyet kârı, dönem net
    kârı, hazır değerler, ticari alacak/borç, stok, cari oran, kaldıraç (Mali Tablo Rasyo Analizi çekirdeği).
İnternete bağlanmaz.

Kullanım:
    python main.py                                           # örnek satış verisiyle dener (Eylül 2026)
    python main.py --satis satislar.xlsx --ay 2026-09
    python main.py --satis satislar.xlsx --ay 2026-09 --mizan mizan_eylul.xlsx
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter, OrderedDict, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ilk = "\n".join(metin.splitlines()[:10])
        satirlar = list(csv.reader(metin.splitlines(), delimiter=max(";\t,", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


ALANLAR = {
    "tarih": ("tarih", "fatura tarihi", "belge tarihi", "sipariş tarihi"),
    "belge": ("fatura no", "belge no", "sipariş no", "evrak no"),
    "musteri": ("müşteri", "cari", "müşteri adı", "cari unvan"),
    "urun": ("ürün", "ürün adı", "stok adı", "malzeme"),
    "kategori": ("kategori", "ürün grubu", "grup"),
    "kanal": ("kanal", "satış kanalı", "bölge", "temsilci"),
    "adet": ("adet", "miktar"),
    "tutar": ("net tutar", "tutar", "satış tutarı", "net satış"),
    "maliyet": ("maliyet", "satış maliyeti", "toplam maliyet"),
    "tur": ("tür", "işlem türü", "belge türü"),
}


def satislari_oku(yol: Path) -> tuple[list[dict], list[str]]:
    s = tablo_oku(yol)
    b = [kucuk(x) for x in s[0]]
    k = {a: next((i for i, x in enumerate(b) if x in adlar), None) for a, adlar in ALANLAR.items()}
    if k["tarih"] is None or k["tutar"] is None:
        raise SystemExit(f"Satış verisinde Tarih ve Net Tutar (KDV hariç) gerekli. Başlıklar: {s[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    kayitlar, uyarilar = [], []
    for r in s[1:]:
        t = tarih(al(r, "tarih"))
        if not t:
            continue
        tutar, maliyet = para(al(r, "tutar")), para(al(r, "maliyet")) if k["maliyet"] is not None else None
        if "iade" in kucuk(al(r, "tur")) and tutar > 0:          # iade satırı pozitif yazılmışsa ters çevir
            tutar, maliyet = -tutar, (-maliyet if maliyet is not None else None)
        kayitlar.append({"tarih": t, "ay": (t.year, t.month), "belge": str(al(r, "belge") or ""), "musteri": str(al(r, "musteri") or "").strip(),
                         "urun": str(al(r, "urun") or "").strip(), "kategori": str(al(r, "kategori") or "").strip() or "(yok)",
                         "kanal": str(al(r, "kanal") or "").strip() or "(yok)", "adet": float(para(al(r, "adet"))), "tutar": tutar, "maliyet": maliyet})
    if k["maliyet"] is None:
        uyarilar.append("Maliyet sütunu yok: brüt kâr ve marj hesaplanmadı")
    return kayitlar, uyarilar


# ----------------------------------------------------------------------------
# KPI hesapları
# ----------------------------------------------------------------------------

def onceki_ay(a: tuple[int, int], n: int = 1) -> tuple[int, int]:
    y, m = a
    m -= n
    while m < 1:
        m += 12
        y -= 1
    return (y, m)


def ozet(kayitlar: list[dict], ilk_musteri: dict[str, tuple]) -> dict:
    tutar = sum((x["tutar"] for x in kayitlar), Decimal(0))
    maliyetler = [x["maliyet"] for x in kayitlar if x["maliyet"] is not None]
    maliyet = sum(maliyetler, Decimal(0)) if maliyetler else None
    belgeler = {x["belge"] for x in kayitlar if x["belge"] and x["tutar"] > 0}
    musteriler = {x["musteri"] for x in kayitlar if x["musteri"] and x["tutar"] > 0}
    aylar = {x["ay"] for x in kayitlar}
    yeni = {m for m in musteriler if ilk_musteri.get(m) in aylar}
    mus_tutar = Counter()
    for x in kayitlar:
        if x["musteri"]:
            mus_tutar[x["musteri"]] += x["tutar"]
    ilk5 = sum((v for _, v in mus_tutar.most_common(5)), Decimal(0))
    return {
        "satis": tutar, "adet": sum(x["adet"] for x in kayitlar), "fatura": len(belgeler) or None,
        "ort_fatura": (tutar / len(belgeler)) if belgeler else None, "aktif_musteri": len(musteriler), "yeni_musteri": len(yeni),
        "brut_kar": None if maliyet is None else tutar - maliyet, "marj": None if maliyet is None or not tutar else (tutar - maliyet) / tutar,
        "ilk5_pay": None if not tutar else ilk5 / tutar,
    }


def degisim(a, b):
    if a is None or b in (None, 0):
        return None
    return float(a / b - 1) if isinstance(a, Decimal) else a / b - 1


def hesapla(kayitlar: list[dict], ay: tuple[int, int]) -> dict:
    ilk_musteri = {}
    for x in sorted(kayitlar, key=lambda x: x["tarih"]):
        if x["musteri"] and x["tutar"] > 0:
            ilk_musteri.setdefault(x["musteri"], x["ay"])
    sec = lambda f: [x for x in kayitlar if f(x["ay"])]  # noqa: E731
    gy = (ay[0] - 1, ay[1])
    donemler = OrderedDict([
        ("Rapor ayı", sec(lambda a: a == ay)),
        ("Önceki ay", sec(lambda a: a == onceki_ay(ay))),
        ("Geçen yıl aynı ay", sec(lambda a: a == gy)),
        ("YTD", sec(lambda a: a[0] == ay[0] and a[1] <= ay[1])),
        ("Geçen yıl YTD", sec(lambda a: a[0] == ay[0] - 1 and a[1] <= ay[1])),
    ])
    kpi = OrderedDict((ad, ozet(v, ilk_musteri)) for ad, v in donemler.items())
    trend = OrderedDict()
    for i in range(12, -1, -1):
        a = onceki_ay(ay, i)
        v = sec(lambda x, a=a: x == a)
        o = ozet(v, ilk_musteri)
        trend[a] = o
    kirilim = {}
    for alan in ("kanal", "kategori", "musteri", "urun"):
        c = defaultdict(lambda: [Decimal(0), Decimal(0), Decimal(0), Decimal(0)])   # ay, gy ay, ytd, gy ytd
        for i, (ad, liste) in enumerate([("ay", donemler["Rapor ayı"]), ("gy", donemler["Geçen yıl aynı ay"]),
                                         ("ytd", donemler["YTD"]), ("gy_ytd", donemler["Geçen yıl YTD"])]):
            for x in liste:
                c[x[alan] or "(yok)"][i] += x["tutar"]
        kirilim[alan] = c
    return {"kpi": kpi, "trend": trend, "kirilim": kirilim, "ay": ay}


def finans_kpi(mizan: Path | None) -> dict | None:
    if not mizan:
        return None
    import rasyo_cekirdek as r
    m, uyarilar = r.mizan_oku(mizan)
    t, u2 = r.tablolar(m)
    oran = r.rasyolar(t, None, Decimal(365))
    bl, gt = t["bilanco"], t["gelir"]
    return {"uyarilar": uyarilar + u2, "satirlar": [
        ("Net satışlar (YTD, mizan)", gt["Net satışlar"]), ("Brüt satış kârı", gt["Brüt satış kârı"]),
        ("Esas faaliyet kârı", gt["Esas faaliyet kârı"]), ("Dönem net kârı", gt["Dönem net kârı"]),
        ("Hazır değerler", bl["Hazır değerler"]), ("Ticari alacaklar", bl["Ticari alacaklar"]),
        ("Ticari borçlar (kısa vadeli)", bl["Ticari borçlar (kısa vadeli)"]), ("Stoklar", bl["Stoklar"]),
        ("Cari oran", oran.get("Cari oran")), ("Kaldıraç oranı", oran.get("Kaldıraç oranı")),
        ("Brüt kâr marjı", oran.get("Brüt kâr marjı")), ("Net kâr marjı", oran.get("Net kâr marjı"))]}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
YESIL = PatternFill("solid", fgColor="E3F5E1")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
PARA = "#,##0"
YUZDE = "0.0%"
DEGISIM = "+0.0%;-0.0%;0.0%"


def f(x):
    return None if x is None else float(x)


def _degisim_hucresi(ws, r, c, v):
    ws.cell(r, c).value = v
    ws.cell(r, c).number_format = DEGISIM
    if v is not None:
        ws.cell(r, c).fill = YESIL if v > 0 else KIRMIZI if v < 0 else PatternFill()


def rapor_yaz(cikti: Path, h: dict, fin: dict | None, uyarilar: list[str]) -> None:
    ay = h["ay"]
    k = h["kpi"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Yönetim Özeti"
    ws.append([f"Yönetim Raporu — {AYLAR[ay[1] - 1]} {ay[0]}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Gösterge", f"{AYLAR[ay[1] - 1]} {ay[0]}", "Önceki ay", "Değişim", "GY aynı ay", "GY'ye göre", "YTD", "GY YTD", "YTD değişim"])
    for c in ws[3]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    satirlar = [("Net satış (TL)", "satis", PARA), ("Satış adedi", "adet", "#,##0"), ("Fatura sayısı", "fatura", "#,##0"),
                ("Ortalama fatura tutarı (TL)", "ort_fatura", PARA), ("Aktif müşteri", "aktif_musteri", "#,##0"), ("Yeni müşteri", "yeni_musteri", "#,##0"),
                ("Brüt kâr (TL)", "brut_kar", PARA), ("Brüt kâr marjı", "marj", YUZDE), ("İlk 5 müşteri payı", "ilk5_pay", YUZDE)]
    for ad, alan, bicim in satirlar:
        v = {p: k[p][alan] for p in k}
        if all(x is None for x in v.values()):
            continue
        r = ws.max_row + 1
        ws.append([ad, f(v["Rapor ayı"]), f(v["Önceki ay"]), None, f(v["Geçen yıl aynı ay"]), None, f(v["YTD"]), f(v["Geçen yıl YTD"]), None])
        for c in (2, 3, 5, 7, 8):
            ws.cell(r, c).number_format = bicim
        if bicim == YUZDE:                               # oranlarda puan farkı
            for c, a, b in ((4, "Rapor ayı", "Önceki ay"), (6, "Rapor ayı", "Geçen yıl aynı ay"), (9, "YTD", "Geçen yıl YTD")):
                if v[a] is not None and v[b] is not None:
                    ws.cell(r, c).value = round(float(v[a] - v[b]) * 100, 1)
                    ws.cell(r, c).number_format = '+0.0" puan";-0.0" puan";0.0" puan"'
        else:
            _degisim_hucresi(ws, r, 4, degisim(v["Rapor ayı"], v["Önceki ay"]))
            _degisim_hucresi(ws, r, 6, degisim(v["Rapor ayı"], v["Geçen yıl aynı ay"]))
            _degisim_hucresi(ws, r, 9, degisim(v["YTD"], v["Geçen yıl YTD"]))
    for j, w in enumerate((30, 16, 14, 10, 14, 11, 16, 16, 12), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    if fin:
        ws.append([])
        ws.append(["Finans (mizan)", "Değer"])
        for c in ws[ws.max_row]:
            c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
        for ad, v in fin["satirlar"]:
            ws.append([ad, f(v)])
            ws.cell(ws.max_row, 2).number_format = YUZDE if ("marj" in ad or "Kaldıraç" in ad) else "0.00" if ad == "Cari oran" else PARA
    ws.append([])
    for u in uyarilar + (fin["uyarilar"] if fin else []):
        ws.append(["Uyarı", u])
    ws.append(["Not", "Net satış KDV hariçtir; iadeler düşülür. Yeni müşteri: veri setinde ilk alımı bu dönemde olan müşteri. "
                      "GY: geçen yıl. Oran değişimleri yüzde puan olarak verilir."])

    t = wb.create_sheet("Trend")
    t.append(["Ay", "Net Satış", "Brüt Kâr Marjı", "Aktif Müşteri"])
    for c in t[1]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI
    for a, o in h["trend"].items():
        t.append([f"{AYLAR[a[1] - 1][:3]} {a[0]}", f(o["satis"]), f(o["marj"]), o["aktif_musteri"]])
        t.cell(t.max_row, 2).number_format = PARA
        t.cell(t.max_row, 3).number_format = YUZDE
    n = len(h["trend"])
    g = BarChart()
    g.title, g.height, g.width = "Net satış — son 13 ay", 8, 22
    g.add_data(Reference(t, min_col=2, min_row=1, max_row=1 + n), titles_from_data=True)
    g.set_categories(Reference(t, min_col=1, min_row=2, max_row=1 + n))
    ws.add_chart(g, "K3")
    if any(o["marj"] is not None for o in h["trend"].values()):
        lc = LineChart()
        lc.title, lc.height, lc.width = "Brüt kâr marjı", 7, 22
        lc.add_data(Reference(t, min_col=3, min_row=1, max_row=1 + n), titles_from_data=True)
        lc.set_categories(Reference(t, min_col=1, min_row=2, max_row=1 + n))
        lc.y_axis.number_format = "0%"
        ws.add_chart(lc, "K21")

    for alan, ad in (("kanal", "Kanal"), ("kategori", "Kategori"), ("musteri", "Müşteri"), ("urun", "Ürün")):
        c = h["kirilim"][alan]
        if len(c) <= 1 and "(yok)" in c:
            continue
        s = wb.create_sheet(ad)
        s.append([ad, "Rapor Ayı", "GY Aynı Ay", "Değişim", "YTD", "GY YTD", "YTD Değişim", "YTD Pay"])
        for x in s[1]:
            x.fill, x.font = BASLIK_DOLGU, BASLIK_YAZI
        toplam = sum((v[2] for v in c.values()), Decimal(0))
        sirali = sorted(c.items(), key=lambda i: -i[1][2])
        for isim, v in (sirali[:30] if alan in ("musteri", "urun") else sirali):
            s.append([isim, f(v[0]), f(v[1]), None, f(v[2]), f(v[3]), None, f(v[2] / toplam) if toplam else None])
            r = s.max_row
            for cc in (2, 3, 5, 6):
                s.cell(r, cc).number_format = PARA
            _degisim_hucresi(s, r, 4, degisim(v[0], v[1]))
            _degisim_hucresi(s, r, 7, degisim(v[2], v[3]))
            s.cell(r, 8).number_format = YUZDE
        s.column_dimensions["A"].width = 30
        for j in range(2, 9):
            s.column_dimensions[get_column_letter(j)].width = 14
        if alan in ("kanal", "kategori") and 1 < len(c) <= 12:
            p = PieChart()
            p.title, p.height, p.width = f"YTD {ad.lower()} dağılımı", 8, 12
            p.add_data(Reference(s, min_col=5, min_row=1, max_row=1 + len(sirali)), titles_from_data=True)
            p.set_categories(Reference(s, min_col=1, min_row=2, max_row=1 + len(sirali)))
            s.add_chart(p, "J2")
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(satis: Path, cikti: Path, ay: tuple[int, int] | None = None, mizan: Path | None = None) -> dict:
    kayitlar, uyarilar = satislari_oku(satis)
    if not kayitlar:
        raise SystemExit("Satış kaydı bulunamadı.")
    ay = ay or max(x["ay"] for x in kayitlar)
    h = hesapla(kayitlar, ay)
    if not h["kpi"]["Geçen yıl YTD"]["satis"]:
        uyarilar.append("Geçen yıl verisi yok: GY karşılaştırmaları boş")
    fin = finans_kpi(mizan)
    rapor_yaz(cikti, h, fin, uyarilar)
    return {**h, "fin": fin, "uyarilar": uyarilar}


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Satış verisinden (ve mizandan) aylık yönetim raporu KPI tablo ve grafiklerini hazırlar.")
    ap.add_argument("--satis", type=Path, default=BURASI / "ornek_veri" / "satislar.csv",
                    help="Satış satırları (.xlsx/.csv): Tarih, Net Tutar (KDV hariç); Fatura No, Müşteri, Ürün, Kategori, Kanal, Adet, Maliyet, Tür")
    ap.add_argument("--ay", help="Rapor ayı (YYYY-AA); varsayılan verideki son ay")
    ap.add_argument("--mizan", type=Path, help="İsteğe bağlı mizan (Tekdüzen, kapanış öncesi) — finans KPI'ları için")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "yonetim_raporu.xlsx")
    a = ap.parse_args(argv)
    ay = tuple(int(x) for x in a.ay.split("-")) if a.ay else None
    s = calistir(a.satis, a.cikti, ay, a.mizan)
    k = s["kpi"]
    tl = lambda x: f"{x:,.0f} TL".replace(",", ".")  # noqa: E731
    d = degisim(k["Rapor ayı"]["satis"], k["Geçen yıl aynı ay"]["satis"])
    print(f"[OK] {AYLAR[s['ay'][1] - 1]} {s['ay'][0]}: net satış {tl(k['Rapor ayı']['satis'])}"
          + (f" (GY'ye göre %{d * 100:+.1f})".replace(".", ",") if d is not None else "")
          + f" · YTD {tl(k['YTD']['satis'])} · aktif müşteri {k['Rapor ayı']['aktif_musteri']}")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
