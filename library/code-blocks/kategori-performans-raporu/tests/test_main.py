import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def u(urun, satis, smm, stok=None, onceki=None, kategori="Kahve", marka="A", adet=None):
    f = lambda v: None if v is None else D(v)   # noqa: E731
    return {"kategori": kategori, "marka": marka, "urun": urun, "adet": f(adet), "satis": D(satis), "smm": D(smm), "stok": f(stok), "onceki": f(onceki)}


def tek(*a, **kw):
    return main.analiz([u("X", *a, **kw)])["urunler"][0]


class AnalizTesti(unittest.TestCase):
    def test_gostergeler(self):
        x = tek(200000, 140000, 70000, 180000)
        self.assertEqual((x["kar"], x["marj"], x["stok_gun"], x["devir"], x["gmroi"], x["buyume"], x["isaret"], x["performans"]),
                         (D(60000), D("30.00"), D("15.00"), D("2.00"), D("0.86"), D("11.11"), [], "İyi"))
        x = tek(100, 80)
        self.assertEqual((x["stok_gun"], x["devir"], x["gmroi"], x["buyume"]), (None, None, None, None))
        self.assertEqual(main.analiz([u("X", 100, 50, 100)], donem_gun=D(90))["urunler"][0]["stok_gun"], D("180.00"))

    def test_isaretler_ve_performans(self):
        self.assertEqual([tek(100, m)["isaret"] for m in (85, "85.01")], [[], ["Düşük marj"]])                     # tam %15 işaretlenmez
        self.assertEqual([tek(100, 50, s)["isaret"] for s in (150, 151)], [[], ["Yavaş dönen stok"]])              # tam 90 gün işaretlenmez
        self.assertEqual([tek(s, 50, onceki=100)["isaret"] for s in (81, 80)], [[], ["Satış düşüyor"]])            # tam %20 düşüş işaretlenir
        x = tek(60000, 54000, 216000, 90000)
        self.assertEqual((x["isaret"], x["performans"]), (["Düşük marj", "Yavaş dönen stok", "Satış düşüyor"], "Düşük"))
        self.assertEqual(tek(100, 50, 151)["performans"], "İzlenecek")
        x = tek(0, 0, 15000, 8000)
        self.assertEqual((x["isaret"], x["performans"], x["marj"], x["abc"]), (["Satış yok"], "Düşük", None, "C"))
        self.assertEqual(main.analiz([u("X", 100, 90)], marj_esik=D(5))["urunler"][0]["isaret"], [])

    def test_gruplar_pay_ve_abc(self):
        l = [u("P1", 300, 240, 120, kategori="Çay", marka="C"), u("P2", 200, 140, 70), u("P3", 120, 96, 48), u("P4", 80, 56, 28, marka="B"), u("P5", 60, 54, 216, marka="B"),
             u("P6", 20, 13, 52, kategori="Çay", marka="C"), u("P7", 0, 0, 15, kategori="Çay", marka="D")]
        r = main.analiz(l, {"kahve": D(4600), "cay": D(8000)})
        self.assertEqual({x["urun"]: x["abc"] for x in r["urunler"]}, {"P1": "A", "P2": "A", "P3": "A", "P4": "B", "P5": "C", "P6": "C", "P7": "C"})   # birikimli %38 / 64 / 79 / 90 / 97 / 100
        k = {x["kategori"]: (x["satis"], x["marj"], x["stok_gun"], x["pay"], x["pazar_payi"], x["urun"]) for x in r["kategoriler"]}
        self.assertEqual(k, {"Kahve": (D(460), D("24.78"), D("31.39"), D("58.97"), D("10.00"), 4), "Çay": (D(320), D("20.94"), D("22.17"), D("41.03"), D("4.00"), 3)})
        m = {(x["kategori"], x["marka"]): (x["satis"], x["pay"]) for x in r["markalar"]}
        self.assertEqual(m, {("Kahve", "A"): (D(320), D("69.57")), ("Kahve", "B"): (D(140), D("30.43")), ("Çay", "C"): (D(320), D("100.00")), ("Çay", "D"): (D(0), D("0.00"))})
        g = r["genel"]
        self.assertEqual((g["satis"], g["kar"], g["marj"], g["stok"], g["stok_gun"]), (D(780), D(181), D("23.21"), D(549), D("27.50")))
        self.assertEqual(next(x["pay"] for x in r["urunler"] if x["urun"] == "P2"), D("43.48"))
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "urun_satislari.csv", cikti, ORNEK / "pazar_verisi.csv")
            self.assertEqual({x["urun"].split()[0]: (x["performans"], x["isaret"]) for x in r["urunler"]},
                             {"Granül": ("Düşük", ["Düşük marj", "Yavaş dönen stok", "Satış düşüyor"]), "Yeşil": ("Düşük", ["Satış yok"]), "Bitki": ("İzlenecek", ["Yavaş dönen stok"]),
                              "Siyah": ("İyi", []), "Filtre": ("İyi", []), "Türk": ("İyi", []), "Kapsül": ("İyi", [])})
            k = {x["kategori"]: (x["satis"], x["marj"], x["buyume"], x["pazar_payi"]) for x in r["kategoriler"]}
            self.assertEqual(k, {"Kahve": (D("460000.00"), D("24.78"), D("1.10"), D("10.00")), "Çay": (D("320000.00"), D("20.94"), D("-5.60"), D("4.00"))})
            self.assertEqual((r["genel"]["marj"], r["genel"]["stok_gun"], r["say"], r["uyarilar"]), (D("23.21"), D("27.50"), {"Düşük": 2, "İzlenecek": 1, "İyi": 4}, []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kategori Özeti", "Marka Bazında", "Ürün Bazında", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, pz = Path(tmp) / "s.csv", Path(tmp) / "p.csv"
            y.write_text("Ürün;Net Satış;SMM;Stok;Önceki Dönem Satış\nA;1.000,00;1.200,00;çok;\nB;yok;5;;\nC;500;300;100;-5\n", encoding="utf-8")
            pz.write_text("Kategori;Pazar Cirosu\nKahve;sıfır\nÇay;1.000\n", encoding="utf-8")
            l, p, uy = main.oku(y, pz)
            self.assertEqual(([(x["urun"], x["satis"], x["stok"], x["onceki"], x["kategori"]) for x in l], p),
                             ([("A", D("1000.00"), None, None, "Belirtilmemiş"), ("C", D(500), D(100), None, "Belirtilmemiş")], {"cay": D(1000)}))
            self.assertEqual(sorted(x["tur"] for x in uy), ["Değer okunamadı", "Değer okunamadı", "Maliyet satıştan büyük", "Pazar cirosu okunamadı", "Satış veya maliyet okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--satislar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--pazar", str(pz), "--donem-gun", "90", "--marj-esik", "25"]), 0)
            self.assertEqual(main.main(arg + ["--donem-gun", "0"]), 2)
            self.assertEqual(main.main(["--satislar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Ürün;Net Satış;SMM\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
