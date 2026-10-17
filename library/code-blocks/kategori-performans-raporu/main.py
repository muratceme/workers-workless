"""
Kategori Performans Raporu — Gerisi Bende kod bloğu
Satın Alma ve Kategori Yönetimi › Kategori Uzmanı

Ürün, marka ve kategori bazında satış, marj, stok devir hızı ve payı raporlar; düşük performanslı ürünleri işaretler.
  Brüt kâr        = net satış − satılan malın maliyeti (SMM);   marj % = brüt kâr ÷ net satış × 100
  Stok gün        = stok (maliyetle) ÷ SMM × dönem gün sayısı (--donem-gun, varsayılan 30)
  Stok devir hızı = SMM ÷ stok                                   (dönem içi; stok dönem sonu veya ortalama stok olabilir)
  GMROI           = brüt kâr ÷ stok
  Büyüme %        = (net satış − önceki dönem satış) ÷ önceki dönem satış × 100
  Kategori payı   = ciro ÷ kategorinin cirosu;   pazar payı = kategori cirosu ÷ verdiğiniz pazar cirosu
  ABC             = ürünler ciroya göre sıralanır; birikimli pay ≤ %80 A, ≤ %95 B, kalanı C
  İşaretler       = Düşük marj (< --marj-esik 15) · Yavaş dönen stok (stok gün > --stok-gun 90) · Satış düşüyor (büyüme ≤ −--dusus-esik 20) · Satış yok
  Performans      = "Satış yok" veya iki ve üzeri işaret: Düşük;  tek işaret: İzlenecek;  işaret yok: İyi
Eşikler kategoriye göre değişir; varsayılanlar genel bir ölçüt değildir. Pazar verisi koda gömülü değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"kategori": ("kategori",), "marka": ("marka",), "urun": ("urun", "urun adi", "sku", "stok kodu"), "adet": ("satis adedi", "adet"),
         "satis": ("net satis", "ciro", "satis tutari", "satis"), "smm": ("smm", "satilan malin maliyeti", "maliyet"), "stok": ("stok", "stok maliyet", "stok degeri", "ortalama stok", "donem sonu stok"),
         "onceki": ("onceki donem satis", "onceki satis", "gecen donem satis", "onceki donem")}
PAZAR_SUTUN = {"kategori": ("kategori",), "pazar": ("pazar cirosu", "pazar buyuklugu", "toplam pazar", "pazar")}
Z = Decimal(0)


def oku(yol: Path, pazar_yol: Path | None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, pazar = [], [], {}
    for r in kayitlar(yol, SUTUN, ("urun", "satis", "smm")):
        u = metin(r["urun"])
        if not u:
            continue
        kim = f"{u} (satır {r['_satir']})"
        v = {a: para(r[a]) for a in ("adet", "satis", "smm", "stok", "onceki")}
        if v["satis"] is None or v["smm"] is None or v["satis"] < 0 or v["smm"] < 0:
            uy.append(uyari("Yüksek", "Satış veya maliyet okunamadı", kim, f"Net satış '{metin(r['satis'])}', SMM '{metin(r['smm'])}'. Ürün rapora alınmadı."))
            continue
        for a, ad in (("stok", "Stok"), ("onceki", "Önceki dönem satış"), ("adet", "Satış adedi")):
            if v[a] is None and metin(r[a]) or (v[a] is not None and v[a] < 0):
                uy.append(uyari("Orta", "Değer okunamadı", kim, f"{ad} '{metin(r[a])}'. Boş sayıldı."))
                v[a] = None
        if v["smm"] > v["satis"] > 0:
            uy.append(uyari("Orta", "Maliyet satıştan büyük", kim, f"SMM {v['smm']:f} > net satış {v['satis']:f}: ürün zararına satılmış görünüyor. KDV dahil / hariç tutarlılığını kontrol edin."))
        l.append({"kategori": metin(r["kategori"]) or "Belirtilmemiş", "marka": metin(r["marka"]) or "Belirtilmemiş", "urun": u} | v)
    if pazar_yol:
        for r in kayitlar(pazar_yol, PAZAR_SUTUN, ("kategori", "pazar")):
            k, p = metin(r["kategori"]), para(r["pazar"])
            if not k:
                continue
            if p is None or p <= 0:
                uy.append(uyari("Yüksek", "Pazar cirosu okunamadı", k, f"{pazar_yol.name} satır {r['_satir']}: '{metin(r['pazar'])}'. Pazar payı hesaplanmadı."))
                continue
            pazar[katla(k)] = p
    return l, pazar, uy


def olcut(x: dict, donem_gun: Decimal) -> dict:
    """satis, smm, stok, onceki alanlarından türetilen göstergeler."""
    kar = x["satis"] - x["smm"]
    stok, smm = x.get("stok"), x["smm"]
    return {"kar": kar, "marj": yuvarla(kar / x["satis"] * 100) if x["satis"] else None, "stok_gun": yuvarla(stok / smm * donem_gun) if stok is not None and smm else None,
            "devir": yuvarla(smm / stok) if stok else None, "gmroi": yuvarla(kar / stok) if stok else None,
            "buyume": yuvarla((x["satis"] - x["onceki"]) / x["onceki"] * 100) if x.get("onceki") else None}


def analiz(urunler: list[dict], pazar: dict | None = None, donem_gun: Decimal = Decimal(30), marj_esik: Decimal = Decimal(15), stok_gun: Decimal = Decimal(90), dusus_esik: Decimal = Decimal(20)) -> dict:
    if not urunler:
        raise ValueError("Raporlanacak ürün yok.")
    pazar = pazar or {}
    kat_ciro = defaultdict(lambda: Z)
    for u in urunler:
        kat_ciro[katla(u["kategori"])] += u["satis"]
    toplam = sum((u["satis"] for u in urunler), Z)
    for u in urunler:
        u |= olcut(u, donem_gun)
        u["pay"] = yuvarla(u["satis"] / kat_ciro[katla(u["kategori"])] * 100) if kat_ciro[katla(u["kategori"])] else Z
        i = []
        if u["satis"] == 0:
            i.append("Satış yok")
        else:
            if u["marj"] is not None and u["marj"] < marj_esik:
                i.append("Düşük marj")
            if u["stok_gun"] is not None and u["stok_gun"] > stok_gun:
                i.append("Yavaş dönen stok")
        if u["buyume"] is not None and u["buyume"] <= -dusus_esik and u["satis"] > 0:
            i.append("Satış düşüyor")
        u["isaret"] = i
        u["performans"] = "Düşük" if "Satış yok" in i or len(i) >= 2 else "İzlenecek" if i else "İyi"
    bir = Z
    for u in sorted(urunler, key=lambda u: (-u["satis"], u["urun"])):
        bir += u["satis"]
        b = bir / toplam * 100 if toplam else Decimal(100)
        u["abc"] = "A" if b <= 80 and u["satis"] > 0 else "B" if b <= 95 and u["satis"] > 0 else "C"

    def grupla(anahtar, ad) -> list[dict]:
        g = {}
        for u in urunler:
            x = g.setdefault(anahtar(u), ad(u) | {"urun": 0, "dusuk": 0, "satis": Z, "smm": Z, "stok": Z, "onceki": Z, "stoklu": 0, "oncekili": 0})
            x["urun"] += 1
            x["dusuk"] += u["performans"] == "Düşük"
            x["satis"] += u["satis"]
            x["smm"] += u["smm"]
            x["stok"] += u["stok"] or Z
            x["onceki"] += u["onceki"] or Z
            x["stoklu"] += u["stok"] is not None
            x["oncekili"] += u["onceki"] is not None
        for x in g.values():
            x |= olcut({"satis": x["satis"], "smm": x["smm"], "stok": x["stok"] if x["stoklu"] else None, "onceki": x["onceki"] if x["oncekili"] else None}, donem_gun)
        return list(g.values())

    kategoriler = grupla(lambda u: katla(u["kategori"]), lambda u: {"kategori": u["kategori"]})
    for k in kategoriler:
        p = pazar.get(katla(k["kategori"]))
        k["pay"] = yuvarla(k["satis"] / toplam * 100) if toplam else Z
        k["pazar"], k["pazar_payi"] = p, yuvarla(k["satis"] / p * 100) if p else None
    markalar = grupla(lambda u: (katla(u["kategori"]), katla(u["marka"])), lambda u: {"kategori": u["kategori"], "marka": u["marka"]})
    for m in markalar:
        c = kat_ciro[katla(m["kategori"])]
        m["pay"] = yuvarla(m["satis"] / c * 100) if c else Z
    genel = grupla(lambda u: 0, lambda u: {})[0]
    return {"urunler": sorted(urunler, key=lambda u: (("Düşük", "İzlenecek", "İyi").index(u["performans"]), -u["satis"])), "kategoriler": sorted(kategoriler, key=lambda k: -k["satis"]),
            "markalar": sorted(markalar, key=lambda m: (m["kategori"], -m["satis"])), "genel": genel, "esik": (donem_gun, marj_esik, stok_gun, dusus_esik),
            "say": {p: sum(u["performans"] == p for u in urunler) for p in ("Düşük", "İzlenecek", "İyi")}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Düşük": KIRMIZI, "İzlenecek": SARI, "İyi": None}
    dg, me, sg, de = r["esik"]
    ortak_bas = ["Net Satış", "Brüt Kâr", "Marj %", "Stok (maliyet)", "Stok Gün", "Stok Devir Hızı", "GMROI", "Büyüme %"]
    deger = lambda x: [x["satis"], x["kar"], x["marj"], x["stok"], x["stok_gun"], x["devir"], x["gmroi"], x["buyume"]]   # noqa: E731
    ws = wb.active
    ws.title = "Kategori Özeti"
    _baslik(ws, ["Kategori", "Ürün Sayısı", "Düşük Performanslı", *ortak_bas, "Ciro Payı %", "Pazar Cirosu", "Pazar Payı %"], (22, 9, 10, 16, 16, 9, 16, 9, 10, 9, 10, 11, 16, 11))
    for k in r["kategoriler"]:
        _satir(ws, [k["kategori"], k["urun"], k["dusuk"] or None, *deger(k), k["pay"], k["pazar"], k["pazar_payi"]], para_s=(4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14),
               dolgu={3: KIRMIZI if k["dusuk"] else None})
    g = r["genel"]
    _satir(ws, ["Toplam", g["urun"], g["dusuk"] or None, *deger(g), 100], para_s=(4, 5, 6, 7, 8, 9, 10, 11, 12))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Dönem {dg.normalize():f} gün. Eşikler: marj < %{me.normalize():f}, stok gün > {sg.normalize():f}, satış düşüşü ≥ %{de.normalize():f}. Stok dönem sonu değeriyse stok gün ve devir hızı yaklaşık değerdir."])
    ws = wb.create_sheet("Marka Bazında")
    _baslik(ws, ["Kategori", "Marka", "Ürün Sayısı", "Düşük Performanslı", *ortak_bas, "Kategori İçi Pay %"], (22, 22, 9, 10, 16, 16, 9, 16, 9, 10, 9, 10, 12))
    for m in r["markalar"]:
        _satir(ws, [m["kategori"], m["marka"], m["urun"], m["dusuk"] or None, *deger(m), m["pay"]], para_s=(5, 6, 7, 8, 9, 10, 11, 12, 13), dolgu={4: KIRMIZI if m["dusuk"] else None})
    ws = wb.create_sheet("Ürün Bazında")
    _baslik(ws, ["Performans", "Kategori", "Marka", "Ürün", "ABC", "Satış Adedi", *ortak_bas, "Kategori İçi Pay %", "İşaretler", "Aksiyon (siz doldurun)"],
            (12, 20, 20, 34, 6, 10, 16, 16, 9, 16, 9, 10, 9, 10, 11, 36, 30))
    for u in r["urunler"]:
        _satir(ws, [u["performans"], u["kategori"], u["marka"], u["urun"], u["abc"], u["adet"], *deger(u), u["pay"], ", ".join(u["isaret"]), ""], para_s=(7, 8, 9, 10, 11, 12, 13, 14, 15),
               dolgu={1: renk[u["performans"]]}, sar=(16,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, pazar_yol: Path | None = None, **esik) -> dict:
    urunler, pazar, uy = oku(yol, pazar_yol)
    r = analiz(urunler, pazar, **esik)
    for k in pazar:
        if k not in {katla(x["kategori"]) for x in r["kategoriler"]}:
            uy.append(uyari("Orta", "Pazar verisinin kategorisi yok", k, "Pazar dosyasında var, satış dosyasında bu kategori yok."))
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ürün, marka ve kategori bazında satış, marj, stok devri ve payı raporlar.")
    p.add_argument("--satislar", type=Path, default=ORNEK / "urun_satislari.csv", help="Kategori, Marka, Ürün, Satış Adedi, Net Satış, SMM, Stok (maliyet), Önceki Dönem Satış")
    p.add_argument("--pazar", type=Path, help="İsteğe bağlı: Kategori, Pazar Cirosu (pazar payı için)")
    p.add_argument("--donem-gun", default="30", help="Verinin kapsadığı gün sayısı (varsayılan 30)")
    p.add_argument("--marj-esik", default="15", help="Bu marj %%'sinin altı 'Düşük marj' (varsayılan 15)")
    p.add_argument("--stok-gun", default="90", help="Bu stok gününün üstü 'Yavaş dönen stok' (varsayılan 90)")
    p.add_argument("--dusus-esik", default="20", help="Önceki döneme göre bu %% ve üzeri düşüş 'Satış düşüyor' (varsayılan 20)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kategori_performans.xlsx")
    a = p.parse_args(argv)
    pazar = a.pazar or (ORNEK / "pazar_verisi.csv" if a.satislar == ORNEK / "urun_satislari.csv" else None)
    for y in (a.satislar, pazar):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    e = [para(x) for x in (a.donem_gun, a.marj_esik, a.stok_gun, a.dusus_esik)]
    if any(x is None or x < 0 for x in e) or e[0] == 0:
        print("[X] --donem-gun pozitif; --marj-esik, --stok-gun ve --dusus-esik sıfır veya pozitif sayı olmalı")
        return 2
    try:
        r = calistir(a.satislar, a.cikti, pazar, donem_gun=e[0], marj_esik=e[1], stok_gun=e[2], dusus_esik=e[3])
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g, s = r["genel"], r["say"]
    print(f"[OK] {g['urun']} ürün · net satış {g['satis']:,.2f} · marj %{g['marj']} · stok gün {g['stok_gun'] if g['stok_gun'] is not None else '—'} · {s['Düşük']} düşük, {s['İzlenecek']} izlenecek")
    for u in r["urunler"]:
        if u["performans"] == "Düşük":
            print(f"[!] {u['kategori']} · {u['marka']} · {u['urun']}: {', '.join(u['isaret'])}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
