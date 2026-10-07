"""
Nakit Akış Tahmini — Workers / Workless kod bloğu
Finans › Finans Uzmanı

Vadeli alacak ve borçlar, alınan/verilen çekler, kredi taksitleri ve tekrarlayan ödemelerden (maaş,
kira, vergi, SGK) haftalık nakit akış tahmini (varsayılan 13 hafta) üretir; kategori bazında giriş-çıkış,
hafta sonu bakiyesi ve minimum nakit eşiğinin altına düşen haftaları gösterir. İnternete bağlanmaz.

Varsayımlar (parametreyle değiştirilebilir):
  - Müşteri tahsilatları vadeden --tahsilat-gecikmesi gün sonra gerçekleşir
  - Başlangıç tarihinden önce vadesi geçmiş alacakların --gecikmis-tahsil-orani kadarı 1. haftada tahsil edilir
  - Vadesi geçmiş borçlar 1. haftada ödenir

Kullanım:
    python main.py                                                    # örnek kalemlerle dener
    python main.py --girdi nakit_kalemleri.xlsx --acilis 1250000 --baslangic 05.10.2026 --minimum 250000
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent


def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def para(x) -> Decimal:
    if x in (None, ""):
        return Decimal(0)
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace("₺", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def tarih(x) -> date | None:
    if x in (None, ""):
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(x).strip(), f).date()
        except ValueError:
            pass
    return None


def ay_ekle(t: date, n: int) -> date:
    y, m = divmod(t.month - 1 + n, 12)
    yil, ay = t.year + y, m + 1
    sonraki = date(yil + (ay == 12), ay % 12 + 1, 1)
    return date(yil, ay, min(t.day, (sonraki - timedelta(days=1)).day))


@dataclass
class Akis:
    tur: str            # "Giriş" | "Çıkış"
    kategori: str
    aciklama: str
    tarih: date
    tutar: Decimal
    not_: str = ""


ALANLAR = {
    "tur": ("tür", "yön", "giriş/çıkış"),
    "kategori": ("kategori", "grup"),
    "aciklama": ("açıklama", "cari", "karşı taraf"),
    "tarih": ("tarih", "vade", "vade tarihi", "beklenen tarih"),
    "tutar": ("tutar",),
    "tekrar": ("tekrar", "periyot", "sıklık"),
    "bitis": ("bitiş", "bitiş tarihi", "son tarih"),
}
GIRIS_KELIMELERI = {"giriş", "tahsilat", "gelir", "giris", "+"}


def kalemleri_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in satirlar[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in ALANLAR.items()}
    eksik = [a for a in ("tur", "tarih", "tutar") if k[a] is None]
    if eksik:
        raise SystemExit(f"Gerekli sütunlar: Tür (Giriş/Çıkış), Tarih, Tutar. Başlıklar: {satirlar[0]}")
    al = lambda r, a: r[k[a]] if k[a] is not None and k[a] < len(r) else None  # noqa: E731
    return [{"tur": "Giriş" if kucuk(al(r, "tur")) in GIRIS_KELIMELERI else "Çıkış",
             "kategori": str(al(r, "kategori") or "Diğer").strip(), "aciklama": str(al(r, "aciklama") or "").strip(),
             "tarih": tarih(al(r, "tarih")), "tutar": abs(para(al(r, "tutar"))),
             "tekrar": kucuk(al(r, "tekrar")), "bitis": tarih(al(r, "bitis"))} for r in satirlar[1:]]


def akislari_uret(kalemler: list[dict], baslangic: date, bitis: date, gecikme: int, gecikmis_oran: Decimal) -> tuple[list[Akis], list[str]]:
    akislar, uyarilar = [], []
    for k in kalemler:
        if k["tarih"] is None or k["tutar"] == 0:
            continue
        tarihler = [k["tarih"]]
        if k["tekrar"] in {"haftalık", "haftalik", "aylık", "aylik"}:
            son = min(k["bitis"] or bitis, bitis)
            n, t = 1, k["tarih"]
            while True:
                t = k["tarih"] + timedelta(weeks=n) if k["tekrar"].startswith("haftal") else ay_ekle(k["tarih"], n)
                if t > son:
                    break
                tarihler.append(t)
                n += 1
        tekrarli = len(tarihler) > 1 or k["tekrar"] in {"haftalık", "haftalik", "aylık", "aylik"}
        for t in tarihler:
            if tekrarli and t < baslangic:
                continue            # tekrarlayan kalemin geçmiş dönemleri tahmine girmez
            tutar, not_ = k["tutar"], ""
            if k["tur"] == "Giriş" and kucuk(k["kategori"]) in {"müşteri", "musteri", "müşteri tahsilatı", "alacak"}:
                t = t + timedelta(days=gecikme)
            if t < baslangic:
                if k["tur"] == "Giriş":
                    tutar = (tutar * gecikmis_oran).quantize(Decimal("0.01"))
                    not_ = f"Vadesi geçmiş ({k['tarih']:%d.%m.%Y}); %{gecikmis_oran * 100:g} tahsil varsayımı"
                else:
                    not_ = f"Vadesi geçmiş borç ({k['tarih']:%d.%m.%Y}); 1. haftada ödenir varsayımı"
                t = baslangic
            if t > bitis:
                continue
            akislar.append(Akis(k["tur"], k["kategori"], k["aciklama"], t, tutar, not_))
    gecmis = [a for a in akislar if a.not_.startswith("Vadesi geçmiş")]
    if gecmis:
        uyarilar.append(f"{len(gecmis)} kalem başlangıç tarihinden önce vadeli; 1. haftaya alındı")
    return akislar, uyarilar


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
PARA = "#,##0;[Red]-#,##0"


def calistir(girdi: Path, acilis: Decimal, baslangic: date, cikti: Path, hafta: int = 13, minimum: Decimal = Decimal(0),
             gecikme: int = 0, gecikmis_oran: Decimal = Decimal("0.5")) -> dict:
    baslangic = baslangic - timedelta(days=baslangic.weekday())          # haftanın pazartesisi
    bitis = baslangic + timedelta(weeks=hafta) - timedelta(days=1)
    akislar, uyarilar = akislari_uret(kalemleri_oku(girdi), baslangic, bitis, gecikme, gecikmis_oran)
    haftalar = [baslangic + timedelta(weeks=i) for i in range(hafta)]
    tablo: dict[tuple[str, str], list[Decimal]] = OrderedDict()
    for tur in ("Giriş", "Çıkış"):
        for kat in sorted({a.kategori for a in akislar if a.tur == tur}):
            tablo[(tur, kat)] = [Decimal(0)] * hafta
    for a in akislar:
        tablo[(a.tur, a.kategori)][(a.tarih - baslangic).days // 7] += a.tutar

    acilislar, kapanislar, girisler, cikislar = [], [], [], []
    bakiye = acilis
    for i in range(hafta):
        g = sum((v[i] for (t, _), v in tablo.items() if t == "Giriş"), Decimal(0))
        c = sum((v[i] for (t, _), v in tablo.items() if t == "Çıkış"), Decimal(0))
        acilislar.append(bakiye)
        bakiye = bakiye + g - c
        girisler.append(g)
        cikislar.append(c)
        kapanislar.append(bakiye)
    acik_haftalar = [i for i, b in enumerate(kapanislar) if b < minimum]

    wb = Workbook()
    ws = wb.active
    ws.title = "Haftalık Nakit Akış"
    ws.append(["Kalem"] + [f"H{i + 1} · {h:%d.%m}" for i, h in enumerate(haftalar)] + ["Toplam"])
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI

    def satir(ad, degerler, kalin=False):
        ws.append([ad] + [float(v) for v in degerler] + ([float(sum(degerler, Decimal(0)))] if not ad.startswith(("Açılış", "Kapanış")) else [None]))
        for h in ws[ws.max_row][1:]:
            h.number_format = PARA
            if kalin:
                h.font = Font(bold=True)
        if kalin:
            ws.cell(ws.max_row, 1).font = Font(bold=True)

    satir("Açılış bakiyesi", acilislar, True)
    for tur in ("Giriş", "Çıkış"):
        for (t, kat), v in tablo.items():
            if t == tur:
                satir(f"  {tur} · {kat}", v)
        satir(f"Toplam {tur.lower()}", girisler if tur == "Giriş" else cikislar, True)
    satir("Net akış", [g - c for g, c in zip(girisler, cikislar)], True)
    satir("Kapanış bakiyesi", kapanislar, True)
    kapanis_satiri = ws.max_row
    for i in acik_haftalar:
        ws.cell(kapanis_satiri, i + 2).fill = KIRMIZI
    ws.column_dimensions["A"].width = 34
    for j in range(2, hafta + 3):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.freeze_panes = "B2"

    yardimci = wb.create_sheet("Grafik Verisi")
    yardimci.append(["Hafta", "Giriş", "Çıkış", "Kapanış", "Minimum"])
    for i, h in enumerate(haftalar):
        yardimci.append([f"H{i + 1}", float(girisler[i]), float(cikislar[i]), float(kapanislar[i]), float(minimum)])
    g = BarChart()
    g.title, g.height, g.width = "Haftalık giriş, çıkış ve kapanış bakiyesi", 9, 26
    g.add_data(Reference(yardimci, min_col=2, max_col=3, min_row=1, max_row=hafta + 1), titles_from_data=True)
    g.set_categories(Reference(yardimci, min_col=1, min_row=2, max_row=hafta + 1))
    c = LineChart()
    c.add_data(Reference(yardimci, min_col=4, max_col=5, min_row=1, max_row=hafta + 1), titles_from_data=True)
    g += c
    ws.add_chart(g, f"A{ws.max_row + 3}")

    d = wb.create_sheet("Kalemler")
    d.append(["Hafta", "Tarih", "Tür", "Kategori", "Açıklama", "Tutar", "Not"])
    for h in d[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for a in sorted(akislar, key=lambda x: (x.tarih, x.tur)):
        d.append([f"H{(a.tarih - baslangic).days // 7 + 1}", a.tarih, a.tur, a.kategori, a.aciklama, float(a.tutar), a.not_])
        d.cell(d.max_row, 2).number_format = "DD.MM.YYYY"
        d.cell(d.max_row, 6).number_format = PARA
    for j, w in enumerate((7, 12, 8, 18, 40, 14, 60), 1):
        d.column_dimensions[get_column_letter(j)].width = w

    b = wb.create_sheet("Varsayımlar")
    for s in [["Başlangıç (pazartesi)", baslangic.strftime("%d.%m.%Y")], ["Süre", f"{hafta} hafta"],
              ["Açılış nakit", float(acilis)], ["Minimum nakit eşiği", float(minimum)],
              ["Müşteri tahsilat gecikmesi", f"{gecikme} gün"], ["Vadesi geçmiş alacak tahsil oranı", f"%{gecikmis_oran * 100:g}"],
              ["Eşiğin altındaki haftalar", ", ".join(f"H{i + 1}" for i in acik_haftalar) or "Yok"]] + [["Uyarı", u] for u in uyarilar]:
        b.append(s)
    b.column_dimensions["A"].width = 34
    b.column_dimensions["B"].width = 80
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    return {"haftalar": haftalar, "kapanis": kapanislar, "giris": girisler, "cikis": cikislar,
            "acik_haftalar": acik_haftalar, "akislar": akislar, "tablo": tablo}


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Haftalık nakit akış tahmini.")
    ap.add_argument("--girdi", type=Path, default=BURASI / "ornek_veri" / "nakit_kalemleri.csv")
    ap.add_argument("--acilis", type=para, default=None, help="Başlangıç nakit (banka + kasa)")
    ap.add_argument("--baslangic", help="Başlangıç tarihi (varsayılan bu hafta)")
    ap.add_argument("--hafta", type=int, default=13)
    ap.add_argument("--minimum", type=para, default=Decimal(0), help="Minimum nakit eşiği")
    ap.add_argument("--tahsilat-gecikmesi", type=int, default=0, help="Müşteri tahsilatlarının vadeye göre ortalama gecikmesi (gün)")
    ap.add_argument("--gecikmis-tahsil-orani", type=float, default=50, help="Vadesi geçmiş alacakların 1. haftada tahsil edileceği varsayılan oran (yüzde)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "nakit_akis.xlsx")
    a = ap.parse_args(argv)
    ornek = a.girdi == BURASI / "ornek_veri" / "nakit_kalemleri.csv"
    bas = tarih(a.baslangic) if a.baslangic else (date(2026, 10, 5) if ornek else date.today())
    acilis = a.acilis if a.acilis is not None else (Decimal(850000) if ornek else Decimal(0))
    minimum = a.minimum or (Decimal(250000) if ornek else Decimal(0))
    s = calistir(a.girdi, acilis, bas, a.cikti, a.hafta, minimum, a.tahsilat_gecikmesi, Decimal(str(a.gecikmis_tahsil_orani)) / 100)
    tl = lambda x: f"{x:,.0f}".replace(",", ".")  # noqa: E731
    en_dusuk = min(range(len(s["kapanis"])), key=lambda i: s["kapanis"][i])
    print(f"[OK] {len(s['haftalar'])} hafta · en düşük bakiye H{en_dusuk + 1}: {tl(s['kapanis'][en_dusuk])} TL · "
          f"dönem sonu {tl(s['kapanis'][-1])} TL")
    print(f"[{'!' if s['acik_haftalar'] else 'OK'}] Eşiğin altındaki haftalar: {', '.join(f'H{i + 1}' for i in s['acik_haftalar']) or 'yok'}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
