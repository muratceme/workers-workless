import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
UZAK = date(2028, 1, 1)
KIS = ((12, 1), (4, 1))


def a(plaka="P", muayene=UZAK, trafik=UZAK, kasko=UZAK, bakim_tarih=None, bakim_km=None, km=None, aralik_km=None, aralik_ay=None, lastik="", kullanici="K"):
    n = lambda x: None if x is None else D(x)   # noqa: E731
    return {"plaka": plaka, "model": "", "kullanici": kullanici, "muayene": muayene, "trafik": trafik, "kasko": kasko, "bakim_tarih": bakim_tarih, "bakim_km": n(bakim_km), "km": n(km), "aralik_km": n(aralik_km), "aralik_ay": aralik_ay,
            "lastik": lastik}


def tek(*, gun=30, km_esik=D(1000), kis=None, **kw):
    r = main.degerlendir([a(**kw)], BUGUN, gun, km_esik, kis)
    return r["araclar"][0], sorted((u["onem"], u["tur"]) for u in r["uyarilar"] if u["tur"] != "Bakım bilgisi yok")


class BelgeTesti(unittest.TestCase):
    def test_belge_durumu(self):
        self.assertEqual([(tek(muayene=t)[0]["muayene_durum"], tek(muayene=t)[0]["muayene_kalan"], tek(muayene=t)[0]["durum"]) for t in (date(2026, 10, 11), BUGUN, date(2026, 11, 11), date(2026, 11, 12))],
                         [("Geçmiş", -1, "Acil"), ("Yaklaşıyor", 0, "Yaklaşan"), ("Yaklaşıyor", 30, "Yaklaşan"), ("Uygun", 31, "Uygun")])
        self.assertEqual((tek(muayene=date(2026, 10, 11))[1], tek(trafik=date(2026, 10, 11))[1], tek(kasko=date(2026, 10, 11))[1]),
                         ([("Yüksek", "Muayene süresi geçmiş")], [("Yüksek", "Trafik sigortası süresi geçmiş")], [("Orta", "Kasko süresi geçmiş")]))
        x, uy = tek(muayene=None, trafik=None, kasko=None)
        self.assertEqual((x["muayene_durum"], x["kasko_durum"], x["durum"], uy), ("Bilgi yok", "Bilgi yok", "Uygun", [("Bilgi", "Kasko bilgisi yok"), ("Orta", "Muayene bilgisi yok"), ("Orta", "Trafik sigortası bilgisi yok")]))
        self.assertEqual(tek(trafik=date(2026, 11, 20), gun=45)[1], [("Orta", "Trafik sigortası süresi yaklaşıyor")])

    def test_bakim(self):
        b = lambda **kw: (lambda x, uy: (x["bakim_durum"], x["sonraki_km"], x["kalan_km"], x["sonraki_tarih"], [t for _, t in uy]))(*tek(**kw))   # noqa: E731
        self.assertEqual(b(bakim_km=80000, km=95400, aralik_km=15000), ("Geçmiş", D(95000), D(-400), None, ["Bakım zamanı geçmiş"]))
        self.assertEqual([b(bakim_km=80000, km=k, aralik_km=15000)[0] for k in (95000, 94000, 93999)], ["Yaklaşıyor", "Yaklaşıyor", "Uygun"])
        self.assertEqual(b(bakim_km=80000, km=93000, aralik_km=15000, km_esik=D(2000))[0], "Yaklaşıyor")
        self.assertEqual([b(bakim_tarih=t, aralik_ay=12)[::3] for t in (date(2025, 10, 11), date(2025, 10, 12), date(2025, 11, 11), date(2025, 11, 12))],
                         [("Geçmiş", date(2026, 10, 11)), ("Yaklaşıyor", date(2026, 10, 12)), ("Yaklaşıyor", date(2026, 11, 11)), ("Uygun", date(2026, 11, 12))])
        self.assertEqual(b(bakim_tarih=date(2026, 4, 20), bakim_km=120000, km=133500, aralik_km=15000, aralik_ay=6), ("Yaklaşıyor", D(135000), D(1500), date(2026, 10, 20), ["Bakım yaklaşıyor"]))
        self.assertEqual(b(bakim_tarih=date(2026, 6, 15), bakim_km=60000, km=59000, aralik_km=15000, aralik_ay=12), ("Uygun", D(75000), None, date(2027, 6, 15), ["Km tutarsız"]))
        self.assertEqual((b(km=12000), b(bakim_km=1000, km=2000), b(bakim_tarih=date(2026, 1, 1))), (("Bilgi yok", None, None, None, []),) * 3)
        self.assertEqual((main.ay_ekle(date(2025, 8, 31), 6), main.ay_ekle(date(2026, 11, 5), 12), main.ay_ekle(date(2026, 10, 31), 1)), (date(2026, 2, 28), date(2027, 11, 5), date(2026, 11, 30)))

    def test_lastik(self):
        k = lambda gun, ay: main.kis_donemi(date(2026, ay, gun), *KIS)   # noqa: E731
        self.assertEqual([k(30, 11), k(1, 12), k(15, 1), k(31, 3), k(1, 4), k(12, 10)], [(False, 1), (True, 0), (True, 320), (True, 245), (False, 244), (False, 50)])
        self.assertEqual(main.kis_donemi(date(2026, 7, 1), (5, 1), (9, 1)), (True, 304))
        l = lambda lastik, tarih=BUGUN, gun=30: [u["tur"] for u in main.degerlendir([a(lastik=lastik)], tarih, gun, D(1000), KIS)["uyarilar"] if "astiğ" in u["tur"]]   # noqa: E731
        self.assertEqual([l("Yaz"), l("Kış"), l("Dört Mevsim"), l(""), l("Yaz", date(2026, 12, 15)), l("Kış", date(2026, 12, 15)), l("Yaz", date(2026, 11, 5)), l("Yaz", gun=50)],
                         [[], ["Kış dönemi dışında kış lastiği"], [], [], ["Kış döneminde yaz lastiği"], [], ["Kış lastiği dönemi yaklaşıyor"], ["Kış lastiği dönemi yaklaşıyor"]])
        self.assertEqual([main.gun_ay(x) for x in ("01.12", "1/4", "15.11.", "31.02", "x", "")], [(12, 1), (4, 1), (11, 15), None, None, None])

    def test_filo(self):
        l = [a("C"), a("B", muayene=date(2026, 10, 20), kasko=date(2026, 9, 1), kullanici="Ali"), a("A", trafik=date(2026, 11, 1)), a("B"), a("D", bakim_tarih=date(2026, 4, 1), aralik_ay=6)]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual(([(x["plaka"], x["durum"]) for x in r["araclar"]], r["say"]), ([("B", "Acil"), ("D", "Acil"), ("A", "Yaklaşan"), ("B", "Uygun"), ("C", "Uygun")], {"Acil": 2, "Yaklaşan": 1, "Uygun": 2}))
        self.assertEqual([(i["tarih"], i["plaka"], i["is"], i["kalan"], i["durum"]) for i in r["isler"]],
                         [(date(2026, 9, 1), "B", "Kasko", -41, "Geçmiş"), (date(2026, 10, 1), "D", "Periyodik bakım", -11, "Geçmiş"), (date(2026, 10, 20), "B", "Muayene", 8, "Yaklaşıyor"),
                          (date(2026, 11, 1), "A", "Trafik sigortası", 20, "Yaklaşıyor")])
        self.assertIn("Plaka tekrar ediyor", [u["tur"] for u in r["uyarilar"]])
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "a.xlsx"
            r = main.calistir(ORNEK / "araclar.csv", cikti, BUGUN, kis=KIS)
            self.assertEqual([(x["plaka"][-3:], x["durum"], x["muayene_durum"], x["trafik_durum"], x["kasko_durum"], x["bakim_durum"], x["kalan_km"]) for x in r["araclar"]],
                             [("102", "Acil", "Geçmiş", "Uygun", "Uygun", "Uygun", D(13800)), ("103", "Acil", "Uygun", "Uygun", "Bilgi yok", "Geçmiş", D(-400)), ("106", "Acil", "Uygun", "Geçmiş", "Uygun", "Uygun", None),
                              ("101", "Yaklaşan", "Uygun", "Yaklaşıyor", "Yaklaşıyor", "Uygun", D(1500)), ("104", "Yaklaşan", "Uygun", "Yaklaşıyor", "Yaklaşıyor", "Yaklaşıyor", D(1500)),
                              ("105", "Uygun", "Uygun", "Uygun", "Uygun", "Bilgi yok", None)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"][-3:]) for u in r["uyarilar"]],
                             [("Yüksek", "Muayene süresi geçmiş", "102"), ("Yüksek", "Trafik sigortası süresi geçmiş", "106"), ("Orta", "Bakım yaklaşıyor", "104"), ("Orta", "Bakım zamanı geçmiş", "103"),
                              ("Orta", "Kasko süresi yaklaşıyor", "101"), ("Orta", "Kasko süresi yaklaşıyor", "104"), ("Orta", "Km tutarsız", "106"), ("Orta", "Trafik sigortası süresi yaklaşıyor", "101"),
                              ("Orta", "Trafik sigortası süresi yaklaşıyor", "104"), ("Bilgi", "Bakım bilgisi yok", "105"), ("Bilgi", "Kasko bilgisi yok", "103"), ("Bilgi", "Kış dönemi dışında kış lastiği", "103")])
            self.assertEqual([(i["tarih"], i["plaka"][-3:], i["is"], i["kalan"]) for i in r["isler"]],
                             [(date(2026, 10, 1), "106", "Trafik sigortası", -11), (date(2026, 10, 5), "102", "Muayene", -7), (date(2026, 10, 20), "101", "Kasko", 8), (date(2026, 10, 20), "101", "Trafik sigortası", 8),
                              (date(2026, 10, 20), "104", "Periyodik bakım", 8), (date(2026, 11, 3), "104", "Kasko", 22), (date(2026, 11, 3), "104", "Trafik sigortası", 22), (date(2026, 11, 5), "103", "Periyodik bakım", 24)])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Araçlar", "Yapılacak İşler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "a.csv"
            y.write_text("Plaka;Araç;Sürücü;Muayene;Sigorta Bitiş;Kasko;Son Bakım;Son Bakım Km;Km;Bakım Periyodu (km);Bakım Periyodu (ay);Lastik Tipi\n34  abc 1;X;Ali;15.03.2027;x;;10.02.2026;45.000;58.500;15000;12;KIŞLIK\n"
                         "06 xy 2;;;;;;;;-5;0;0;lastik\n;;;;;;;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["plaka"], x["kullanici"], x["muayene"], x["trafik"], x["bakim_km"], x["km"], x["aralik_km"], x["aralik_ay"], x["lastik"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("34 ABC 1", "Ali", date(2027, 3, 15), None, D(45000), D(58500), D(15000), 12, "Kış"), ("06 XY 2", "", None, None, None, None, None, None, "lastik")], [("Tarih okunamadı", "34 ABC 1")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--araclar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "45", "--km-esik", "2.000", "--kis-baslangic", "01.12", "--kis-bitis", "01.04"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + x) for x in (["--gun", "-1"], ["--km-esik", "x"], ["--kis-baslangic", "01.12"], ["--kis-baslangic", "01.12", "--kis-bitis", "32.04"], ["--kis-baslangic", "01.12", "--kis-bitis", "1.12"],
                                                            ["--tarih", "x"])], [2] * 6)
            self.assertEqual(main.main(["--araclar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Marka;Km\nA;1\n", "Plaka;Km\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
