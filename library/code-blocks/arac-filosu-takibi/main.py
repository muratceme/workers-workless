"""
Araç Filosu Takibi — Gerisi Bende kod bloğu
İdari İşler › İdari İşler Sorumlusu

Şirket araçlarının muayene, trafik sigortası, kasko, bakım ve lastik durumunu izler; yaklaşan ve geçmiş işleri listeler.
  Belge durumu (muayene, trafik sigortası, kasko)
                      = bitiş --tarih gününden önceyse Geçmiş · --gun (30) gün veya daha az kaldıysa Yaklaşıyor · değilse Uygun
  Sonraki bakım       = son bakım km + bakım aralığı (km)  ve  son bakım tarihi + bakım aralığı (ay); hangisi önce dolarsa
  Bakım durumu        = Geçmiş: güncel km sonraki bakım km'sini aşmış veya sonraki bakım tarihi geçmiş
                        Yaklaşıyor: kalan km ≤ --km-esik (1.000) veya bakım tarihine ≤ --gun gün
  Lastik              = --kis-baslangic ve --kis-bitis (GG.AA) verilirse: dönem içinde yaz lastiği, dönem dışında kış lastiği işaretlenir
  Araç durumu         = herhangi bir kalem Geçmiş ise Acil · Yaklaşıyor ise Yaklaşan · değilse Uygun
Tarihler ve km sizin kaydınızdır. Yasal süreleri, kış lastiği dönemini ve bakım aralıklarını kod bilmez; siz girersiniz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"plaka": ("plaka", "arac plakasi"), "model": ("marka model", "model", "marka", "arac"), "kullanici": ("kullanici", "zimmetli", "surucu", "birim"), "muayene": ("muayene bitis", "muayene", "muayene tarihi", "muayene gecerlilik"),
         "trafik": ("trafik sigortasi bitis", "trafik sigortasi", "zorunlu sigorta bitis", "sigorta bitis"), "kasko": ("kasko bitis", "kasko"), "bakim_tarih": ("son bakim tarihi", "son bakim"), "bakim_km": ("son bakim km", "son bakim kilometresi"),
         "km": ("guncel km", "km", "kilometre"), "aralik_km": ("bakim araligi km", "bakim periyodu km"), "aralik_ay": ("bakim araligi ay", "bakim periyodu ay"), "lastik": ("lastik", "lastik tipi", "takili lastik")}
BELGE = (("muayene", "Muayene", "Yüksek"), ("trafik", "Trafik sigortası", "Yüksek"), ("kasko", "Kasko", "Orta"))
LASTIK = {"yaz": "Yaz", "yazlik": "Yaz", "kis": "Kış", "kislik": "Kış", "dort mevsim": "Dört Mevsim", "4 mevsim": "Dört Mevsim", "": ""}
D_RENK = {"Geçmiş": KIRMIZI, "Yaklaşıyor": SARI, "Uygun": YESIL, "Acil": KIRMIZI, "Yaklaşan": SARI}


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def gun_ay(s: str) -> tuple[int, int] | None:
    m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})\.?", (s or "").strip())
    try:
        return (int(m.group(2)), date(2024, int(m.group(2)), int(m.group(1))).day) if m else None
    except ValueError:
        return None


def kis_donemi(bugun: date, bas: tuple[int, int], bit: tuple[int, int]) -> tuple[bool, int]:
    """(bugün kış döneminde mi, dönem başlangıcına kalan gün) döndürür. Dönem başlangıç dahil, bitiş hariçtir."""
    b = (bugun.month, bugun.day)
    icinde = bas <= b < bit if bas < bit else (b >= bas or b < bit)
    sonraki = date(bugun.year + (b > bas), bas[0], bas[1])
    return icinde, (sonraki - bugun).days


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("plaka",)):
        plaka = " ".join(metin(r["plaka"]).upper().split())
        if not plaka:
            continue
        for alan, ad in (("muayene", "Muayene bitiş"), ("trafik", "Trafik sigortası bitiş"), ("kasko", "Kasko bitiş"), ("bakim_tarih", "Son bakım tarihi")):
            if metin(r[alan]) and tarih(r[alan]) is None:
                uy.append(uyari("Orta", "Tarih okunamadı", plaka, f"{ad} '{metin(r[alan])}' okunamadı; bilgi yok sayıldı."))
        sayi = {a: (v if v is not None and v >= 0 else None) for a in ("bakim_km", "km", "aralik_km", "aralik_ay") for v in [para(r[a])]}
        ls = katla(r["lastik"])
        l.append({"plaka": plaka, "model": metin(r["model"]), "kullanici": metin(r["kullanici"]), "muayene": tarih(r["muayene"]), "trafik": tarih(r["trafik"]), "kasko": tarih(r["kasko"]), "bakim_tarih": tarih(r["bakim_tarih"]),
                  "bakim_km": sayi["bakim_km"], "km": sayi["km"], "aralik_km": sayi["aralik_km"] or None, "aralik_ay": int(sayi["aralik_ay"]) if sayi["aralik_ay"] else None, "lastik": LASTIK.get(ls, metin(r["lastik"]))})
    return l, uy


def degerlendir(araclar: list[dict], bugun: date, gun: int = 30, km_esik: Decimal = Decimal(1000), kis: tuple | None = None) -> dict:
    if not araclar:
        raise ValueError("Listede araç yok.")
    uy, isler, gorulen = [], [], set()
    for a in araclar:
        p = a["plaka"]
        if p in gorulen:
            uy.append(uyari("Orta", "Plaka tekrar ediyor", p, "Aynı plaka birden çok satırda var."))
        gorulen.add(p)
        durumlar = []
        for alan, ad, onem in BELGE:
            t = a[alan]
            if t is None:
                a[alan + "_durum"], a[alan + "_kalan"] = "Bilgi yok", None
                uy.append(uyari("Bilgi" if alan == "kasko" else "Orta", f"{ad} bilgisi yok", p, f"{ad} bitiş tarihi yazılmamış." + (" Kasko yoksa boş bırakabilirsiniz." if alan == "kasko" else " Belgeyi kontrol edip tarihi girin.")))
                continue
            kalan = (t - bugun).days
            a[alan + "_kalan"], a[alan + "_durum"] = kalan, "Geçmiş" if kalan < 0 else "Yaklaşıyor" if kalan <= gun else "Uygun"
            durumlar.append(a[alan + "_durum"])
            if kalan < 0:
                uy.append(uyari(onem, f"{ad} süresi geçmiş", p, f"{ad} {t:%d.%m.%Y} tarihinde bitti ({-kalan} gün önce)." + (" Araç bu hâliyle trafiğe çıkarılmamalıdır." if onem == "Yüksek" else "")))
                isler.append({"plaka": p, "is": ad, "tarih": t, "kalan": kalan, "durum": "Geçmiş", "not": a["kullanici"]})
            elif kalan <= gun:
                uy.append(uyari("Orta", f"{ad} süresi yaklaşıyor", p, f"{ad} {t:%d.%m.%Y} tarihinde bitiyor ({kalan} gün)."))
                isler.append({"plaka": p, "is": ad, "tarih": t, "kalan": kalan, "durum": "Yaklaşıyor", "not": a["kullanici"]})
        a["sonraki_km"] = a["bakim_km"] + a["aralik_km"] if a["bakim_km"] is not None and a["aralik_km"] else None
        a["sonraki_tarih"] = ay_ekle(a["bakim_tarih"], a["aralik_ay"]) if a["bakim_tarih"] and a["aralik_ay"] else None
        a["kalan_km"] = None
        if a["km"] is not None and a["bakim_km"] is not None and a["km"] < a["bakim_km"]:
            uy.append(uyari("Orta", "Km tutarsız", p, f"Güncel km ({a['km']:,.0f}) son bakım km'sinden ({a['bakim_km']:,.0f}) küçük. Kaydı düzeltin; km'ye göre bakım hesabı yapılmadı."))
        elif a["sonraki_km"] is not None and a["km"] is not None:
            a["kalan_km"] = a["sonraki_km"] - a["km"]
        kalan_gun = (a["sonraki_tarih"] - bugun).days if a["sonraki_tarih"] else None
        if a["kalan_km"] is None and kalan_gun is None:
            a["bakim_durum"] = "Bilgi yok"
            uy.append(uyari("Bilgi", "Bakım bilgisi yok", p, "Son bakım ve bakım aralığı yazılmadığı için sonraki bakım hesaplanamadı."))
        else:
            gecmis = (a["kalan_km"] is not None and a["kalan_km"] < 0) or (kalan_gun is not None and kalan_gun < 0)
            yakin = (a["kalan_km"] is not None and a["kalan_km"] <= km_esik) or (kalan_gun is not None and kalan_gun <= gun)
            a["bakim_durum"] = "Geçmiş" if gecmis else "Yaklaşıyor" if yakin else "Uygun"
            durumlar.append(a["bakim_durum"])
            ayrinti = " · ".join(x for x in (f"sonraki bakım {a['sonraki_km']:,.0f} km, kalan {a['kalan_km']:,.0f} km" if a["kalan_km"] is not None else "",
                                             f"sonraki bakım tarihi {a['sonraki_tarih']:%d.%m.%Y} ({kalan_gun} gün)" if kalan_gun is not None else "") if x)
            if gecmis or yakin:
                uy.append(uyari("Orta", "Bakım zamanı geçmiş" if gecmis else "Bakım yaklaşıyor", p, ayrinti[0].upper() + ayrinti[1:] + "."))
                isler.append({"plaka": p, "is": "Periyodik bakım", "tarih": a["sonraki_tarih"], "kalan": kalan_gun, "durum": a["bakim_durum"], "not": ayrinti})
        if kis and a["lastik"] in ("Yaz", "Kış"):
            icinde, kalan = kis_donemi(bugun, *kis)
            if icinde and a["lastik"] == "Yaz":
                uy.append(uyari("Orta", "Kış döneminde yaz lastiği", p, "Belirttiğiniz kış lastiği döneminde araçta yaz lastiği görünüyor. Zorunluluk kapsamındaysa değiştirin."))
            elif not icinde and a["lastik"] == "Kış":
                uy.append(uyari("Bilgi", "Kış dönemi dışında kış lastiği", p, "Araçta kış lastiği takılı görünüyor; mevsime göre değişimi planlayın."))
            elif not icinde and kalan <= gun:
                uy.append(uyari("Bilgi", "Kış lastiği dönemi yaklaşıyor", p, f"Belirttiğiniz dönemin başlamasına {kalan} gün var; lastik değişimini planlayın."))
        a["durum"] = "Acil" if "Geçmiş" in durumlar else "Yaklaşan" if "Yaklaşıyor" in durumlar else "Uygun"
    sira = {"Acil": 0, "Yaklaşan": 1, "Uygun": 2}
    araclar.sort(key=lambda a: (sira[a["durum"]], a["plaka"]))
    isler.sort(key=lambda i: (i["tarih"] or date.max, i["plaka"], i["is"]))
    return {"araclar": araclar, "isler": isler, "say": {d: sum(a["durum"] == d for a in araclar) for d in sira}, "bugun": bugun, "gun": gun, "kis": kis, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Araçlar"
    _baslik(ws, ["Durum", "Plaka", "Marka / Model", "Kullanıcı", "Muayene Bitiş", "Muayene", "Trafik Sigortası Bitiş", "Trafik Sigortası", "Kasko Bitiş", "Kasko", "Güncel Km", "Sonraki Bakım Km", "Kalan Km", "Sonraki Bakım Tarihi", "Bakım",
                 "Lastik"], (10, 13, 22, 24, 12, 12, 12, 12, 12, 12, 11, 11, 10, 12, 12, 12))
    for a in r["araclar"]:
        _satir(ws, [a["durum"], a["plaka"], a["model"], a["kullanici"], a["muayene"], a["muayene_durum"], a["trafik"], a["trafik_durum"], a["kasko"], a["kasko_durum"], a["km"], a["sonraki_km"], a["kalan_km"], a["sonraki_tarih"],
                    a["bakim_durum"], a["lastik"]], tarih_s=(5, 7, 9, 14),
               dolgu={1: D_RENK.get(a["durum"]), 6: D_RENK.get(a["muayene_durum"]), 8: D_RENK.get(a["trafik_durum"]), 10: D_RENK.get(a["kasko_durum"]), 15: D_RENK.get(a["bakim_durum"])})
        for j in (11, 12, 13):
            ws.cell(ws.max_row, j).number_format = "#,##0"
    ws.append([])
    ws.append([f"Durum tarihi {r['bugun']:%d.%m.%Y}; {r['gun']} gün veya daha az kalanlar 'Yaklaşıyor' sayılır." + ("" if r["kis"] else " Kış lastiği dönemi verilmedi; lastik denetimi yapılmadı.")])
    ws = wb.create_sheet("Yapılacak İşler")
    _baslik(ws, ["Tarih", "Kalan Gün", "Durum", "Plaka", "İş", "Not", "Yapıldı (siz doldurun)"], (12, 9, 12, 13, 20, 60, 20))
    for i in r["isler"]:
        _satir(ws, [i["tarih"], i["kalan"], i["durum"], i["plaka"], i["is"], i["not"], ""], tarih_s=(1,), dolgu={3: D_RENK.get(i["durum"])}, sar=(6,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 30, km_esik: Decimal = Decimal(1000), kis: tuple | None = None) -> dict:
    araclar, uy = oku(yol)
    r = degerlendir(araclar, bugun, gun, km_esik, kis)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Şirket araçlarının muayene, sigorta, kasko, bakım ve lastik durumunu izler.")
    p.add_argument("--araclar", type=Path, help="Plaka, Marka Model, Kullanıcı, Muayene Bitiş, Trafik Sigortası Bitiş, Kasko Bitiş, Son Bakım Tarihi, Son Bakım Km, Güncel Km, Bakım Aralığı (km), Bakım Aralığı (ay), Lastik")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=30, help="Kaç gün önceden uyarılsın (varsayılan 30)")
    p.add_argument("--km-esik", default="1000", help="Bakıma bu kadar km veya daha az kaldıysa uyarılır (varsayılan 1000)")
    p.add_argument("--kis-baslangic", help="Kış lastiği döneminin başlangıcı (GG.AA); --kis-bitis ile birlikte verilir")
    p.add_argument("--kis-bitis", help="Kış lastiği döneminin bitişi (GG.AA)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "arac_filosu_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.araclar is None
    yol = ORNEK / "araclar.csv" if ornek else a.araclar
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    kb, kt = (a.kis_baslangic or "01.12", a.kis_bitis or "01.04") if ornek else (a.kis_baslangic, a.kis_bitis)
    km_esik, kis = para(a.km_esik), (gun_ay(kb), gun_ay(kt)) if kb or kt else None
    if bugun is None or not 0 <= a.gun <= 365 or km_esik is None or km_esik < 0 or (kis and (None in kis or kis[0] == kis[1])):
        print("[X] --tarih GG.AA.YYYY; --gun 0–365; --km-esik sıfır veya pozitif; --kis-baslangic ve --kis-bitis birlikte ve GG.AA biçiminde (örn. 01.12 ve 01.04) olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun, km_esik, kis)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['araclar'])} araç · Acil {s['Acil']}, Yaklaşan {s['Yaklaşan']}, Uygun {s['Uygun']} · {len(r['isler'])} yapılacak iş")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
