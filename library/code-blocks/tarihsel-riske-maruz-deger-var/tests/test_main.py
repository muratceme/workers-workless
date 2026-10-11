import math
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def gunler(n):
    return [date(2025, 1, 1) + timedelta(days=i) for i in range(n)]


def fiyat_serisi(getiriler, p0=100.0):
    p = [p0]
    for r in getiriler:
        p.append(p[-1] * (1 + r))
    return p


def poz(ad, faktor, deger):
    return {"ad": ad, "faktor": faktor, "deger": float(deger), "portfoy": ""}


class FormulTesti(unittest.TestCase):
    def test_var_es_sirasi(self):
        kz = [-float(i) for i in range(1, 101)]                                     # zararlar 1..100
        self.assertEqual(main.var_es(kz, 99)[:3], (99.0, 99.5, 2))                  # k = ⌊100 × 0,01⌋ + 1 = 2
        self.assertEqual(main.var_es(kz, 95)[:3], (95.0, 97.5, 6))
        self.assertEqual(main.var_es([-float(i) for i in range(1, 251)], 99)[2], 3)  # 250 gün → 3. en kötü
        self.assertEqual(main.var_es([-float(i) for i in range(1, 501)], 99)[2], 6)
        self.assertEqual(main.var_es([1.0, 2.0, 3.0] * 20, 99)[:2], (0.0, 0.0))     # hiç zarar yoksa VaR 0

    def test_getiri(self):
        g = main.getiriler({"X": [100.0, 110.0, 99.0]})["X"]
        self.assertAlmostEqual(g[0], 0.10)
        self.assertAlmostEqual(g[1], -0.10)

    def test_tek_pozisyon(self):
        r_ = [0.001] * 40
        r_[10], r_[20], r_[30] = -0.05, -0.03, -0.02
        r = main.hesapla([poz("A", "X", 1_000_000)], gunler(41), {"X": fiyat_serisi(r_)}, guven=95, pencere=40, elde_tutma=10)
        self.assertEqual((r["n"], r["k"]), (40, 3))                                  # ⌊40 × 0,05⌋ + 1 = 3
        self.assertAlmostEqual(r["var"], 20_000, places=4)
        self.assertAlmostEqual(r["es"], (50_000 + 30_000 + 20_000) / 3, places=4)
        self.assertAlmostEqual(r["var_h"], 20_000 * math.sqrt(10), places=4)
        self.assertEqual(r["var_gunu"], gunler(41)[31])                              # 30. getiri 31. günün kapanışında

    def test_korunma_ve_katkilar(self):
        rx = [0.002 * ((i % 5) - 2) for i in range(60)]
        ry = [0.003 * (((i * 3) % 7) - 3) for i in range(60)]
        f = {"X": fiyat_serisi(rx), "Y": fiyat_serisi(ry)}
        tam = main.hesapla([poz("Uzun", "X", 1000), poz("Kısa", "X", -1000)], gunler(61), f, pencere=60)
        self.assertAlmostEqual(tam["var"], 0.0, places=9)                            # tam korunma
        self.assertAlmostEqual(tam["cesitlendirme"], sum(s["bireysel_var"] for s in tam["satirlar"]), places=9)
        r = main.hesapla([poz("A", "X", 5000), poz("B", "Y", 3000), poz("C", "X", -1000)], gunler(61), f, pencere=60)
        self.assertAlmostEqual(sum(s["es_katki"] for s in r["satirlar"]), r["es"], places=9)
        self.assertAlmostEqual(sum(s["var_senaryosu"] for s in r["satirlar"]), r["var"], places=9)
        self.assertGreaterEqual(r["es"], r["var"])
        self.assertGreaterEqual(r["cesitlendirme"], -1e-9)
        self.assertEqual((r["toplam"], r["brut"]), (7000.0, 9000.0))

    def test_geriye_donuk_test(self):
        r_ = [0.001] * 60
        r_[40], r_[50], r_[55] = -0.05, -0.03, -0.06
        r = main.hesapla([poz("A", "X", 1000)], gunler(61), {"X": fiyat_serisi(r_)}, guven=99, pencere=30)
        t = r["test"]
        self.assertEqual((t["gun"], [a["gun"] for a in t["asim"]], t["bolge"]), (30, [gunler(61)[41], gunler(61)[56]], None))   # −%3, penceredeki −%5'i aşmaz
        self.assertAlmostEqual(t["asim"][1]["var"], 50.0, places=6)
        kisa = main.hesapla([poz("A", "X", 1000)], gunler(41), {"X": fiyat_serisi([0.001] * 40)}, pencere=250)
        self.assertIsNone(kisa["test"])
        self.assertEqual(sorted(u["tur"] for u in kisa["uyarilar"]), ["Fiyat geçmişi kısa", "Geriye dönük test yapılamadı", "Yoğunlaşma"])

    def test_veri_uyarilari(self):
        r_ = [0.001] * 20 + [0.0] * 6 + [0.35] + [0.001] * 20
        r = main.hesapla([poz("A", "X", 1000), poz("B", "X", 900)], gunler(48), {"X": fiyat_serisi(r_)}, pencere=40)
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"] if u["onem"] == "Orta"), ["Değişmeyen fiyat", "Uç getiri"])
        self.assertEqual(main.en_uzun_sifir([0, 0, 1, 0, 0, 0, 2]), 3)
        with self.assertRaises(ValueError):
            main.hesapla([poz("A", "X", 1)], gunler(20), {"X": fiyat_serisi([0.0] * 19)})


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "v.xlsx"
            r = main.calistir(ORNEK / "pozisyonlar.csv", ORNEK / "fiyatlar.csv", cikti, elde_tutma=10)
            self.assertEqual((len(r["satirlar"]), r["n"], r["k"], r["toplam"], r["brut"]), (6, 250, 3, 117_500_000.0, 166_500_000.0))
            self.assertAlmostEqual(sum(s["es_katki"] for s in r["satirlar"]), r["es"], places=4)
            self.assertTrue(0 < r["var"] <= r["es"] < r["brut"] * 0.05)
            self.assertEqual((r["test"]["gun"], r["test"]["bolge"]), (250, "Yeşil"))
            kisa = next(s for s in r["satirlar"] if s["ad"].startswith("Endeks vadeli"))
            self.assertLess(kisa["es_katki"], 0)                                     # kısa pozisyon en kötü günlerde kazandırır
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Pozisyonlar", "Senaryolar", "Geriye Dönük Test", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            p, f = Path(tmp) / "p.csv", Path(tmp) / "f.csv"
            p.write_text("Enstrüman;Risk Faktörü;Piyasa Değeri\nA;X;1.000,00\nB;Y;-500,00\nC;X;sıfır\n", encoding="utf-8")
            satir = ["Tarih;X;Y"] + [f"{d:%d.%m.%Y};{100 + i % 7};{50 + i % 3}" for i, d in enumerate(gunler(45))]
            satir[10] = satir[10].rsplit(";", 1)[0] + ";"                            # bir günde Y yok
            f.write_text("\n".join(satir) + "\n", encoding="utf-8")
            poz_, g, fy, uy = main.oku(p, f)
            self.assertEqual(([x["deger"] for x in poz_], len(g), len(fy["Y"])), ([1000.0, -500.0], 44, 44))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Eksik fiyatlı günler", "Pozisyon değeri okunamadı"])
            arg = ["--pozisyonlar", str(p), "--fiyatlar", str(f), "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg + ["--pencere", "40"]), 0)
            self.assertEqual(main.main(arg + ["--guven", "100"]), 2)
            self.assertEqual(main.main(["--fiyatlar", str(Path(tmp) / "yok.csv")]), 1)
            p.write_text("Enstrüman;Risk Faktörü;Piyasa Değeri\nA;Z;1000\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)                                      # Z için fiyat sütunu yok


if __name__ == "__main__":
    unittest.main()
