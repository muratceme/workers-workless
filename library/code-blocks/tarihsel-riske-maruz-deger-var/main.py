"""
Tarihsel Riske Maruz Değer (VaR) — Gerisi Bende kod bloğu
İç Kontrol ve Risk Yönetimi › Risk Yönetimi Uzmanı

Portföy pozisyonları ve geçmiş fiyat serilerinden tarihsel simülasyonla günlük VaR ve beklenen kaybı (ES) hesaplar.
  Günlük getiri        r(t) = P(t) ÷ P(t−1) − 1                     (her risk faktörü için)
  Senaryo kâr / zararı K/Z(t) = Σ pozisyon değeri × r(t)            (bugünkü pozisyonlar geçmişin her gününe uygulanır)
  VaR (güven α, N senaryo): k = ⌊N × (1 − α)⌋ + 1; en kötü k'ıncı senaryonun zararı.   Ör. N = 250, α = %99 → 3. en kötü gün.
  Beklenen kayıp (ES)  = en kötü k senaryonun ortalama zararı (VaR senaryosu dahil); ES ≥ VaR.
  Elde tutma süresi h gün: VaR(h) = VaR(1) × √h   (karekök kuralı; getirilerin bağımsız olduğu varsayımına dayanan yaklaşım)
  Bireysel VaR         = pozisyon tek başına portföy olsaydı VaR'ı.  Çeşitlendirme etkisi = Σ bireysel VaR − portföy VaR.
  ES katkısı           = pozisyonun, portföyün en kötü k senaryosundaki ortalama zararı (katkıların toplamı ES'e eşittir).
  Geriye dönük test    : her gün için önceki `pencere` günün senaryolarıyla VaR hesaplanır, ertesi günün K/Z'si ile
                         karşılaştırılır; zarar VaR'ı aşarsa "aşım". Pozisyonlar sabit varsayılır (varsayımsal K/Z).
Getiriler yalnız tüm risk faktörlerinin fiyatı olan günlerden hesaplanır. İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 6 pozisyon, 520 iş günü fiyat
    python main.py --pozisyonlar p.xlsx --fiyatlar f.xlsx --guven 99 --pencere 250 --elde-tutma 10
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


POZISYON_SUTUNLARI = {"ad": ("enstruman", "pozisyon", "kiymet", "ad"), "faktor": ("risk faktoru", "fiyat serisi", "faktor", "sembol"),
                      "deger": ("piyasa degeri", "piyasa degeri tl", "deger", "pozisyon degeri"), "portfoy": ("portfoy", "masa", "hesap")}


def oku(p_yol: Path, f_yol: Path) -> tuple[list[dict], list[date], dict, list[dict]]:
    uy, poz = [], []
    for r in kayitlar(p_yol, POZISYON_SUTUNLARI, ("faktor", "deger")):
        if not metin(r["faktor"]):
            continue
        d = para(r["deger"])
        if d is None or d == 0:
            uy.append(uyari("Yüksek", "Pozisyon değeri okunamadı", r["ad"] or r["faktor"], f"Satır {r['_satir']}: '{metin(r['deger'])}'. Pozisyon hesaba katılmadı."))
            continue
        poz.append({"ad": metin(r["ad"]) or metin(r["faktor"]), "faktor": metin(r["faktor"]), "deger": float(d), "portfoy": metin(r["portfoy"])})
    if not poz:
        raise ValueError(f"{p_yol.name}: pozisyon bulunamadı")
    t = tablo_oku(f_yol)
    bi = next((i for i, r in enumerate(t[:15]) if katla(r[0]) in ("tarih", "gun", "date")), None)
    if bi is None:
        raise ValueError(f"{f_yol.name}: ilk sütunu 'Tarih' olan başlık satırı bulunamadı")
    sutun = {katla(c): j for j, c in enumerate(t[bi]) if j and metin(c)}
    eksik = sorted({p["faktor"] for p in poz if katla(p["faktor"]) not in sutun})
    if eksik:
        raise ValueError(f"{f_yol.name}: şu risk faktörlerinin fiyat sütunu yok: {', '.join(eksik)}")
    faktorler = list(dict.fromkeys(p["faktor"] for p in poz))
    satirlar, atlanan = [], 0
    for r in t[bi + 1:]:
        g = tarih(r[0])
        v = [para(r[sutun[katla(f)]]) if sutun[katla(f)] < len(r) else None for f in faktorler]
        if not g:
            continue
        if any(x is None or x <= 0 for x in v):
            atlanan += 1
            continue
        satirlar.append((g, [float(x) for x in v]))
    satirlar.sort()
    if atlanan:
        uy.append(uyari("Orta", "Eksik fiyatlı günler", f_yol.name, f"{atlanan} günde en az bir risk faktörünün fiyatı yok veya sıfır / eksi; bu günler atlandı "
                                                                    "(atlanan günün ertesindeki getiri birden çok günü kapsar)."))
    fiyat = {f: [s[1][i] for s in satirlar] for i, f in enumerate(faktorler)}
    return poz, [s[0] for s in satirlar], fiyat, uy


def getiriler(fiyat: dict) -> dict:
    return {f: [b / a - 1 for a, b in zip(p, p[1:])] for f, p in fiyat.items()}


def var_es(kz: list[float], guven: float) -> tuple[float, float, int, list[int]]:
    """(VaR, ES, k, en kötü k senaryonun indeksleri) — VaR ve ES pozitif sayı olarak zarar."""
    n = len(kz)
    k = int(math.floor(n * (1 - guven / 100) + 1e-9)) + 1
    sirali = sorted(range(n), key=lambda i: kz[i])[:k]
    return max(-kz[sirali[-1]], 0.0), max(-sum(kz[i] for i in sirali) / k, 0.0), k, sirali


def hesapla(poz: list[dict], gunler: list[date], fiyat: dict, guven: float = 99.0, pencere: int = 250, elde_tutma: int = 1) -> dict:
    uy, r = [], getiriler(fiyat)
    n_tum = len(gunler) - 1
    if n_tum < 30:
        raise ValueError(f"En az 31 günlük ortak fiyat gerekli ({len(gunler)} gün var)")
    if n_tum < pencere:
        uy.append(uyari("Yüksek", "Fiyat geçmişi kısa", "Tüm portföy", f"{n_tum} getiri var, pencere {pencere}. VaR mevcut {n_tum} senaryoyla hesaplandı; sonuç daha az güvenilir."))
    n = min(pencere, n_tum)
    for f, seri in r.items():
        son = seri[-n:]
        uc = [(gunler[len(seri) - n + i + 1], x) for i, x in enumerate(son) if abs(x) > 0.20]
        if uc:
            uy.append(uyari("Orta", "Uç getiri", f, "Günlük %20'yi aşan getiri: " + ", ".join(f"{g:%d.%m.%Y} %{x * 100:.1f}" for g, x in uc[:5]) +
                            ". Bölünme, temettü veya veri hatası olabilir; düzeltilmiş fiyat kullanın."))
        sabit = en_uzun_sifir(son)
        if sabit >= 5:
            uy.append(uyari("Orta", "Değişmeyen fiyat", f, f"Arka arkaya {sabit} gün fiyat değişmemiş. Bayat fiyat riski olduğundan düşük gösterir."))
    senaryo_gun = gunler[-n:]
    poz_kz = [[p["deger"] * x for x in r[p["faktor"]][-n:]] for p in poz]
    kz = [sum(s) for s in zip(*poz_kz)]
    var, es, k, kotu = var_es(kz, guven)
    toplam, brut = sum(p["deger"] for p in poz), sum(abs(p["deger"]) for p in poz)
    kok = math.sqrt(elde_tutma)
    satirlar = []
    for p, pk in zip(poz, poz_kz):
        bv, be, _, _ = var_es(pk, guven)
        satirlar.append(p | {"agirlik": p["deger"] / brut * 100, "bireysel_var": bv, "bireysel_es": be, "var_senaryosu": -pk[kotu[-1]], "es_katki": -sum(pk[i] for i in kotu) / k,
                             "oynaklik": _std(r[p["faktor"]][-n:]) * 100})
    for s in satirlar:
        s["es_pay"] = s["es_katki"] / es * 100 if es else 0.0
    # geriye dönük test
    test = None
    if n_tum >= pencere + 20:
        tum_kz = [sum(p["deger"] * r[p["faktor"]][i] for p in poz) for i in range(n_tum)]
        gun_say = min(250, n_tum - pencere)
        asimlar = []
        for i in range(n_tum - gun_say, n_tum):
            v = var_es(tum_kz[i - pencere:i], guven)[0]
            if -tum_kz[i] > v:
                asimlar.append({"gun": gunler[i + 1], "kz": tum_kz[i], "var": v})
        beklenen = gun_say * (1 - guven / 100)
        bolge = None
        if gun_say == 250 and guven == 99.0:
            bolge = "Yeşil" if len(asimlar) <= 4 else "Sarı" if len(asimlar) <= 9 else "Kırmızı"
        test = {"gun": gun_say, "asim": asimlar, "beklenen": beklenen, "bolge": bolge}
    else:
        uy.append(uyari("Bilgi", "Geriye dönük test yapılamadı", "Tüm portföy", f"Test için en az pencere + 20 = {pencere + 20} getiri gerekir ({n_tum} var)."))
    if brut and max(abs(p["deger"]) for p in poz) / brut > 0.5:
        b = max(poz, key=lambda p: abs(p["deger"]))
        uy.append(uyari("Bilgi", "Yoğunlaşma", b["ad"], f"Brüt pozisyonun %{abs(b['deger']) / brut * 100:.0f} kadarı tek pozisyonda."))
    return {"var": var, "es": es, "k": k, "n": n, "guven": guven, "pencere": pencere, "elde": elde_tutma, "var_h": var * kok, "es_h": es * kok, "toplam": toplam, "brut": brut,
            "satirlar": satirlar, "cesitlendirme": sum(s["bireysel_var"] for s in satirlar) - var, "kz": kz, "gunler": senaryo_gun, "kotu": kotu,
            "var_gunu": senaryo_gun[kotu[-1]], "test": test, "uyarilar": uy, "donem": (gunler[-n - 1], gunler[-1])}


def en_uzun_sifir(seri: list[float]) -> int:
    en, say = 0, 0
    for x in seri:
        say = say + 1 if x == 0 else 0
        en = max(en, say)
    return en


def _std(x: list[float]) -> float:
    if len(x) < 2:
        return 0.0
    m = sum(x) / len(x)
    return math.sqrt(sum((v - m) ** 2 for v in x) / (len(x) - 1))


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer", "Açıklama"], (40, 20, 90))
    h = r["elde"]
    satir = [("Portföy net değeri", r["toplam"], "Uzun ve kısa pozisyonların toplamı"), ("Portföy brüt değeri", r["brut"], "Mutlak değerlerin toplamı"),
             (f"VaR — %{r['guven']:g}, 1 gün", r["var"], f"{r['n']} senaryonun en kötü {r['k']}. günü ({r['var_gunu']:%d.%m.%Y})"),
             ("VaR / brüt değer %", r["var"] / r["brut"] * 100 if r["brut"] else 0, ""),
             (f"Beklenen kayıp (ES) — %{r['guven']:g}, 1 gün", r["es"], f"En kötü {r['k']} senaryonun ortalaması"),
             ("Çeşitlendirme etkisi", r["cesitlendirme"], "Bireysel VaR'ların toplamı − portföy VaR")]
    if h > 1:
        satir += [(f"VaR — {h} gün", r["var_h"], f"1 günlük VaR × √{h} (karekök kuralı)"), (f"ES — {h} gün", r["es_h"], f"1 günlük ES × √{h}")]
    satir += [("Senaryo dönemi", f"{r['donem'][0]:%d.%m.%Y} – {r['donem'][1]:%d.%m.%Y}", f"{r['n']} günlük getiri")]
    t = r["test"]
    if t:
        satir += [("Geriye dönük test — aşım sayısı", len(t["asim"]), f"Son {t['gun']} gün; beklenen yaklaşık {t['beklenen']:.1f}" + (f" · bölge: {t['bolge']}" if t["bolge"] else ""))]
    for k, v, a in satir:
        n = _satir(oz, [k, round(v, 2) if isinstance(v, float) else v, a], para_s=(2,) if isinstance(v, float) else ())
        if k.startswith("Geriye") and t["bolge"]:
            oz.cell(n, 2).fill = PatternFill("solid", fgColor={"Yeşil": YESIL, "Sarı": SARI, "Kırmızı": KIRMIZI}[t["bolge"]])
    oz.append([])
    oz.append(["Tarihsel simülasyon geçmişte yaşanmamış bir şoku göremez. Sonuç; pencere uzunluğuna, fiyat kalitesine ve pozisyonların doğrusal olduğu varsayımına bağlıdır."])

    pz = wb.create_sheet("Pozisyonlar")
    _baslik(pz, ["Enstrüman", "Risk Faktörü", "Portföy", "Piyasa Değeri", "Brüt Ağırlık %", "Günlük Oynaklık %", "Bireysel VaR", "Bireysel ES", "VaR Günündeki Zarar",
                 "ES Katkısı", "ES Payı %"], (28, 16, 14, 18, 10, 10, 16, 16, 16, 16, 10))
    for s in sorted(r["satirlar"], key=lambda s: -s["es_katki"]):
        _satir(pz, [s["ad"], s["faktor"], s["portfoy"], s["deger"], round(s["agirlik"], 1), round(s["oynaklik"], 2), round(s["bireysel_var"], 2), round(s["bireysel_es"], 2),
                    round(s["var_senaryosu"], 2), round(s["es_katki"], 2), round(s["es_pay"], 1)], para_s=(4, 7, 8, 9, 10),
               dolgu={10: KIRMIZI if s["es_pay"] > 40 else MAVI if s["es_katki"] < 0 else None})
    pz.append([])
    pz.append(["ES katkısı eksi olan pozisyon, portföyün en kötü günlerinde kazandırmıştır (riski azaltır). Katkıların toplamı portföy ES'ine eşittir."])

    sn = wb.create_sheet("Senaryolar")
    _baslik(sn, ["Sıra (en kötüden)", "Tarih", "Portföy K/Z", "K/Z ÷ Brüt %", "Not"], (10, 12, 18, 12, 30))
    for i, j in enumerate(sorted(range(r["n"]), key=lambda j: r["kz"][j]), 1):
        _satir(sn, [i, r["gunler"][j], round(r["kz"][j], 2), round(r["kz"][j] / r["brut"] * 100, 3) if r["brut"] else 0, "VaR senaryosu" if i == r["k"] else "ES'e giriyor" if i < r["k"] else ""],
               para_s=(3,), tarih_s=(2,), dolgu={5: KIRMIZI if i == r["k"] else SARI if i < r["k"] else None})
    if t:
        bt = wb.create_sheet("Geriye Dönük Test")
        _baslik(bt, ["Tarih", "Gerçekleşen K/Z (varsayımsal)", "Bir Önceki Günün VaR'ı", "Aşım Oranı"], (12, 22, 22, 12))
        for a in t["asim"]:
            _satir(bt, [a["gun"], round(a["kz"], 2), round(a["var"], 2), round(-a["kz"] / a["var"], 2) if a["var"] else None], para_s=(2, 3), tarih_s=(1,))
        if not t["asim"]:
            bt.append(["Aşım yok."])
        bt.append([])
        bt.append([f"Son {t['gun']} günde {len(t['asim'])} aşım; %{r['guven']:g} güvende beklenen yaklaşık {t['beklenen']:.1f}. Pozisyonlar dönem boyunca sabit varsayılmıştır."])
        if t["bolge"]:
            bt.append(["Bölge (250 gün, %99; Basel geriye dönük test çerçevesi): 0–4 aşım Yeşil, 5–9 Sarı, 10 ve üzeri Kırmızı."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(p_yol: Path, f_yol: Path, cikti: Path, guven: float = 99.0, pencere: int = 250, elde_tutma: int = 1) -> dict:
    poz, gunler, fiyat, uy = oku(p_yol, f_yol)
    r = hesapla(poz, gunler, fiyat, guven, pencere, elde_tutma)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Tarihsel simülasyonla günlük VaR ve beklenen kaybı (ES) hesaplar.")
    p.add_argument("--pozisyonlar", type=Path, default=ORNEK / "pozisyonlar.csv", help="Enstrüman, Risk Faktörü, Piyasa Değeri (kısa pozisyon eksi), Portföy")
    p.add_argument("--fiyatlar", type=Path, default=ORNEK / "fiyatlar.csv", help="Tarih + her risk faktörü için bir sütun (günlük kapanış)")
    p.add_argument("--guven", type=float, default=99.0, help="Güven düzeyi %% (varsayılan 99)")
    p.add_argument("--pencere", type=int, default=250, help="Senaryo sayısı / gün (varsayılan 250)")
    p.add_argument("--elde-tutma", type=int, default=1, help="Elde tutma süresi, gün (varsayılan 1; 10 için karekök kuralı)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "var_raporu.xlsx")
    a = p.parse_args(argv)
    for y in (a.pozisyonlar, a.fiyatlar):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    if not (50 <= a.guven < 100) or a.pencere < 30 or a.elde_tutma < 1:
        print("[X] --guven 50 ile 100 arasında, --pencere en az 30, --elde-tutma en az 1 olmalı")
        return 2
    try:
        r = calistir(a.pozisyonlar, a.fiyatlar, a.cikti, a.guven, a.pencere, a.elde_tutma)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    b = lambda x: f"{x:,.0f}".replace(",", ".")                                    # noqa: E731
    print(f"[OK] {len(r['satirlar'])} pozisyon · {r['n']} senaryo · VaR %{r['guven']:g} 1 gün: {b(r['var'])} · ES: {b(r['es'])} · çeşitlendirme etkisi {b(r['cesitlendirme'])}")
    if r["test"]:
        print(f"[OK] Geriye dönük test: son {r['test']['gun']} günde {len(r['test']['asim'])} aşım" + (f" ({r['test']['bolge']})" if r["test"]["bolge"] else ""))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
