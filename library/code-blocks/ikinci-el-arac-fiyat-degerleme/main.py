"""
İkinci El Araç Fiyat Değerleme — Gerisi Bende kod bloğu
Satış › İkinci El Satış Danışmanı

Benzer ilan verisinden değerlenecek araç için piyasa fiyat aralığını ve alım fiyatı önerisini üretir.
  Benzer ilan        = aynı marka ve model;  model yılı farkı ≤ --yil-fark (1);  km, aracın km'sinin ±--km-fark % (30) içinde;  yakıt ve vites iki tarafta da yazılıysa aynı
  Az ilan varsa      = benzer ilan sayısı --asgari (3) altındaysa km koşulu kaldırılıp yeniden aranır; yine azsa "Yetersiz veri"
  Aykırı fiyat       = en az 4 ilan varsa, [Q1 − 1,5 × (Q3 − Q1), Q3 + 1,5 × (Q3 − Q1)] dışındaki ilanlar hesaba katılmaz
  Piyasa değeri      = kalan ilan fiyatlarının medyanı;   piyasa aralığı = Q1 – Q3            (çeyrekler doğrusal ara değerle)
  Alım önerisi       = medyan × (1 − --marj % ÷ 100), --yuvarla (1.000) katına aşağı yuvarlanır   (marj: pazarlık payı, hazırlık gideri ve kâr; varsayılan %10)
İlan fiyatı satış fiyatı değildir; istenen fiyattır. Hasar, boya, değişen, bakım geçmişi ve donanım farkını kod bilmez; ekspertiz sonucuna göre siz düzeltirsiniz.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
from decimal import ROUND_FLOOR

ILAN = {"marka": ("marka",), "model": ("model", "seri"), "yil": ("yil", "model yili"), "km": ("km", "kilometre"), "fiyat": ("fiyat", "ilan fiyati"), "yakit": ("yakit", "yakit tipi"), "vites": ("vites", "vites tipi"),
        "paket": ("paket", "donanim", "versiyon")}
ARAC = {"ad": ("arac", "plaka", "arac adi", "no"), "marka": ILAN["marka"], "model": ILAN["model"], "yil": ILAN["yil"], "km": ILAN["km"], "yakit": ILAN["yakit"], "vites": ILAN["vites"], "paket": ILAN["paket"],
        "hasar": ("hasar tutari", "hasar kaydi", "tramer")}
Z = Decimal(0)


def yil_coz(x) -> int | None:
    v = para(x)
    return int(v) if v is not None and v == v.to_integral_value() and 1950 <= v <= 2100 else None


def oku(ilan_yol: Path, arac_yol: Path) -> tuple[list[dict], list[dict], list[dict]]:
    uy, ilanlar, araclar = [], [], []
    for r in kayitlar(ilan_yol, ILAN, ("marka", "model", "yil", "fiyat")):
        if not metin(r["marka"]) and not metin(r["fiyat"]):
            continue
        y, km, f = yil_coz(r["yil"]), (para(r["km"]) if metin(r["km"]) else None), para(r["fiyat"])
        if not metin(r["marka"]) or not metin(r["model"]) or y is None or f is None or f <= 0 or (metin(r["km"]) and (km is None or km < 0)):
            uy.append(uyari("Orta", "İlan okunamadı", f"{metin(r['marka'])} {metin(r['model'])} (satır {r['_satir']})", f"Yıl '{metin(r['yil'])}', km '{metin(r['km'])}', fiyat '{metin(r['fiyat'])}'. İlan hesaba katılmadı."))
            continue
        ilanlar.append({"marka": metin(r["marka"]), "model": metin(r["model"]), "yil": y, "km": km, "fiyat": f, "yakit": katla(r["yakit"]), "vites": katla(r["vites"]), "paket": metin(r["paket"])})
    for n, r in enumerate(kayitlar(arac_yol, ARAC, ("marka", "model", "yil")), 1):
        if not metin(r["marka"]) and not metin(r["model"]):
            continue
        ad = metin(r["ad"]) or f"Araç {n}"
        y, km, h = yil_coz(r["yil"]), (para(r["km"]) if metin(r["km"]) else None), (para(r["hasar"]) if metin(r["hasar"]) else Z)
        if not metin(r["marka"]) or not metin(r["model"]) or y is None or (metin(r["km"]) and (km is None or km < 0)) or h is None or h < 0:
            uy.append(uyari("Yüksek", "Araç okunamadı", f"{ad} (satır {r['_satir']})", f"Marka, model, yıl '{metin(r['yil'])}', km '{metin(r['km'])}' veya hasar tutarı '{metin(r['hasar'])}' hatalı. Araç değerlenmedi."))
            continue
        araclar.append({"ad": ad, "marka": metin(r["marka"]), "model": metin(r["model"]), "yil": y, "km": km, "yakit": katla(r["yakit"]), "vites": katla(r["vites"]), "paket": metin(r["paket"]), "hasar": h})
    return ilanlar, araclar, uy


def ceyrek(sirali: list[Decimal], p: Decimal) -> Decimal:
    konum = (len(sirali) - 1) * p
    alt = int(konum)
    return sirali[alt] if alt + 1 >= len(sirali) else sirali[alt] + (sirali[alt + 1] - sirali[alt]) * (konum - alt)


def degerle(arac: dict, ilanlar: list[dict], yil_fark: int = 1, km_fark: Decimal = Decimal(30), asgari: int = 3, marj: Decimal = Decimal(10), adim: Decimal = Decimal(1000)) -> dict:
    uy, kim = [], arac["ad"]
    aday = [i for i in ilanlar if katla(i["marka"]) == katla(arac["marka"]) and katla(i["model"]) == katla(arac["model"]) and abs(i["yil"] - arac["yil"]) <= yil_fark
            and not (i["yakit"] and arac["yakit"] and i["yakit"] != arac["yakit"]) and not (i["vites"] and arac["vites"] and i["vites"] != arac["vites"])]
    benzer, yontem = aday, "Marka, model, yıl"
    if arac["km"] is not None:
        alt, ust = arac["km"] * (1 - km_fark / 100), arac["km"] * (1 + km_fark / 100)
        kmli = [i for i in aday if i["km"] is not None and alt <= i["km"] <= ust]
        if len(kmli) >= asgari:
            benzer, yontem = kmli, f"Marka, model, yıl, km ±%{km_fark.normalize():f}"
        elif len(aday) >= asgari:
            uy.append(uyari("Orta", "Km koşulu kaldırıldı", kim, f"Km aralığında yalnız {len(kmli)} ilan vardı; {len(aday)} ilanla km'ye bakılmadan değerlendi. Aracın km'si benzerlerden farklıysa fiyatı elle düzeltin."))
    sonuc = {"ad": kim, "marka": arac["marka"], "model": arac["model"], "yil": arac["yil"], "km": arac["km"], "paket": arac["paket"], "hasar": arac["hasar"], "yontem": yontem, "ilan": len(benzer), "ayrik": 0,
             "min": None, "q1": None, "medyan": None, "q3": None, "max": None, "alim": None, "km_medyan": None, "km_fark": None, "durum": "Değerlendi", "benzerler": [], "uyarilar": uy}
    if len(benzer) < asgari:
        sonuc["durum"] = "Yetersiz veri"
        uy.append(uyari("Yüksek", "Yetersiz veri", kim, f"{arac['marka']} {arac['model']} {arac['yil']} için yalnız {len(benzer)} benzer ilan var (en az {asgari} gerekir). Daha fazla ilan toplayın veya --yil-fark değerini artırın."))
        return sonuc
    fiyatlar = sorted(i["fiyat"] for i in benzer)
    if len(fiyatlar) >= 4:
        q1, q3 = ceyrek(fiyatlar, Decimal("0.25")), ceyrek(fiyatlar, Decimal("0.75"))
        a, u = q1 - (q3 - q1) * Decimal("1.5"), q3 + (q3 - q1) * Decimal("1.5")
        temiz = [i for i in benzer if a <= i["fiyat"] <= u]
        if len(temiz) >= asgari:
            sonuc["ayrik"], benzer = len(benzer) - len(temiz), temiz
            fiyatlar = sorted(i["fiyat"] for i in benzer)
    med = ceyrek(fiyatlar, Decimal("0.5"))
    kmler = sorted(i["km"] for i in benzer if i["km"] is not None)
    km_med = ceyrek(kmler, Decimal("0.5")) if kmler else None
    alim = med * (1 - marj / 100)
    sonuc |= {"ilan": len(benzer), "min": fiyatlar[0], "q1": yuvarla(ceyrek(fiyatlar, Decimal("0.25"))), "medyan": yuvarla(med), "q3": yuvarla(ceyrek(fiyatlar, Decimal("0.75"))), "max": fiyatlar[-1],
              "alim": (alim / adim).quantize(Decimal(1), rounding=ROUND_FLOOR) * adim if adim > 0 else yuvarla(alim), "km_medyan": yuvarla(km_med, "1") if km_med is not None else None,
              "km_fark": yuvarla(arac["km"] - km_med, "1") if km_med is not None and arac["km"] is not None else None,
              "benzerler": sorted(({"arac": kim} | i for i in benzer), key=lambda i: i["fiyat"])}
    if arac["hasar"]:
        uy.append(uyari("Orta", "Hasar kaydı var", kim, f"Araçta {arac['hasar']:,.2f} hasar kaydı yazılı. Öneri hasarsız emsallere göredir; ekspertize göre fiyatı düşürün."))
    if len(benzer) < 5:
        uy.append(uyari("Bilgi", "Az sayıda emsal", kim, f"Değerleme {len(benzer)} ilana dayanıyor; aralık güvenilir olmayabilir."))
    return sonuc


def rapor_yaz(cikti: Path, sonuclar: list[dict], ayar: dict, uy: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Değerleme"
    _baslik(ws, ["Araç", "Marka", "Model", "Yıl", "Km", "Paket", "Durum", "Emsal Sayısı", "Çıkarılan Aykırı", "En Düşük", "Alt Çeyrek", "Piyasa Değeri (medyan)", "Üst Çeyrek", "En Yüksek", "Alım Önerisi", "Emsal Medyan Km", "Km Farkı",
                 "Emsal Ölçütü", "Karar (siz doldurun)"], (16, 14, 18, 7, 11, 16, 14, 8, 9, 14, 14, 16, 14, 14, 15, 12, 11, 30, 24), "B2")
    for s in sonuclar:
        _satir(ws, [s["ad"], s["marka"], s["model"], s["yil"], s["km"], s["paket"], s["durum"], s["ilan"], s["ayrik"] or None, s["min"], s["q1"], s["medyan"], s["q3"], s["max"], s["alim"], s["km_medyan"], s["km_fark"], s["yontem"], ""],
               para_s=(10, 11, 12, 13, 14, 15), dolgu={7: KIRMIZI if s["durum"] != "Değerlendi" else YESIL, 12: MAVI, 15: SARI})
        for j in (5, 16, 17):
            ws.cell(ws.max_row, j).number_format = "#,##0"
    ws.append([])
    ws.append([f"Piyasa değeri ilan fiyatlarının medyanıdır; ilan fiyatı satış fiyatı değildir. Alım önerisi = medyan × (1 − %{ayar['marj'].normalize():f}). Hasar, boya, değişen ve donanım farkı hesaba katılmaz."])
    ws = wb.create_sheet("Emsal İlanlar")
    _baslik(ws, ["Değerlenen Araç", "Marka", "Model", "Yıl", "Km", "Fiyat", "Paket"], (16, 14, 18, 7, 11, 15, 18))
    for s in sonuclar:
        for i in s["benzerler"]:
            _satir(ws, [i["arac"], i["marka"], i["model"], i["yil"], i["km"], i["fiyat"], i["paket"]], para_s=(6,))
            ws.cell(ws.max_row, 5).number_format = "#,##0"
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(ilan_yol: Path, arac_yol: Path, cikti: Path, **ayar) -> dict:
    ilanlar, araclar, uy = oku(ilan_yol, arac_yol)
    if not ilanlar:
        raise ValueError("İlan dosyasında geçerli satır yok.")
    if not araclar:
        raise ValueError("Değerlenecek araç yok.")
    sonuclar = [degerle(a, ilanlar, **ayar) for a in araclar]
    for s in sonuclar:
        uy += s.pop("uyarilar")
    rapor_yaz(cikti, sonuclar, {"marj": Decimal(10)} | ayar, uy)
    return {"araclar": sonuclar, "uyarilar": sirala(uy)}


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Benzer ilanlardan piyasa fiyat aralığı ve alım fiyatı önerisi üretir.")
    p.add_argument("--ilanlar", type=Path, default=ORNEK / "ilanlar.csv", help="Marka, Model, Yıl, Km, Fiyat, Yakıt, Vites, Paket")
    p.add_argument("--araclar", type=Path, help="Araç (plaka veya ad), Marka, Model, Yıl, Km, Yakıt, Vites, Paket, Hasar Tutarı")
    p.add_argument("--yil-fark", type=int, default=1, help="Emsal için kabul edilen model yılı farkı (varsayılan 1)")
    p.add_argument("--km-fark", default="30", help="Emsal için kabul edilen km farkı, ±%% (varsayılan 30)")
    p.add_argument("--asgari", type=int, default=3, help="Değerleme için gereken en az emsal sayısı (varsayılan 3)")
    p.add_argument("--marj", default="10", help="Alım önerisinde medyandan düşülecek pay, %% (varsayılan 10)")
    p.add_argument("--yuvarla", default="1000", help="Alım önerisi bu değerin katına aşağı yuvarlanır (varsayılan 1000)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "ikinci_el_degerleme.xlsx")
    a = p.parse_args(argv)
    arac = a.araclar or (ORNEK / "araclar.csv" if a.ilanlar == ORNEK / "ilanlar.csv" else None)
    kf, mj, yv = para(a.km_fark), para(a.marj), para(a.yuvarla)
    if arac is None or a.yil_fark < 0 or a.asgari < 1 or kf is None or not 0 < kf <= 100 or mj is None or not 0 <= mj < 100 or yv is None or yv < 0:
        print("[X] --araclar dosyası verilmeli; --yil-fark ≥ 0, --asgari ≥ 1, --km-fark 0–100, --marj 0–100, --yuvarla sıfır veya pozitif olmalı")
        return 2
    for y in (a.ilanlar, arac):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(a.ilanlar, arac, a.cikti, yil_fark=a.yil_fark, km_fark=kf, asgari=a.asgari, marj=mj, adim=yv)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    for s in r["araclar"]:
        if s["durum"] == "Değerlendi":
            print(f"[OK] {s['ad']} · {s['marka']} {s['model']} {s['yil']}: piyasa {s['q1']:,.0f} – {s['q3']:,.0f}, medyan {s['medyan']:,.0f} ({s['ilan']} emsal) · alım önerisi {s['alim']:,.0f}")
        else:
            print(f"[!] {s['ad']} · {s['marka']} {s['model']} {s['yil']}: yetersiz veri ({s['ilan']} emsal)")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
