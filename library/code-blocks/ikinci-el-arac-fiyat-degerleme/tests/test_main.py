import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def il(yil, km, fiyat, model="A", yakit="", vites=""):
    return {"marka": "M", "model": model, "yil": yil, "km": None if km is None else D(km), "fiyat": D(fiyat), "yakit": yakit, "vites": vites, "paket": ""}


def ar(yil=2020, km=85000, model="A", yakit="", vites="", hasar=0):
    return {"ad": "X", "marka": "m", "model": model.lower(), "yil": yil, "km": None if km is None else D(km), "yakit": yakit, "vites": vites, "paket": "", "hasar": D(hasar)}


ILAN = lambda: [il(2020, 80000, 900000), il(2020, 90000, 880000), il(2021, 70000, 950000), il(2019, 100000, 820000), il(2020, 85000, 2000000), il(2020, 200000, 600000), il(2017, 80000, 700000),   # noqa: E731
                il(2022, 40000, 1400000, "B")]


class DegerlemeTesti(unittest.TestCase):
    def test_ceyrek(self):
        d = [D(x) for x in (820, 880, 900, 950)]
        self.assertEqual([main.ceyrek(d, D(p)) for p in ("0.25", "0.5", "0.75", "0", "1")], [D(865), D(890), D("912.5"), D(820), D(950)])
        self.assertEqual((main.ceyrek([D(5)], D("0.5")), main.ceyrek([D(1), D(2), D(10)], D("0.5"))), (D(5), D(2)))

    def test_emsal_ve_aykiri(self):
        s = main.degerle(ar(), ILAN())
        self.assertEqual((s["durum"], s["ilan"], s["ayrik"], s["min"], s["q1"], s["medyan"], s["q3"], s["max"], s["alim"], s["km_medyan"], s["km_fark"], s["yontem"], [u["tur"] for u in s["uyarilar"]]),
                         ("Değerlendi", 4, 1, D(820000), D("865000.00"), D("890000.00"), D("912500.00"), D(950000), D(801000), D(85000), D(0), "Marka, model, yıl, km ±%30", ["Az sayıda emsal"]))
        self.assertEqual([i["fiyat"] for i in s["benzerler"]], [D(820000), D(880000), D(900000), D(950000)])
        s = main.degerle(ar(), ILAN(), marj=D(15), adim=D(0))
        self.assertEqual(s["alim"], D("756500.00"))
        s = main.degerle(ar(yil=2019), ILAN(), yil_fark=0, asgari=1)
        self.assertEqual((s["ilan"], s["medyan"], s["alim"]), (1, D("820000.00"), D(738000)))

    def test_km_gevsetme_ve_yetersiz_veri(self):
        s = main.degerle(ar(km=150000, hasar=45000), ILAN())
        self.assertEqual((s["durum"], s["ilan"], s["ayrik"], s["medyan"], s["yontem"], s["km_fark"], sorted(u["tur"] for u in s["uyarilar"])), ("Değerlendi", 4, 2, D("890000.00"), "Marka, model, yıl", D(65000), ["Az sayıda emsal", "Hasar kaydı var", "Km koşulu kaldırıldı"]))
        s = main.degerle(ar(yil=2022, km=35000, model="B"), ILAN())
        self.assertEqual((s["durum"], s["ilan"], s["medyan"], s["alim"], [u["tur"] for u in s["uyarilar"]]), ("Yetersiz veri", 1, None, None, ["Yetersiz veri"]))
        s = main.degerle(ar(km=None), ILAN())
        self.assertEqual((s["ilan"], s["ayrik"], s["km_fark"], s["yontem"]), (4, 2, None, "Marka, model, yıl"))
        l = [il(2020, 80000, 900000, yakit="benzin", vites="otomatik"), il(2020, 81000, 800000, yakit="dizel", vites="otomatik"), il(2020, 82000, 700000, yakit="benzin", vites="manuel"), il(2020, 83000, 600000)]
        self.assertEqual((main.degerle(ar(yakit="benzin", vites="otomatik"), l, asgari=1)["ilan"], main.degerle(ar(), l, asgari=1)["ilan"], main.degerle(ar(yakit="dizel"), l, asgari=1)["ilan"]), (2, 4, 2))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "d.xlsx"
            r = main.calistir(ORNEK / "ilanlar.csv", ORNEK / "araclar.csv", cikti)
            self.assertEqual([(s["ad"], s["durum"], s["ilan"], s["medyan"], s["alim"]) for s in r["araclar"]],
                             [("Stok-001", "Değerlendi", 4, D("890000.00"), D(801000)), ("Stok-002", "Değerlendi", 4, D("890000.00"), D(801000)), ("Stok-003", "Yetersiz veri", 1, None, None)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "Yetersiz veri", "Stok-003"), ("Orta", "Hasar kaydı var", "Stok-002"), ("Orta", "Km koşulu kaldırıldı", "Stok-002"),
                                                                                       ("Bilgi", "Az sayıda emsal", "Stok-001"), ("Bilgi", "Az sayıda emsal", "Stok-002")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Değerleme", "Emsal İlanlar", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            iy, ay = Path(tmp) / "i.csv", Path(tmp) / "a.csv"
            iy.write_text("Marka;Model;Model Yılı;Kilometre;İlan Fiyatı\nM;A;2020;85.000;1.250.000\nM;A;20;1;5\nM;A;2020;x;5\nM;;2020;1;5\nM;A;2020;;900.000\n", encoding="utf-8")
            ay.write_text("Plaka;Marka;Model;Yıl;Km\n34 ABC 01;M;A;2020;85000\n34 ABC 02;M;A;iki bin;1\n", encoding="utf-8")
            ilanlar, araclar, uy = main.oku(iy, ay)
            self.assertEqual(([(i["yil"], i["km"], i["fiyat"]) for i in ilanlar], [(a["ad"], a["km"], a["hasar"]) for a in araclar], sorted(u["tur"] for u in uy)),
                             ([(2020, D(85000), D(1250000)), (2020, None, D(900000))], [("34 ABC 01", D(85000), D(0))], ["Araç okunamadı"] + ["İlan okunamadı"] * 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--ilanlar", str(iy), "--araclar", str(ay), "--cikti", c]
            self.assertEqual(main.main(arg + ["--asgari", "1", "--yil-fark", "2", "--km-fark", "20", "--marj", "12,5", "--yuvarla", "500"]), 0)
            for ek in (["--asgari", "0"], ["--km-fark", "0"], ["--marj", "100"], ["--yuvarla", "x"]):
                self.assertEqual(main.main(arg + ek), 2)
            self.assertEqual(main.main(["--ilanlar", str(iy), "--cikti", c]), 2)
            self.assertEqual(main.main(["--ilanlar", str(iy), "--araclar", str(Path(tmp) / "yok.csv")]), 1)
            ay.write_text("Plaka;Marka;Model;Yıl\nP;M;A;x\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
