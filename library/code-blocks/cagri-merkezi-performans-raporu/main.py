"""
Çağrı Merkezi Performans Raporu — Workers / Workless kod bloğu
Müşteri Hizmetleri ve Çağrı Merkezi › Çağrı Merkezi Yöneticisi

Çağrı kayıtlarından temel performans göstergelerini hesaplar:
  - Kısa terk: --kisa-terk (5) saniyeden önce kapatılan çağrı gelen ve terk sayısına katılmaz.
  - Karşılama oranı = cevaplanan / gelen; terk oranı = terk / gelen.
  - Servis seviyesi (SL) = T saniye içinde cevaplanan / (cevaplanan + T'den sonra terk edilen); T = --sl-sure (20).
    T içinde terk edilen çağrı paydaya girmez.
  - ASA = cevaplanan çağrılarda ortalama bekleme; AHT = (görüşme + beklemeye alma + çağrı sonrası iş) / cevaplanan.
  - İlk temasta çözüm (FCR), tekrar arama yöntemiyle: cevaplanan çağrıdan sonra aynı numara --tekrar-gun (7) gün içinde
    yeniden aramadıysa çözülmüş sayılır. Veri bitişine --tekrar-gun günden yakın çağrılar paydaya girmez. "Çözüldü"
    sütunu varsa temsilci beyanı ayrıca gösterilir.
  - 30 dakikalık aralıklar: hafta içi günlerin ortalama gelen çağrısı, AHT ve Erlang C ile hedef SL için gereken
    temsilci (+ --kayip payı); aynı aralıkta çağrı cevaplayan temsilci sayısıyla karşılaştırma.
  - Temsilci: cevaplanan, AHT, beklemeye alma ve çağrı sonrası iş payı, kısa görüşme (< --kisa-gorusme sn) oranı.
Arayan numaralar raporda maskelenir (son iki hane). İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 15 iş günü, 3 kuyruk, 12 temsilci
    python main.py --cagrilar cagrilar.xlsx --sl-sure 20 --sl-hedef 80 --kayip 30
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]
CEVAP = {"cevaplandi", "cevaplanan", "answered", "basarili", "tamamlandi", "gorusuldu"}
TERK = {"terk", "terk edildi", "abandoned", "kayip", "kacan", "cevapsiz"}
EVET = {"evet", "e", "1", "cozuldu", "true"}

SUTUNLAR = {"id": ("cagri id", "cagri no", "id"), "bas": ("baslangic", "baslangic zamani", "tarih saat", "cagri zamani"), "tarih": ("tarih",), "saat": ("saat",),
            "kuyruk": ("kuyruk", "hat", "skill", "grup"), "temsilci": ("temsilci", "agent", "operator"), "bekleme": ("bekleme", "bekleme sn", "bekleme suresi"),
            "gorusme": ("gorusme", "gorusme sn", "konusma suresi", "gorusme suresi"), "hold": ("beklemeye alma", "beklemeye alma sn", "hold"),
            "acw": ("cagri sonrasi is", "cagri sonrasi is sn", "acw"), "durum": ("durum", "sonuc"), "arayan": ("arayan", "arayan no", "telefon", "ani"),
            "kategori": ("kategori", "konu"), "cozuldu": ("cozuldu", "cozum")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def zaman(x, saat_x=None) -> datetime | None:
    if isinstance(x, datetime):
        return x
    if isinstance(x, date) and saat_x is not None:
        s = saat_x if isinstance(saat_x, time) else datetime.strptime(str(saat_x).strip()[:8].ljust(8, ":00")[:8], "%H:%M:%S").time()
        return datetime.combine(x, s)
    s = str(x or "").strip()
    if saat_x is not None:
        s = f"{s[:10]} {str(saat_x).strip()}"
    for f in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def sure(x) -> float | None:
    """Saniye: 95, '95', '01:35', '00:01:35'."""
    if x in (None, ""):
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    if isinstance(x, time):
        return x.hour * 3600 + x.minute * 60 + x.second
    s = str(x).strip().replace(",", ".")
    if ":" in s:
        p = [float(v) for v in s.split(":")]
        return sum(v * 60 ** i for i, v in enumerate(reversed(p)))
    try:
        return float(s)
    except ValueError:
        return None


def yz(v: float | None, basamak: int = 0) -> str:
    """0.619 → '%62' / '%61,9' (Türkçe)."""
    return "—" if v is None else ("%" + f"{v * 100:.{basamak}f}").replace(".", ",")


def sy(v: float, basamak: int = 1) -> str:
    return f"{v:.{basamak}f}".replace(".", ",")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def maske(no: str) -> str:
    d = re.sub(r"\D", "", no)
    return ("*" * max(0, len(d) - 2) + d[-2:]) if d else ""


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Erlang C
# ----------------------------------------------------------------------------

def erlang_c(trafik: float, n: int) -> float:
    """Bekleme olasılığı. Erlang B özyinelemesiyle sayısal olarak kararlı hesap."""
    if n <= trafik:
        return 1.0
    b = 1.0
    for k in range(1, n + 1):
        b = trafik * b / (k + trafik * b)
    return n * b / (n - trafik * (1 - b))


def servis_seviyesi(cagri: float, aht: float, n: int, hedef_sn: float, aralik_sn: float = 1800) -> float:
    a = cagri * aht / aralik_sn
    if n <= a:
        return 0.0
    return 1 - erlang_c(a, n) * math.exp(-(n - a) * hedef_sn / aht)


def gereken_temsilci(cagri: float, aht: float, hedef_sn: float, sl_hedef: float, aralik_sn: float = 1800) -> int:
    if cagri <= 0 or aht <= 0:
        return 0
    a = cagri * aht / aralik_sn
    n = max(1, math.floor(a) + 1)
    while servis_seviyesi(cagri, aht, n, hedef_sn, aralik_sn) < sl_hedef:
        n += 1
    return n


# ----------------------------------------------------------------------------
# Veri ve hesap
# ----------------------------------------------------------------------------

@dataclass
class Cagri:
    satir: int
    id: str
    bas: datetime
    kuyruk: str
    temsilci: str
    bekleme: float
    gorusme: float
    hold: float
    acw: float
    durum: str             # cevap | terk | kisa_terk | diger
    arayan: str
    kategori: str
    cozuldu: bool | None
    fcr: bool | None = None

    @property
    def ele_alma(self) -> float:
        return self.gorusme + self.hold + self.acw


def oku(yol: Path, kisa_terk: float) -> tuple[list[Cagri], list[dict]]:
    cagrilar, uy, bilinmeyen = [], [], 0
    satirlar = kayitlar(yol, SUTUNLAR, ("durum",))
    cozum_var = any("cozuldu" in r for r in satirlar[:1])
    for r in satirlar:
        bas = zaman(r.get("bas")) if r.get("bas") not in (None, "") else zaman(r.get("tarih"), r.get("saat"))
        if not bas:
            uy.append({"onem": "Orta", "tur": "Zaman okunamadı", "kim": metin(r.get("id")) or f"satır {r['_satir']}", "aciklama": "Çağrı alınmadı"})
            continue
        b, g, h, a = sure(r.get("bekleme")), sure(r.get("gorusme")), sure(r.get("hold")), sure(r.get("acw"))
        if None in (b, g, h, a):
            uy.append({"onem": "Orta", "tur": "Süre okunamadı", "kim": metin(r.get("id")) or f"satır {r['_satir']}", "aciklama": "Çağrı alınmadı"})
            continue
        d = katla(r.get("durum"))
        if d in CEVAP or (d not in TERK and g > 0 and metin(r.get("temsilci"))):
            durum = "cevap"
        elif d in TERK:
            durum = "kisa_terk" if b < kisa_terk else "terk"
        else:
            durum, bilinmeyen = "diger", bilinmeyen + 1
        cz = katla(r.get("cozuldu"))
        cagrilar.append(Cagri(r["_satir"], metin(r.get("id")), bas, metin(r.get("kuyruk")) or "(genel)", metin(r.get("temsilci")), b, g, h, a, durum,
                              re.sub(r"\D", "", metin(r.get("arayan"))), metin(r.get("kategori")), (cz in EVET) if cozum_var and cz else None))
    if bilinmeyen:
        uy.append({"onem": "Orta", "tur": "Bilinmeyen durum", "kim": "", "aciklama": f"{bilinmeyen} çağrının durumu Cevaplandı / Terk olarak okunamadı; göstergelere katılmadı"})
    return cagrilar, uy


def ozet(lst: list[Cagri], sl_sure: float) -> dict:
    cev = [c for c in lst if c.durum == "cevap"]
    terk = [c for c in lst if c.durum == "terk"]
    gelen = len(cev) + len(terk)
    sl_pay = sum(c.bekleme <= sl_sure for c in cev)
    sl_payda = len(cev) + sum(c.bekleme > sl_sure for c in terk)
    fcr = [c for c in cev if c.fcr is not None]
    beyan = [c for c in cev if c.cozuldu is not None]
    return {"gelen": gelen, "cevap": len(cev), "terk": len(terk), "kisa_terk": sum(c.durum == "kisa_terk" for c in lst),
            "karsilama": len(cev) / gelen if gelen else None, "terk_orani": len(terk) / gelen if gelen else None,
            "sl": sl_pay / sl_payda if sl_payda else None, "asa": statistics.fmean(c.bekleme for c in cev) if cev else None,
            "aht": statistics.fmean(c.ele_alma for c in cev) if cev else None, "gorusme": statistics.fmean(c.gorusme for c in cev) if cev else None,
            "hold": statistics.fmean(c.hold for c in cev) if cev else None, "acw": statistics.fmean(c.acw for c in cev) if cev else None,
            "terk_bekleme": statistics.fmean(c.bekleme for c in terk) if terk else None, "en_uzun": max((c.bekleme for c in cev), default=None),
            "fcr": sum(c.fcr for c in fcr) / len(fcr) if fcr else None, "fcr_n": len(fcr),
            "fcr_beyan": sum(c.cozuldu for c in beyan) / len(beyan) if beyan else None}


def fcr_isaretle(cagrilar: list[Cagri], tekrar_gun: int) -> None:
    son = max(c.bas for c in cagrilar)
    numara = defaultdict(list)
    for c in cagrilar:
        if c.arayan and c.durum in ("cevap", "terk"):
            numara[c.arayan].append(c)
    for lst in numara.values():
        lst.sort(key=lambda c: c.bas)
        for i, c in enumerate(lst):
            if c.durum != "cevap" or c.bas > son - timedelta(days=tekrar_gun):
                continue
            sonraki = lst[i + 1] if i + 1 < len(lst) else None
            c.fcr = not (sonraki and sonraki.bas - c.bas <= timedelta(days=tekrar_gun))


def aralik(dt: datetime) -> time:
    return time(dt.hour, 30 if dt.minute >= 30 else 0)


def analiz_et(cagrilar: list[Cagri], *, sl_sure: float = 20, sl_hedef: float = 0.80, tekrar_gun: int = 7, kayip: float = 0.30, kisa_gorusme: float = 10,
              terk_hedef: float = 0.05) -> dict:
    fcr_isaretle(cagrilar, tekrar_gun)
    uy = []
    genel = ozet(cagrilar, sl_sure)
    gunluk = {}
    for g in sorted({c.bas.date() for c in cagrilar}):
        gunluk[g] = ozet([c for c in cagrilar if c.bas.date() == g], sl_sure)
        o = gunluk[g]
        if o["sl"] is not None and o["sl"] < sl_hedef:
            uy.append({"onem": "Orta", "tur": "Servis seviyesi hedef altı", "kim": f"{g:%d.%m.%Y} {GUNLER[g.weekday()]}",
                       "aciklama": f"SL {yz(o['sl'], 1)} (hedef {yz(sl_hedef)}, {sl_sure:g} sn); gelen {o['gelen']}, terk {o['terk']}"})
    kuyruk = {k: ozet([c for c in cagrilar if c.kuyruk == k], sl_sure) for k in sorted({c.kuyruk for c in cagrilar})}
    for k, o in kuyruk.items():
        if o["terk_orani"] is not None and o["terk_orani"] > terk_hedef:
            uy.append({"onem": "Orta", "tur": "Terk oranı yüksek", "kim": k, "aciklama": f"Terk {yz(o['terk_orani'], 1)} (hedef ≤ {yz(terk_hedef)})"})

    # 30 dakikalık aralıklar (hafta içi ortalaması)
    hi = [c for c in cagrilar if c.bas.weekday() < 5]
    gun_sayisi = len({c.bas.date() for c in hi}) or 1
    araliklar = []
    for a in sorted({aralik(c.bas) for c in hi}):
        lst = [c for c in hi if aralik(c.bas) == a]
        o = ozet(lst, sl_sure)
        ort_gelen = o["gelen"] / gun_sayisi
        gerek = gereken_temsilci(ort_gelen, o["aht"] or 0, sl_sure, sl_hedef)
        planlanacak = math.ceil(gerek / (1 - kayip)) if gerek else 0
        fiili = statistics.fmean(len({c.temsilci for c in lst if c.durum == "cevap" and c.bas.date() == g and c.temsilci}) for g in {c.bas.date() for c in hi})
        araliklar.append({"aralik": a, "ort_gelen": ort_gelen, **o, "gerek": gerek, "planlanacak": planlanacak, "fiili": fiili,
                          "doluluk": (ort_gelen * (o["aht"] or 0) / 1800) / gerek if gerek else None})
        if gerek and fiili < gerek and o["sl"] is not None and o["sl"] < sl_hedef:
            uy.append({"onem": "Orta", "tur": "Aralıkta personel yetersiz", "kim": f"{a:%H:%M}",
                       "aciklama": f"Ortalama {sy(ort_gelen)} çağrı, AHT {o['aht']:.0f} sn → en az {gerek} temsilci gerekir (kayıp payıyla {planlanacak}); "
                                   f"cevaplayan temsilci ort. {sy(fiili)}; SL {yz(o['sl'])}"})

    temsilciler = []
    cev = [c for c in cagrilar if c.durum == "cevap" and c.temsilci]
    aht_lar = []
    for t in sorted({c.temsilci for c in cev}):
        lst = [c for c in cev if c.temsilci == t]
        toplam = sum(c.ele_alma for c in lst)
        x = {"temsilci": t, "cevap": len(lst), "aht": toplam / len(lst), "gorusme": statistics.fmean(c.gorusme for c in lst),
             "hold_pay": sum(c.hold for c in lst) / toplam if toplam else 0, "acw_pay": sum(c.acw for c in lst) / toplam if toplam else 0,
             "kisa": sum(c.gorusme < kisa_gorusme for c in lst), "fcr": None}
        f = [c for c in lst if c.fcr is not None]
        x["fcr"] = sum(c.fcr for c in f) / len(f) if f else None
        x["kisa_oran"] = x["kisa"] / len(lst)
        temsilciler.append(x)
        aht_lar.append(x["aht"])
    medyan = statistics.median(aht_lar) if aht_lar else 0
    for x in temsilciler:
        x["aht_kat"] = x["aht"] / medyan if medyan else None
        if x["kisa_oran"] > 0.05 and x["kisa"] >= 3:
            uy.append({"onem": "Orta", "tur": "Kısa görüşme", "kim": x["temsilci"],
                       "aciklama": f"{x['kisa']} görüşme {kisa_gorusme:g} saniyeden kısa ({yz(x['kisa_oran'])}); çağrı kayıtlarını dinleyin (erken kapatma / hat sorunu)"})
        if x["aht_kat"] and x["aht_kat"] > 1.5:
            uy.append({"onem": "Bilgi", "tur": "Yüksek AHT", "kim": x["temsilci"], "aciklama": f"AHT {x['aht']:.0f} sn, ekip medyanının {sy(x['aht_kat'])} katı; koçluk / eğitim ihtiyacına bakın"})

    # Tekrar arayanlar
    numara = defaultdict(list)
    for c in cagrilar:
        if c.arayan and c.durum in ("cevap", "terk"):
            numara[c.arayan].append(c)
    tekrar = []
    for no, lst in numara.items():
        lst.sort(key=lambda c: c.bas)
        en_cok = max(sum(1 for d in lst if timedelta(0) <= d.bas - c.bas <= timedelta(days=tekrar_gun)) for c in lst)
        if en_cok >= 3:
            tekrar.append({"arayan": maske(no), "cagri": len(lst), "pencere": en_cok, "ilk": lst[0].bas, "son": lst[-1].bas,
                           "kategoriler": ", ".join(sorted({c.kategori for c in lst if c.kategori})), "kuyruklar": ", ".join(sorted({c.kuyruk for c in lst}))})
    tekrar.sort(key=lambda x: (-x["pencere"], -x["cagri"]))
    if genel["sl"] is not None and genel["sl"] < sl_hedef:
        uy.append({"onem": "Yüksek", "tur": "Dönem SL hedef altı", "kim": "Genel", "aciklama": f"SL {yz(genel['sl'], 1)}, hedef {yz(sl_hedef)}"})
    if genel["terk_orani"] is not None and genel["terk_orani"] > terk_hedef:
        uy.append({"onem": "Yüksek", "tur": "Dönem terk oranı yüksek", "kim": "Genel", "aciklama": f"Terk {yz(genel['terk_orani'], 1)}, hedef ≤ {yz(terk_hedef)}"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"cagrilar": cagrilar, "genel": genel, "gunluk": gunluk, "kuyruk": kuyruk, "araliklar": araliklar, "temsilciler": temsilciler, "tekrar": tekrar,
            "uyarilar": uy, "parametre": {"sl_sure": sl_sure, "sl_hedef": sl_hedef, "tekrar_gun": tekrar_gun, "kayip": kayip, "terk_hedef": terk_hedef},
            "gun_sayisi": gun_sayisi}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
YF, SF = "0.0%", "0"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


GOSTERGE = ["gelen", "cevap", "terk", "karsilama", "terk_orani", "sl", "asa", "aht", "fcr"]
GOSTERGE_ADI = ["Gelen", "Cevaplanan", "Terk", "Karşılama", "Terk Oranı", "SL", "ASA (sn)", "AHT (sn)", "FCR (tekrar arama)"]
GOSTERGE_BICIM = [None, None, None, YF, YF, YF, SF, SF, YF]


def _satir(ws, ilk, o, hedef=None):
    ws.append(ilk + [o[k] for k in GOSTERGE])
    n = len(ilk)
    for j, b in enumerate(GOSTERGE_BICIM, n + 1):
        if b:
            ws.cell(ws.max_row, j).number_format = b
    if hedef and o["sl"] is not None:
        ws.cell(ws.max_row, n + 6).fill = PatternFill("solid", fgColor="E3F4E1" if o["sl"] >= hedef else "FDE2E1")


def rapor_yaz(cikti: Path, s: dict) -> None:
    p = s["parametre"]
    g = s["genel"]
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer", "Tanım"], (30, 14, 90))
    tarihler = sorted(s["gunluk"])
    for ad, v, b, tanim in [
        ("Dönem", f"{tarihler[0]:%d.%m.%Y} – {tarihler[-1]:%d.%m.%Y}", None, f"{len(tarihler)} gün"),
        ("Gelen çağrı", g["gelen"], None, f"Kısa terk ({g['kisa_terk']} çağrı) hariç"),
        ("Cevaplanan", g["cevap"], None, ""), ("Terk", g["terk"], None, ""),
        ("Karşılama oranı", g["karsilama"], YF, "Cevaplanan / gelen"), ("Terk oranı", g["terk_orani"], YF, f"Hedef ≤ %{p['terk_hedef'] * 100:.0f}"),
        ("Servis seviyesi", g["sl"], YF, f"{p['sl_sure']:g} sn içinde cevaplanan / (cevaplanan + {p['sl_sure']:g} sn'den sonra terk); hedef %{p['sl_hedef'] * 100:.0f}"),
        ("ASA (sn)", g["asa"], SF, "Cevaplanan çağrılarda ortalama bekleme"), ("En uzun bekleme (sn)", g["en_uzun"], SF, ""),
        ("Terk edenlerin ort. beklemesi (sn)", g["terk_bekleme"], SF, ""),
        ("AHT (sn)", g["aht"], SF, f"Görüşme {g['gorusme']:.0f} + beklemeye alma {g['hold']:.0f} + çağrı sonrası iş {g['acw']:.0f}" if g["aht"] else ""),
        ("FCR (tekrar arama)", g["fcr"], YF, f"{p['tekrar_gun']} gün içinde aynı numaradan yeni çağrı gelmeyen cevaplanmış çağrılar; {g['fcr_n']} çağrı değerlendirildi"),
        ("FCR (temsilci beyanı)", g["fcr_beyan"], YF, "'Çözüldü' sütunu varsa")]:
        oz.append([ad, v, tanim])
        if b:
            oz.cell(oz.max_row, 2).number_format = b
    oz.cell(8, 2).fill = PatternFill("solid", fgColor="E3F4E1" if g["sl"] is not None and g["sl"] >= p["sl_hedef"] else "FDE2E1")
    oz.append([])
    _baslik(oz, ["Kuyruk"] + GOSTERGE_ADI, ())
    for k, o in s["kuyruk"].items():
        _satir(oz, [k], o, p["sl_hedef"])

    gn = wb.create_sheet("Günlük")
    _baslik(gn, ["Tarih", "Gün"] + GOSTERGE_ADI, (11, 10, 8, 10, 7, 10, 9, 8, 8, 8, 10))
    for d, o in s["gunluk"].items():
        _satir(gn, [d, GUNLER[d.weekday()]], o, p["sl_hedef"])
        gn.cell(gn.max_row, 1).number_format = "DD.MM.YYYY"

    ar = wb.create_sheet("Aralıklar")
    _baslik(ar, ["Aralık", "Ort. Gelen (gün)", "SL", "ASA (sn)", "AHT (sn)", "Terk Oranı", "Gereken Temsilci (Erlang C)", "Planlanacak (kayıp payıyla)",
                 "Cevaplayan Temsilci (ort.)", "Fark", "Doluluk"], (8, 10, 8, 8, 8, 8, 12, 12, 12, 7, 8))
    for x in s["araliklar"]:
        ar.append([x["aralik"].strftime("%H:%M"), round(x["ort_gelen"], 1), x["sl"], x["asa"], x["aht"], x["terk_orani"], x["gerek"], x["planlanacak"],
                   round(x["fiili"], 1), round(x["fiili"] - x["gerek"], 1), x["doluluk"]])
        for j, b in ((3, YF), (4, SF), (5, SF), (6, YF), (11, YF)):
            ar.cell(ar.max_row, j).number_format = b
        if x["fiili"] < x["gerek"]:
            ar.cell(ar.max_row, 10).fill = PatternFill("solid", fgColor="FDE2E1")
    son = ar.max_row
    ar.append([])
    ar.append([f"Hafta içi {s['gun_sayisi']} günün ortalaması. Gereken temsilci: hedef SL %{p['sl_hedef'] * 100:.0f} / {p['sl_sure']:g} sn için Erlang C; "
               f"planlanacak = gereken / (1 − %{p['kayip'] * 100:.0f} kayıp: mola, eğitim, devamsızlık)."])
    ar.append(["Cevaplayan temsilci, o aralıkta en az bir çağrı cevaplayan temsilci sayısıdır; vardiyadaki gerçek temsilci sayısı değildir."])
    if son > 1:
        gr = BarChart()
        gr.title, gr.height, gr.width = "Aralık bazında ortalama gelen çağrı", 8, 22
        gr.add_data(Reference(ar, min_col=2, min_row=1, max_row=son), titles_from_data=True)
        gr.set_categories(Reference(ar, min_col=1, min_row=2, max_row=son))
        ar.add_chart(gr, f"A{son + 5}")
        ln = LineChart()
        ln.title, ln.height, ln.width = "Gereken ve cevaplayan temsilci", 8, 22
        ln.add_data(Reference(ar, min_col=7, min_row=1, max_row=son), titles_from_data=True)
        ln.add_data(Reference(ar, min_col=9, min_row=1, max_row=son), titles_from_data=True)
        ln.set_categories(Reference(ar, min_col=1, min_row=2, max_row=son))
        ar.add_chart(ln, f"M{son + 5}")

    te = wb.create_sheet("Temsilciler")
    _baslik(te, ["Temsilci", "Cevaplanan", "AHT (sn)", "Medyana Oranı", "Ort. Görüşme (sn)", "Beklemeye Alma Payı", "Çağrı Sonrası İş Payı",
                 f"Kısa Görüşme", "Kısa Görüşme Oranı", "FCR"], (18, 10, 9, 9, 10, 10, 10, 9, 9, 8))
    for x in sorted(s["temsilciler"], key=lambda x: -x["cevap"]):
        te.append([x["temsilci"], x["cevap"], round(x["aht"]), round(x["aht_kat"], 2) if x["aht_kat"] else None, round(x["gorusme"]), x["hold_pay"], x["acw_pay"],
                   x["kisa"], x["kisa_oran"], x["fcr"]])
        for j in (6, 7, 9, 10):
            te.cell(te.max_row, j).number_format = YF
        if x["aht_kat"] and x["aht_kat"] > 1.5:
            te.cell(te.max_row, 4).fill = PatternFill("solid", fgColor="FFF4CE")
        if x["kisa_oran"] > 0.05 and x["kisa"] >= 3:
            te.cell(te.max_row, 9).fill = PatternFill("solid", fgColor="FFF4CE")

    tk = wb.create_sheet("Tekrar Arayanlar")
    _baslik(tk, ["Arayan (maskeli)", "Toplam Çağrı", f"{p['tekrar_gun']} Gün İçinde En Çok", "İlk", "Son", "Kategoriler", "Kuyruklar"], (16, 9, 12, 16, 16, 30, 20))
    for x in s["tekrar"]:
        tk.append([x["arayan"], x["cagri"], x["pencere"], x["ilk"], x["son"], x["kategoriler"], x["kuyruklar"]])
        tk.cell(tk.max_row, 4).number_format = tk.cell(tk.max_row, 5).number_format = "DD.MM.YYYY HH:MM"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kim / Ne Zaman", "Açıklama"], (9, 26, 22, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(yol: Path, cikti: Path, *, kisa_terk: float = 5, **kw) -> dict:
    cagrilar, uy = oku(yol, kisa_terk)
    if not cagrilar:
        raise ValueError(f"{yol.name}: çağrı bulunamadı")
    s = analiz_et(cagrilar, **kw)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(uy + s["uyarilar"], key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Çağrı verisinden karşılama, servis seviyesi, AHT, FCR ve personel ihtiyacını hesaplar.")
    p.add_argument("--cagrilar", type=Path, default=ORNEK / "cagrilar.csv",
                   help="Çağrı ID, Başlangıç, Kuyruk, Temsilci, Bekleme, Görüşme, Beklemeye Alma, Çağrı Sonrası İş (sn veya ss:dd:sn), Durum, Arayan, "
                        "Kategori, Çözüldü")
    p.add_argument("--sl-sure", type=float, default=20, help="Servis seviyesi süresi T (sn; varsayılan 20)")
    p.add_argument("--sl-hedef", type=float, default=80, help="Servis seviyesi hedefi %% (varsayılan 80)")
    p.add_argument("--kisa-terk", type=float, default=5, help="Bu kadar saniyeden önce kapatılan çağrı sayılmaz (varsayılan 5)")
    p.add_argument("--tekrar-gun", type=int, default=7, help="FCR: kaç gün içinde tekrar arama çözülmemiş sayılır (varsayılan 7)")
    p.add_argument("--kayip", type=float, default=30, help="Personel planında kayıp payı %% (mola, eğitim, devamsızlık; varsayılan 30)")
    p.add_argument("--terk-hedef", type=float, default=5, help="Terk oranı hedefi %% (varsayılan 5)")
    p.add_argument("--kisa-gorusme", type=float, default=10, help="Bu kadar saniyeden kısa görüşme işaretlenir (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "cagri_merkezi_raporu.xlsx")
    a = p.parse_args(argv)
    if not a.cagrilar.exists():
        print(f"[X] Dosya bulunamadı: {a.cagrilar}")
        return 1
    if not (0 < a.sl_hedef < 100 and 0 <= a.kayip < 100):
        print("[X] --sl-hedef ve --kayip 0–100 arasında olmalı")
        return 2
    try:
        s = calistir(a.cagrilar, a.cikti, kisa_terk=a.kisa_terk, sl_sure=a.sl_sure, sl_hedef=a.sl_hedef / 100, tekrar_gun=a.tekrar_gun, kayip=a.kayip / 100,
                     kisa_gorusme=a.kisa_gorusme, terk_hedef=a.terk_hedef / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = s["genel"]
    print(f"[OK] Gelen {g['gelen']} · cevaplanan {g['cevap']} · karşılama {yz(g['karsilama'], 1)} · terk {yz(g['terk_orani'], 1)} · SL {yz(g['sl'], 1)} · "
          f"ASA {g['asa'] or 0:.0f} sn · AHT {g['aht'] or 0:.0f} sn · FCR {yz(g['fcr'], 1)}")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
