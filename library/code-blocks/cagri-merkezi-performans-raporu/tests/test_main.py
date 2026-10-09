import csv
import math
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class ErlangTesti(unittest.TestCase):
    def test_bilinen_degerler(self):
        self.assertAlmostEqual(main.erlang_c(0.5, 1), 0.5)                      # M/M/1: P(bekleme) = ρ
        self.assertAlmostEqual(main.erlang_c(1.0, 2), 1 / 3)                    # M/M/2, A = 1
        self.assertEqual(main.erlang_c(3.0, 3), 1.0)                            # N ≤ A: kuyruk sınırsız büyür
        # M/M/1: SL = 1 − ρ·e^(−(1−ρ)·T/AHT); 30 dk'da 9 çağrı × 100 sn = 0,5 Erlang
        self.assertAlmostEqual(main.servis_seviyesi(9, 100, 1, 20), 1 - 0.5 * math.exp(-0.5 * 20 / 100))

    def test_gereken_temsilci(self):
        n = main.gereken_temsilci(30, 300, 20, 0.80)                            # 5 Erlang
        self.assertGreaterEqual(main.servis_seviyesi(30, 300, n, 20), 0.80)
        self.assertLess(main.servis_seviyesi(30, 300, n - 1, 20), 0.80)         # en küçük sayı
        self.assertEqual(main.gereken_temsilci(0, 300, 20, 0.8), 0)


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "c.xlsx"
        cls.s = main.calistir(ORNEK / "cagrilar.csv", cls.cikti)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_bagimsiz_hesap(self):
        cevap = terk = kisa = sl_pay = sl_terk = 0
        with open(ORNEK / "cagrilar.csv", encoding="utf-8") as f:
            for r in csv.DictReader(f, delimiter=";"):
                b = float(r["Bekleme (sn)"])
                if r["Durum"] == "Cevaplandı":
                    cevap += 1
                    sl_pay += b <= 20
                elif b < 5:
                    kisa += 1
                else:
                    terk += 1
                    sl_terk += b > 20
        g = self.s["genel"]
        self.assertEqual((g["cevap"], g["terk"], g["kisa_terk"], g["gelen"]), (cevap, terk, kisa, cevap + terk))
        self.assertAlmostEqual(g["sl"], sl_pay / (cevap + sl_terk))

    def test_fcr_ve_tekrar(self):
        son = max(c.bas for c in self.s["cagrilar"])
        degerlendirilen = [c for c in self.s["cagrilar"] if c.fcr is not None]
        self.assertTrue(all((son - c.bas).days >= 7 for c in degerlendirilen))      # veri sonuna yakın çağrılar dışarıda
        tekrar = [c for c in self.s["cagrilar"] if c.arayan == "05551234567"]
        self.assertEqual([c.fcr for c in tekrar], [False, False, False, True])       # 16.09, 17.09 (×2) → 7 gün içinde yeniden aradı; 21.09 son
        self.assertIn("*********67", {x["arayan"] for x in self.s["tekrar"]})

    def test_temsilci_ve_uyarilar(self):
        t = {x["temsilci"]: x for x in self.s["temsilciler"]}
        self.assertGreater(t["Temsilci 07"]["aht_kat"], 1.5)
        self.assertGreater(t["Temsilci 11"]["kisa_oran"], 0.05)
        turler = {(u["tur"], u["kim"]) for u in self.s["uyarilar"]}
        self.assertTrue({("Kısa görüşme", "Temsilci 11"), ("Yüksek AHT", "Temsilci 07"), ("Servis seviyesi hedef altı", "14.09.2026 Pazartesi")} <= turler)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Günlük", "Aralıklar", "Temsilciler", "Tekrar Arayanlar", "Uyarılar"])
        self.assertNotIn("05551234567", str([c.value for r in wb["Tekrar Arayanlar"].iter_rows() for c in r]))


class KuralTesti(unittest.TestCase):
    def test_sure_ve_maske(self):
        self.assertEqual(main.sure("01:35"), 95)
        self.assertEqual(main.sure("00:01:35"), 95)
        self.assertEqual(main.sure("12,5"), 12.5)
        self.assertEqual(main.maske("0532 111 22 33"), "*********33")

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--sl-hedef", "120"]), 2)
            self.assertEqual(main.main(["--cagrilar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
