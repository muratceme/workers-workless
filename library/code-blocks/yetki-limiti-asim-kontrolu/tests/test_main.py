import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
YETKI = {("ali", "kasko"): {"personel": "Ali", "unvan": "", "brans": "Kasko", "bedel": D(1000), "indirim": D(10), "pb": "TRY"},
         ("veli", "*"): {"personel": "Veli", "unvan": "", "brans": "Tümü", "bedel": D(5000), "indirim": None, "pb": "TRY"}}


def t(no, kim, brans, bedel, ind=None, sig="", gun=1, pb="TRY"):
    return {"no": no, "tarih": date(2026, 9, 1) + timedelta(days=gun - 1), "brans": brans, "onaylayan": kim, "sigortali": sig, "bedel": D(bedel), "indirim": None if ind is None else D(ind), "prim": None, "pb": pb}


class KontrolTesti(unittest.TestCase):
    def test_kontrol(self):
        l = [t("1", "Ali", "Kasko", 900), t("2", "Ali", "Kasko", 1100), t("3", "Ali", "Kasko", 500, 12), t("4", "Ali", "Kasko", 500, 9.5), t("5", "Ali", "Konut", 100),
             t("6", "Veli", "Yangın", 4000, 50), t("7", "Veli", "Yangın", 100, pb="EUR"), t("8", "Ali", "kasko", 600, sig="X", gun=1), t("9", "Ali", "Kasko", 600, sig="X", gun=20),
             t("10", "Ali", "Kasko", 600, sig="Y", gun=1), t("11", "Ali", "Kasko", 600, sig="Y", gun=40)]
        r = main.kontrol(l, YETKI)
        s = {x["no"]: x["sonuc"] for x in r["teklifler"]}
        self.assertEqual(s, {"1": "Limite yakın", "2": "Limit aşımı", "3": "Limit aşımı", "4": "Limite yakın", "5": "Yetki yok", "6": "Uygun", "7": "Kontrol edilemedi",
                             "8": "Uygun", "9": "Uygun", "10": "Uygun", "11": "Uygun"})
        self.assertEqual([x["no"] for x in r["teklifler"]][:3], ["2", "3", "5"])
        self.assertEqual([(b["sigortali"], b["teklifler"], b["toplam"]) for b in r["bolme"]], [("X", ["8", "9"], D(1200))])
        self.assertEqual([(p["ad"], p["teklif"], p["asim"], p["yetkisiz"]) for p in r["personel"]], [("Ali", 9, 2, 1), ("Veli", 2, 0, 0)])
        self.assertEqual(sorted(u["tur"] for u in r["uyarilar"]), ["Bedel limiti aşıldı", "Bölme şüphesi", "Para birimi farklı", "Yetki tanımı yok", "İndirim limiti aşıldı"])
        self.assertEqual(main.kontrol([t("8", "Ali", "Kasko", 600, sig="X", gun=1), t("9", "Ali", "Kasko", 600, sig="X", gun=20)], YETKI, gun=10)["bolme"], [])
        with self.assertRaises(ValueError):
            main.kontrol([], YETKI)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "yetkiler.csv", ORNEK / "teklifler.csv", cikti)
            l = r["teklifler"]
            self.assertEqual((len(l), sum(1 for x in l if x["sonuc"] == "Limit aşımı"), sum(1 for x in l if x["sonuc"] == "Yetki yok"), sum(1 for x in l if x["sonuc"] == "Limite yakın")),
                             (134, 12, 1, 15))
            self.assertEqual([b["sigortali"] for b in r["bolme"]], ["Firma 158", "Firma Z"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Teklif Kontrolü", "Personel", "Bölme Şüphesi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            yy, ty = Path(tmp) / "y.csv", Path(tmp) / "t.csv"
            yy.write_text("Kullanıcı;Ürün;Limit;İndirim Limiti\nA;;1.000;10\nB;Kasko;x;\nC;Kasko;5;150\n", encoding="utf-8")
            ty.write_text("Poliçe No;Teknik Personel;Bedel;Uygulanan İndirim;Tarih\n1;A;900;5;01.09.2026\n2;A;x;;\n3;;5;;\n4;A;5;;x\n", encoding="utf-8")
            y, uy = main.yetki_oku(yy)
            self.assertEqual((list(y), len(uy)), ([("a", "*")], 2))
            l, uy = main.teklif_oku(ty)
            self.assertEqual(([x["no"] for x in l], len(uy)), (["1"], 3))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--yetkiler", str(yy), "--teklifler", str(ty), "--cikti", c]
            self.assertEqual(main.main(arg + ["--yakin", "80", "--gun", "60"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--yakin", "0"], ["--gun", "0"])], [2, 2])
            self.assertEqual(main.main(["--yetkiler", str(Path(tmp) / "yok.csv")]), 1)
            ty.write_text("Poliçe No;Teknik Personel;Bedel\n1;;\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
