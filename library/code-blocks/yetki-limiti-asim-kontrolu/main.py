"""
Yetki Limiti Aşım Kontrolü — Gerisi Bende kod bloğu
Teknik (Risk Kabul / Underwriting) › Teknik Müdür
Onaylanan tekliflerin teknik personel yetki limitlerine (branş bazında azami sigorta bedeli ve azami indirim) uygunluğunu kontrol eder.
  Yetki satırı      = onaylayan + branş; branşa özel satır yoksa onaylayanın "Tümü" satırı
  Bedel aşımı       = sigorta bedeli > azami bedel;  indirim aşımı = uygulanan indirim % > azami indirim %
  Limite yakın      = bedel ya da indirim limitin --yakin (%90) ve üstünde
  Bölme şüphesi     = aynı sigortalı + branş, aynı onaylayan, --gun (30) gün içindeki tekliflerin bedel toplamı limiti aşıyor (tek tek aşmıyor)
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
YETKI = {"personel": ("personel", "ad soyad", "kullanici", "onaylayan"), "unvan": ("unvan", "gorev"), "brans": ("brans", "urun"), "bedel": ("azami sigorta bedeli", "azami bedel", "bedel limiti", "limit"),
         "indirim": ("azami indirim", "indirim limiti", "azami indirim yuzde"), "pb": ("para birimi", "doviz")}
TEKLIF = {"no": ("teklif no", "police no", "no"), "tarih": ("onay tarihi", "tarih"), "brans": ("brans", "urun"), "onaylayan": ("onaylayan", "onaylayan kullanici", "teknik personel"),
          "sigortali": ("sigortali", "musteri", "sigorta ettiren"), "bedel": ("sigorta bedeli", "bedel", "toplam bedel"), "indirim": ("indirim", "indirim orani", "uygulanan indirim"),
          "prim": ("prim", "net prim", "brut prim"), "pb": ("para birimi", "doviz")}
SIFIR = Decimal(0)


def yetki_oku(yol: Path) -> tuple[dict, list[dict]]:
    uy, d = [], {}
    for r in kayitlar(yol, YETKI, ("personel", "bedel")):
        p = metin(r["personel"])
        if not p:
            continue
        b = para(r["bedel"])
        i = para(r["indirim"]) if metin(r["indirim"]) else None
        if b is None or b < 0 or (metin(r["indirim"]) and (i is None or not 0 <= i <= 100)):
            uy.append(uyari("Yüksek", "Yetki satırı okunamadı", p, f"Satır {r['_satir']}: azami bedel ya da indirim okunamadı."))
            continue
        br = katla(r["brans"]) if katla(r["brans"]) not in ("", "tumu", "hepsi", "tum branslar") else "*"
        d[(katla(p), br)] = {"personel": p, "unvan": metin(r["unvan"]), "brans": metin(r["brans"]) or "Tümü", "bedel": b, "indirim": i, "pb": metin(r["pb"]).upper() or "TRY"}
    return d, uy


def teklif_oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, TEKLIF, ("onaylayan", "bedel")):
        no = metin(r["no"]) or f"Satır {r['_satir']}"
        if not metin(r["onaylayan"]) and not metin(r["bedel"]):
            continue
        b, t = para(r["bedel"]), tarih(r["tarih"]) if metin(r["tarih"]) else None
        i, pr = (para(r[a]) if metin(r[a]) else None for a in ("indirim", "prim"))
        if b is None or b < 0 or not metin(r["onaylayan"]) or (metin(r["tarih"]) and t is None) or (metin(r["indirim"]) and i is None) or (metin(r["prim"]) and pr is None):
            uy.append(uyari("Yüksek", "Teklif satırı okunamadı", no, f"Satır {r['_satir']}: onaylayan, bedel, tarih, indirim ya da prim okunamadı."))
            continue
        l.append({"no": no, "tarih": t, "brans": metin(r["brans"]) or "Belirtilmemiş", "onaylayan": metin(r["onaylayan"]), "sigortali": metin(r["sigortali"]), "bedel": b, "indirim": i,
                  "prim": pr, "pb": metin(r["pb"]).upper() or "TRY"})
    return l, uy


def kontrol(teklifler: list[dict], yetki: dict, yakin: Decimal = Decimal(90), gun: int = 30) -> dict:
    if not teklifler:
        raise ValueError("Okunabilen teklif yok.")
    uy = []
    for t in teklifler:
        y = yetki.get((katla(t["onaylayan"]), katla(t["brans"]))) or yetki.get((katla(t["onaylayan"]), "*"))
        t["yetki"] = y
        t["sonuc"], t["notlar"] = "Uygun", []
        kim = f"{t['no']} · {t['onaylayan']}"
        if y is None:
            t["sonuc"] = "Yetki yok"
            uy.append(uyari("Yüksek", "Yetki tanımı yok", kim, f"{t['brans']} branşında onaylayan için yetki satırı yok."))
            continue
        if y["pb"] != t["pb"]:
            t["sonuc"] = "Kontrol edilemedi"
            uy.append(uyari("Orta", "Para birimi farklı", kim, f"Teklif {t['pb']}, yetki limiti {y['pb']}; kurla çevirip elle kontrol edin."))
            continue
        t["bedel_oran"] = yuvarla(t["bedel"] * 100 / y["bedel"]) if y["bedel"] else None
        if t["bedel"] > y["bedel"]:
            t["notlar"].append("Bedel aşımı")
            uy.append(uyari("Yüksek", "Bedel limiti aşıldı", kim, f"{t['brans']}: bedel {t['bedel']:,.2f} {t['pb']}, limit {y['bedel']:,.2f} (%{t['bedel_oran']}); üst yetki onayı gerekirdi."))
        elif t["bedel_oran"] is not None and t["bedel_oran"] >= yakin:
            t["notlar"].append("Bedel limite yakın")
        if t["indirim"] is not None and y["indirim"] is not None:
            if t["indirim"] > y["indirim"]:
                t["notlar"].append("İndirim aşımı")
                uy.append(uyari("Yüksek", "İndirim limiti aşıldı", kim, f"{t['brans']}: indirim %{t['indirim']}, limit %{y['indirim']}."))
            elif y["indirim"] and t["indirim"] * 100 >= y["indirim"] * yakin:
                t["notlar"].append("İndirim limite yakın")
        if any("aşımı" in n for n in t["notlar"]):
            t["sonuc"] = "Limit aşımı"
        elif t["notlar"]:
            t["sonuc"] = "Limite yakın"
    gruplar = {}
    for t in teklifler:
        if t["yetki"] and t["sigortali"] and t["tarih"] and t["sonuc"] in ("Uygun", "Limite yakın"):
            gruplar.setdefault((katla(t["sigortali"]), katla(t["brans"]), katla(t["onaylayan"])), []).append(t)
    bolme = []
    for l in gruplar.values():
        if len(l) < 2:
            continue
        l.sort(key=lambda t: t["tarih"])
        for i, t in enumerate(l):
            pencere = [x for x in l[i:] if (x["tarih"] - t["tarih"]).days <= gun]
            top = sum((x["bedel"] for x in pencere), SIFIR)
            if len(pencere) > 1 and top > t["yetki"]["bedel"]:
                bolme.append({"sigortali": t["sigortali"], "brans": t["brans"], "onaylayan": t["onaylayan"], "teklifler": [x["no"] for x in pencere], "toplam": top, "limit": t["yetki"]["bedel"]})
                uy.append(uyari("Orta", "Bölme şüphesi", f"{t['sigortali']} · {t['brans']}", f"{t['onaylayan']} {gun} gün içinde {len(pencere)} teklif onayladı ({', '.join(x['no'] for x in pencere)}); "
                                f"toplam bedel {top:,.2f}, limit {t['yetki']['bedel']:,.2f}. Riskin bölünüp bölünmediğini inceleyin."))
                break
    personel = {}
    for t in teklifler:
        p = personel.setdefault(t["onaylayan"], {"ad": t["onaylayan"], "teklif": 0, "asim": 0, "yakin": 0, "yetkisiz": 0})
        p["teklif"] += 1
        p["asim"] += t["sonuc"] == "Limit aşımı"
        p["yakin"] += t["sonuc"] == "Limite yakın"
        p["yetkisiz"] += t["sonuc"] == "Yetki yok"
    sira = {"Limit aşımı": 0, "Yetki yok": 1, "Kontrol edilemedi": 2, "Limite yakın": 3, "Uygun": 4}
    return {"teklifler": sorted(teklifler, key=lambda t: (sira[t["sonuc"]], t["tarih"] or date.min, t["no"])), "bolme": bolme,
            "personel": sorted(personel.values(), key=lambda p: (-p["asim"], -p["yetkisiz"], katla(p["ad"]))), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teklif Kontrolü"
    _baslik(ws, ["Teklif No", "Onay Tarihi", "Branş", "Onaylayan", "Sigortalı", "Sigorta Bedeli", "Bedel Limiti", "Limit Kullanımı %", "İndirim %", "İndirim Limiti %", "Para Birimi", "Sonuç", "Not"],
            (12, 11, 14, 16, 22, 15, 15, 9, 8, 9, 7, 15, 30), "B2")
    renk = {"Limit aşımı": KIRMIZI, "Yetki yok": KIRMIZI, "Kontrol edilemedi": SARI, "Limite yakın": SARI, "Uygun": YESIL}
    for t in r["teklifler"]:
        y = t["yetki"] or {}
        _satir(ws, [t["no"], t["tarih"], t["brans"], t["onaylayan"], t["sigortali"], t["bedel"], y.get("bedel"), t.get("bedel_oran"), t["indirim"], y.get("indirim"), t["pb"], t["sonuc"],
                    ", ".join(t["notlar"])], para_s=(6, 7, 8, 9, 10), tarih_s=(2,), dolgu={12: renk[t["sonuc"]]})
    ws = wb.create_sheet("Personel")
    _baslik(ws, ["Onaylayan", "Teklif", "Limit Aşımı", "Yetkisiz", "Limite Yakın"], (20, 8, 10, 9, 10))
    for p in r["personel"]:
        _satir(ws, [p["ad"], p["teklif"], p["asim"], p["yetkisiz"], p["yakin"]], dolgu={3: KIRMIZI if p["asim"] else None})
    ws = wb.create_sheet("Bölme Şüphesi")
    _baslik(ws, ["Sigortalı", "Branş", "Onaylayan", "Teklifler", "Toplam Bedel", "Limit"], (22, 14, 16, 30, 15, 15))
    for b in r["bolme"]:
        _satir(ws, [b["sigortali"], b["brans"], b["onaylayan"], ", ".join(b["teklifler"]), b["toplam"], b["limit"]], para_s=(5, 6))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yetki_yol: Path, teklif_yol: Path, cikti: Path, yakin: Decimal = Decimal(90), gun: int = 30) -> dict:
    y, uy1 = yetki_oku(yetki_yol)
    t, uy2 = teklif_oku(teklif_yol)
    r = kontrol(t, y, yakin, gun)
    r["uyarilar"] = sirala(uy1 + uy2 + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Onaylanan tekliflerin teknik personel yetki limitlerine uygunluğunu kontrol eder.")
    p.add_argument("--yetkiler", type=Path, help="Personel, Unvan, Branş (boş / Tümü), Azami Sigorta Bedeli, Azami İndirim %%, Para Birimi")
    p.add_argument("--teklifler", type=Path, help="Teklif No, Onay Tarihi, Branş, Onaylayan, Sigortalı, Sigorta Bedeli, İndirim %%, Prim, Para Birimi")
    p.add_argument("--yakin", default="90", help="Limitin yüzde kaçı 'limite yakın' sayılır (varsayılan 90)")
    p.add_argument("--gun", type=int, default=30, help="Bölme şüphesi penceresi, gün (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "yetki_limiti_kontrolu.xlsx")
    a = p.parse_args(argv)
    yy, ty = a.yetkiler or ORNEK / "yetkiler.csv", a.teklifler or ORNEK / "teklifler.csv"
    yk = para(a.yakin)
    if yk is None or not 0 < yk <= 100 or a.gun < 1:
        print("[X] --yakin 1-100, --gun en az 1 olmalı")
        return 2
    for y in (yy, ty):
        if not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yy, ty, a.cikti, yk, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    l = r["teklifler"]
    print(f"[OK] {len(l)} teklif · limit aşımı {sum(1 for t in l if t['sonuc'] == 'Limit aşımı')} · yetkisiz {sum(1 for t in l if t['sonuc'] == 'Yetki yok')} · "
          f"limite yakın {sum(1 for t in l if t['sonuc'] == 'Limite yakın')} · bölme şüphesi {len(r['bolme'])}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
