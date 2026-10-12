"""
Envanter Kayıp Risk Skoru — Workers / Workless kod bloğu
Kayıp Önleme › Kayıp Önleme Uzmanı

Sayım farkı, fire, iade ve iptal göstergelerinden mağaza bazında kayıp risk skoru üretir, denetim önceliği çıkarır.
  Oran %          = gösterge tutarı ÷ net satış × 100          (sayım farkında yalnız eksik dikkate alınır; fazla 0 sayılır)
  Gösterge puanı  = oran ÷ eşik × 50, en çok 100               (oran eşikteyse 50, eşiğin iki katı ve üzerindeyse 100)
  Risk skoru      = Σ(ağırlık × puan) ÷ Σ(ağırlık)             (verisi olmayan gösterge hesaptan çıkar, ağırlıklar yeniden dağıtılır)
  Risk sınıfı     = skor ≥ --yuksek (70): Yüksek;  ≥ --orta (40): Orta;  altı: Düşük
  Denetim önceliği= skora göre büyükten küçüğe sıra
Eşikler (sayım %1, fire %1, iade %5, iptal %2) ve ağırlıklar (40, 20, 20, 20) varsayılandır; sektör standardı değildir, kendi verinize göre ayarlayın.
Skor bir önceliklendirme aracıdır; hırsızlık veya suistimal tespiti değildir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"magaza": ("magaza", "magaza adi", "sube"), "bolge": ("bolge", "il"), "satis": ("net satis", "satis", "ciro"), "sayim": ("sayim farki", "sayim farki tutari", "envanter farki", "sayim"),
         "fire": ("fire", "fire tutari", "imha"), "iade": ("iade", "iade tutari", "musteri iadesi"), "iptal": ("iptal", "iptal tutari", "fis iptal", "fis iptal tutari")}
GOSTERGE = (("sayim", "Sayım eksiği"), ("fire", "Fire"), ("iade", "İade"), ("iptal", "İptal"))
VARSAYILAN_ESIK = {"sayim": Decimal(1), "fire": Decimal(1), "iade": Decimal(5), "iptal": Decimal(2)}
VARSAYILAN_AGIRLIK = {"sayim": Decimal(40), "fire": Decimal(20), "iade": Decimal(20), "iptal": Decimal(20)}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("magaza", "satis")):
        m = metin(r["magaza"])
        if not m:
            continue
        kim = f"{m} (satır {r['_satir']})"
        satis = para(r["satis"])
        if satis is None or satis <= 0:
            uy.append(uyari("Yüksek", "Net satış okunamadı", kim, f"Net satış '{metin(r['satis'])}'. Mağaza skorlanmadı."))
            continue
        x = {"magaza": m, "bolge": metin(r["bolge"]), "satis": satis, "sayim_fazlasi": Z}
        for a, ad in GOSTERGE:
            v = para(r[a])
            if v is None:
                if metin(r[a]):
                    uy.append(uyari("Orta", "Değer okunamadı", kim, f"{ad} '{metin(r[a])}'. Bu gösterge mağaza için hesaba katılmadı."))
            elif a == "sayim":
                if v > 0:
                    x["sayim_fazlasi"] = v
                    uy.append(uyari("Bilgi", "Sayım fazlası", kim, f"Sayım farkı +{v:f}: fazla çıkmış. Kayıp 0 sayıldı; fazla da kayıt hatasına (yanlış ürün okutma, giriş hatası) işaret edebilir."))
                v = max(Z, -v)
            elif v < 0:
                v = -v
            x[a] = v
        eksik = [ad for a, ad in GOSTERGE if x.get(a) is None]
        if len(eksik) == 4:
            uy.append(uyari("Yüksek", "Gösterge yok", kim, "Sayım farkı, fire, iade ve iptal sütunlarının hiçbiri dolu değil. Mağaza skorlanmadı."))
            continue
        if eksik:
            uy.append(uyari("Orta", "Eksik gösterge", kim, f"{', '.join(eksik)} verisi yok; skor kalan göstergelerle hesaplandı. Mağazalar arası karşılaştırma tam değildir."))
        l.append(x)
    return l, uy


def skorla(magazalar: list[dict], esik: dict | None = None, agirlik: dict | None = None, yuksek: Decimal = Decimal(70), orta: Decimal = Decimal(40)) -> dict:
    if not magazalar:
        raise ValueError("Skorlanacak mağaza yok.")
    esik, agirlik = VARSAYILAN_ESIK | (esik or {}), VARSAYILAN_AGIRLIK | (agirlik or {})
    zincir = {}
    for a, _ in GOSTERGE:
        dolu = [m for m in magazalar if m.get(a) is not None]
        zincir[a] = yuvarla(sum((m[a] for m in dolu), Z) / sum((m["satis"] for m in dolu), Z) * 100) if dolu else None
    for m in magazalar:
        m["oran"], m["puan"], pay, payda = {}, {}, Z, Z
        for a, _ in GOSTERGE:
            if m.get(a) is None or not agirlik[a]:
                m["oran"][a] = yuvarla(m[a] / m["satis"] * 100) if m.get(a) is not None else None
                m["puan"][a] = None
                continue
            oran = m[a] / m["satis"] * 100
            m["oran"][a], m["puan"][a] = yuvarla(oran), yuvarla(min(Decimal(100), oran / esik[a] * 50))
            pay += agirlik[a] * m["puan"][a]
            payda += agirlik[a]
        m["skor"] = yuvarla(pay / payda) if payda else Z
        m["sinif"] = "Yüksek" if m["skor"] >= yuksek else "Orta" if m["skor"] >= orta else "Düşük"
        m["asan"] = [ad for a, ad in GOSTERGE if m["oran"][a] is not None and m["oran"][a] > esik[a]]
        katki = {ad: agirlik[a] * m["puan"][a] for a, ad in GOSTERGE if m["puan"][a]}
        m["baskin"] = max(katki, key=katki.get) if katki else ""
    sirali = sorted(magazalar, key=lambda m: (-m["skor"], m["magaza"]))
    for i, m in enumerate(sirali, 1):
        m["oncelik"] = i
    return {"magazalar": sirali, "zincir": zincir, "esik": esik, "agirlik": agirlik, "sinir": (yuksek, orta), "say": {s: sum(m["sinif"] == s for m in sirali) for s in ("Yüksek", "Orta", "Düşük")}}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb = Workbook()
    renk = {"Yüksek": KIRMIZI, "Orta": SARI, "Düşük": None}
    ws = wb.active
    ws.title = "Denetim Önceliği"
    _baslik(ws, ["Öncelik", "Mağaza", "Bölge", "Risk Skoru", "Risk Sınıfı", "Baskın Gösterge", "Eşiği Aşan Göstergeler", "Net Satış", *[f"{ad} %" for _, ad in GOSTERGE], *[f"{ad} Puanı" for _, ad in GOSTERGE],
                 "Denetim Notu (siz doldurun)"], (8, 24, 14, 9, 10, 16, 30, 16, 10, 8, 8, 8, 10, 8, 8, 8, 30))
    for m in r["magazalar"]:
        _satir(ws, [m["oncelik"], m["magaza"], m["bolge"], m["skor"], m["sinif"], m["baskin"], ", ".join(m["asan"]), m["satis"], *[m["oran"][a] for a, _ in GOSTERGE], *[m["puan"][a] for a, _ in GOSTERGE], ""],
               para_s=(4, 8, 9, 10, 11, 12, 13, 14, 15, 16), dolgu={5: renk[m["sinif"]]} | {9 + i: KIRMIZI for i, (a, ad) in enumerate(GOSTERGE) if ad in m["asan"]}, sar=(7,))
    ws = wb.create_sheet("Göstergeler")
    _baslik(ws, ["Gösterge", "Eşik % (puan 50)", "Ağırlık", "Zincir Ortalaması %", "Eşiği Aşan Mağaza Sayısı"], (20, 14, 9, 14, 14))
    for a, ad in GOSTERGE:
        _satir(ws, [ad, r["esik"][a], r["agirlik"][a], r["zincir"][a], sum(ad in m["asan"] for m in r["magazalar"])], para_s=(2, 4))
    ws.append([])
    y, o = r["sinir"]
    ws.append([f"Gösterge puanı = oran ÷ eşik × 50 (en çok 100). Risk skoru ağırlıklı ortalamadır. Skor ≥ {y.normalize():f}: Yüksek; ≥ {o.normalize():f}: Orta. Eşik ve ağırlıklar varsayılandır; sektör standardı değildir."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, **ayar) -> dict:
    magazalar, uy = oku(yol)
    r = skorla(magazalar, **ayar)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Mağaza bazında kayıp risk skoru ve denetim önceliği üretir.")
    p.add_argument("--magazalar", type=Path, default=ORNEK / "magaza_gostergeleri.csv", help="Mağaza, Bölge, Net Satış, Sayım Farkı (eksik negatif), Fire, İade, İptal (tutar)")
    for a, ad in GOSTERGE:
        p.add_argument(f"--{a}-esik", default=str(VARSAYILAN_ESIK[a]), help=f"{ad} için kabul edilebilir üst sınır, satışın %%'si (varsayılan {VARSAYILAN_ESIK[a]})")
    p.add_argument("--agirlik", default="40,20,20,20", help="Sayım, fire, iade, iptal ağırlıkları (varsayılan 40,20,20,20)")
    p.add_argument("--yuksek", default="70", help="Yüksek risk skor sınırı (varsayılan 70)")
    p.add_argument("--orta", default="40", help="Orta risk skor sınırı (varsayılan 40)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "kayip_risk_skoru.xlsx")
    a = p.parse_args(argv)
    esik = {g: para(getattr(a, f"{g}_esik")) for g, _ in GOSTERGE}
    ag = [para(x) for x in a.agirlik.replace(";", ",").split(",")]
    y, o = para(a.yuksek), para(a.orta)
    if any(e is None or e <= 0 for e in esik.values()) or len(ag) != 4 or any(x is None or x < 0 for x in ag) or not sum(x for x in ag if x) or y is None or o is None or not 0 < o < y <= 100:
        print("[X] Eşikler pozitif; --agirlik dört sayı (ör. 40,20,20,20, en az biri pozitif); 0 < --orta < --yuksek ≤ 100 olmalı")
        return 2
    if not a.magazalar.exists():
        print(f"[X] Dosya bulunamadı: {a.magazalar}")
        return 1
    try:
        r = calistir(a.magazalar, a.cikti, esik=esik, agirlik=dict(zip((g for g, _ in GOSTERGE), ag)), yuksek=y, orta=o)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['magazalar'])} mağaza · {s['Yüksek']} yüksek, {s['Orta']} orta, {s['Düşük']} düşük risk")
    for m in r["magazalar"]:
        if m["sinif"] == "Yüksek":
            print(f"[!] {m['oncelik']}. {m['magaza']}: skor {m['skor']} · baskın gösterge {m['baskin']} · eşiği aşan: {', '.join(m['asan']) or '—'}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
