import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def m(ad, satis, sayim=None, fire=None, iade=None, iptal=None):
    return {"magaza": ad, "bolge": "", "satis": D(satis), "sayim_fazlasi": D(0)} | {k: D(v) for k, v in (("sayim", sayim), ("fire", fire), ("iade", iade), ("iptal", iptal)) if v is not None}


def tek(*a, **kw):
    return main.skorla([m("X", *a)], **kw)["magazalar"][0]


class SkorTesti(unittest.TestCase):
    def test_puan_ve_skor(self):
        x = tek(1000000, 25000, 8000, 60000, 30000)
        self.assertEqual((x["oran"], x["puan"]), ({"sayim": D("2.50"), "fire": D("0.80"), "iade": D("6.00"), "iptal": D("3.00")},
                                                  {"sayim": D("100.00"), "fire": D("40.00"), "iade": D("60.00"), "iptal": D("75.00")}))
        self.assertEqual((x["skor"], x["sinif"], x["baskin"], x["asan"]), (D("75.00"), "Yüksek", "Sayım eksiği", ["Sayım eksiği", "İade", "İptal"]))
        x = tek(100, 1, 1, 5, 2)                                                     # her gösterge tam eşikte: puan 50, eşik aşılmış sayılmaz
        self.assertEqual((x["skor"], x["sinif"], x["asan"]), (D("50.00"), "Orta", []))
        self.assertEqual([tek(1000, s, 0, 0, 0)["sinif"] for s in (19, 20)], ["Düşük", "Orta"])          # %2 sayım eksiği → puan 100 × 0,4 = 40
        self.assertEqual(tek(600000, 3000, None, 18000, 6000)["skor"], D("26.25"))                      # fire yok: (10 + 6 + 5) ÷ 0,8
        self.assertEqual(tek(100, 0, 0, 0, 0)["baskin"], "")

    def test_ayarlar_zincir_ve_siralama(self):
        self.assertEqual(tek(1000, 10, 10, 50, 20, agirlik={"sayim": D(100), "fire": D(0), "iade": D(0), "iptal": D(0)}, esik={"sayim": D(2)})["skor"], D("25.00"))
        self.assertEqual(tek(100, 1, 1, 5, 2, orta=D(51))["sinif"], "Düşük")
        r = main.skorla([m("A", 800000, 4000, 4000, 24000, 8000), m("B", 1000000, 25000, 8000, 60000, 30000), m("C", 200000, 0, None, 10000, 0)])
        self.assertEqual([(x["oncelik"], x["magaza"], x["skor"]) for x in r["magazalar"]], [(1, "B", D("75.00")), (2, "A", D("26.00")), (3, "C", D("12.50"))])
        self.assertEqual((r["zincir"], r["say"]), ({"sayim": D("1.45"), "fire": D("0.67"), "iade": D("4.70"), "iptal": D("1.90")}, {"Yüksek": 1, "Orta": 0, "Düşük": 2}))
        with self.assertRaises(ValueError):
            main.skorla([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "magaza_gostergeleri.csv", cikti)
            self.assertEqual([(x["magaza"], x["skor"], x["sinif"], x["baskin"]) for x in r["magazalar"]],
                             [("Kadıköy", D("75.00"), "Yüksek", "Sayım eksiği"), ("Beşiktaş", D("48.00"), "Orta", "Sayım eksiği"), ("Bakırköy", D("45.00"), "Orta", "İade"),
                              ("Pendik", D("26.25"), "Düşük", "Sayım eksiği"), ("Ümraniye", D("26.00"), "Düşük", "Sayım eksiği")])
            self.assertEqual((r["zincir"]["sayim"], [u["tur"] for u in r["uyarilar"]]), (D("1.07"), ["Eksik gösterge", "Sayım fazlası"]))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Denetim Önceliği", "Göstergeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "m.csv"
            y.write_text("Mağaza;Net Satış;Sayım Farkı;Fire;İade;İptal\nA;1.000;-10;5;-50;çok\nB;0;1;1;1;1\nC;1.000;;;;\nD;1.000;20;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["magaza"], x.get("sayim"), x.get("iade"), x.get("iptal"), x["sayim_fazlasi"]) for x in l], [("A", D(10), D(50), None, D(0)), ("D", D(0), None, None, D(20))])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Değer okunamadı", "Eksik gösterge", "Eksik gösterge", "Gösterge yok", "Net satış okunamadı", "Sayım fazlası"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--magazalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--sayim-esik", "1,5", "--agirlik", "50,10,20,20", "--yuksek", "60", "--orta", "30"]), 0)
            self.assertEqual(main.main(arg + ["--agirlik", "50,50"]), 2)
            self.assertEqual(main.main(arg + ["--fire-esik", "0"]), 2)
            self.assertEqual(main.main(arg + ["--orta", "80"]), 2)
            self.assertEqual(main.main(["--magazalar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Mağaza;Net Satış\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
