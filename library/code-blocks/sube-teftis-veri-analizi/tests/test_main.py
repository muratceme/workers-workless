import sys
import tempfile
import unittest
from collections import Counter
from datetime import time
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "s.xlsx"
        cls.s = main.calistir(ORNEK / "islemler.csv", ORNEK / "personel.csv", cls.cikti)
        cls.tur = {}
        for x in cls.s["bulgular"]:
            cls.tur.setdefault(x["tur"], []).append(x)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sayilar(self):
        self.assertEqual(len(self.s["islemler"]), 160)
        self.assertEqual(Counter(x["onem"] for x in self.s["bulgular"]), {"Yüksek": 3, "Orta": 8, "Bilgi": 1})

    def test_yuksek(self):
        self.assertEqual(self.tur["Personelin kendi hesabı"][0]["islemler"], ["İ0067"])
        self.assertEqual(self.tur["Limit aşımı — onay yok"][0]["personel"], "P002")
        self.assertEqual(self.tur["Kendi işlemini onaylama"][0]["personel"], "P006")

    def test_orta(self):
        self.assertEqual([x["islemler"] for x in self.tur["Personel yakınının hesabı"]], [["İ0040"], ["İ0113"]])
        self.assertEqual(len(self.tur["Mesai dışı işlem"]), 2)
        self.assertIn("26.09.2026 (hafta sonu)", self.tur["Tatil / hafta sonu işlemi"][0]["aciklama"])
        self.assertIn("P005: 34 işlemden 6 iptal (%17,6); şube geneli %5,0", self.tur["Sık iptal"][0]["aciklama"])
        self.assertEqual(self.tur["İptal sonrası farklı tutarla yeniden işlem"][0]["islemler"], ["İ0091", "İ0092"])
        b = self.tur["Bölünmüş nakit işlem"][0]
        self.assertEqual((b["musteri"], b["islemler"]), ("10050", ["İ0105", "İ0107", "İ0109"]))
        self.assertIn("toplam 133.000 TL", b["aciklama"])
        self.assertEqual(self.tur["Personel-müşteri yoğunlaşması"][0]["musteri"], "10042")

    def test_personel_ozeti(self):
        self.assertEqual([(x["sicil"], x["puan"]) for x in self.s["ozet"]][:2], [("P002", 7), ("P004", 6)])

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Özet", "Bulgular", "Personel Özeti", "İşlemler", "Parametreler"])
        self.assertEqual(wb["Bulgular"]["G1"].value, "Müfettiş Değerlendirmesi")
        self.assertEqual(wb["Bulgular"].max_row, 13)
        isl = {r[0]: r[11] for r in wb["İşlemler"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(isl["İ0091"], "Sık iptal; İptal sonrası farklı tutarla yeniden işlem")

    def test_parametreler(self):
        ay = main.Ayarlar(mesai_bas=time(7, 30), mesai_bit=time(20, 0), cumartesi_acik=True, bolunmus_esik=D(150000))
        s = main.calistir(ORNEK / "islemler.csv", ORNEK / "personel.csv", self.cikti, ay)
        turler = {x["tur"] for x in s["bulgular"]}
        self.assertNotIn("Mesai dışı işlem", turler)
        self.assertNotIn("Tatil / hafta sonu işlemi", turler)
        self.assertNotIn("Bölünmüş nakit işlem", turler)


class KuralTesti(unittest.TestCase):
    def test_tatil_ve_onaylayan_limiti(self):
        with tempfile.TemporaryDirectory() as tmp:
            i = Path(tmp) / "i.csv"
            i.write_text("ŞUBE İŞLEM DÖKÜMÜ\nİşlem No;Tarih;Saat;Personel;Müşteri No;İşlem Türü;Tutar;Durum;Onaylayan\n"
                         "1;29.10.2026;10:00;A1;900;Havale;5.000,00;Tamamlandı;\n"
                         "2;28.10.2026;11:00;A1;901;EFT;300.000,00;Tamamlandı;B1\n"
                         "3;28.10.2026;11:30;X9;902;EFT;1.000,00;Tamamlandı;\n", encoding="utf-8")
            p = Path(tmp) / "p.csv"
            p.write_text("Sicil;Ad Soyad;İşlem Limiti;Müşteri No\nA1;Deneme Bir;50000;800\nB1;Deneme İki;200000;\n", encoding="utf-8")
            s = main.calistir(i, p, Path(tmp) / "o.xlsx")
            t = {x["tur"]: x for x in s["bulgular"]}
            self.assertIn("resmî tatil", t["Tatil / hafta sonu işlemi"]["aciklama"])
            self.assertIn("onaylayan B1 limiti 200.000 TL", t["Onaylayanın limiti yetersiz"]["aciklama"])
            self.assertIn("X9 personel listesinde yok", t["Tanımsız personel"]["aciklama"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--mesai", "sabah", "--cikti", str(Path(tmp) / "x.xlsx")]), 2)


if __name__ == "__main__":
    unittest.main()
