"""
Tedarikçi Performans Değerlendirme — Workers / Workless kod bloğu
Satın Alma › Satın Alma Uzmanı

Teslimat kayıtlarından tedarikçi karnesi hazırlar:
  - Sipariş satırı = Sipariş No + Malzeme; kısmi teslimatlar toplanır. Satır, teslim edilen miktar sipariş miktarının
    (1 − miktar toleransı) kadarına ulaştığı teslimatta tamamlanmış sayılır. Eksik teslim edilmiş satır "Satır Durumu"
    Açık değilse son teslimatla kapanmış sayılır (miktar uyumsuz).
  - Zamanında teslim: tamamlanma tarihi ≤ termin + --tolerans-gun. Termini geçmiş ve tamamlanmamış açık satır da
    geciken sayılır (gecikme rapor tarihine kadar). Termini gelmemiş açık satır değerlendirmeye girmez.
  - Kalite: retsiz teslimat (lot) oranı; ayrıca ret oranı ve PPM (milyonda ret).
  - Miktar uyumu: teslim miktarı sipariş miktarının ±%--miktar-tol içinde.
  - Fiyat uyumu: fatura birim fiyatı ≤ sipariş birim fiyatı × (1 + %--fiyat-tol); fazla faturalanan tutar.
  - Toplam puan = kriter puanlarının ağırlıklı ortalaması (varsayılan teslim 40, kalite 35, fiyat 15, miktar 10;
    verisi olmayan kriter ağırlıktan çıkarılır). Sınıf A/B/C/D (90 / 75 / 60).
  - Değerlendirilen satır --min-satir altındaysa sınıf verilmez ("Yetersiz veri").
  - Çeyreklik trend (termin tarihine göre) ve bir önceki çeyreğe göre --dusus puan düşüş uyarısı.
Rapor: karne, çeyreklik trend, geciken siparişler, kalite retleri, fiyat farkları, sipariş satırları, tedarikçiye
gönderilebilecek karne metinleri, uyarılar. İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 6 tedarikçi, Ocak–Eylül 2026
    python main.py --teslimatlar teslimatlar.xlsx --bugun 09.10.2026 --agirlik teslim=50 kalite=30 fiyat=10 miktar=10
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
YUZ = Decimal(100)
KRITERLER = ["teslim", "kalite", "fiyat", "miktar"]
KRITER_ADI = {"teslim": "Zamanında Teslim", "kalite": "Kalite (retsiz lot)", "fiyat": "Fiyat Uyumu", "miktar": "Miktar Uyumu"}
VARSAYILAN_AGIRLIK = {"teslim": Decimal(40), "kalite": Decimal(35), "fiyat": Decimal(15), "miktar": Decimal(10)}

SUTUNLAR = {"siparis": ("siparis no", "siparis numarasi", "po no"), "tedarikci": ("tedarikci", "satici", "firma"), "malzeme": ("malzeme", "malzeme kodu", "urun"),
            "siparis_tarihi": ("siparis tarihi",), "termin": ("termin tarihi", "termin", "istenen teslim"), "teslim_tarihi": ("teslim tarihi", "giris tarihi", "irsaliye tarihi"),
            "siparis_miktari": ("siparis miktari",), "teslim_miktari": ("teslim miktari", "gelen miktar", "teslim edilen"),
            "ret": ("ret miktari", "red miktari", "iade miktari"), "kabul": ("kabul miktari",),
            "siparis_fiyati": ("siparis birim fiyati", "siparis fiyati"), "fatura_fiyati": ("fatura birim fiyati", "fatura fiyati"),
            "satir_durumu": ("satir durumu", "kapandi", "kapali")}
ACIK_DEGERLER = {"acik", "hayir", "devam", "bekliyor"}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.0f}".replace(",", ".")


def yuzde(x: Decimal | None) -> str:
    return "—" if x is None else f"%{x:.0f}"


def ondalik(x: Decimal | None) -> str:
    return "—" if x is None else f"{x:.1f}".replace(".", ",")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def ceyrek(d: date) -> str:
    return f"{d.year}-Ç{(d.month - 1) // 3 + 1}"


# ----------------------------------------------------------------------------
# Veri
# ----------------------------------------------------------------------------

@dataclass
class Teslimat:
    satir: int
    tarih: date
    miktar: Decimal
    ret: Decimal


@dataclass
class SiparisSatiri:
    siparis: str
    tedarikci: str
    malzeme: str
    siparis_tarihi: date | None
    termin: date
    miktar: Decimal
    siparis_fiyati: Decimal | None
    fatura_fiyati: Decimal | None = None
    acik_isaretli: bool | None = None   # Satır Durumu sütunu: True = açık, False = kapalı, None = belirtilmemiş
    teslimatlar: list[Teslimat] = field(default_factory=list)
    # hesaplanan
    durum: str = ""            # tamam | acik_gecikmis | acik
    tamam_tarihi: date | None = None
    zamaninda: bool | None = None
    gecikme: int = 0
    miktar_uyumlu: bool | None = None
    fiyat_uyumlu: bool | None = None
    fiyat_farki: Decimal = SIFIR
    notlar: list[str] = field(default_factory=list)

    @property
    def teslim_toplam(self) -> Decimal:
        return sum((t.miktar for t in self.teslimatlar), SIFIR)

    @property
    def ret_toplam(self) -> Decimal:
        return sum((t.ret for t in self.teslimatlar), SIFIR)


def oku(yol: Path) -> tuple[list[SiparisSatiri], list[dict]]:
    satirlar: dict[tuple, SiparisSatiri] = {}
    uy = []
    for r in kayitlar(yol, SUTUNLAR, ("tedarikci", "termin", "siparis_miktari")):
        ted = metin(r.get("tedarikci"))
        if not ted:
            continue
        termin, sm = tarih(r.get("termin")), para(r.get("siparis_miktari"))
        if not termin or not sm or sm <= 0:
            uy.append({"onem": "Orta", "tur": "Okunamayan satır", "kim": ted, "aciklama": f"Satır {r['_satir']}: termin veya sipariş miktarı okunamadı; alınmadı"})
            continue
        sno = metin(r.get("siparis")) or f"S{r['_satir']}"
        mal = metin(r.get("malzeme"))
        anahtar = (katla(ted), sno, katla(mal))
        s = satirlar.get(anahtar)
        if s is None:
            s = satirlar[anahtar] = SiparisSatiri(sno, ted, mal, tarih(r.get("siparis_tarihi")), termin, sm, para(r.get("siparis_fiyati")))
        elif s.miktar != sm or s.termin != termin:
            s.notlar.append(f"satır {r['_satir']}: sipariş miktarı / termin ilk satırdan farklı; ilk satır esas alındı")
        if metin(r.get("satir_durumu")):
            s.acik_isaretli = katla(r.get("satir_durumu")) in ACIK_DEGERLER
        ff = para(r.get("fatura_fiyati"))
        if ff is not None:
            s.fatura_fiyati = ff if s.fatura_fiyati is None else max(s.fatura_fiyati, ff)
        td, tm = tarih(r.get("teslim_tarihi")), para(r.get("teslim_miktari"))
        if td and tm and tm > 0:
            ret = para(r.get("ret"))
            if ret is None and para(r.get("kabul")) is not None:
                ret = max(SIFIR, tm - para(r.get("kabul")))
            s.teslimatlar.append(Teslimat(r["_satir"], td, tm, ret or SIFIR))
            if ret and ret > tm:
                uy.append({"onem": "Orta", "tur": "Ret > teslim", "kim": ted, "aciklama": f"{sno} satır {r['_satir']}: ret miktarı teslim miktarından büyük"})
    return list(satirlar.values()), uy


# ----------------------------------------------------------------------------
# Değerlendirme
# ----------------------------------------------------------------------------

def satir_degerlendir(s: SiparisSatiri, bugun: date, tolerans_gun: int, miktar_tol: Decimal, fiyat_tol: Decimal) -> None:
    s.teslimatlar.sort(key=lambda t: (t.tarih, t.satir))
    esik = s.miktar * (1 - miktar_tol)
    kum = SIFIR
    for t in s.teslimatlar:
        kum += t.miktar
        if kum >= esik:
            s.tamam_tarihi = t.tarih
            break
    if not s.tamam_tarihi and s.teslimatlar and not s.acik_isaretli:
        # Kısmi teslim edilmiş ve "Açık" işaretlenmemiş satır son teslimatla kapanmış sayılır (eksik kapama)
        s.tamam_tarihi = s.teslimatlar[-1].tarih
        s.notlar.append("eksik teslimle kapandı" + ("" if s.acik_isaretli is False else " sayıldı (açıksa Satır Durumu = Açık yazın)"))
    if s.tamam_tarihi:
        s.durum = "tamam"
        s.gecikme = max(0, (s.tamam_tarihi - s.termin).days)
        s.zamaninda = (s.tamam_tarihi - s.termin).days <= tolerans_gun
    elif s.termin < bugun:
        s.durum = "acik_gecikmis"
        s.gecikme = (bugun - s.termin).days
        s.zamaninda = s.gecikme <= tolerans_gun
    else:
        s.durum = "acik"
    if s.durum != "acik":
        fark = (s.teslim_toplam - s.miktar) / s.miktar
        s.miktar_uyumlu = abs(fark) <= miktar_tol
        if not s.miktar_uyumlu:
            s.notlar.append(f"teslim {s.teslim_toplam:g} / sipariş {s.miktar:g} ({fark * 100:+.1f}%)")
    if s.siparis_fiyati is not None and s.fatura_fiyati is not None and s.teslimatlar:
        s.fiyat_uyumlu = s.fatura_fiyati <= s.siparis_fiyati * (1 + fiyat_tol)
        if s.fatura_fiyati > s.siparis_fiyati:
            s.fiyat_farki = ((s.fatura_fiyati - s.siparis_fiyati) * s.teslim_toplam).quantize(K2, ROUND_HALF_UP)


def oran(pay: int, payda: int) -> Decimal | None:
    return (Decimal(pay) * YUZ / payda).quantize(Decimal("0.1"), ROUND_HALF_UP) if payda else None


def puanla(satirlar: list[SiparisSatiri], agirlik: dict[str, Decimal]) -> dict:
    deg = [s for s in satirlar if s.durum != "acik"]
    lotlar = [t for s in satirlar for t in s.teslimatlar]
    fiyatli = [s for s in satirlar if s.fiyat_uyumlu is not None]
    teslim_m = sum((t.miktar for t in lotlar), SIFIR)
    ret_m = sum((t.ret for t in lotlar), SIFIR)
    gec = [s for s in deg if not s.zamaninda]
    p = {"teslim": oran(sum(bool(s.zamaninda) for s in deg), len(deg)),
         "kalite": oran(sum(t.ret == 0 for t in lotlar), len(lotlar)),
         "fiyat": oran(sum(bool(s.fiyat_uyumlu) for s in fiyatli), len(fiyatli)),
         "miktar": oran(sum(bool(s.miktar_uyumlu) for s in deg), len(deg))}
    w = {k: agirlik[k] for k in KRITERLER if p[k] is not None and agirlik.get(k, SIFIR) > 0}
    toplam = (sum(p[k] * w[k] for k in w) / sum(w.values())).quantize(Decimal("0.1"), ROUND_HALF_UP) if w else None
    return {"puan": p, "toplam": toplam, "satir": len(deg), "acik": sum(s.durum == "acik" for s in satirlar), "lot": len(lotlar),
            "ret_lot": sum(t.ret > 0 for t in lotlar), "teslim_miktar": teslim_m, "ret_miktar": ret_m,
            "ppm": (ret_m * 1000000 / teslim_m).quantize(Decimal(1), ROUND_HALF_UP) if teslim_m else None,
            "ort_gecikme": (Decimal(sum(s.gecikme for s in gec)) / len(gec)).quantize(Decimal("0.1"), ROUND_HALF_UP) if gec else None,
            "acik_gecikmis": sum(s.durum == "acik_gecikmis" for s in satirlar), "fiyat_farki": sum((s.fiyat_farki for s in satirlar), SIFIR),
            "eksik_kriter": [KRITER_ADI[k] for k in KRITERLER if p[k] is None]}


def sinif(puan: Decimal | None, esikler: tuple[Decimal, Decimal, Decimal], yeterli: bool) -> str:
    if puan is None or not yeterli:
        return "Yetersiz veri"
    a, b, c = esikler
    return "A" if puan >= a else "B" if puan >= b else "C" if puan >= c else "D"


def karne_metni(ted: str, k: dict, donem: str) -> str:
    p = k["puan"]
    parca = [f"{donem} döneminde {k['satir']} sipariş satırınız değerlendirildi. Zamanında teslim oranı {yuzde(p['teslim'])}"
             + (f"; geciken siparişlerde ortalama gecikme {ondalik(k['ort_gecikme'])} gün." if k["ort_gecikme"] is not None else ".")]
    if p["kalite"] is not None:
        parca.append(f"Ret görülen teslimat: {k['ret_lot']} / {k['lot']}; ret oranı {k['ppm']:,} PPM.".replace(",", "."))
    if p["fiyat"] is not None:
        parca.append(f"Fiyat uyumu {yuzde(p['fiyat'])}" + (f"; sipariş fiyatının üzerinde faturalanan tutar {tl(k['fiyat_farki'])} TL." if k["fiyat_farki"] else "."))
    if p["miktar"] is not None:
        parca.append(f"Miktar uyumu {yuzde(p['miktar'])}.")
    if k["acik_gecikmis"]:
        parca.append(f"Termini geçmiş, teslimi tamamlanmamış {k['acik_gecikmis']} sipariş satırı bulunmaktadır.")
    parca.append(f"Toplam puanınız {ondalik(k['toplam'])} / 100" + (f", sınıfınız {k['sinif']}." if k["sinif"] in "ABCD" else "."))
    return f"Sayın {ted} yetkilisi,\n" + " ".join(parca)


def degerlendir(satirlar: list[SiparisSatiri], bugun: date, *, tolerans_gun: int = 0, miktar_tol: Decimal = Decimal("0.05"),
                fiyat_tol: Decimal = Decimal("0.005"), agirlik: dict | None = None, esikler=(Decimal(90), Decimal(75), Decimal(60)), min_satir: int = 3,
                dusus: Decimal = Decimal(10)) -> dict:
    agirlik = agirlik or VARSAYILAN_AGIRLIK
    for s in satirlar:
        satir_degerlendir(s, bugun, tolerans_gun, miktar_tol, fiyat_tol)
    deg = [s for s in satirlar if s.durum != "acik"]
    tarihler = [s.termin for s in deg] or [bugun]
    donem = f"{min(tarihler):%d.%m.%Y} – {max(tarihler):%d.%m.%Y}"
    ted_satir = defaultdict(list)
    for s in satirlar:
        ted_satir[s.tedarikci].append(s)
    karne, trend, uy = {}, {}, []
    ceyrekler = sorted({ceyrek(s.termin) for s in deg})
    for ted, lst in sorted(ted_satir.items(), key=lambda i: katla(i[0])):
        k = puanla(lst, agirlik)
        k["sinif"] = sinif(k["toplam"], esikler, k["satir"] >= min_satir)
        k["metin"] = karne_metni(ted, k, donem)
        karne[ted] = k
        trend[ted] = {c: puanla([s for s in lst if s.durum != "acik" and ceyrek(s.termin) == c], agirlik)["toplam"] for c in ceyrekler}
        dolu = [(c, v) for c, v in trend[ted].items() if v is not None]
        if len(dolu) >= 2 and dolu[-2][1] - dolu[-1][1] >= dusus:
            uy.append({"onem": "Orta", "tur": "Performans düşüşü", "kim": ted,
                       "aciklama": f"{dolu[-2][0]} puanı {ondalik(dolu[-2][1])} → {dolu[-1][0]} puanı {ondalik(dolu[-1][1])} "
                                   f"({ondalik(dolu[-1][1] - dolu[-2][1])})"})
        if k["sinif"] == "D":
            uy.append({"onem": "Yüksek", "tur": "D sınıfı tedarikçi", "kim": ted,
                       "aciklama": f"Toplam puan {ondalik(k['toplam'])}; düzeltici faaliyet planı isteyin ve alternatif tedarikçi değerlendirin"})
        elif k["sinif"] == "Yetersiz veri":
            uy.append({"onem": "Bilgi", "tur": "Yetersiz veri", "kim": ted, "aciklama": f"{k['satir']} değerlendirilen satır (< {min_satir}); sınıf verilmedi"})
        if k["acik_gecikmis"]:
            uy.append({"onem": "Yüksek", "tur": "Termini geçmiş açık sipariş", "kim": ted,
                       "aciklama": f"{k['acik_gecikmis']} satır: " + ", ".join(f"{s.siparis} {s.malzeme} ({s.gecikme} gün)" for s in lst if s.durum == "acik_gecikmis")})
        if k["puan"]["kalite"] is not None and k["puan"]["kalite"] < 90:
            uy.append({"onem": "Orta", "tur": "Kalite retleri", "kim": ted,
                       "aciklama": f"Ret görülen teslimat {k['ret_lot']} / {k['lot']} (retsiz lot %{ondalik(k['puan']['kalite'])}, {k['ppm']} PPM); "
                                   "8D / DÖF isteyin"})
        if k["fiyat_farki"] > 0:
            uy.append({"onem": "Orta", "tur": "Fazla faturalama", "kim": ted,
                       "aciklama": f"Sipariş fiyatı üzerinde faturalanan toplam {tl(k['fiyat_farki'])} TL; fiyat farkı faturası / iade talep edin"})
        if k["eksik_kriter"] and k["satir"]:
            uy.append({"onem": "Bilgi", "tur": "Eksik kriter", "kim": ted, "aciklama": f"Verisi olmayan kriter ağırlıktan çıkarıldı: {', '.join(k['eksik_kriter'])}"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], katla(u["kim"])))
    return {"satirlar": satirlar, "karne": karne, "trend": trend, "ceyrekler": ceyrekler, "uyarilar": uy, "donem": donem, "bugun": bugun,
            "agirlik": agirlik, "esikler": esikler, "parametre": {"tolerans_gun": tolerans_gun, "miktar_tol": miktar_tol, "fiyat_tol": fiyat_tol, "min_satir": min_satir}}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
SINIF_RENK = {"A": "E3F4E1", "B": "E8F0FE", "C": "FFF4CE", "D": "FDE2E1"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


def _f(x):
    return None if x is None else float(x)


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    ka = wb.active
    ka.title = "Karne"
    _baslik(ka, ["Tedarikçi", "Değerlendirilen Satır", "Termini Gelmemiş", "Zamanında Teslim %", "Ort. Gecikme (gün)", "Kalite (retsiz lot) %", "PPM",
                 "Fiyat Uyumu %", "Fazla Faturalanan (TL)", "Miktar Uyumu %", "Toplam Puan", "Sınıf", "Aksiyon"],
            (26, 11, 10, 11, 10, 11, 9, 10, 13, 10, 10, 12, 30))
    for ted, k in sorted(s["karne"].items(), key=lambda i: -(i[1]["toplam"] or 0)):
        p = k["puan"]
        ka.append([ted, k["satir"], k["acik"], _f(p["teslim"]), _f(k["ort_gecikme"]), _f(p["kalite"]), _f(k["ppm"]), _f(p["fiyat"]), float(k["fiyat_farki"]),
                   _f(p["miktar"]), _f(k["toplam"]), k["sinif"], ""])
        ka.cell(ka.max_row, 9).number_format = PF
        ka.cell(ka.max_row, 12).fill = PatternFill("solid", fgColor=SINIF_RENK.get(k["sinif"], "FFFFFF"))
        ka.cell(ka.max_row, 13).fill = PatternFill("solid", fgColor="FFF4CE")
    ka.append([])
    ka.append([f"Dönem (termin): {s['donem']} · rapor tarihi {s['bugun']:%d.%m.%Y}"])
    ka.append(["Ağırlıklar: " + ", ".join(f"{KRITER_ADI[k]} {s['agirlik'].get(k, 0):g}" for k in KRITERLER)
               + f" · Sınıf: A ≥ {s['esikler'][0]:g}, B ≥ {s['esikler'][1]:g}, C ≥ {s['esikler'][2]:g}, D altı"])
    pr = s["parametre"]
    ka.append([f"Zamanında: termin + {pr['tolerans_gun']} gün · miktar toleransı ±%{pr['miktar_tol'] * 100:g} · fiyat toleransı %{pr['fiyat_tol'] * 100:g} · "
               f"sınıf için en az {pr['min_satir']} satır"])

    tr = wb.create_sheet("Çeyreklik Trend")
    _baslik(tr, ["Tedarikçi"] + s["ceyrekler"], [26] + [11] * len(s["ceyrekler"]))
    for ted, d in sorted(s["trend"].items(), key=lambda i: katla(i[0])):
        tr.append([ted] + [_f(d[c]) for c in s["ceyrekler"]])
    if s["ceyrekler"]:
        g = LineChart()
        g.title, g.height, g.width = "Çeyreklik toplam puan", 8, 18
        g.add_data(Reference(tr, min_col=1, max_col=len(s["ceyrekler"]) + 1, min_row=2, max_row=tr.max_row), titles_from_data=True, from_rows=True)
        g.set_categories(Reference(tr, min_col=2, max_col=len(s["ceyrekler"]) + 1, min_row=1))
        tr.add_chart(g, f"A{tr.max_row + 3}")

    gs = wb.create_sheet("Geciken Siparişler")
    _baslik(gs, ["Tedarikçi", "Sipariş No", "Malzeme", "Termin", "Tamamlanma", "Gecikme (gün)", "Durum", "Sipariş Miktarı", "Teslim Edilen"],
            (26, 12, 28, 11, 11, 10, 22, 11, 11))
    for x in sorted((x for x in s["satirlar"] if x.zamaninda is False), key=lambda x: (-x.gecikme, katla(x.tedarikci))):
        gs.append([x.tedarikci, x.siparis, x.malzeme, x.termin, x.tamam_tarihi, x.gecikme,
                   "Açık — termini geçmiş" if x.durum == "acik_gecikmis" else "Geç tamamlandı", float(x.miktar), float(x.teslim_toplam)])
        gs.cell(gs.max_row, 4).number_format = gs.cell(gs.max_row, 5).number_format = "DD.MM.YYYY"
        if x.durum == "acik_gecikmis":
            gs.cell(gs.max_row, 7).fill = PatternFill("solid", fgColor="FDE2E1")

    kr = wb.create_sheet("Kalite Retleri")
    _baslik(kr, ["Tedarikçi", "Sipariş No", "Malzeme", "Teslim Tarihi", "Teslim Miktarı", "Ret Miktarı", "Ret Oranı", "DÖF / 8D No"], (26, 12, 28, 11, 11, 10, 9, 14))
    for x in s["satirlar"]:
        for t in x.teslimatlar:
            if t.ret > 0:
                kr.append([x.tedarikci, x.siparis, x.malzeme, t.tarih, float(t.miktar), float(t.ret), float(t.ret / t.miktar), ""])
                kr.cell(kr.max_row, 4).number_format = "DD.MM.YYYY"
                kr.cell(kr.max_row, 7).number_format = "0.0%"
                kr.cell(kr.max_row, 8).fill = PatternFill("solid", fgColor="FFF4CE")

    ff = wb.create_sheet("Fiyat Farkları")
    _baslik(ff, ["Tedarikçi", "Sipariş No", "Malzeme", "Sipariş Fiyatı", "Fatura Fiyatı", "Fark %", "Teslim Miktarı", "Fazla Faturalanan (TL)", "Tolerans İçinde"],
            (26, 12, 28, 11, 11, 8, 11, 13, 10))
    for x in s["satirlar"]:
        if x.fiyat_farki > 0:
            ff.append([x.tedarikci, x.siparis, x.malzeme, float(x.siparis_fiyati), float(x.fatura_fiyati), float(x.fatura_fiyati / x.siparis_fiyati - 1),
                       float(x.teslim_toplam), float(x.fiyat_farki), "Evet" if x.fiyat_uyumlu else "Hayır"])
            ff.cell(ff.max_row, 6).number_format = "0.0%"
            for j in (4, 5, 8):
                ff.cell(ff.max_row, j).number_format = PF

    ss = wb.create_sheet("Sipariş Satırları")
    _baslik(ss, ["Tedarikçi", "Sipariş No", "Malzeme", "Sipariş Tarihi", "Termin", "Sipariş Miktarı", "Teslim Edilen", "Ret", "Teslimat Sayısı", "Tamamlanma",
                 "Durum", "Zamanında", "Gecikme", "Miktar Uyumlu", "Fiyat Uyumlu", "Not"], (24, 11, 26, 11, 11, 10, 10, 8, 8, 11, 14, 9, 8, 9, 9, 36))
    for x in sorted(s["satirlar"], key=lambda x: (katla(x.tedarikci), x.termin, x.siparis)):
        e = {True: "Evet", False: "Hayır", None: "—"}
        ss.append([x.tedarikci, x.siparis, x.malzeme, x.siparis_tarihi, x.termin, float(x.miktar), float(x.teslim_toplam), float(x.ret_toplam), len(x.teslimatlar),
                   x.tamam_tarihi, {"tamam": "Tamamlandı", "acik_gecikmis": "Açık, gecikmiş", "acik": "Açık, termini gelmedi"}[x.durum], e[x.zamaninda], x.gecikme,
                   e[x.miktar_uyumlu], e[x.fiyat_uyumlu], "; ".join(x.notlar)])
        for j in (4, 5, 10):
            ss.cell(ss.max_row, j).number_format = "DD.MM.YYYY"
    ss.auto_filter.ref = f"A1:P{ss.max_row}"

    km = wb.create_sheet("Karne Metinleri")
    _baslik(km, ["Tedarikçi", "Sınıf", "Metin"], (26, 12, 110))
    for ted, k in sorted(s["karne"].items(), key=lambda i: katla(i[0])):
        km.append([ted, k["sinif"], k["metin"]])
        km.cell(km.max_row, 3).alignment = UST

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Tedarikçi", "Açıklama"], (9, 26, 26, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(yol: Path, cikti: Path, bugun: date, **kw) -> dict:
    satirlar, uy = oku(yol)
    if not satirlar:
        raise ValueError(f"{yol.name}: sipariş satırı bulunamadı")
    s = degerlendir(satirlar, bugun, **kw)
    s["uyarilar"] = uy + s["uyarilar"]
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Teslim zamanı, kalite reddi, miktar ve fiyat uyumundan tedarikçi karnesi hazırlar.")
    p.add_argument("--teslimatlar", type=Path, default=ORNEK / "teslimatlar.csv",
                   help="Sipariş No, Tedarikçi, Malzeme, Sipariş Tarihi, Termin Tarihi, Teslim Tarihi, Sipariş Miktarı, Teslim Miktarı, Ret Miktarı, "
                        "Sipariş Birim Fiyatı, Fatura Birim Fiyatı, Satır Durumu (Açık/Kapalı) — kısmi teslimatlar ayrı satır")
    p.add_argument("--bugun", help="Rapor tarihi GG.AA.YYYY (açık siparişlerin gecikmesi; örnek veride 09.10.2026)")
    p.add_argument("--tolerans-gun", type=int, default=0, help="Termin + bu kadar gün zamanında sayılır (varsayılan 0)")
    p.add_argument("--miktar-tol", type=float, default=5, help="Miktar uyumu toleransı %% (varsayılan 5)")
    p.add_argument("--fiyat-tol", type=float, default=0.5, help="Fiyat uyumu toleransı %% (varsayılan 0,5)")
    p.add_argument("--agirlik", nargs="*", default=[], metavar="KRİTER=AĞIRLIK", help="teslim, kalite, fiyat, miktar (varsayılan 40 / 35 / 15 / 10)")
    p.add_argument("--siniflar", default="90,75,60", help="A, B, C alt sınırları (varsayılan 90,75,60)")
    p.add_argument("--min-satir", type=int, default=3, help="Sınıf verilmesi için en az değerlendirilen satır (varsayılan 3)")
    p.add_argument("--dusus", type=float, default=10, help="Önceki çeyreğe göre bu kadar puan düşüşte uyarı (varsayılan 10)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tedarikci_karnesi.xlsx")
    a = p.parse_args(argv)
    if not a.teslimatlar.exists():
        print(f"[X] Dosya bulunamadı: {a.teslimatlar}")
        return 1
    ornek = a.teslimatlar == ORNEK / "teslimatlar.csv"
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    agirlik = dict(VARSAYILAN_AGIRLIK)
    try:
        for x in a.agirlik:
            k, _, v = x.partition("=")
            if katla(k) not in KRITERLER or para(v) is None:
                raise ValueError(x)
            agirlik[katla(k)] = para(v)
        esikler = tuple(Decimal(x.strip()) for x in a.siniflar.split(","))
        if len(esikler) != 3:
            raise ValueError(a.siniflar)
    except (ValueError, InvalidOperation) as h:
        print(f"[X] Geçersiz parametre: {h} (ör. --agirlik teslim=50 kalite=30, --siniflar 90,75,60)")
        return 2
    try:
        s = calistir(a.teslimatlar, a.cikti, bugun, tolerans_gun=a.tolerans_gun, miktar_tol=Decimal(str(a.miktar_tol)) / 100,
                     fiyat_tol=Decimal(str(a.fiyat_tol)) / 100, agirlik=agirlik, esikler=esikler, min_satir=a.min_satir, dusus=Decimal(str(a.dusus)))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['karne'])} tedarikçi · {sum(k['satir'] for k in s['karne'].values())} değerlendirilen sipariş satırı · dönem {s['donem']}")
    for ted, k in sorted(s["karne"].items(), key=lambda i: -(i[1]["toplam"] or 0)):
        p_ = k["puan"]
        print(f"     {k['sinif']:<13} {k['toplam'] if k['toplam'] is not None else '—':>5}  {ted}  (teslim {yuzde(p_['teslim'])}, kalite {yuzde(p_['kalite'])}, "
              f"fiyat {yuzde(p_['fiyat'])}, miktar {yuzde(p_['miktar'])})")
    for u in s["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
