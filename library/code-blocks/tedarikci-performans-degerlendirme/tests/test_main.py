import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "t.xlsx"
        cls.s = main.calistir(ORNEK / "teslimatlar.csv", cls.cikti, date(2026, 10, 9))
        cls.k = cls.s["karne"]
        cls.u = {(u["tur"], u["kim"]) for u in cls.s["uyarilar"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_elle_hesaplanan_tedarikci(self):
        # Epsilon: 12 satır (10 kapalı + 2 açık gecikmiş); zamanında 2 (PO-4554 eksik kapandı, PO-4555); miktar uyumlu 5;
        # 14 lot, 2'sinde ret; fiyatlı 11 satır, 6'sı uyumlu; fazla faturalama 560 + 1024 + 237,44 + 280 + 211,2
        e = self.k["Epsilon Elektrik"]
        self.assertEqual((e["satir"], e["lot"], e["ret_lot"]), (12, 14, 2))
        self.assertEqual((e["puan"]["teslim"], e["puan"]["miktar"], e["puan"]["kalite"], e["puan"]["fiyat"]), (D("16.7"), D("41.7"), D("85.7"), D("54.5")))
        self.assertEqual(e["fiyat_farki"], D("2312.64"))
        self.assertEqual((e["toplam"], e["sinif"]), (D("49.0"), "D"))

    def test_satir_mantigi(self):
        sat = {x.siparis: x for x in self.s["satirlar"]}
        self.assertEqual((sat["PO-4549"].tamam_tarihi, sat["PO-4549"].zamaninda), (date(2026, 1, 24), False))     # kısmi teslimle tamamlandı
        self.assertEqual((sat["PO-4554"].durum, sat["PO-4554"].zamaninda, sat["PO-4554"].miktar_uyumlu), ("tamam", True, False))   # eksik kapama
        self.assertEqual((sat["PO-4602"].durum, sat["PO-4602"].gecikme), ("acik_gecikmis", 9))              # Açık işaretli, kısmi
        self.assertEqual(sat["PO-4603"].durum, "acik")                                                          # termini gelmedi
        self.assertEqual(self.k["Örnek Metal A.Ş."]["acik"], 1)

    def test_siniflar_ve_uyarilar(self):
        self.assertEqual(self.k["Zeta Hırdavat"]["sinif"], "Yetersiz veri")
        self.assertEqual(self.k["Örnek Metal A.Ş."]["sinif"], "A")
        self.assertTrue({("D sınıfı tedarikçi", "Epsilon Elektrik"), ("Termini geçmiş açık sipariş", "Epsilon Elektrik"), ("Kalite retleri", "Gama Ambalaj San."),
                         ("Fazla faturalama", "Delta Kimya Tic."), ("Performans düşüşü", "Beta Plastik Ltd."), ("Yetersiz veri", "Zeta Hırdavat")} <= self.u)
        self.assertIn("sınıfınız D", self.k["Epsilon Elektrik"]["metin"])

    def test_excel(self):
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Karne", "Çeyreklik Trend", "Geciken Siparişler", "Kalite Retleri", "Fiyat Farkları", "Sipariş Satırları",
                                         "Karne Metinleri", "Uyarılar"])
        self.assertEqual(wb["Karne"]["A2"].value, "Zeta Hırdavat")      # puana göre sıralı (100)
        self.assertEqual(wb["Çeyreklik Trend"]["B1"].value, "2026-Ç1")


class KuralTesti(unittest.TestCase):
    def satir(self, teslimatlar, termin=date(2026, 3, 10), miktar=100, acik=None):
        s = main.SiparisSatiri("P1", "T", "M", None, termin, D(miktar), D(10), D("10.04"), acik_isaretli=acik)
        s.teslimatlar = [main.Teslimat(i, t, D(m), D(r)) for i, (t, m, r) in enumerate(teslimatlar)]
        main.satir_degerlendir(s, date(2026, 4, 1), 0, D("0.05"), D("0.005"))
        return s

    def test_tolerans_sinirlari(self):
        s = self.satir([(date(2026, 3, 10), 95, 0)])            # tam %95 → tamamlandı, termin günü → zamanında
        self.assertEqual((s.durum, s.zamaninda, s.miktar_uyumlu), ("tamam", True, True))
        self.assertTrue(s.fiyat_uyumlu)                         # %0,4 fark tolerans içinde
        self.assertEqual(s.fiyat_farki, D("3.80"))
        s = self.satir([(date(2026, 3, 11), 100, 0)])
        self.assertFalse(s.zamaninda)
        s = self.satir([], termin=date(2026, 3, 1))
        self.assertEqual((s.durum, s.gecikme), ("acik_gecikmis", 31))

    def test_agirlik_yeniden_dagitimi(self):
        s = self.satir([(date(2026, 3, 10), 100, 0)])
        s.fiyat_uyumlu = None
        k = main.puanla([s], {"teslim": D(40), "kalite": D(35), "fiyat": D(15), "miktar": D(10)})
        self.assertEqual((k["toplam"], k["eksik_kriter"]), (D("100.0"), ["Fiyat Uyumu"]))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx"), "--agirlik", "teslim=50"]), 0)
            self.assertEqual(main.main(["--agirlik", "hiz=5"]), 2)
            self.assertEqual(main.main(["--bugun", "x"]), 2)
            self.assertEqual(main.main(["--teslimatlar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
