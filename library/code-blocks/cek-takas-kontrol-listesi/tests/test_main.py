import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
V, G = "Takasa verilecek", "Takastan gelen"


def c(no, once=1, rakam=100, yazi="ayni", karekod=True, ciro=True, imza=True, farkli=False, yon=V, banka="A", keside="x"):
    """once: keşide tarihi kaç gün önce."""
    return {"no": no, "yon": yon, "banka": banka, "keside": BUGUN - timedelta(days=once) if keside == "x" else keside, "rakam": D(str(rakam)), "yazi": D(str(rakam)) if yazi == "ayni" else (None if yazi is None else D(str(yazi))),
            "karekod": karekod, "ciro": ciro, "imza": imza, "farkli": farkli}


def s(*cekler, **kw):
    return {x["no"]: (x["sonuc"], x["bulgu"]) for x in main.kontrol_et(list(cekler), BUGUN, **kw)["cekler"]}


class KontrolTesti(unittest.TestCase):
    def test_tarih_kontrolleri(self):
        r = s(c("bugün", 0), c("yarın", -1), c("on", 10), c("onbir", 11), c("tarihsiz", keside=None), c("ay", keside=date(2026, 9, 12), farkli=True), c("ay+1", keside=date(2026, 9, 11), farkli=True),
              c("ay_sonu", keside=date(2026, 8, 31), farkli=True))
        self.assertEqual(r, {"bugün": ("Uygun", ""), "yarın": ("İşleme alma", "İleri tarihli çek"), "on": ("Uygun", ""), "onbir": ("İncele", "İbraz süresi geçmiş"), "tarihsiz": ("İşleme alma", "Keşide tarihi yok"), "ay": ("Uygun", ""),
                             "ay+1": ("İncele", "İbraz süresi geçmiş"), "ay_sonu": ("İncele", "İbraz süresi geçmiş")})
        x = {v["no"]: v["ibraz_son"] for v in main.kontrol_et([c("a", 11), c("b", keside=date(2026, 8, 31), farkli=True), c("c", keside=None)], BUGUN)["cekler"]}
        self.assertEqual(x, {"a": date(2026, 10, 11), "b": date(2026, 9, 30), "c": None})
        self.assertEqual((s(c("a", 11), ibraz_gun=11)["a"][0], s(c("a", keside=date(2026, 8, 11), farkli=True), ibraz_ay=2)["a"][0], s(c("a", keside=date(2026, 8, 13), farkli=True), ibraz_ay=2)["a"][0]), ("Uygun", "İncele", "Uygun"))
        self.assertEqual(main.ay_ekle(date(2027, 1, 31), 1), date(2027, 2, 28))

    def test_diger_kontroller(self):
        r = s(c("tutar", yazi=101), c("yazısız", yazi=None), c("karekod", karekod=False), c("ciro", ciro=False), c("imza", imza=False), c("boş", karekod=None, ciro=None, imza=None), c("hepsi", 11, yazi=5, karekod=False, imza=None),
              c("M1"), c("m 1"), c("M1", banka="B"), c("M2", yon=G), c("M2", yon=G, banka="a"))
        self.assertEqual(r, {"tutar": ("İncele", "Tutar tutarsız"), "yazısız": ("İncele", "Kontrol işaretlenmemiş"), "karekod": ("İşleme alma", "Karekod yok / okunamadı"), "ciro": ("İşleme alma", "Ciro silsilesi kopuk"),
                             "imza": ("İşleme alma", "İmza yok"), "boş": ("İncele", "Kontrol işaretlenmemiş"), "hepsi": ("İşleme alma", "İbraz süresi geçmiş; Tutar tutarsız; Karekod yok / okunamadı; Kontrol işaretlenmemiş"),
                             "M1": ("Uygun", ""), "m 1": ("İşleme alma", "Mükerrer çek"), "M2": ("İşleme alma", "Mükerrer çek")})
        r = main.kontrol_et([c("1", rakam=100), c("2", rakam=50, imza=False), c("3", rakam=30, yazi=31), c("4", rakam=20, yon=G, karekod=None)], BUGUN)
        self.assertEqual(r["ozet"], [{"yon": V, "adet": 3, "tutar": D(180), "İşleme alma": 1, "İncele": 1, "Uygun": 1, "sorunlu_tutar": D(50)}, {"yon": G, "adet": 1, "tutar": D(20), "İşleme alma": 0, "İncele": 1, "Uygun": 0, "sorunlu_tutar": D(0)}])
        self.assertEqual((r["say"], [(b["no"], b["duzey"], b["tur"]) for b in r["bulgular"]], [(u["onem"], u["kim"]) for u in r["uyarilar"]]),
                         ({"İşleme alma": 1, "İncele": 2, "Uygun": 1}, [("2", "Engel", "İmza yok"), ("3", "Dikkat", "Tutar tutarsız"), ("4", "Eksik", "Kontrol işaretlenmemiş")], [("Yüksek", "2 · A"), ("Orta", "3 · A"), ("Bilgi", "4 · A")]))
        with self.assertRaises(ValueError):
            main.kontrol_et([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "cekler.csv", cikti, BUGUN)
            self.assertEqual([(x["no"], x["sonuc"], x["bulgu"]) for x in r["cekler"]],
                             [("1000101", "Uygun", ""), ("1000102", "İşleme alma", "İleri tarihli çek"), ("1000103", "İncele", "İbraz süresi geçmiş"), ("1000104", "Uygun", ""), ("1000105", "İncele", "Tutar tutarsız"),
                              ("1000106", "İşleme alma", "Karekod yok / okunamadı"), ("2000201", "İşleme alma", "Ciro silsilesi kopuk"), ("2000202", "İşleme alma", "İmza yok"), ("2000203", "Uygun", ""), ("2000204", "İşleme alma", "Mükerrer çek"),
                              ("2000204", "İşleme alma", "Mükerrer çek"), ("2000205", "İncele", "Kontrol işaretlenmemiş")])
            self.assertEqual(r["ozet"], [{"yon": V, "adet": 6, "tutar": D(552500), "İşleme alma": 2, "İncele": 2, "Uygun": 2, "sorunlu_tutar": D(290000)},
                                         {"yon": G, "adet": 6, "tutar": D(268000), "İşleme alma": 4, "İncele": 1, "Uygun": 1, "sorunlu_tutar": D(132000)}])
            self.assertEqual((r["say"], r["cekler"][2]["ibraz_son"], r["cekler"][3]["ibraz_son"], len(r["bulgular"])), ({"İşleme alma": 6, "İncele": 3, "Uygun": 3}, date(2026, 10, 8), date(2026, 10, 28), 11))
            self.assertEqual([(u["onem"], u["tur"]) for u in r["uyarilar"]][:6] + [(u["onem"], u["tur"]) for u in r["uyarilar"]][-5:],
                             [("Yüksek", "Ciro silsilesi kopuk"), ("Yüksek", "Karekod yok / okunamadı"), ("Yüksek", "Mükerrer çek"), ("Yüksek", "Mükerrer çek"), ("Yüksek", "İleri tarihli çek"), ("Yüksek", "İmza yok"),
                              ("Orta", "Tutar tutarsız"), ("Orta", "İbraz süresi geçmiş"), ("Bilgi", "Kontrol işaretlenmemiş"), ("Bilgi", "Kontrol işaretlenmemiş"), ("Bilgi", "Kontrol işaretlenmemiş")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kontrol Listesi", "Özet", "Bulgular", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "c.csv"
            y.write_text("No;İşlem;Muhatap Banka;Çek Tarihi;Tutar;Yazıyla Tutar;Karekod;Ciro;İmza;Başka Yerde Ödenecek\n77;GELEN;X;05.10.2026;1.250,50;1.250,50;OKUNDU;düzgün;evet;EVET\n78;giden;;x;5;;okunmadı;eksik;hayır;\n79;iade;;;5;;;;;\n"
                         "80;gelen;;;0;;;;;\n81;gelen;;;5;x;;;;\n;gelen;;;5;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual((l, [(u["tur"], u["kim"]) for u in uy]),
                             ([{"no": "77", "yon": G, "banka": "X", "keside": date(2026, 10, 5), "rakam": D("1250.50"), "yazi": D("1250.50"), "karekod": True, "ciro": True, "imza": True, "farkli": True},
                               {"no": "78", "yon": V, "banka": "", "keside": None, "rakam": D(5), "yazi": None, "karekod": False, "ciro": False, "imza": False, "farkli": False}], [("Çek satırı okunamadı", n) for n in ("79", "80", "81")]))
            c_ = str(Path(tmp) / "x.xlsx")
            arg = ["--cekler", str(y), "--cikti", c_]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--ibraz-gun", "8", "--ibraz-ay", "2"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--ibraz-gun", "0"], ["--ibraz-ay", "13"])], [2] * 3)
            self.assertEqual(main.main(["--cekler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Çek No;Tutar\nA;1\n", "Çek No;Keşide Tarihi;Tutar\n;05.10.2026;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c_]), 0)


if __name__ == "__main__":
    unittest.main()
