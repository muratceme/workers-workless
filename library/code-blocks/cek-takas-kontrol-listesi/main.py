"""
Çek Takas Kontrol Listesi — Gerisi Bende kod bloğu
Şube Bankacılığı › Operasyon Yetkilisi

Takasa verilecek ve takastan gelen çeklerde keşide tarihi, ibraz süresi, karekod, ciro silsilesi, imza ve tutar tutarlılığını kontrol eder.
  İleri tarihli        = keşide tarihi işlem tarihinden sonra                                    → İşleme alma
  İbraz süresi son gün = keşide tarihi + 10 gün (aynı yerde ödenecek) / + 1 ay (başka yerde)   → geçmişse İncele
  Tutar tutarlılığı    = rakamla yazılan tutar ile yazıyla yazılan tutar aynı olmalı            → farklıysa İncele
  Karekod              = "Yok" ya da "Okunamadı"                                                → İşleme alma
  Ciro silsilesi       = "Kopuk"                                                                → İşleme alma
  İmza                 = "Yok"                                                                  → İşleme alma
  Mükerrer             = aynı banka ve çek numarası listede birden çok kez                      → İşleme alma
  İşaretlenmemiş kontrol (karekod, ciro, imza boş)                                              → İncele
Kod çeki görmez; operasyon yetkilisinin fiziki kontrolde işaretlediği sütunları değerlendirir. Süreler komut satırından değişir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
import calendar

SUTUN = {"no": ("cek no", "cek numarasi", "seri no", "no"), "yon": ("yon", "islem", "tur"), "banka": ("banka", "muhatap banka", "muhatap"), "keside": ("keside tarihi", "duzenleme tarihi", "cek tarihi"),
         "rakam": ("tutar rakam", "tutar", "rakamla tutar"), "yazi": ("tutar yazi", "yaziyla tutar", "yazi ile tutar"), "karekod": ("karekod",), "ciro": ("ciro silsilesi", "ciro"), "imza": ("imza", "kesideci imzasi"),
         "farkli": ("farkli yer", "baska yerde odenecek", "farkli yerde odeme")}
YON = {"takasa verilecek": "Takasa verilecek", "verilecek": "Takasa verilecek", "giden": "Takasa verilecek", "takastan gelen": "Takastan gelen", "gelen": "Takastan gelen"}
OLUMLU, OLUMSUZ = {"var", "evet", "tam", "duzgun", "okundu", "uygun"}, {"yok", "hayir", "kopuk", "eksik", "okunamadi", "okunmadi", "uygun degil"}
KONTROL = (("karekod", "Karekod", "Karekod yok / okunamadı", "Karekod yok ya da okunamadı."), ("ciro", "Ciro silsilesi", "Ciro silsilesi kopuk", "Ciro silsilesi kopuk ya da eksik işaretlenmiş."),
           ("imza", "İmza", "İmza yok", "Keşideci imzası yok işaretlenmiş."))
SONUC_RENK = {"İşleme alma": KIRMIZI, "İncele": SARI, "Uygun": YESIL}
DUZEY = {"Engel": "Yüksek", "Dikkat": "Orta", "Eksik": "Bilgi"}
Z = Decimal(0)


def ay_ekle(t: date, n: int) -> date:
    y, a = divmod(t.year * 12 + t.month - 1 + n, 12)
    return date(y, a + 1, min(t.day, calendar.monthrange(y, a + 1)[1]))


def _isaret(x) -> bool | None:
    k = katla(x)
    return True if k in OLUMLU else False if k in OLUMSUZ else None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "keside", "rakam")):
        no = metin(r["no"])
        if not no:
            continue
        yon, rakam, yazi = YON.get(katla(r["yon"])), para(r["rakam"]), para(r["yazi"])
        if yon is None or rakam is None or rakam <= 0 or (metin(r["yazi"]) and yazi is None):
            uy.append(uyari("Yüksek", "Çek satırı okunamadı", no, f"Satır {r['_satir']}: yön '{metin(r['yon'])}' (Takasa Verilecek / Takastan Gelen olmalı), tutar '{metin(r['rakam'])}', yazıyla tutar '{metin(r['yazi'])}'. Satır alınmadı."))
            continue
        l.append({"no": no, "yon": yon, "banka": metin(r["banka"]), "keside": tarih(r["keside"]), "rakam": rakam, "yazi": yazi, "karekod": _isaret(r["karekod"]), "ciro": _isaret(r["ciro"]), "imza": _isaret(r["imza"]),
                  "farkli": _isaret(r["farkli"]) is True})
    return l, uy


def kontrol_et(cekler: list[dict], bugun: date, ibraz_gun: int = 10, ibraz_ay: int = 1) -> dict:
    if not cekler:
        raise ValueError("Kontrol edilecek çek yok.")
    sayim = defaultdict(int)
    for c in cekler:
        sayim[(katla(c["banka"]), katla(c["no"]).replace(" ", ""))] += 1
    uy, bulgular = [], []
    for c in cekler:
        kim, b = f"{c['no']} · {c['banka']}".strip(" ·"), []
        c["ibraz_son"] = (ay_ekle(c["keside"], ibraz_ay) if c["farkli"] else c["keside"] + timedelta(days=ibraz_gun)) if c["keside"] else None
        if c["keside"] is None:
            b.append(("Engel", "Keşide tarihi yok", "Keşide tarihi yazılmamış ya da okunamadı."))
        elif c["keside"] > bugun:
            b.append(("Engel", "İleri tarihli çek", f"Keşide tarihi {c['keside']:%d.%m.%Y}, işlem tarihi {bugun:%d.%m.%Y}. Keşide tarihinden önce ibraz konusundaki yürürlükteki düzenlemeyi kontrol edin."))
        elif bugun > c["ibraz_son"]:
            b.append(("Dikkat", "İbraz süresi geçmiş", f"Keşide tarihi {c['keside']:%d.%m.%Y}; ibraz için son gün {c['ibraz_son']:%d.%m.%Y} ({(bugun - c['ibraz_son']).days} gün geçti)."))
        if c["yazi"] is None:
            b.append(("Eksik", "Kontrol işaretlenmemiş", "Yazıyla yazılan tutar girilmemiş; tutar tutarlılığı kontrol edilemedi."))
        elif c["yazi"] != c["rakam"]:
            b.append(("Dikkat", "Tutar tutarsız", f"Rakamla {c['rakam']:,.2f}, yazıyla {c['yazi']:,.2f}. Hangi tutarın esas alınacağını mevzuata ve banka uygulamanıza göre belirleyin."))
        for alan, ad, tur, aciklama in KONTROL:
            if c[alan] is False:
                b.append(("Engel", tur, aciklama))
            elif c[alan] is None:
                b.append(("Eksik", "Kontrol işaretlenmemiş", f"{ad} sütunu boş; kontrol yapılıp yapılmadığı bilinmiyor."))
        if sayim[(katla(c["banka"]), katla(c["no"]).replace(" ", ""))] > 1:
            b.append(("Engel", "Mükerrer çek", "Aynı banka ve çek numarası listede birden çok kez geçiyor."))
        duzeyler = {x[0] for x in b}
        c["sonuc"] = "İşleme alma" if "Engel" in duzeyler else "İncele" if duzeyler else "Uygun"
        c["bulgu"] = "; ".join(dict.fromkeys(x[1] for x in b))
        for duzey, tur, aciklama in b:
            bulgular.append({"no": c["no"], "banka": c["banka"], "yon": c["yon"], "duzey": duzey, "tur": tur, "aciklama": aciklama})
            uy.append(uyari(DUZEY[duzey], tur, kim, aciklama))
    ozet = [{"yon": y, "adet": sum(1 for c in cekler if c["yon"] == y), "tutar": sum((c["rakam"] for c in cekler if c["yon"] == y), Z), **{s: sum(1 for c in cekler if c["yon"] == y and c["sonuc"] == s) for s in SONUC_RENK},
             "sorunlu_tutar": sum((c["rakam"] for c in cekler if c["yon"] == y and c["sonuc"] == "İşleme alma"), Z)} for y in ("Takasa verilecek", "Takastan gelen")]
    return {"cekler": cekler, "bulgular": bulgular, "ozet": ozet, "say": {s: sum(1 for c in cekler if c["sonuc"] == s) for s in SONUC_RENK}, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontrol Listesi"
    isaret = lambda v: "Var" if v is True else "Sorun" if v is False else "—"   # noqa: E731
    _baslik(ws, ["Sonuç", "Yön", "Çek No", "Banka", "Keşide Tarihi", "İbraz Son Gün", "Tutar (rakam)", "Tutar (yazı)", "Karekod", "Ciro Silsilesi", "İmza", "Bulgu", "Kontrol Eden / Paraf"], (14, 18, 14, 20, 12, 12, 15, 15, 9, 10, 8, 46, 20))
    for c in sorted(r["cekler"], key=lambda c: (c["yon"], list(SONUC_RENK).index(c["sonuc"]))):
        _satir(ws, [c["sonuc"], c["yon"], c["no"], c["banka"], c["keside"], c["ibraz_son"], c["rakam"], c["yazi"], isaret(c["karekod"]), isaret(c["ciro"]), isaret(c["imza"]), c["bulgu"], None], para_s=(7, 8), tarih_s=(5, 6),
               dolgu={1: SONUC_RENK[c["sonuc"]]}, sar=(12,))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Yön", "Çek", "Tutar", "Uygun", "İncele", "İşleme Alma", "İşleme Alınmayacak Tutar"], (20, 8, 16, 9, 9, 11, 18))
    for o in r["ozet"]:
        _satir(ws, [o["yon"], o["adet"], o["tutar"], o["Uygun"], o["İncele"], o["İşleme alma"], o["sorunlu_tutar"]], para_s=(3, 7))
    ws.append([])
    ws.append([f"İşlem tarihi: {bugun:%d.%m.%Y}"])
    ws = wb.create_sheet("Bulgular")
    _baslik(ws, ["Düzey", "Bulgu", "Çek No", "Banka", "Yön", "Açıklama"], (10, 28, 14, 20, 18, 90))
    for b in r["bulgular"]:
        _satir(ws, [b["duzey"], b["tur"], b["no"], b["banka"], b["yon"], b["aciklama"]], dolgu={1: RENK[DUZEY[b["duzey"]]]}, sar=(6,))
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, ibraz_gun: int = 10, ibraz_ay: int = 1) -> dict:
    cekler, uy = oku(yol)
    r = kontrol_et(cekler, bugun, ibraz_gun, ibraz_ay)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Takasa verilecek ve takastan gelen çekler için kontrol listesi üretir.")
    p.add_argument("--cekler", type=Path, help="Çek No, Yön, Banka, Keşide Tarihi, Tutar (rakam), Tutar (yazı), Karekod, Ciro Silsilesi, İmza, Farklı Yer")
    p.add_argument("--tarih", help="İşlem tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--ibraz-gun", type=int, default=10, help="Düzenlendiği yerde ödenecek çekte ibraz süresi, gün (varsayılan 10)")
    p.add_argument("--ibraz-ay", type=int, default=1, help="Başka yerde ödenecek çekte ibraz süresi, ay (varsayılan 1)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "cek_takas_kontrol_listesi.xlsx")
    a = p.parse_args(argv)
    ornek = a.cekler is None
    yol = ORNEK / "cekler.csv" if ornek else a.cekler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or not 1 <= a.ibraz_gun <= 365 or not 1 <= a.ibraz_ay <= 12:
        print("[X] --tarih GG.AA.YYYY; --ibraz-gun 1–365; --ibraz-ay 1–12 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.ibraz_gun, a.ibraz_ay)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {len(r['cekler'])} çek · Uygun {s['Uygun']}, İncele {s['İncele']}, İşleme alma {s['İşleme alma']}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
