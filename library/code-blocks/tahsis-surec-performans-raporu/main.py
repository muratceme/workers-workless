"""
Tahsis Süreç Performans Raporu — Gerisi Bende kod bloğu
Krediler Tahsis › Kredi Tahsis Müdürü

Kredi tekliflerinin bekleme sürelerini, onay / ret oranlarını ve iade nedenlerini raporlar.
  Süre (gün)       = karar tarihi − giriş tarihi   (takvim günü; bekleyen teklifte durum tarihi − giriş tarihi)
  Sonuçlanan       = Onay + Ret + İade   (Bekliyor hariç)
  Onay / Ret / İade % = ilgili adet ÷ sonuçlanan × 100;   tutar bazında onay % = onaylanan tutar ÷ sonuçlanan tutar × 100
  Ortalama ve medyan süre = sonuçlanan tekliflerin süreleri üzerinden
  Süre hedefi içinde % = süresi --sla (5) gün ve altında olan sonuçlanan teklif ÷ sonuçlanan × 100
  İade oranı uyarısı = en az 3 sonuçlanan teklifi olan şubede iade oranı --iade-esik (%30) ve üstünde
Süre hedefi ve eşikler yönetim tercihidir. İade edilip yeniden gelen teklif ayrı numarayla girilirse ayrı teklif sayılır.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"no": ("teklif no", "teklif numarasi", "basvuru no", "no"), "sube": ("sube", "birim"), "segment": ("segment", "musteri segmenti"), "tutar": ("tutar", "teklif tutari", "talep tutari"),
         "giris": ("giris tarihi", "teklif tarihi", "basvuru tarihi"), "karar": ("karar tarihi", "sonuc tarihi"), "sonuc": ("karar", "sonuc", "durum"), "neden": ("iade nedeni", "neden", "ret nedeni"), "analist": ("analist", "tahsis uzmani", "uzman")}
KARAR = {"onay": "Onay", "onaylandi": "Onay", "kabul": "Onay", "ret": "Ret", "red": "Ret", "reddedildi": "Ret", "iade": "İade", "iade edildi": "İade", "": "Bekliyor", "bekliyor": "Bekliyor", "incelemede": "Bekliyor", "degerlendirmede": "Bekliyor"}
KIRILIM = (("Şubeler", "sube", "Şube"), ("Segmentler", "segment", "Segment"), ("Analistler", "analist", "Analist"))
Z, YUZ = Decimal(0), Decimal(100)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("no", "giris")):
        no = metin(r["no"])
        if not no:
            continue
        giris, karar_t, sonuc, tutar = tarih(r["giris"]), tarih(r["karar"]), KARAR.get(katla(r["sonuc"])), para(r["tutar"]) if metin(r["tutar"]) else Z
        hata = ("giriş tarihi okunamadı" if giris is None else "karar 'Onay', 'Ret', 'İade' ya da 'Bekliyor' olmalı" if sonuc is None else "tutar okunamadı" if tutar is None or tutar < 0
                else "sonuçlanmış teklifin karar tarihi yok" if sonuc != "Bekliyor" and karar_t is None else "karar tarihi giriş tarihinden önce" if karar_t and karar_t < giris else None)
        if hata:
            uy.append(uyari("Yüksek", "Teklif okunamadı", no, f"Satır {r['_satir']}: {hata}. Satır alınmadı."))
            continue
        l.append({"no": no, "sube": metin(r["sube"]) or "(şube yok)", "segment": metin(r["segment"]) or "(segment yok)", "tutar": tutar, "giris": giris, "karar_t": karar_t if sonuc != "Bekliyor" else None, "sonuc": sonuc,
                  "neden": metin(r["neden"]), "analist": metin(r["analist"]) or "(atanmamış)"})
    return l, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * YUZ / Decimal(payda)) if payda else None


def _medyan(d: list[int]) -> Decimal:
    d = sorted(d)
    return yuvarla(Decimal(d[len(d) // 2]) if len(d) % 2 else Decimal(d[len(d) // 2 - 1] + d[len(d) // 2]) / 2)


def _ozet(liste: list[dict], sla: int) -> dict:
    s = [t for t in liste if t["sonuc"] != "Bekliyor"]
    say = {k: sum(1 for t in s if t["sonuc"] == k) for k in ("Onay", "Ret", "İade")}
    sureler = [t["sure"] for t in s]
    return {"teklif": len(liste), "tutar": sum((t["tutar"] for t in liste), Z), "sonuclanan": len(s), "onay": say["Onay"], "ret": say["Ret"], "iade": say["İade"], "bekleyen": len(liste) - len(s),
            "bekleyen_tutar": sum((t["tutar"] for t in liste if t["sonuc"] == "Bekliyor"), Z), "onay_oran": _oran(say["Onay"], len(s)), "ret_oran": _oran(say["Ret"], len(s)), "iade_oran": _oran(say["İade"], len(s)),
            "onay_tutar_oran": _oran(sum((t["tutar"] for t in s if t["sonuc"] == "Onay"), Z), sum((t["tutar"] for t in s), Z)), "ort_sure": yuvarla(Decimal(sum(sureler)) / len(sureler)) if sureler else None,
            "medyan_sure": _medyan(sureler) if sureler else None, "sla_oran": _oran(sum(1 for x in sureler if x <= sla), len(sureler)), "geciken_bekleyen": sum(1 for t in liste if t["sonuc"] == "Bekliyor" and t["sure"] > sla)}


def raporla(teklifler: list[dict], bugun: date, sla: int = 5, iade_esik: Decimal = Decimal(30)) -> dict:
    if not teklifler:
        raise ValueError("Raporlanacak teklif yok.")
    uy = []
    for t in teklifler:
        t["sure"] = max(((t["karar_t"] or bugun) - t["giris"]).days, 0)
        if t["sonuc"] == "Bekliyor" and t["sure"] > sla:
            uy.append(uyari("Yüksek" if t["sure"] > 2 * sla else "Orta", "Bekleyen teklif süreyi aştı", f"{t['no']} · {t['sube']}", f"{t['giris']:%d.%m.%Y} tarihinde girildi, {t['sure']} gündür bekliyor (hedef {sla} gün). Tutar {t['tutar']:,.2f}; "
                            f"analist: {t['analist']}."))
        if t["sonuc"] == "İade" and not t["neden"]:
            uy.append(uyari("Bilgi", "İade nedeni yazılmamış", f"{t['no']} · {t['sube']}", "İade edilen teklifin nedeni boş; neden dağılımında '(belirtilmemiş)' olarak yer aldı."))
    kirilimlar = {}
    for baslik, alan, _ in KIRILIM:
        grup = {}
        for t in teklifler:
            grup.setdefault(katla(t[alan]), []).append(t)
        kirilimlar[baslik] = [{"ad": g[0][alan], **_ozet(g, sla)} for _, g in sorted(grup.items())]
    for s in kirilimlar["Şubeler"]:
        if s["sonuclanan"] >= 3 and s["iade_oran"] >= iade_esik:
            uy.append(uyari("Orta", "İade oranı yüksek", f"Şube · {s['ad']}", f"Sonuçlanan {s['sonuclanan']} teklifin {s['iade']} tanesi iade edildi (%{s['iade_oran']}; eşik %{iade_esik:f}). İade nedenlerine bakın."))
    nedenler = {}
    for t in teklifler:
        if t["sonuc"] == "İade":
            n = nedenler.setdefault(katla(t["neden"]), {"neden": t["neden"] or "(belirtilmemiş)", "adet": 0, "tutar": Z})
            n["adet"] += 1
            n["tutar"] += t["tutar"]
    toplam_iade = sum(n["adet"] for n in nedenler.values())
    for n in nedenler.values():
        n["pay"] = _oran(n["adet"], toplam_iade)
    bekleyen = sorted((t for t in teklifler if t["sonuc"] == "Bekliyor"), key=lambda t: (-t["sure"], -t["tutar"], katla(t["no"])))
    return {"teklifler": teklifler, "ozet": _ozet(teklifler, sla), "kirilimlar": kirilimlar, "nedenler": sorted(nedenler.values(), key=lambda n: (-n["adet"], -n["tutar"], katla(n["neden"]))), "bekleyen": bekleyen, "sla": sla, "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    o, sla = r["ozet"], r["sla"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 18))
    for etiket, deger, para_mi in (("Teklif sayısı", o["teklif"], False), ("Teklif tutarı", o["tutar"], True), ("Sonuçlanan", o["sonuclanan"], False), ("Onay", o["onay"], False), ("Ret", o["ret"], False), ("İade", o["iade"], False),
                                   ("Onay %", o["onay_oran"], True), ("Ret %", o["ret_oran"], True), ("İade %", o["iade_oran"], True), ("Onay % (tutar)", o["onay_tutar_oran"], True), ("Ortalama süre (gün)", o["ort_sure"], True),
                                   ("Medyan süre (gün)", o["medyan_sure"], True), (f"Süre hedefi ({sla} gün) içinde %", o["sla_oran"], True), ("Bekleyen", o["bekleyen"], False), ("Bekleyen tutar", o["bekleyen_tutar"], True),
                                   ("Süre hedefini aşan bekleyen", o["geciken_bekleyen"], False)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}. Süreler takvim günüdür."])
    for baslik, _, ad in KIRILIM:
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Teklif", "Tutar", "Sonuçlanan", "Onay", "Ret", "İade", "Onay %", "Ret %", "İade %", "Ort. Süre", "Medyan Süre", "Hedef İçinde %", "Bekleyen", "Hedefi Aşan Bekleyen"], (24, 8, 16, 11, 8, 8, 8, 9, 9, 9, 10, 10, 11, 10, 12))
        for s in r["kirilimlar"][baslik]:
            _satir(ws, [s["ad"], s["teklif"], s["tutar"], s["sonuclanan"], s["onay"], s["ret"], s["iade"], s["onay_oran"], s["ret_oran"], s["iade_oran"], s["ort_sure"], s["medyan_sure"], s["sla_oran"], s["bekleyen"], s["geciken_bekleyen"] or None],
                   para_s=(3, 8, 9, 10, 11, 12, 13), dolgu={15: KIRMIZI if s["geciken_bekleyen"] else None})
    ws = wb.create_sheet("İade Nedenleri")
    _baslik(ws, ["Neden", "Adet", "Tutar", "Pay %"], (36, 8, 16, 9))
    for n in r["nedenler"]:
        _satir(ws, [n["neden"], n["adet"], n["tutar"], n["pay"]], para_s=(3, 4))
    ws = wb.create_sheet("Bekleyenler")
    _baslik(ws, ["Teklif No", "Şube", "Segment", "Tutar", "Giriş Tarihi", "Bekleme (gün)", "Analist"], (14, 20, 14, 16, 12, 10, 22))
    for t in r["bekleyen"]:
        _satir(ws, [t["no"], t["sube"], t["segment"], t["tutar"], t["giris"], t["sure"], t["analist"]], para_s=(4,), tarih_s=(5,), dolgu={6: KIRMIZI if t["sure"] > 2 * sla else SARI if t["sure"] > sla else None})
    ws = wb.create_sheet("Teklifler")
    _baslik(ws, ["Teklif No", "Şube", "Segment", "Tutar", "Giriş Tarihi", "Karar Tarihi", "Karar", "Süre (gün)", "İade Nedeni", "Analist"], (14, 20, 14, 16, 12, 12, 10, 9, 30, 22))
    for t in r["teklifler"]:
        _satir(ws, [t["no"], t["sube"], t["segment"], t["tutar"], t["giris"], t["karar_t"], t["sonuc"], t["sure"], t["neden"], t["analist"]], para_s=(4,), tarih_s=(5, 6), dolgu={7: {"Onay": YESIL, "Ret": KIRMIZI, "İade": SARI}.get(t["sonuc"])})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, sla: int = 5, iade_esik: Decimal = Decimal(30)) -> dict:
    teklifler, uy = oku(yol)
    r = raporla(teklifler, bugun, sla, iade_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Kredi tekliflerinin bekleme sürelerini, onay / ret oranlarını ve iade nedenlerini raporlar.")
    p.add_argument("--teklifler", type=Path, help="Teklif No, Şube, Segment, Tutar, Giriş Tarihi, Karar Tarihi, Karar (Onay / Ret / İade / Bekliyor), İade Nedeni, Analist")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--sla", type=int, default=5, help="Süre hedefi, takvim günü (varsayılan 5)")
    p.add_argument("--iade-esik", default="30", help="Şube iade oranı bu yüzde ve üstündeyse uyarılır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "tahsis_surec_performans.xlsx")
    a = p.parse_args(argv)
    ornek = a.teklifler is None
    yol = ORNEK / "teklifler.csv" if ornek else a.teklifler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    esik = para(a.iade_esik)
    if bugun is None or a.sla < 1 or esik is None or not 0 < esik <= 100:
        print("[X] --tarih GG.AA.YYYY; --sla en az 1; --iade-esik 0–100 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.sla, esik)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['teklif']} teklif · sonuçlanan {o['sonuclanan']}" + (f" (onay %{o['onay_oran']}, ret %{o['ret_oran']}, iade %{o['iade_oran']}) · ortalama {o['ort_sure']} gün, medyan {o['medyan_sure']} gün · hedef içinde %{o['sla_oran']}" if o["sonuclanan"] else "")
          + f" · bekleyen {o['bekleyen']} ({o['geciken_bekleyen']} tanesi hedefi aştı)")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
