import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
GIRIS = date(2026, 9, 1)


def t(no, sonuc="Onay", sure=2, tutar=100, sube="S", segment="G", analist="A", neden=""):
    """sure: sonuçlanan teklifte karar süresi; bekleyende kaç gündür beklediği."""
    bekliyor = sonuc == "Bekliyor"
    giris = BUGUN - timedelta(days=sure) if bekliyor else GIRIS
    return {"no": no, "sube": sube, "segment": segment, "tutar": D(str(tutar)), "giris": giris, "karar_t": None if bekliyor else giris + timedelta(days=sure), "sonuc": sonuc, "neden": neden, "analist": analist}


class RaporTesti(unittest.TestCase):
    def test_sure_ve_oranlar(self):
        l = [t("1", "Onay", 1, 300), t("2", "Onay", 5, 100), t("3", "Ret", 6, 400), t("4", "İade", 0, 200, neden="Eksik"), t("5", "Bekliyor", 5, 1000), t("6", "Bekliyor", 6, 50), t("7", "Bekliyor", 10, 60), t("8", "Bekliyor", 11, 70)]
        r = main.raporla(l, BUGUN)
        self.assertEqual(r["ozet"], {"teklif": 8, "tutar": D(2180), "sonuclanan": 4, "onay": 2, "ret": 1, "iade": 1, "bekleyen": 4, "bekleyen_tutar": D(1180), "onay_oran": D("50.00"), "ret_oran": D("25.00"), "iade_oran": D("25.00"),
                                     "onay_tutar_oran": D("40.00"), "ort_sure": D("3.00"), "medyan_sure": D("3.00"), "sla_oran": D("75.00"), "geciken_bekleyen": 3})
        self.assertEqual(([x["no"] for x in r["bekleyen"]], [(u["onem"], u["kim"]) for u in r["uyarilar"]]), (["8", "7", "6", "5"], [("Orta", "6 · S"), ("Orta", "7 · S"), ("Yüksek", "8 · S")]))
        o = main.raporla([t("1", sure=3), t("2", sure=9), t("3", sure=4)], BUGUN, sla=3)["ozet"]
        self.assertEqual((o["ort_sure"], o["medyan_sure"], o["sla_oran"]), (D("5.33"), D("4.00"), D("33.33")))
        o = main.raporla([t("1", "Bekliyor", 1)], BUGUN)["ozet"]
        self.assertEqual((o["sonuclanan"], o["onay_oran"], o["ort_sure"], o["medyan_sure"], o["sla_oran"], o["onay_tutar_oran"]), (0, None, None, None, None, None))
        self.assertEqual((main._medyan([7]), main._medyan([1, 2]), main._medyan([5, 1, 3, 2])), (D("7.00"), D("1.50"), D("2.50")))
        with self.assertRaises(ValueError):
            main.raporla([], BUGUN)

    def test_kirilim_ve_iade(self):
        l = [t("1", "Onay", 2, sube="A", segment="K", analist="X"), t("2", "İade", 4, sube="a", segment="T", analist="Y", neden="Teminat"), t("3", "Ret", 6, sube="A", segment="K", analist="X"), t("4", "İade", 1, 500, sube="B", analist="Y"),
             t("5", "İade", 1, sube="B", analist="Y", neden="teminat"), t("6", "Bekliyor", 2, sube="B", analist="X"), t("7", "İade", 3, sube="C", neden="Eksik"), t("8", "İade", 3, sube="C", neden="Eksik"), t("9", "Onay", 3, sube="C"),
             t("10", "Onay", 3, sube="C"), t("11", "Onay", 3, sube="C"), t("12", "Onay", 3, sube="C"), t("13", "Onay", 3, sube="C")]
        r = main.raporla(l, BUGUN)
        s = lambda ad: [(x["ad"], x["teklif"], x["sonuclanan"], x["onay"], x["ret"], x["iade"], x["bekleyen"], x["iade_oran"], x["ort_sure"]) for x in r["kirilimlar"][ad]]   # noqa: E731
        self.assertEqual(s("Şubeler"), [("A", 3, 3, 1, 1, 1, 0, D("33.33"), D("4.00")), ("B", 3, 2, 0, 0, 2, 1, D("100.00"), D("1.00")), ("C", 7, 7, 5, 0, 2, 0, D("28.57"), D("3.00"))])
        self.assertEqual([x[:3] for x in s("Segmentler")] + [x[:3] for x in s("Analistler")], [("G", 10, 9), ("K", 2, 2), ("T", 1, 1), ("A", 7, 7), ("X", 3, 2), ("Y", 3, 3)])
        self.assertEqual(r["nedenler"], [{"neden": "Eksik", "adet": 2, "tutar": D(200), "pay": D("40.00")}, {"neden": "Teminat", "adet": 2, "tutar": D(200), "pay": D("40.00")},
                                         {"neden": "(belirtilmemiş)", "adet": 1, "tutar": D(500), "pay": D("20.00")}])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "İade nedeni yazılmamış", "4 · B"), ("Orta", "İade oranı yüksek", "Şube · A")])
        self.assertEqual(sorted(u["kim"] for u in main.raporla(l, BUGUN, iade_esik=D("28.57"))["uyarilar"] if u["tur"] == "İade oranı yüksek"), ["Şube · A", "Şube · C"])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "teklifler.csv", cikti, BUGUN)
            self.assertEqual(r["ozet"], {"teklif": 13, "tutar": D(17800000), "sonuclanan": 10, "onay": 5, "ret": 2, "iade": 3, "bekleyen": 3, "bekleyen_tutar": D(4600000), "onay_oran": D("50.00"), "ret_oran": D("20.00"), "iade_oran": D("30.00"),
                                         "onay_tutar_oran": D("55.30"), "ort_sure": D("4.70"), "medyan_sure": D("3.50"), "sla_oran": D("70.00"), "geciken_bekleyen": 2})
            s = lambda ad: [(x["ad"], x["teklif"], x["sonuclanan"], x["onay"], x["ret"], x["iade"], x["ort_sure"], x["sla_oran"], x["bekleyen"]) for x in r["kirilimlar"][ad]]   # noqa: E731
            self.assertEqual(s("Şubeler"), [("Çarşı", 5, 4, 1, 1, 2, D("2.75"), D("100.00"), 1), ("Merkez", 4, 3, 2, 1, 0, D("6.33"), D("33.33"), 1), ("Sanayi", 4, 3, 2, 0, 1, D("5.67"), D("66.67"), 1)])
            self.assertEqual(s("Segmentler"), [("KOBİ", 9, 7, 4, 1, 2, D("3.43"), D("85.71"), 2), ("Ticari", 4, 3, 1, 1, 1, D("7.67"), D("33.33"), 1)])
            self.assertEqual(s("Analistler"), [("Ali Can", 8, 6, 3, 1, 2, D("3.33"), D("83.33"), 2), ("Ece Tan", 5, 4, 2, 1, 1, D("6.75"), D("50.00"), 1)])
            self.assertEqual(r["nedenler"], [{"neden": "Eksik mali veri", "adet": 2, "tutar": D(550000), "pay": D("66.67")}, {"neden": "Teminat yetersiz", "adet": 1, "tutar": D(2000000), "pay": D("33.33")}])
            self.assertEqual([(x["no"], x["sure"]) for x in r["bekleyen"]], [("TK-11", 14), ("TK-12", 6), ("TK-13", 3)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Bekleyen teklif süreyi aştı", "TK-11 · Merkez"), ("Orta", "Bekleyen teklif süreyi aştı", "TK-12 · Sanayi"), ("Orta", "İade oranı yüksek", "Şube · Sanayi"), ("Orta", "İade oranı yüksek", "Şube · Çarşı")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Şubeler", "Segmentler", "Analistler", "İade Nedenleri", "Bekleyenler", "Teklifler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Başvuru No;Birim;Talep Tutarı;Başvuru Tarihi;Sonuç Tarihi;Durum;Ret Nedeni;Uzman\nB1;;1.250.000,50;01.09.2026;05.09.2026;ONAYLANDI;;\nB2;Ş;;02.09.2026;09.09.2026;İncelemede;;U\nB3;;5;x;;;;\n"
                         "B4;;5;01.09.2026;;Askıda;;\nB5;;x;01.09.2026;;;;\nB6;;5;01.09.2026;;Ret;;\nB7;;5;05.09.2026;01.09.2026;Ret;;\n;;5;01.09.2026;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["no"], x["sube"], x["segment"], x["tutar"], x["giris"], x["karar_t"], x["sonuc"], x["analist"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("B1", "(şube yok)", "(segment yok)", D("1250000.50"), date(2026, 9, 1), date(2026, 9, 5), "Onay", "(atanmamış)"), ("B2", "Ş", "(segment yok)", D(0), date(2026, 9, 2), None, "Bekliyor", "U")],
                              [("Teklif okunamadı", f"B{n}") for n in (3, 4, 5, 6, 7)]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--teklifler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--sla", "3", "--iade-esik", "25"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--sla", "0"], ["--iade-esik", "0"], ["--iade-esik", "101"])], [2] * 4)
            self.assertEqual(main.main(["--teklifler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Teklif No;Şube\nA;B\n", "Teklif No;Giriş Tarihi\n;01.09.2026\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
