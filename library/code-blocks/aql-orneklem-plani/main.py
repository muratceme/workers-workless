"""
AQL Örneklem Planı — Workers / Workless kod bloğu
Üretim › Kalite Kontrol Elemanı · Tekstil › Kalite Kontrol

Parti büyüklüğü, muayene seviyesi ve AQL değerinden tek örneklemeli normal muayene planını
(örneklem sayısı, kabul ve ret sayısı) bulur; muayene sonuçları verilirse partiyi KABUL / RET
olarak değerlendirir. Kritik, majör ve minör hatalar için ayrı AQL kullanılabilir.

Tablolar: MIL-STD-105E Tablo I ve II-A (kamu malı); ISO 2859-1 Tablo 1 ve 2-A ile eşdeğerdir.
Kapsam: AQL 0,010 – 10 (yüzde kusurlu), normal muayene, tek örnekleme.

Kullanım:
    python main.py --parti 1200 --aql 2.5                    # tek plan
    python main.py --parti 3000 --majör 2.5 --minör 4.0 --kritik 0 --seviye II
    python main.py --girdi muayeneler.xlsx                   # parti listesi + bulunan hatalar
"""
from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent

# Tablo I — örneklem kod harfleri (parti büyüklüğü üst sınırı → S-1, S-2, S-3, S-4, I, II, III)
KOD_HARFLERI = [
    (8,        "AAAAAAB"), (15,      "AAAAABC"), (25,     "AABBBCD"), (50,      "ABBCCDE"),
    (90,       "BBCCCEF"), (150,     "BBCDDFG"), (280,    "BCDEEGH"), (500,     "BCDEFHJ"),
    (1200,     "CCEFGJK"), (3200,    "CDEGHKL"), (10000,  "CDFGJLM"), (35000,   "CDFHKMN"),
    (150000,   "DEGJLNP"), (500000,  "DEGJMPQ"), (None,   "DEHKNQR"),
]
SEVIYELER = ["S-1", "S-2", "S-3", "S-4", "I", "II", "III"]
HARFLER = "ABCDEFGHJKLMNPQR"
ORNEKLEM = dict(zip(HARFLER, [2, 3, 5, 8, 13, 20, 32, 50, 80, 125, 200, 315, 500, 800, 1250, 2000]))
AQL_DEGERLERI = [Decimal(x) for x in ("0.010", "0.015", "0.025", "0.040", "0.065", "0.10", "0.15", "0.25",
                                      "0.40", "0.65", "1.0", "1.5", "2.5", "4.0", "6.5", "10")]
# Tablo II-A köşegeni: 0/1'den sonra ↑ ve ↓, sonra bu kabul/ret çiftleri
CIFTLER = [(0, 1), (1, 2), (2, 3), (3, 4), (5, 6), (7, 8), (10, 11), (14, 15), (21, 22)]
YUKARI, ASAGI = "↑", "↓"


def kod_harfi(parti: int, seviye: str = "II") -> str:
    if parti < 2:
        raise ValueError("Parti büyüklüğü en az 2 olmalı")
    if seviye not in SEVIYELER:
        raise ValueError(f"Muayene seviyesi {', '.join(SEVIYELER)} olmalı")
    for ust, harfler in KOD_HARFLERI:
        if ust is None or parti <= ust:
            return harfler[SEVIYELER.index(seviye)]


def _hucre(satir: int, sutun: int):
    """Tablo II-A hücresi: (Ac, Re) ya da ok. satir: A=0..R=15, sutun: AQL 0.010=0..10=15."""
    j = sutun - (14 - satir)
    if j < 0:
        return ASAGI
    if j == 0:
        return CIFTLER[0]
    if j == 1:
        return YUKARI
    if j == 2:
        return ASAGI
    if j - 2 < len(CIFTLER):
        return CIFTLER[j - 2]
    return YUKARI


@dataclass
class Plan:
    parti: int
    seviye: str
    aql: Decimal
    kod_harfi: str
    kullanilan_harf: str
    orneklem: int
    kabul: int
    ret: int
    yuzde_yuz: bool
    not_: str = ""


def plan_bul(parti: int, aql, seviye: str = "II") -> Plan:
    aql = Decimal(str(aql))
    if aql not in AQL_DEGERLERI:
        raise ValueError(f"AQL {aql} desteklenmiyor. Kullanılabilir: {', '.join(str(a) for a in AQL_DEGERLERI)}")
    harf = kod_harfi(parti, seviye)
    satir, sutun = HARFLER.index(harf), AQL_DEGERLERI.index(aql)
    notlar = []
    for _ in range(len(HARFLER)):
        h = _hucre(satir, sutun)
        if h == ASAGI and satir == len(HARFLER) - 1:
            h = YUKARI          # R satırının altında plan yoktur
        if h == ASAGI:
            satir += 1
        elif h == YUKARI:
            satir -= 1
        else:
            break
    kullanilan = HARFLER[satir]
    if kullanilan != harf:
        notlar.append(f"Tablodaki ok nedeniyle {harf} yerine {kullanilan} harfinin planı kullanıldı")
    n = ORNEKLEM[kullanilan]
    yuzde_yuz = n >= parti
    if yuzde_yuz:
        notlar.append("Örneklem parti büyüklüğüne eşit veya büyük: %100 muayene yapın")
        n = parti
    return Plan(parti, seviye, aql, harf, kullanilan, n, h[0], h[1], yuzde_yuz, "; ".join(notlar))


def karar(hata_sayisi: int | None, plan: Plan) -> str:
    if hata_sayisi is None:
        return ""
    return "KABUL" if hata_sayisi <= plan.kabul else "RET"


# ----------------------------------------------------------------------------
# Toplu muayene listesi
# ----------------------------------------------------------------------------

def kucuk(s) -> str:
    return str(s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def tam_sayi(x) -> int | None:
    if x in (None, ""):
        return None
    return int(Decimal(str(x).replace(",", ".")))


def ondalik(x) -> Decimal | None:
    if x in (None, ""):
        return None
    return Decimal(str(x).replace(",", ".")).normalize()


SUTUNLAR = {
    "ref": ("sipariş", "sipariş no", "parti", "parti no", "lot", "referans"),
    "parti": ("parti büyüklüğü", "parti adedi", "adet", "lot büyüklüğü", "miktar"),
    "seviye": ("seviye", "muayene seviyesi"),
    "aql_kritik": ("kritik aql", "aql kritik"),
    "aql_major": ("majör aql", "major aql", "aql majör", "aql major"),
    "aql_minor": ("minör aql", "minor aql", "aql minör", "aql minor"),
    "kritik": ("kritik hata", "kritik"),
    "major": ("majör hata", "major hata", "majör", "major"),
    "minor": ("minör hata", "minor hata", "minör", "minor"),
}


def liste_oku(yol: Path) -> list[dict]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        metin = yol.read_text(encoding="utf-8-sig")
        ayirici = csv.Sniffer().sniff(metin.splitlines()[0], delimiters=",;\t").delimiter
        satirlar = list(csv.reader(metin.splitlines(), delimiter=ayirici))
    satirlar = [r for r in satirlar if any(c not in (None, "") for c in r)]
    b = [kucuk(x) for x in satirlar[0]]
    k = {alan: next((i for i, x in enumerate(b) if x in adlar), None) for alan, adlar in SUTUNLAR.items()}
    if k["parti"] is None:
        raise SystemExit(f"'Parti Büyüklüğü' sütunu bulunamadı. Başlıklar: {satirlar[0]}")
    return [{alan: (r[i] if i is not None and i < len(r) else None) for alan, i in k.items()} for r in satirlar[1:]]


def degerlendir(k: dict, varsayilan: dict) -> dict:
    parti = tam_sayi(k["parti"])
    seviye = str(k.get("seviye") or varsayilan["seviye"]).strip().upper().replace("S", "S-").replace("S--", "S-")
    sonuc = {"ref": k.get("ref"), "parti": parti, "seviye": seviye, "kararlar": {}, "planlar": {}}
    for tur in ("kritik", "major", "minor"):
        aql = ondalik(k.get(f"aql_{tur}"))
        aql = varsayilan[tur] if aql is None else aql
        if aql is None:
            continue
        bulunan = tam_sayi(k.get(tur))
        if aql == 0:   # sıfır tolerans: hiç hata kabul edilmez (yaygın kritik hata uygulaması)
            ref_plan = plan_bul(parti, varsayilan.get("major") or Decimal("2.5"), seviye)
            p = Plan(parti, seviye, Decimal(0), ref_plan.kod_harfi, ref_plan.kullanilan_harf, ref_plan.orneklem, 0, 1,
                     ref_plan.yuzde_yuz, "Sıfır tolerans: tek hata reddedilir")
        else:
            p = plan_bul(parti, aql, seviye)
        sonuc["planlar"][tur] = p
        sonuc["kararlar"][tur] = karar(bulunan, p)
        sonuc[f"bulunan_{tur}"] = bulunan
    kararlar = [x for x in sonuc["kararlar"].values() if x]
    sonuc["genel"] = ("RET" if "RET" in kararlar else "KABUL") if kararlar else "Plan"
    return sonuc


BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"KABUL": "B7E4C7", "RET": "FF8A8A"}
TUR_AD = {"kritik": "Kritik", "major": "Majör", "minor": "Minör"}


def rapor_yaz(sonuclar: list[dict], cikti: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Muayene Planı"
    basliklar = ["Referans", "Parti", "Seviye", "Kod Harfi", "Örneklem"]
    for tur in ("kritik", "major", "minor"):
        basliklar += [f"{TUR_AD[tur]} AQL", f"{TUR_AD[tur]} Ac/Re", f"{TUR_AD[tur]} Bulunan", f"{TUR_AD[tur]} Karar"]
    basliklar += ["Genel Karar", "Notlar"]
    ws.append(basliklar)
    for h in ws[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for s in sonuclar:
        planlar = s["planlar"]
        ilk = next(iter(planlar.values()), None)
        satir = [s["ref"], s["parti"], s["seviye"], ilk.kod_harfi if ilk else "", ilk.orneklem if ilk else ""]
        notlar = []
        for tur in ("kritik", "major", "minor"):
            p = planlar.get(tur)
            if p:
                satir += [float(p.aql), f"{p.kabul}/{p.ret}" + (f" (n={p.orneklem})" if ilk and p.orneklem != ilk.orneklem else ""),
                          s.get(f"bulunan_{tur}"), s["kararlar"][tur]]
                if p.not_:
                    notlar.append(f"{TUR_AD[tur]}: {p.not_}")
            else:
                satir += ["", "", "", ""]
        satir += [s["genel"], "; ".join(notlar)]
        ws.append(satir)
        for j, v in enumerate(satir, 1):
            if v in RENK:
                ws.cell(ws.max_row, j).fill = PatternFill("solid", fgColor=RENK[v])
    for j, b in enumerate(basliklar, 1):
        ws.column_dimensions[get_column_letter(j)].width = 50 if b == "Notlar" else max(10, len(b) + 2)
    ws.freeze_panes = "B2"

    bilgi = wb.create_sheet("Bilgi")
    for satir in [
        ["Tablolar", "MIL-STD-105E Tablo I ve II-A (kamu malı); ISO 2859-1 Tablo 1 ve 2-A ile eşdeğer"],
        ["Kapsam", "Normal muayene, tek örnekleme, AQL 0,010 – 10"],
        ["Karar", "Bulunan hata ≤ Ac → KABUL; ≥ Re → RET"],
        ["Ok kuralı", "Tabloda ok varsa okun gösterdiği ilk plan kullanılır; örneklem parti büyüklüğüne eşit/büyükse %100 muayene"],
        ["Sıfır tolerans", "AQL 0 girilirse (yaygın kritik hata uygulaması) tek hata partiyi reddeder"],
        ["Seviyeler", "Genel: I, II (varsayılan), III · Özel: S-1, S-2, S-3, S-4"],
    ]:
        bilgi.append(satir)
    bilgi.column_dimensions["A"].width = 18
    bilgi.column_dimensions["B"].width = 110
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="AQL tek örnekleme planı (ISO 2859-1 / MIL-STD-105E, normal muayene).")
    ap.add_argument("--parti", type=int, help="Parti büyüklüğü (adet)")
    ap.add_argument("--aql", type=Decimal, help="Tek bir AQL değeri")
    ap.add_argument("--kritik", type=Decimal, default=None, help="Kritik hata AQL (0 = sıfır tolerans)")
    ap.add_argument("--majör", "--major", dest="major", type=Decimal, default=None, help="Majör hata AQL")
    ap.add_argument("--minör", "--minor", dest="minor", type=Decimal, default=None, help="Minör hata AQL")
    ap.add_argument("--seviye", default="II", choices=SEVIYELER)
    ap.add_argument("--girdi", type=Path, help="Muayene listesi (.xlsx/.csv)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "aql_plani.xlsx")
    a = ap.parse_args(argv)

    if a.parti and a.aql is not None:
        p = plan_bul(a.parti, a.aql, a.seviye)
        print(f"[OK] Parti {a.parti}, seviye {a.seviye}, AQL {a.aql}: kod harfi {p.kod_harfi} → örneklem {p.orneklem}, "
              f"kabul {p.kabul}, ret {p.ret}" + (f" ({p.not_})" if p.not_ else ""))
        return
    varsayilan = {"seviye": a.seviye, "kritik": a.kritik, "major": a.major, "minor": a.minor}
    if a.parti:
        kayitlar = [{"ref": "-", "parti": a.parti}]
    else:
        kayitlar = liste_oku(a.girdi or BURASI / "ornek_veri" / "muayeneler.csv")
        if all(v is None for k, v in varsayilan.items() if k != "seviye"):
            varsayilan.update(kritik=Decimal(0), major=Decimal("2.5"), minor=Decimal("4.0"))
    sonuclar = [degerlendir(k, varsayilan) for k in kayitlar]
    rapor_yaz(sonuclar, a.cikti)
    for s in sonuclar:
        print(f"[OK] {s['ref']}: parti {s['parti']} → {s['genel']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
