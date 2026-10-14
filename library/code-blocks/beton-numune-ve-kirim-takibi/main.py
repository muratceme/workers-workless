"""
Beton Numune ve Kırım Takibi — Workers / Workless kod bloğu
Şantiye › Saha Mühendisi
Beton döküm kayıtları ve 7 / 28 günlük numune kırım sonuçlarını takip edip sınıf dayanımının altında kalanları işaretler.
  Karakteristik dayanım (fck) = beton sınıfından (C30/37): silindir numunede ilk sayı, küp numunede ikinci sayı
  Deney sonucu                = aynı döküm ve aynı yaştaki numunelerin ortalaması (MPa)
  Bireysel kriter             = 28 günlük her sonuç ≥ fck − --eksi (4)
  Grup kriteri                = aynı sınıftaki ardışık --grup (3) sonucun ortalaması ≥ fck + --ek (4)   (gruplar örtüşmez, döküm tarihine göre)
  7 gün oranı                 = 7 günlük sonuç ÷ fck × 100 (bilgi amaçlı; uygunluk 28 günlük sonuçla değerlendirilir)
  Numune farkı                = (en büyük − en küçük) ÷ ortalama × 100; --fark (%15) üstü bilgi notu
  Eksik sonuç                 = dökümün üzerinden 28 + --tolerans (2) gün geçmiş ama 28 günlük sonuç yok
Varsayılan kriterler TS EN 206'daki başlangıç üretimi uygunluk kriteridir; projenizde geçerli kriteri ve yönetmeliği doğrulayın.
"""
from __future__ import annotations

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"dokum": ("dokum no", "dokum", "irsaliye no", "parti no"), "tarih": ("dokum tarihi", "tarih"), "yer": ("yer", "yer eleman", "eleman", "blok", "konum", "yapi elemani"),
         "sinif": ("beton sinifi", "sinif", "beton"), "tip": ("numune tipi", "numune", "numune sekli"), "numune": ("numune no", "numune kodu", "kup no"),
         "yas": ("kirim yasi", "yas", "gun", "kirim gunu"), "dayanim": ("dayanim", "basinc dayanimi", "kirim sonucu", "mpa")}
SIFIR = Decimal(0)


def sinif_oku(x) -> tuple[Decimal, Decimal] | None:
    m = re.fullmatch(r"c?\s*(\d+(?:[.,]\d+)?)\s*/\s*(\d+(?:[.,]\d+)?)", metin(x).lower().replace(" ", ""))
    return (Decimal(m.group(1).replace(",", ".")), Decimal(m.group(2).replace(",", "."))) if m else None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("dokum", "tarih", "sinif")):
        dokum = metin(r["dokum"])
        if not dokum:
            continue
        t, sinif = tarih(r["tarih"]), sinif_oku(r["sinif"])
        tip = "Silindir" if "silindir" in katla(r["tip"]) else "Küp"
        yas, dayanim = (para(r["yas"]) if metin(r["yas"]) else None), (para(r["dayanim"]) if metin(r["dayanim"]) else None)
        if t is None or sinif is None or (metin(r["yas"]) and (yas is None or yas <= 0)) or (metin(r["dayanim"]) and (dayanim is None or dayanim <= 0)) or ((yas is None) != (dayanim is None)):
            uy.append(uyari("Yüksek", "Satır okunamadı", dokum, f"Satır {r['_satir']}: döküm tarihi, beton sınıfı ('{metin(r['sinif'])}'), kırım yaşı ya da dayanım okunamadı (yaş ve dayanım birlikte yazılmalı). Satır alınmadı."))
            continue
        l.append({"dokum": dokum, "tarih": t, "yer": metin(r["yer"]), "sinif": metin(r["sinif"]).upper().replace(" ", ""), "fck": sinif[0] if tip == "Silindir" else sinif[1], "tip": tip,
                  "numune": metin(r["numune"]), "yas": int(yas) if yas is not None else None, "dayanim": dayanim})
    return l, uy


def degerlendir(kayit: list[dict], gun: date, ek: Decimal = Decimal(4), eksi: Decimal = Decimal(4), grup: int = 3, fark: Decimal = Decimal(15), tolerans: int = 2) -> dict:
    if not kayit:
        raise ValueError("Döküm kaydı yok.")
    uy, dokumler = [], {}
    for k in kayit:
        d = dokumler.setdefault(katla(k["dokum"]), {"dokum": k["dokum"], "tarih": k["tarih"], "yer": k["yer"], "sinif": k["sinif"], "tip": k["tip"], "fck": k["fck"], "numuneler": defaultdict(list)})
        if (k["sinif"], k["tip"]) != (d["sinif"], d["tip"]):
            uy.append(uyari("Orta", "Döküm bilgisi tutarsız", k["dokum"], "Aynı dökümde farklı beton sınıfı ya da numune tipi yazılmış; ilk satırdaki kullanıldı."))
        if k["yas"] is not None:
            d["numuneler"][k["yas"]].append(k["dayanim"])
    sonuclar = []
    for d in sorted(dokumler.values(), key=lambda d: (d["tarih"], katla(d["dokum"]))):
        d["sonuc"] = {}
        for yas, l in sorted(d["numuneler"].items()):
            ort = sum(l, SIFIR) / len(l)
            s = {"dokum": d["dokum"], "tarih": d["tarih"], "yer": d["yer"], "sinif": d["sinif"], "tip": d["tip"], "fck": d["fck"], "yas": yas, "adet": len(l), "ortalama": yuvarla(ort, "0.1"),
                 "fark": yuvarla((max(l) - min(l)) * 100 / ort) if len(l) > 1 else None, "oran": yuvarla(ort * 100 / d["fck"])}
            s["bireysel"] = None if yas != 28 else ("Uygun" if ort >= d["fck"] - eksi else "Uygun değil")
            d["sonuc"][yas] = s
            sonuclar.append(s)
            if s["bireysel"] == "Uygun değil":
                uy.append(uyari("Yüksek", "Bireysel sonuç sınırın altında", d["dokum"], f"{d['yer']} · {d['sinif']} {d['tip'].lower()}: 28 gün {s['ortalama']} MPa < fck − {eksi} = {d['fck'] - eksi} MPa. "
                                                                                       "Yapı denetim ve proje müellifiyle değerlendirin (karot vb.)."))
            if s["fark"] is not None and s["fark"] > fark:
                uy.append(uyari("Bilgi", "Numuneler arası fark yüksek", f"{d['dokum']} · {yas} gün", f"Numuneler {', '.join(f'{x}' for x in l)} MPa; fark ortalamaya göre %{s['fark']} (eşik %{fark})."))
        if 28 not in d["sonuc"] and (gun - d["tarih"]).days >= 28 + tolerans:
            uy.append(uyari("Orta", "28 günlük sonuç eksik", d["dokum"], f"{d['tarih']:%d.%m.%Y} dökümü ({d['yer']}) {(gun - d['tarih']).days} gün önce yapıldı; 28 günlük kırım sonucu yok."))
    gruplar = []
    sinif = {}
    for s in sonuclar:
        if s["yas"] == 28:
            sinif.setdefault((s["sinif"], s["tip"]), []).append(s)
    for (sn, tip), l in sorted(sinif.items()):
        for i in range(0, len(l), grup):
            g = l[i:i + grup]
            ort = yuvarla(sum((s["ortalama"] for s in g), SIFIR) / len(g), "0.1")
            tam = len(g) == grup
            durum = ("Uygun" if ort >= g[0]["fck"] + ek else "Uygun değil") if tam else f"Eksik grup ({len(g)}/{grup})"
            gruplar.append({"sinif": sn, "tip": tip, "fck": g[0]["fck"], "dokumler": ", ".join(s["dokum"] for s in g), "ilk": g[0]["tarih"], "son": g[-1]["tarih"], "ortalama": ort, "sinir": g[0]["fck"] + ek, "durum": durum})
            if durum == "Uygun değil":
                uy.append(uyari("Yüksek", "Grup ortalaması yetersiz", f"{sn} {tip.lower()}", f"Dökümler {gruplar[-1]['dokumler']}: ortalama {ort} MPa < fck + {ek} = {g[0]['fck'] + ek} MPa."))
    eksik = [d for d in dokumler.values() if 28 not in d["sonuc"]]
    return {"sonuclar": sonuclar, "gruplar": gruplar, "dokum_sayisi": len(dokumler), "eksik": len(eksik), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, gun: date, ek: Decimal, eksi: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kırım Sonuçları"
    _baslik(ws, ["Döküm No", "Döküm Tarihi", "Yer / Eleman", "Sınıf", "Numune", "fck (MPa)", "Yaş (gün)", "Numune Sayısı", "Ortalama (MPa)", "fck'ye Oran %", "Numune Farkı %", "Bireysel Kriter"],
            (12, 11, 24, 9, 9, 9, 8, 9, 10, 10, 10, 13), "B2")
    for s in r["sonuclar"]:
        _satir(ws, [s["dokum"], s["tarih"], s["yer"], s["sinif"], s["tip"], s["fck"], s["yas"], s["adet"], s["ortalama"], s["oran"], s["fark"], s["bireysel"]], para_s=(10, 11), tarih_s=(2,),
               dolgu={12: None if s["bireysel"] is None else YESIL if s["bireysel"] == "Uygun" else KIRMIZI})
    ws.append([])
    ws.append([f"Bireysel: 28 gün sonucu ≥ fck − {eksi}. Grup: ardışık sonuçların ortalaması ≥ fck + {ek}. Rapor tarihi {gun:%d.%m.%Y}."])
    ws = wb.create_sheet("Grup Değerlendirmesi")
    _baslik(ws, ["Sınıf", "Numune", "fck (MPa)", "Dökümler", "İlk Döküm", "Son Döküm", "Grup Ortalaması", "Sınır (fck + ek)", "Durum"], (9, 9, 9, 30, 11, 11, 10, 10, 16))
    for g in r["gruplar"]:
        _satir(ws, [g["sinif"], g["tip"], g["fck"], g["dokumler"], g["ilk"], g["son"], g["ortalama"], g["sinir"], g["durum"]], tarih_s=(5, 6),
               dolgu={9: YESIL if g["durum"] == "Uygun" else KIRMIZI if g["durum"] == "Uygun değil" else MAVI})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, gun: date, ek: Decimal = Decimal(4), eksi: Decimal = Decimal(4), grup: int = 3, fark: Decimal = Decimal(15), tolerans: int = 2) -> dict:
    kayit, uy = oku(yol)
    r = degerlendir(kayit, gun, ek, eksi, grup, fark, tolerans)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, gun, ek, eksi)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Beton döküm ve 7/28 günlük numune kırım sonuçlarını takip eder, sınıf dayanımının altında kalanları işaretler.")
    p.add_argument("--numuneler", type=Path, help="Döküm No, Döküm Tarihi, Yer, Beton Sınıfı (C30/37), Numune Tipi (Küp / Silindir), Numune No, Kırım Yaşı, Dayanım (MPa)")
    p.add_argument("--tarih", help="Rapor tarihi GG.AA.YYYY (varsayılan bugün; örnek veride 15.10.2026)")
    p.add_argument("--ek", default="4", help="Grup ortalaması için fck üzerine eklenecek değer, MPa (varsayılan 4)")
    p.add_argument("--eksi", default="4", help="Bireysel sonuç için fck'den düşülecek değer, MPa (varsayılan 4)")
    p.add_argument("--grup", type=int, default=3, help="Grup kriterindeki ardışık sonuç sayısı (varsayılan 3)")
    p.add_argument("--fark", default="15", help="Numuneler arası fark bu yüzdeyi aşarsa bilgi verilir (varsayılan 15)")
    p.add_argument("--tolerans", type=int, default=2, help="28 günlük sonucun beklenmesi için ek gün (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "beton_kirim_takibi.xlsx")
    a = p.parse_args(argv)
    yol = a.numuneler or ORNEK / "beton_numuneleri.csv"
    gun = tarih(a.tarih) if a.tarih else (date(2026, 10, 15) if a.numuneler is None else date.today())
    ek, eksi, fark = para(a.ek), para(a.eksi), para(a.fark)
    if gun is None or None in (ek, eksi, fark) or ek < 0 or eksi < 0 or fark < 0 or a.grup < 1 or a.tolerans < 0:
        print("[X] --tarih GG.AA.YYYY, --ek / --eksi / --fark sıfır ya da üstü, --grup en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, gun, ek, eksi, a.grup, fark, a.tolerans)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    y28 = [s for s in r["sonuclar"] if s["yas"] == 28]
    print(f"[OK] {r['dokum_sayisi']} döküm · {len(y28)} adet 28 günlük sonuç · bireysel uygun değil {sum(1 for s in y28 if s['bireysel'] == 'Uygun değil')}"
          f" · grup uygun değil {sum(1 for g in r['gruplar'] if g['durum'] == 'Uygun değil')} · 28 günlük sonucu olmayan {r['eksik']}")
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
