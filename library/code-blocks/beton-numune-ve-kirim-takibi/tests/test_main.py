import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
GUN = date(2026, 10, 15)


def n(dokum, gun, sinif="C30/37", tip="Küp", yas=None, dayanim=None, ay=9):
    s = main.sinif_oku(sinif)
    return {"dokum": dokum, "tarih": date(2026, ay, gun), "yer": "Y", "sinif": sinif, "fck": s[0] if tip == "Silindir" else s[1], "tip": tip, "numune": "",
            "yas": yas, "dayanim": None if dayanim is None else D(dayanim)}


def veri():
    return [n("D1", 1, yas=28, dayanim=42), n("D1", 1, yas=28, dayanim=44), n("D1", 1, yas=28, dayanim=43), n("D1", 1, yas=7, dayanim=30), n("D1", 1, yas=7, dayanim=31),
            n("D2", 5, yas=28, dayanim=38), n("D2", 5, yas=28, dayanim=40), n("D3", 10, yas=28, dayanim=33), n("D3", 10, yas=28, dayanim=35),
            n("D4", 20, "C25/30", "Silindir", 28, 20), n("D5", 1)]


class DegerlendirTesti(unittest.TestCase):
    def test_kriterler(self):
        r = main.degerlendir(veri(), GUN)
        self.assertEqual([(s["dokum"], s["yas"], s["fck"], s["adet"], s["ortalama"], s["oran"], s["fark"], s["bireysel"]) for s in r["sonuclar"]],
                         [("D1", 7, D(37), 2, D("30.5"), D("82.43"), D("3.28"), None), ("D1", 28, D(37), 3, D("43.0"), D("116.22"), D("4.65"), "Uygun"),
                          ("D2", 28, D(37), 2, D("39.0"), D("105.41"), D("5.13"), "Uygun"), ("D3", 28, D(37), 2, D("34.0"), D("91.89"), D("5.88"), "Uygun"),
                          ("D4", 28, D(25), 1, D("20.0"), D("80.00"), None, "Uygun değil")])
        self.assertEqual([(g["sinif"], g["dokumler"], g["ortalama"], g["sinir"], g["durum"]) for g in r["gruplar"]],
                         [("C25/30", "D4", D("20.0"), D(29), "Eksik grup (1/3)"), ("C30/37", "D1, D2, D3", D("38.7"), D(41), "Uygun değil")])
        self.assertEqual((r["dokum_sayisi"], r["eksik"]), (5, 1))
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Orta", "28 günlük sonuç eksik", "D5"), ("Yüksek", "Bireysel sonuç sınırın altında", "D4"), ("Yüksek", "Grup ortalaması yetersiz", "C30/37 küp")])

    def test_secenekler(self):
        r = main.degerlendir(veri(), date(2026, 9, 30), ek=D(1), eksi=D(5), grup=2, fark=D(5))
        self.assertEqual([(g["dokumler"], g["durum"]) for g in r["gruplar"]], [("D4", "Eksik grup (1/2)"), ("D1, D2", "Uygun"), ("D3", "Eksik grup (1/2)")])
        self.assertEqual(sorted(u["kim"] for u in r["uyarilar"] if u["tur"] == "Numuneler arası fark yüksek"), ["D2 · 28 gün", "D3 · 28 gün"])
        self.assertEqual([main.sinif_oku(x) for x in ("C30/37", "c 25 / 30", "C30", "")], [(D(30), D(37)), (D(25), D(30)), None, None])
        with self.assertRaises(ValueError):
            main.degerlendir([], GUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "beton_numuneleri.csv", cikti, GUN)
            self.assertEqual([(g["sinif"], g["ortalama"], g["durum"]) for g in r["gruplar"]], [("C16/20", D("20.1"), "Eksik grup (1/3)"), ("C30/37", D("43.6"), "Uygun"), ("C30/37", D("37.0"), "Uygun değil")])
            self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("Grup ortalaması yetersiz", "C30/37 küp"), ("28 günlük sonuç eksik", "BD-108"), ("Numuneler arası fark yüksek", "BD-106 · 28 gün")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Kırım Sonuçları", "Grup Değerlendirmesi", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "b.csv"
            y.write_text("Döküm No;Tarih;Beton Sınıfı;Numune Tipi;Yaş;Basınç Dayanımı\nA;01.09.2026;C30/37;Silindir;28;35,5\nA;01.09.2026;C30/37;Silindir;;\nB;x;C30/37;;;\nC;01.09.2026;X;;;\n"
                         "D;01.09.2026;C30/37;;28;\nE;01.09.2026;C30/37;;0;30\n;01.09.2026;C30/37;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["dokum"], x["fck"], x["tip"], x["yas"], x["dayanim"]) for x in l], len(uy)), ([("A", D(30), "Silindir", 28, D("35.5")), ("A", D(30), "Silindir", None, None)], 4))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--numuneler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "15.10.2026", "--ek", "1,48", "--eksi", "4", "--grup", "15", "--fark", "10", "--tolerans", "0"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--ek", "-1"], ["--grup", "0"])], [2] * 3)
            self.assertEqual(main.main(["--numuneler", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Döküm No;Tarih\nA;01.09.2026\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
