"""
Fazla Mesai Hesaplama — Workers / Workless kod bloğu
İnsan Kaynakları › Bordro ve Özlük İşleri Uzmanı

Günlük giriş-çıkış kayıtlarından (PDKS dökümü) 4857 sayılı İş Kanunu'na göre fazla çalışma ücretini hesaplar:
  - Günlük çalışma = çıkış − giriş − ara dinlenmesi; mola verilmemişse md. 68'deki asgari süre düşülür
    (4 saate kadar 15 dk, 7,5 saate kadar 30 dk, daha fazlası 1 saat). Gece yarısını geçen vardiyalar desteklenir.
  - Haftalık (Pazartesi–Pazar) toplam 45 saati aşan süre fazla çalışma: saat ücreti × 1,5 (md. 41).
    Sözleşmede haftalık süre 45'in altındaysa, sözleşme süresi ile 45 saat arası fazla sürelerle çalışma: × 1,25.
  - Denkleştirme (md. 63) seçilirse fazla çalışma, N haftalık dönemin ortalaması üzerinden hesaplanır.
  - Genel tatilde çalışılan gün için ayrıca bir günlük ücret (md. 47); arife günü 13.00'ten sonra çalışma yarım gün.
  - Serbest zaman karşılığı (fazla çalışmanın her saati için 1 saat 30 dk, fazla sürelerle çalışmanın 1 saat 15 dk),
    günlük 11 saat ve yıllık 270 saat sınırı kontrolü.
İnternete bağlanmaz.

Kullanım:
    python main.py                                         # örnek PDKS dökümüyle dener
    python main.py --kayitlar pdks.xlsx --personel personel.xlsx
    python main.py --kayitlar pdks.xlsx --personel personel.xlsx --denklestirme 8 --ubgt-haric
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import puantaj_cekirdek as pc

BURASI = Path(__file__).resolve().parent
SIFIR = Decimal(0)
HAFTALIK = Decimal(45)
GUNLUK_AZAMI = Decimal(11)
AYLIK_SAAT = Decimal(225)          # 30 gün × 7,5 saat: aylık ücretten saat ücretine geçiş
FM_ZAM, FSC_ZAM = Decimal("1.5"), Decimal("1.25")


def katla(s) -> str:
    return pc.katla(s)


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def saat_coz(x) -> time | None:
    if isinstance(x, datetime):
        return x.time()
    if isinstance(x, time):
        return x
    if isinstance(x, float) and 0 <= x < 1:                  # Excel saat kesri
        dk = round(x * 24 * 60)
        return time(dk // 60 % 24, dk % 60)
    m = re.match(r"^\s*(\d{1,2})[:.](\d{2})", str(x or ""))
    return time(int(m.group(1)) % 24, int(m.group(2))) if m else None


def ara_dinlenme(brut_saat: Decimal) -> Decimal:
    """4857 md. 68 asgari ara dinlenmesi (saat)."""
    if brut_saat <= 0:
        return SIFIR
    if brut_saat <= 4:
        return Decimal("0.25")
    if brut_saat <= Decimal("7.5"):
        return Decimal("0.5")
    return Decimal(1)


def s2(x: Decimal) -> Decimal:
    return x.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# ----------------------------------------------------------------------------
# Okuma
# ----------------------------------------------------------------------------

def kayitlari_oku(yol: Path) -> list[dict]:
    s = pc.tablo_oku(yol)
    b = [katla(x) for x in s[0]]
    bul = lambda *a: next((b.index(katla(x)) for x in a if katla(x) in b), None)  # noqa: E731
    i = {"sicil": bul("sicil", "sicil no", "personel no"), "ad": bul("ad soyad", "adı soyadı", "personel", "çalışan"),
         "tarih": bul("tarih", "gün"), "giris": bul("giriş", "giriş saati", "başlangıç"), "cikis": bul("çıkış", "çıkış saati", "bitiş"),
         "mola": bul("mola dk", "mola", "ara dinlenme dk", "ara dinlenme"), "saat": bul("çalışma saati", "net çalışma", "saat")}
    if i["tarih"] is None or (i["sicil"] is None and i["ad"] is None):
        raise SystemExit(f"Kayıt dosyasında Sicil (veya Ad Soyad) ve Tarih sütunları gerekli. Başlıklar: {s[0]}")
    if i["saat"] is None and (i["giris"] is None or i["cikis"] is None):
        raise SystemExit("Kayıt dosyasında Giriş ve Çıkış saatleri ya da Çalışma Saati sütunu gerekli.")
    al = lambda r, k: r[i[k]] if i[k] is not None and i[k] < len(r) else None  # noqa: E731
    kayitlar = []
    for n, r in enumerate(s[1:], 2):
        t = pc.tarih(al(r, "tarih"))
        sicil = str(al(r, "sicil") or al(r, "ad") or "").strip()
        if not t or not sicil:
            continue
        k = {"satir": n, "sicil": sicil, "ad": str(al(r, "ad") or sicil).strip(), "tarih": t, "giris": None, "cikis": None,
             "mola": None, "brut": None, "net": None, "not": ""}
        if i["giris"] is not None and al(r, "giris") not in (None, ""):
            g, c = saat_coz(al(r, "giris")), saat_coz(al(r, "cikis"))
            if not g or not c:
                k["not"] = "giriş/çıkış okunamadı"
                kayitlar.append(k)
                continue
            bas = datetime.combine(t, g)
            son = datetime.combine(t, c)
            if son <= bas:
                son += timedelta(days=1)                       # gece yarısını geçen vardiya
            k["giris"], k["cikis"] = bas, son
            k["brut"] = Decimal((son - bas).seconds) / 3600
            mola = sayi(al(r, "mola"))
            if mola is None:
                k["mola"] = ara_dinlenme(k["brut"])
                k["not"] = "mola verilmedi: md. 68 asgari ara dinlenmesi düşüldü"
            else:
                k["mola"] = mola / 60
                yasal = ara_dinlenme(k["brut"] - k["mola"])
                if k["mola"] < yasal:
                    k["not"] = f"mola ({mola:g} dk) md. 68 asgarisinin ({yasal * 60:g} dk) altında"
            k["net"] = max(SIFIR, k["brut"] - k["mola"])
        else:
            k["net"] = sayi(al(r, "saat")) or SIFIR
        kayitlar.append(k)
    return kayitlar


def personel_oku(yol: Path | None) -> dict[str, dict]:
    if not yol:
        return {}
    s = pc.tablo_oku(yol)
    b = [katla(x) for x in s[0]]
    bul = lambda *a: next((b.index(katla(x)) for x in a if katla(x) in b), None)  # noqa: E731
    i_s = bul("sicil", "sicil no", "personel no", "ad soyad")
    i_ucret = bul("aylık brüt ücret", "brüt ücret", "aylık ücret", "ücret")
    i_saat_u = bul("saat ücreti", "brüt saat ücreti")
    i_soz = bul("haftalık sözleşme saati", "haftalık çalışma saati", "sözleşme saati")
    i_yfm = bul("yıl başından fm", "yıllık fm", "kümülatif fm", "yıl başından fazla mesai")
    if i_s is None:
        raise SystemExit(f"Personel dosyasında Sicil sütunu bulunamadı. Başlıklar: {s[0]}")
    al = lambda r, i: r[i] if i is not None and i < len(r) else None  # noqa: E731
    return {str(al(r, i_s)).strip(): {"aylik": sayi(al(r, i_ucret)), "saat_u": sayi(al(r, i_saat_u)),
                                       "soz": sayi(al(r, i_soz)) or HAFTALIK, "yil_fm": sayi(al(r, i_yfm)) or SIFIR}
            for r in s[1:] if al(r, i_s)}


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

def gece_payi(k: dict) -> Decimal:
    """Vardiyanın 20.00–06.00 arasına düşen kısmı (saat)."""
    if not k["giris"] or not k["cikis"]:
        return SIFIR
    toplam, t = 0, k["giris"]
    while t < k["cikis"]:
        if t.hour >= 20 or t.hour < 6:
            toplam += 1
        t += timedelta(minutes=1)
    return Decimal(toplam) / 60


def pazartesi(d: date) -> date:
    return d - timedelta(days=d.weekday())


def ubgt_payi(k: dict, tam: dict, yarim: dict) -> Decimal:
    """Genel tatilde çalışma: tam gün 1, arifede 13.00'ten sonra çalışma 0,5 gün."""
    if k["tarih"] in tam and k["net"]:
        return Decimal(1)
    if k["tarih"] in yarim and k["cikis"] and k["cikis"] > datetime.combine(k["tarih"], time(13, 0)):
        return Decimal("0.5")
    return SIFIR


def ubgt_saati(k: dict, tam: dict, yarim: dict) -> Decimal:
    if k["tarih"] in tam:
        return k["net"] or SIFIR
    if k["tarih"] in yarim and k["giris"] and k["cikis"]:
        on3 = datetime.combine(k["tarih"], time(13, 0))
        if k["cikis"] > on3:
            return Decimal((k["cikis"] - max(k["giris"], on3)).seconds) / 3600
    return SIFIR


def calistir(kayit_yolu: Path, cikti: Path, personel_yolu: Path | None = None, denklestirme: int = 0,
             ubgt_haric: bool = False, yillik_sinir: float = 270.0) -> dict:
    kayitlar = kayitlari_oku(kayit_yolu)
    personel = personel_oku(personel_yolu)
    yillar = {k["tarih"].year for k in kayitlar}
    tam, yarim = {}, {}
    for y in yillar:
        t, h = pc.tatiller(y)
        tam.update(t)
        yarim.update(h)
    uyarilar = [f"{y} yılı dini bayram tarihleri tanımlı değil." for y in sorted(yillar) if y not in pc.DINI_BAYRAMLAR]
    gruplar: dict[str, list[dict]] = defaultdict(list)
    for k in kayitlar:
        gruplar[k["sicil"]].append(k)
    haftalar, ozet, gunluk_ihlal = [], [], []
    for sicil, ks in sorted(gruplar.items()):
        p = personel.get(sicil, {})
        soz = p.get("soz", HAFTALIK)
        saat_u = p.get("saat_u") or (p["aylik"] / AYLIK_SAAT if p.get("aylik") else None)
        gunluk_u = (p["aylik"] / 30) if p.get("aylik") else (saat_u * Decimal("7.5") if saat_u else None)
        hafta: dict[date, dict] = {}
        for k in sorted(ks, key=lambda x: x["tarih"]):
            if k["net"] is None:
                continue
            u_gun, u_saat = ubgt_payi(k, tam, yarim), ubgt_saati(k, tam, yarim)
            h = hafta.setdefault(pazartesi(k["tarih"]), {"sicil": sicil, "ad": k["ad"], "pzt": pazartesi(k["tarih"]),
                                                        "toplam": SIFIR, "ubgt_saat": SIFIR, "ubgt_gun": SIFIR, "gun": 0})
            h["toplam"] += k["net"]
            h["ubgt_saat"] += u_saat
            h["ubgt_gun"] += u_gun
            h["gun"] += 1 if k["net"] else 0
            if k["net"] > GUNLUK_AZAMI:
                gunluk_ihlal.append({"sicil": sicil, "ad": k["ad"], "tarih": k["tarih"], "saat": k["net"], "tur": "11 saat"})
            elif k["brut"] and k["net"] > Decimal("7.5") and gece_payi(k) * 2 >= k["brut"]:
                gunluk_ihlal.append({"sicil": sicil, "ad": k["ad"], "tarih": k["tarih"], "saat": k["net"], "tur": "gece 7,5 saat"})
        sirali = [hafta[w] for w in sorted(hafta)]
        for h in sirali:
            h["sayilan"] = h["toplam"] - (h["ubgt_saat"] if ubgt_haric else SIFIR)
        if denklestirme and denklestirme > 1:
            # N haftalık bloklar: blok toplamı N × 45'i aşan kısım fazla çalışma; bloğun son haftasına yazılır
            for i in range(0, len(sirali), denklestirme):
                blok = sirali[i:i + denklestirme]
                n = Decimal(len(blok))
                top = sum((h["sayilan"] for h in blok), SIFIR)
                for h in blok:
                    h["fm"] = h["fsc"] = SIFIR
                blok[-1]["fm"] = max(SIFIR, top - HAFTALIK * n)
                blok[-1]["fsc"] = max(SIFIR, min(top, HAFTALIK * n) - soz * n) if soz < HAFTALIK else SIFIR
        else:
            for h in sirali:
                h["fm"] = max(SIFIR, h["sayilan"] - HAFTALIK)
                h["fsc"] = max(SIFIR, min(h["sayilan"], HAFTALIK) - soz) if soz < HAFTALIK else SIFIR
        kum = p.get("yil_fm", SIFIR)
        for h in sirali:
            kum += h["fm"]
            h["kumulatif_fm"] = kum
            h["fm_ucret"] = s2(h["fm"] * saat_u * FM_ZAM) if saat_u else None
            h["fsc_ucret"] = s2(h["fsc"] * saat_u * FSC_ZAM) if saat_u else None
            h["ubgt_ucret"] = s2(h["ubgt_gun"] * gunluk_u) if gunluk_u else None
            h["serbest"] = h["fm"] * FM_ZAM + h["fsc"] * FSC_ZAM
            haftalar.append(h)
        top = lambda a: sum((h[a] for h in sirali if h[a] is not None), SIFIR)  # noqa: E731
        o = {"sicil": sicil, "ad": sirali[0]["ad"] if sirali else sicil, "saat_u": saat_u, "soz": soz, "toplam": top("toplam"),
             "fm": top("fm"), "fsc": top("fsc"), "ubgt_gun": top("ubgt_gun"), "fm_ucret": top("fm_ucret") if saat_u else None,
             "fsc_ucret": top("fsc_ucret") if saat_u else None, "ubgt_ucret": top("ubgt_ucret") if gunluk_u else None,
             "serbest": top("serbest"), "yil_fm": kum, "yillik_asim": kum > Decimal(str(yillik_sinir))}
        o["toplam_ucret"] = sum((o[a] for a in ("fm_ucret", "fsc_ucret", "ubgt_ucret") if o[a] is not None), SIFIR) if saat_u else None
        if not saat_u:
            uyarilar.append(f"{sicil}: ücret bilgisi yok, yalnız saatler hesaplandı")
        ozet.append(o)
    _rapor(ozet, haftalar, kayitlar, gunluk_ihlal, uyarilar, denklestirme, ubgt_haric, yillik_sinir, cikti)
    return {"ozet": ozet, "haftalar": haftalar, "kayitlar": kayitlar, "gunluk_ihlal": gunluk_ihlal, "uyarilar": uyarilar}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="FDE2E1")
PARA = "#,##0.00"
SAAT = "#,##0.00"


def _baslik(ws, satir=1):
    for c in ws[satir]:
        c.fill, c.font = BASLIK_DOLGU, BASLIK_YAZI


def _f(x):
    return None if x is None else float(x)


def _rapor(ozet, haftalar, kayitlar, gunluk_ihlal, uyarilar, denklestirme, ubgt_haric, yillik_sinir, cikti):
    wb = Workbook()
    o = wb.active
    o.title = "Personel Özeti"
    o.append(["Sicil", "Ad Soyad", "Haftalık Sözleşme Saati", "Saat Ücreti", "Toplam Çalışma (saat)", "Fazla Çalışma (saat, ×1,5)",
              "Fazla Sürelerle Çalışma (saat, ×1,25)", "Genel Tatil Çalışması (gün)", "Fazla Çalışma Ücreti", "Fazla Sürelerle Ç. Ücreti",
              "Genel Tatil Ücreti", "Toplam Brüt Ek Ödeme", "Serbest Zaman Karşılığı (saat)", "Yıl Başından Fazla Çalışma",
              f"{yillik_sinir:g} Saat Aşımı"])
    _baslik(o)
    for x in ozet:
        o.append([x["sicil"], x["ad"], _f(x["soz"]), _f(x["saat_u"]), _f(x["toplam"]), _f(x["fm"]), _f(x["fsc"]), _f(x["ubgt_gun"]),
                  _f(x["fm_ucret"]), _f(x["fsc_ucret"]), _f(x["ubgt_ucret"]), _f(x["toplam_ucret"]), _f(x["serbest"]),
                  _f(x["yil_fm"]), "EVET" if x["yillik_asim"] else ""])
        for c in (4, 9, 10, 11, 12):
            o.cell(o.max_row, c).number_format = PARA
        for c in (5, 6, 7, 13, 14):
            o.cell(o.max_row, c).number_format = SAAT
        if x["yillik_asim"]:
            o.cell(o.max_row, 15).fill = KIRMIZI
    for j, w in enumerate((10, 22, 11, 11, 12, 13, 15, 12, 14, 14, 13, 15, 14, 13, 10), 1):
        o.column_dimensions[get_column_letter(j)].width = w
    o.freeze_panes = "C2"

    h = wb.create_sheet("Haftalık")
    h.append(["Sicil", "Ad Soyad", "Hafta (Pazartesi)", "Çalışılan Gün", "Toplam Saat", "Genel Tatil Saati", "Sayılan Saat",
              "Fazla Çalışma", "Fazla Sürelerle Ç.", "Genel Tatil (gün)", "FÇ Ücreti", "FSÇ Ücreti", "Genel Tatil Ücreti",
              "Kümülatif FÇ (yıl)"])
    _baslik(h)
    for x in haftalar:
        h.append([x["sicil"], x["ad"], x["pzt"], x["gun"], _f(x["toplam"]), _f(x["ubgt_saat"]) or None, _f(x["sayilan"]), _f(x["fm"]),
                  _f(x["fsc"]), _f(x["ubgt_gun"]) or None, _f(x["fm_ucret"]), _f(x["fsc_ucret"]), _f(x["ubgt_ucret"]) or None,
                  _f(x["kumulatif_fm"])])
        h.cell(h.max_row, 3).number_format = "DD.MM.YYYY"
        for c in (11, 12, 13):
            h.cell(h.max_row, c).number_format = PARA
        for c in (5, 6, 7, 8, 9, 14):
            h.cell(h.max_row, c).number_format = SAAT
        if x["gun"] >= 7:
            h.cell(h.max_row, 4).fill = KIRMIZI
    for j, w in enumerate((10, 22, 13, 9, 10, 11, 10, 11, 12, 10, 12, 12, 12, 12), 1):
        h.column_dimensions[get_column_letter(j)].width = w
    h.freeze_panes = "D2"
    h.auto_filter.ref = h.dimensions

    g = wb.create_sheet("Günlük Kayıtlar")
    g.append(["Sicil", "Ad Soyad", "Tarih", "Giriş", "Çıkış", "Brüt Süre", "Ara Dinlenme", "Net Çalışma", "Not"])
    _baslik(g)
    for k in sorted(kayitlar, key=lambda x: (x["sicil"], x["tarih"])):
        g.append([k["sicil"], k["ad"], k["tarih"], k["giris"].strftime("%H:%M") if k["giris"] else None,
                  k["cikis"].strftime("%H:%M") if k["cikis"] else None, _f(k["brut"]), _f(k["mola"]), _f(k["net"]), k["not"]])
        g.cell(g.max_row, 3).number_format = "DD.MM.YYYY"
        for c in (6, 7, 8):
            g.cell(g.max_row, c).number_format = SAAT
        if k["net"] is not None and k["net"] > GUNLUK_AZAMI:
            g.cell(g.max_row, 8).fill = KIRMIZI
    for j, w in enumerate((10, 22, 12, 8, 8, 10, 11, 11, 55), 1):
        g.column_dimensions[get_column_letter(j)].width = w
    g.freeze_panes = "D2"
    g.auto_filter.ref = g.dimensions

    b = wb.create_sheet("Bilgi")
    satirlar = [["Saat ücreti", "aylık brüt ücret ÷ 225 (30 gün × 7,5 saat); personel dosyasında saat ücreti verilmişse o"],
                ["Fazla çalışma", "haftalık 45 saati aşan süre; saat ücreti × 1,5 (İş K. md. 41)"],
                ["Fazla sürelerle çalışma", "sözleşmedeki haftalık süre 45'in altındaysa, o süre ile 45 saat arası; × 1,25"],
                ["Denkleştirme", f"{denklestirme} haftalık bloklar, ortalama 45 saat (md. 63)" if denklestirme and denklestirme > 1 else "uygulanmadı"],
                ["Genel tatil", "çalışılan her genel tatil günü için ayrıca bir günlük ücret (aylık ÷ 30) (md. 47); arifede 13.00 "
                                "sonrası çalışma yarım gün. Genel tatil saatleri haftalık toplamdan "
                                + ("ÇIKARILDI (--ubgt-haric)" if ubgt_haric else "çıkarılmadı (varsayılan)")],
                ["Ara dinlenmesi", "mola verilmemişse md. 68 asgarisi düşülür: ≤ 4 saat 15 dk, ≤ 7,5 saat 30 dk, > 7,5 saat 1 saat"],
                ["Serbest zaman", "işçi isterse ücret yerine: fazla çalışmanın her saati için 1 saat 30 dk, fazla sürelerle "
                                  "çalışmanın her saati için 1 saat 15 dk, 6 ay içinde (md. 41)"],
                ["Sınırlar", f"günlük 11 saat (md. 63), gece çalışması 7,5 saat (md. 69), yıllık {yillik_sinir:g} saat fazla çalışma "
                             "(md. 41); 7 gün çalışılan haftalar kırmızı"],
                ["Vardiya", "gece yarısını geçen vardiya, başladığı günün kaydı sayılır (hafta ve genel tatil ataması bu güne göre)"],
                ["Not", "Genel tatil saatlerinin haftalık 45 saat hesabına dahil edilip edilmeyeceği uygulamada tartışmalıdır; "
                        "işyeri uygulamanıza ve hukuk danışmanınızın görüşüne göre --ubgt-haric seçeneğini belirleyin. "
                        "Hesaplanan tutarlar brüttür; SGK primi, gelir ve damga vergisi bordroda hesaplanır."]]
    for i in gunluk_ihlal:
        satirlar.append([f"{i['tur']} aşımı", f"{i['sicil']} {i['ad']} · {i['tarih']:%d.%m.%Y} · {float(i['saat']):.2f} saat"
                         + (" (vardiyanın yarısından fazlası 20.00–06.00 arasında; md. 69: gece çalışması 7,5 saati geçemez, "
                            "kanundaki istisnalar hariç)" if i["tur"].startswith("gece") else "")])
    for u in uyarilar:
        satirlar.append(["Uyarı", u])
    for s in satirlar:
        b.append(s)
    b.column_dimensions["A"].width = 24
    b.column_dimensions["B"].width = 130
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main(argv: list[str] | None = None) -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ornek = BURASI / "ornek_veri"
    ap = argparse.ArgumentParser(description="PDKS giriş-çıkış kayıtlarından 4857 sayılı Kanun'a göre fazla çalışma ücretini hesaplar.")
    ap.add_argument("--kayitlar", type=Path, default=ornek / "pdks.csv",
                    help="Günlük kayıtlar (.xlsx/.csv): Sicil, Ad Soyad, Tarih, Giriş, Çıkış [, Mola (dk)] — veya Çalışma Saati")
    ap.add_argument("--personel", type=Path, help="Personel: Sicil, Aylık Brüt Ücret [, Haftalık Sözleşme Saati, Yıl Başından FM]")
    ap.add_argument("--denklestirme", type=int, default=0, help="Denkleştirme dönemi, hafta (ör. 8 ≈ 2 ay); 0 = uygulanmaz")
    ap.add_argument("--ubgt-haric", action="store_true", help="Genel tatil saatlerini haftalık 45 saat hesabından çıkar")
    ap.add_argument("--yillik-sinir", type=float, default=270.0, help="Yıllık fazla çalışma sınırı (varsayılan 270)")
    ap.add_argument("--cikti", type=Path, default=Path("cikti") / "fazla_mesai.xlsx")
    a = ap.parse_args(argv)
    if a.kayitlar == ornek / "pdks.csv" and not a.personel:
        a.personel = ornek / "personel.csv"
    s = calistir(a.kayitlar, a.cikti, a.personel, a.denklestirme, a.ubgt_haric, a.yillik_sinir)
    for o in s["ozet"]:
        ucret = f" · ek brüt {o['toplam_ucret']:,.2f} TL".replace(",", "X").replace(".", ",").replace("X", ".") if o["toplam_ucret"] is not None else ""
        print(f"[OK] {o['sicil']} {o['ad']}: FÇ {o['fm']:g} sa · FSÇ {o['fsc']:g} sa · genel tatil {o['ubgt_gun']:g} gün{ucret}"
              + (" · YILLIK SINIR AŞILDI" if o["yillik_asim"] else ""))
    for i in s["gunluk_ihlal"]:
        print(f"[!] {i['tur']} aşımı: {i['sicil']} {i['tarih']:%d.%m.%Y} {float(i['saat']):.2f} saat")
    for u in s["uyarilar"]:
        print(f"[!] {u}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")


if __name__ == "__main__":
    sys.exit(main())
