"""
Standart-Fiili Maliyet Sapma Analizi — Workers / Workless kod bloğu
Muhasebe › Maliyet Muhasebesi Uzmanı

Standart maliyet ile fiili maliyet arasındaki farkı klasik sapma bileşenlerine ayırır (+ aleyhte, − lehte):
  - Direkt ilk madde ve malzeme (711 / 712):
      fiyat sapması  = (fiili fiyat − standart fiyat) × fiili miktar
      miktar sapması = (fiili miktar − standart miktar) × standart fiyat
    Standart miktar = Σ ürün (fiili üretim × birim standart miktar).
  - Direkt işçilik (721 / 722):
      ücret sapması = (fiili ücret − standart ücret) × fiili saat
      süre sapması  = (fiili saat − standart saat) × standart ücret
  - Genel üretim giderleri, üç sapma yöntemi, iş merkezi bazında (731 / 732 / 733):
      standart GÜG oranı = değişken oran + bütçelenen sabit GÜG ÷ normal kapasite
      esnek bütçe (fiili saat) = bütçelenen sabit GÜG + fiili saat × değişken oran
      bütçe sapması      = fiili GÜG − esnek bütçe (fiili saat)
      kapasite sapması   = esnek bütçe (fiili saat) − fiili saat × standart GÜG oranı
      verimlilik sapması = (fiili saat − standart saat) × standart GÜG oranı
    Yük tabanı: iş merkezinin direkt işçilik saati.
  - Toplam sapma = fiili maliyet − standart maliyet (fiili üretim için); bileşenlerin toplamına eşittir.
Rapor: sapma özeti (hesap kodlarıyla), malzeme, işçilik, GÜG, ürün standart maliyet kartı, uyarılar.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 2 ürün, 4 malzeme, 2 iş merkezi
    python main.py --standartlar s.xlsx --uretim u.xlsx --fiili f.xlsx --gug gug.xlsx --esik 5
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)
HESAPLAR = {"fiyat": ("711", "Direkt İlk Madde ve Malzeme Fiyat Farkı"), "miktar": ("712", "Direkt İlk Madde ve Malzeme Miktar Farkı"),
            "ucret": ("721", "Direkt İşçilik Ücret Farkları"), "sure": ("722", "Direkt İşçilik Süre Farkları"),
            "butce": ("731", "Genel Üretim Giderleri Bütçe Farkları"), "verimlilik": ("732", "Genel Üretim Giderleri Verimlilik Farkları"),
            "kapasite": ("733", "Genel Üretim Giderleri Kapasite Farkları")}

STD_SUTUNLARI = {"urun": ("urun", "mamul", "urun kodu"), "tur": ("tur", "kalem turu", "maliyet turu"), "kalem": ("kalem", "malzeme", "is merkezi"),
                 "birim": ("birim",), "miktar": ("standart miktar", "birim miktar", "miktar"), "fiyat": ("standart fiyat", "standart ucret", "fiyat")}
URETIM_SUTUNLARI = {"urun": ("urun", "mamul", "urun kodu"), "miktar": ("uretim miktari", "fiili uretim", "miktar")}
FIILI_SUTUNLARI = {"tur": ("tur", "kalem turu", "maliyet turu"), "kalem": ("kalem", "malzeme", "is merkezi"), "miktar": ("fiili miktar", "miktar", "saat"),
                   "tutar": ("fiili tutar", "tutar")}
GUG_SUTUNLARI = {"merkez": ("is merkezi", "gider yeri", "kalem"), "degisken": ("degisken oran tl saat", "degisken oran"),
                 "sabit": ("butcelenen sabit gug", "sabit gug", "butce sabit"), "kapasite": ("normal kapasite saat", "normal kapasite", "kapasite")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def sayi(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tur_bul(x) -> str:
    k = katla(x)
    if k.startswith(("malz", "ilk madde", "dimm", "hammadde")):
        return "malzeme"
    if k.startswith(("isci", "direkt isci", "dis", "emek")):
        return "iscilik"
    if k.startswith(("gug", "genel uretim")):
        return "gug"
    return ""


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


# ----------------------------------------------------------------------------
# Hesap
# ----------------------------------------------------------------------------

@dataclass
class Sapma:
    tur: str                  # malzeme | iscilik
    kalem: str
    birim: str
    sm: Decimal               # standart miktar / saat (fiili üretim için)
    sf: Decimal               # standart fiyat / ücret
    fm: Decimal | None
    ft: Decimal | None        # fiili tutar

    @property
    def ff(self) -> Decimal | None:
        return self.ft / self.fm if self.ft is not None and self.fm else None

    @property
    def standart(self) -> Decimal:
        return self.sm * self.sf

    @property
    def fiyat(self) -> Decimal:
        return (self.ft - self.fm * self.sf) if self.ft is not None and self.fm is not None else SIFIR

    @property
    def miktar(self) -> Decimal:
        return ((self.fm - self.sm) * self.sf) if self.fm is not None else -self.standart

    @property
    def toplam(self) -> Decimal:
        return (self.ft or SIFIR) - self.standart


@dataclass
class GugSapma:
    merkez: str
    degisken: Decimal
    sabit: Decimal
    kapasite: Decimal
    ss: Decimal
    fs: Decimal
    fiili: Decimal

    @property
    def oran(self) -> Decimal:
        return self.degisken + (self.sabit / self.kapasite if self.kapasite else SIFIR)

    @property
    def esnek(self) -> Decimal:
        return self.sabit + self.fs * self.degisken

    @property
    def standart(self) -> Decimal:
        return self.ss * self.oran

    @property
    def butce(self) -> Decimal:
        return self.fiili - self.esnek

    @property
    def kapasite_sapmasi(self) -> Decimal:
        return self.esnek - self.fs * self.oran

    @property
    def verimlilik(self) -> Decimal:
        return (self.fs - self.ss) * self.oran

    @property
    def toplam(self) -> Decimal:
        return self.fiili - self.standart


def analiz_et(std: list[dict], uretim: dict[str, Decimal], fiili: list[dict], gug: list[dict], esik: Decimal) -> dict:
    uy = []
    # Standart miktar ve fiyat (fiili üretim için)
    sm = defaultdict(lambda: SIFIR)
    sdeger = defaultdict(lambda: SIFIR)
    birim, urun_kart = {}, defaultdict(list)
    for r in std:
        tur, kalem, u = tur_bul(r.get("tur")), metin(r.get("kalem")), metin(r.get("urun"))
        m, f = sayi(r.get("miktar")), sayi(r.get("fiyat"))
        if tur not in ("malzeme", "iscilik") or not kalem or m is None or f is None:
            uy.append({"onem": "Orta", "tur": "Okunamayan standart", "kalem": kalem, "aciklama": f"Standart satırı {r['_satir']} atlandı"})
            continue
        anahtar = (tur, katla(kalem))
        birim.setdefault(anahtar, (kalem, metin(r.get("birim"))))
        urun_kart[u].append((tur, kalem, metin(r.get("birim")), m, f))
        adet = uretim.get(katla(u))
        if adet is None:
            continue
        sm[anahtar] += adet * m
        sdeger[anahtar] += adet * m * f
    for u in urun_kart:
        if katla(u) not in uretim:
            uy.append({"onem": "Bilgi", "tur": "Üretim yok", "kalem": u, "aciklama": "Standart kartı var, dönemde üretim kaydı yok"})
    for u in uretim:
        if u not in {katla(x) for x in urun_kart}:
            uy.append({"onem": "Yüksek", "tur": "Standart kart yok", "kalem": u, "aciklama": "Üretilen ürünün standart maliyet kartı yok; sapmaya dahil edilmedi"})
    fi = {}
    gug_fiili = defaultdict(lambda: SIFIR)
    for r in fiili:
        tur, kalem = tur_bul(r.get("tur")), metin(r.get("kalem"))
        if tur == "gug":
            gug_fiili[katla(kalem)] += sayi(r.get("tutar")) or SIFIR
            continue
        if tur not in ("malzeme", "iscilik"):
            continue
        a = (tur, katla(kalem))
        m0, t0 = fi.get(a, (SIFIR, SIFIR))
        fi[a] = (m0 + (sayi(r.get("miktar")) or SIFIR), t0 + (sayi(r.get("tutar")) or SIFIR))
        birim.setdefault(a, (kalem, ""))
    sapmalar = []
    for a in sorted(set(sm) | set(fi), key=lambda x: (x[0], x[1])):
        ad, b = birim[a]
        s_m = sm.get(a, SIFIR)
        s_f = sdeger[a] / s_m if s_m else SIFIR
        fm, ft = fi.get(a, (None, None))
        x = Sapma(a[0], ad, b, s_m, s_f, fm, ft)
        sapmalar.append(x)
        if a not in fi:
            uy.append({"onem": "Yüksek", "tur": "Fiili kayıt yok", "kalem": ad, "aciklama": f"Standartta {s_m:g} {b} var, fiili tüketim / saat kaydı yok"})
        elif not s_m:
            uy.append({"onem": "Orta", "tur": "Standart dışı kalem", "kalem": ad, "aciklama": f"Standart kartta yok; fiili {tl(ft)} TL tamamı sapma"})
        elif x.standart and abs(x.toplam) / x.standart > esik:
            o = abs(x.toplam) / x.standart
            uy.append({"onem": "Yüksek" if o > esik * 3 else "Orta", "tur": "Sapma eşiği", "kalem": ad,
                       "aciklama": f"Toplam sapma {tl(x.toplam)} TL ({'aleyhte' if x.toplam > 0 else 'lehte'}, standardın %{str(round(o * 100, 1)).replace('.', ',')}'i): "
                       f"{'fiyat' if x.tur == 'malzeme' else 'ücret'} {tl(x.fiyat)}, {'miktar' if x.tur == 'malzeme' else 'süre'} {tl(x.miktar)}"})
    # GÜG
    gs = []
    for r in gug:
        m = metin(r.get("merkez"))
        d, s, k = sayi(r.get("degisken")) or SIFIR, sayi(r.get("sabit")) or SIFIR, sayi(r.get("kapasite")) or SIFIR
        a = ("iscilik", katla(m))
        if not k:
            uy.append({"onem": "Yüksek", "tur": "Normal kapasite yok", "kalem": m, "aciklama": "Sabit GÜG oranı hesaplanamadı"})
        if katla(m) not in gug_fiili:
            uy.append({"onem": "Orta", "tur": "Fiili GÜG yok", "kalem": m, "aciklama": "Bu iş merkezi için fiili GÜG tutarı verilmedi; 0 kabul edildi"})
        g = GugSapma(m, d, s, k, sm.get(a, SIFIR), (fi.get(a) or (SIFIR,))[0], gug_fiili.get(katla(m), SIFIR))
        gs.append(g)
        if g.standart and abs(g.toplam) / g.standart > esik:
            o = abs(g.toplam) / g.standart
            uy.append({"onem": "Yüksek" if o > esik * 3 else "Orta", "tur": "Sapma eşiği", "kalem": f"GÜG · {m}",
                       "aciklama": f"Toplam GÜG sapması {tl(g.toplam)} TL (standardın %{str(round(o * 100, 1)).replace('.', ',')}'i): bütçe {tl(g.butce)}, kapasite "
                       f"{tl(g.kapasite_sapmasi)}, verimlilik {tl(g.verimlilik)}"})
        if k and g.fs < k * Decimal("0.8"):
            uy.append({"onem": "Bilgi", "tur": "Düşük kapasite kullanımı", "kalem": m, "aciklama": f"Fiili saat {g.fs:g}, normal kapasite {k:g} "
                       f"(%{g.fs / k * 100:.0f}); sabit giderin kullanılmayan kısmı kapasite sapmasına düşer"})
    for k in gug_fiili:
        if k not in {katla(g.merkez) for g in gs}:
            uy.append({"onem": "Orta", "tur": "GÜG oranı yok", "kalem": k, "aciklama": "Fiili GÜG var ama iş merkezi GÜG tablosunda tanımlı değil"})
    ozet = {
        "fiyat": sum((x.fiyat for x in sapmalar if x.tur == "malzeme"), SIFIR), "miktar": sum((x.miktar for x in sapmalar if x.tur == "malzeme"), SIFIR),
        "ucret": sum((x.fiyat for x in sapmalar if x.tur == "iscilik"), SIFIR), "sure": sum((x.miktar for x in sapmalar if x.tur == "iscilik"), SIFIR),
        "butce": sum((g.butce for g in gs), SIFIR), "verimlilik": sum((g.verimlilik for g in gs), SIFIR),
        "kapasite": sum((g.kapasite_sapmasi for g in gs), SIFIR)}
    standart = sum((x.standart for x in sapmalar), SIFIR) + sum((g.standart for g in gs), SIFIR)
    fiili_t = sum((x.ft or SIFIR for x in sapmalar), SIFIR) + sum((g.fiili for g in gs), SIFIR)
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kalem"]))
    return {"sapmalar": sapmalar, "gug": gs, "ozet": ozet, "standart": standart, "fiili": fiili_t, "uyarilar": uy, "kart": urun_kart, "uretim": uretim}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
ALEYHTE, LEHTE = PatternFill("solid", fgColor="FDE2E1"), PatternFill("solid", fgColor="E3F4E1")
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0.00;[Red]-#,##0.00"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def _renk(h, v):
    if v:
        h.fill = ALEYHTE if v > 0 else LEHTE


def rapor_yaz(cikti: Path, s: dict) -> None:
    wb = Workbook()
    oz = wb.active
    oz.title = "Sapma Özeti"
    _baslik(oz, ["Hesap", "Sapma", "Tutar (TL)", "Yön"], (8, 46, 16, 10))
    for anahtar, (kod, ad) in HESAPLAR.items():
        v = s["ozet"][anahtar]
        oz.append([kod, ad, float(v), "Aleyhte" if v > 0 else "Lehte" if v < 0 else "—"])
        oz.cell(oz.max_row, 3).number_format = PF
        _renk(oz.cell(oz.max_row, 4), v)
    top = sum(s["ozet"].values(), SIFIR)
    for a, v in (("Toplam sapma", top), ("Standart maliyet (fiili üretim)", s["standart"]), ("Fiili maliyet", s["fiili"])):
        oz.append(["", a, float(v)])
        oz.cell(oz.max_row, 3).number_format = PF
        oz.cell(oz.max_row, 2).font = Font(bold=True)
    oz.append(["", "Kontrol: fiili − standart − toplam sapma", float(s["fiili"] - s["standart"] - top)])
    oz.cell(oz.max_row, 3).number_format = PF
    g = BarChart()
    g.type, g.title, g.height, g.width = "bar", "Sapmalar (+ aleyhte, − lehte)", 7, 16
    g.add_data(Reference(oz, min_col=3, min_row=1, max_row=8), titles_from_data=True)
    g.set_categories(Reference(oz, min_col=2, min_row=2, max_row=8))
    g.legend = None
    oz.add_chart(g, "F2")
    oz.append([])
    oz.append(["", "Pozitif tutar aleyhte (fiili > standart), negatif tutar lehtedir. Hesap kodları Tek Düzen Hesap Planı 7/A seçeneğine göredir."])

    for tur, baslik, fad, mad in (("malzeme", "Malzeme", "Fiyat Sapması (711)", "Miktar Sapması (712)"),
                                  ("iscilik", "İşçilik", "Ücret Sapması (721)", "Süre Sapması (722)")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, ["Kalem", "Birim", "Standart Miktar", "Standart Fiyat", "Standart Tutar", "Fiili Miktar", "Fiili Fiyat", "Fiili Tutar", fad, mad,
                     "Toplam Sapma", "Sapma %"], (20, 7, 13, 12, 15, 13, 12, 15, 15, 15, 15, 9))
        for x in s["sapmalar"]:
            if x.tur != tur:
                continue
            ws.append([x.kalem, x.birim, float(x.sm), float(x.sf), float(x.standart), None if x.fm is None else float(x.fm),
                       None if x.ff is None else float(x.ff), None if x.ft is None else float(x.ft), float(x.fiyat), float(x.miktar), float(x.toplam),
                       float(x.toplam / x.standart) if x.standart else None])
            r = ws.max_row
            for j in range(3, 12):
                ws.cell(r, j).number_format = PF if j >= 9 else "#,##0.00##"
            ws.cell(r, 12).number_format = "0.0%"
            for j in (9, 10, 11):
                _renk(ws.cell(r, j), x.fiyat if j == 9 else x.miktar if j == 10 else x.toplam)

    gw = wb.create_sheet("GÜG")
    _baslik(gw, ["İş Merkezi", "Değişken Oran", "Bütçelenen Sabit", "Normal Kapasite", "Standart GÜG Oranı", "Standart Saat", "Fiili Saat",
                 "Esnek Bütçe (fiili saat)", "Standart GÜG", "Fiili GÜG", "Bütçe (731)", "Verimlilik (732)", "Kapasite (733)", "Toplam"],
            (14, 11, 14, 11, 12, 11, 11, 15, 15, 15, 13, 13, 13, 13))
    for g in s["gug"]:
        gw.append([g.merkez, float(g.degisken), float(g.sabit), float(g.kapasite), float(g.oran), float(g.ss), float(g.fs), float(g.esnek), float(g.standart),
                   float(g.fiili), float(g.butce), float(g.verimlilik), float(g.kapasite_sapmasi), float(g.toplam)])
        r = gw.max_row
        for j in range(2, 15):
            gw.cell(r, j).number_format = PF
        for j, v in ((11, g.butce), (12, g.verimlilik), (13, g.kapasite_sapmasi), (14, g.toplam)):
            _renk(gw.cell(r, j), v)

    kt = wb.create_sheet("Standart Maliyet Kartı")
    _baslik(kt, ["Ürün", "Tür", "Kalem", "Birim", "Birim Standart Miktar", "Standart Fiyat", "Birim Standart Tutar", "Fiili Üretim"], (14, 10, 20, 7, 12, 12, 13, 11))
    oranlar = {katla(g.merkez): g.oran for g in s["gug"]}
    for u, satirlar in s["kart"].items():
        top_u = SIFIR
        for tur, kalem, b, m, f in satirlar:
            kt.append([u, "Malzeme" if tur == "malzeme" else "İşçilik", kalem, b, float(m), float(f), float(m * f), float(s["uretim"].get(katla(u), 0))])
            top_u += m * f
            if tur == "iscilik" and katla(kalem) in oranlar:
                kt.append([u, "GÜG", kalem, "saat", float(m), float(oranlar[katla(kalem)]), float(m * oranlar[katla(kalem)])])
                top_u += m * oranlar[katla(kalem)]
        kt.append([u, "Toplam", "", "", None, None, float(top_u)])
        kt.cell(kt.max_row, 7).font = Font(bold=True)
        for row in kt.iter_rows(min_row=2, min_col=5, max_col=7):
            for c in row:
                c.number_format = "#,##0.00##"

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kalem", "Açıklama", "Açıklama / Aksiyon"], (9, 22, 18, 90, 24))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kalem"], u["aciklama"], ""])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
        uy.cell(uy.max_row, 5).fill = PatternFill("solid", fgColor="FFF4CE")
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def calistir(std_yolu: Path, uretim_yolu: Path, fiili_yolu: Path, cikti: Path, gug_yolu: Path | None = None, esik: Decimal = Decimal("0.05")) -> dict:
    std = kayitlar(std_yolu, STD_SUTUNLARI, ("urun", "tur", "kalem", "miktar", "fiyat"))
    uretim = {}
    for r in kayitlar(uretim_yolu, URETIM_SUTUNLARI, ("urun", "miktar")):
        m = sayi(r.get("miktar"))
        if metin(r.get("urun")) and m is not None:
            uretim[katla(r["urun"])] = uretim.get(katla(r["urun"]), SIFIR) + m
    fiili = kayitlar(fiili_yolu, FIILI_SUTUNLARI, ("tur", "kalem", "tutar"))
    gug = kayitlar(gug_yolu, GUG_SUTUNLARI, ("merkez", "degisken")) if gug_yolu else []
    s = analiz_et(std, uretim, fiili, gug, esik)
    rapor_yaz(cikti, s)
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Standart ve fiili maliyet farkını fiyat / miktar, ücret / süre ve GÜG bütçe / verimlilik / kapasite sapmalarına ayırır.")
    p.add_argument("--standartlar", type=Path, default=ORNEK / "standartlar.csv", help="Ürün, Tür (Malzeme/İşçilik), Kalem, Birim, Standart Miktar, Standart Fiyat")
    p.add_argument("--uretim", type=Path, default=ORNEK / "uretim.csv", help="Ürün, Üretim Miktarı (dönem fiili üretimi)")
    p.add_argument("--fiili", type=Path, default=ORNEK / "fiili.csv", help="Tür (Malzeme/İşçilik/GÜG), Kalem, Fiili Miktar, Fiili Tutar")
    p.add_argument("--gug", type=Path, help="İş Merkezi, Değişken Oran (TL/saat), Bütçelenen Sabit GÜG, Normal Kapasite (saat)")
    p.add_argument("--esik", type=float, default=5, help="Kalem sapması standardın bu %%'sini aşarsa uyarı (varsayılan 5)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "maliyet_sapma_analizi.xlsx")
    a = p.parse_args(argv)
    gug = a.gug or (ORNEK / "gug.csv" if a.standartlar == ORNEK / "standartlar.csv" else None)
    for y in (a.standartlar, a.uretim, a.fiili, gug):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        s = calistir(a.standartlar, a.uretim, a.fiili, a.cikti, gug, Decimal(str(a.esik)) / 100)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    top = sum(s["ozet"].values(), SIFIR)
    print(f"[OK] Standart {tl(s['standart'])} TL · fiili {tl(s['fiili'])} TL · toplam sapma {tl(top)} TL ({'aleyhte' if top > 0 else 'lehte'})")
    for k, (kod, ad) in HESAPLAR.items():
        print(f"     {kod} {ad}: {tl(s['ozet'][k])}")
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kalem']} {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
