import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class OrnekVeriTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.cikti = Path(cls.tmp.name) / "s.xlsx"
        cls.s = main.calistir(ORNEK / "standartlar.csv", ORNEK / "uretim.csv", ORNEK / "fiili.csv", cls.cikti, ORNEK / "gug.csv")
        cls.x = {x.kalem: x for x in cls.s["sapmalar"]}
        cls.g = {g.merkez: g for g in cls.s["gug"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_malzeme(self):
        y = self.x["Yonga levha"]                       # SM = 1000 × 1,8 + 2400 × 0,5 = 3000
        self.assertEqual((y.sm, y.fiyat, y.miktar), (D("3000.0"), D("37800"), D("63000")))
        k = self.x["Kumaş"]
        self.assertEqual((k.fiyat, k.miktar), (D("18500"), D("-11200")))
        self.assertEqual(self.x["Zımpara"].toplam, D("3200"))

    def test_iscilik_ve_gug(self):
        self.assertEqual((self.x["Kesim"].fiyat, self.x["Kesim"].miktar), (D("9300"), D("15000")))
        g = self.g["Kesim"]                             # oran 150 + 180.000 / 1.200 = 300; SS 880; FS 930
        self.assertEqual((g.oran, g.butce, g.kapasite_sapmasi, g.verimlilik), (D("300"), D("3500"), D("40500"), D("15000")))
        g = self.g["Montaj"]
        self.assertEqual((g.butce, g.kapasite_sapmasi, g.verimlilik, g.toplam), (D("-3400"), D("-36400"), D("-9200"), D("-49000")))

    def test_ozet_mutabakati(self):
        o = self.s["ozet"]
        self.assertEqual((o["fiyat"], o["miktar"], o["ucret"], o["sure"], o["butce"], o["verimlilik"], o["kapasite"]),
                         (D("52930"), D("57100"), D("18100"), D("3800"), D("100"), D("5800"), D("4100")))
        self.assertEqual(self.s["fiili"] - self.s["standart"], sum(o.values()))
        self.assertEqual(self.s["standart"], D("3374200"))

    def test_uyarilar_ve_excel(self):
        u = {(x["tur"], x["kalem"]) for x in self.s["uyarilar"]}
        self.assertTrue({("Standart dışı kalem", "Zımpara"), ("Sapma eşiği", "Yonga levha"), ("Sapma eşiği", "GÜG · Kesim")} <= u)
        self.assertNotIn(("Sapma eşiği", "Kenar bandı"), u)
        wb = load_workbook(self.cikti)
        self.assertEqual(wb.sheetnames, ["Sapma Özeti", "Malzeme", "İşçilik", "GÜG", "Standart Maliyet Kartı", "Uyarılar"])
        self.assertEqual(wb["Sapma Özeti"]["C12"].value, 0)


class KuralTesti(unittest.TestCase):
    def test_eksik_kayitlar(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp)
            (t / "u.csv").write_text("Ürün;Üretim Miktarı\nMasa;100\nDolap;5\n", encoding="utf-8")
            (t / "f.csv").write_text("Tür;Kalem;Fiili Miktar;Fiili Tutar\nMalzeme;Yonga levha;180;75600\n", encoding="utf-8")
            s = main.calistir(ORNEK / "standartlar.csv", t / "u.csv", t / "f.csv", t / "o.xlsx")
            u = {(x["tur"], x["kalem"]) for x in s["uyarilar"]}
            self.assertTrue({("Standart kart yok", "dolap"), ("Fiili kayıt yok", "Kenar bandı"), ("Üretim yok", "Sandalye")} <= u)
            y = next(x for x in s["sapmalar"] if x.kalem == "Yonga levha")
            self.assertEqual((y.sm, y.toplam), (D("180.0"), D("0.0")))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--fiili", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
