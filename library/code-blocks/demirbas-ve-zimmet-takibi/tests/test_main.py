import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def d(no, konum="Oda 1", zimmetli="Ali", durum="Aktif", deger=100):
    return {"no": no, "ad": "X", "kategori": "K", "konum": konum, "zimmetli": zimmetli, "durum": durum, "aktif": durum == "Aktif", "alis": None, "deger": None if deger is None else D(deger)}


def s(no, konum="Oda 1"):
    return {"no": no, "konum": konum, "sayan": "S", "tarih": None, "nota": "n"}


class KarsilastirTesti(unittest.TestCase):
    def test_sonuclar(self):
        l = [d("D-1"), d("D-2"), d("D-3"), d("D-4"), d("D-5", durum="Hurda"), d("D-6", durum="Satıldı"), d("D-7", konum=""), d("D-8"), d("D-9"), d("D-10")]
        r = main.karsilastir(l, [s("d 1", "ODA-1"), s("D-2", "Oda 2"), s("D-4"), s("D-4", "Oda 3"), s("D-5", "Depo"), s("D-7", "Oda 9"), s("D-8", ""), s("D-9"), s("D-9", "oda 1"), s("X-1", "Depo"), s("X-1", "Depo")])
        self.assertEqual([x["sonuc"] for x in r["demirbaslar"]],
                         ["Yerinde", "Farklı konumda", "Bulunamadı", "Mükerrer sayım", "Kayıttan çıkmış, sayıldı", None, "Farklı konumda", "Yerinde", "Yerinde", "Bulunamadı"])
        self.assertEqual((r["aktif"], r["deger"], r["say"], r["oran"], r["kayip_deger"]),
                         (8, D(800), {"Yerinde": 3, "Farklı konumda": 2, "Bulunamadı": 2, "Mükerrer sayım": 1, "Kayıttan çıkmış, sayıldı": 1}, D("75.00"), D(200)))
        self.assertEqual(r["kayit_disi"], [{"no": "X-1", "konum": "Depo", "sayan": "S", "nota": "n"}])
        self.assertEqual(r["demirbaslar"][3]["sayim_konum"], "Oda 1, Oda 3")
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         sorted([("Orta", "Farklı konumda bulundu", "D-2 · X"), ("Yüksek", "Demirbaş bulunamadı", "D-3 · X"), ("Yüksek", "Mükerrer sayım", "D-4 · X"), ("Orta", "Kayıttan çıkmış demirbaş sayıldı", "D-5 · X"),
                                 ("Bilgi", "Konum kaydı yok", "D-7 · X"), ("Orta", "Farklı konumda bulundu", "D-7 · X"), ("Yüksek", "Demirbaş bulunamadı", "D-10 · X"), ("Orta", "Kayıt dışı demirbaş", "X-1")]))

    def test_ozet_ve_sayimsiz(self):
        l = [d("1", "A", "Ali", deger=100), d("2", "a", "Ece", deger=50), d("3", "B", "", deger=None), d("4", "B", "Ali", deger=30), d("5", "B", "Ali", "Kayıp", 999)]
        r = main.karsilastir(l, [s("1", "A"), s("2", "B"), s("3", "B")])
        k = lambda liste: [(g["ad"], g["adet"], g["deger"], g["Yerinde"], g["Farklı konumda"], g["Bulunamadı"], g["kayip_deger"]) for g in liste]   # noqa: E731
        self.assertEqual(k(r["konumlar"]), [("A", 2, D(150), 1, 1, 0, D(0)), ("B", 2, D(30), 1, 0, 1, D(30))])
        self.assertEqual(k(r["zimmet"]), [("Ali", 2, D(130), 1, 0, 1, D(30)), ("Ece", 1, D(50), 0, 1, 0, D(0)), ("(zimmetsiz)", 1, D(0), 1, 0, 0, D(0))])
        self.assertEqual((r["oran"], [u["tur"] for u in r["uyarilar"] if u["tur"] == "Zimmetsiz demirbaş"]), (D("75.00"), ["Zimmetsiz demirbaş"]))
        r = main.karsilastir([d("1"), d("2", zimmetli="")])
        self.assertEqual((r["sayim_var"], r["oran"], [x["sonuc"] for x in r["demirbaslar"]], [u["tur"] for u in r["uyarilar"]], r["kayit_disi"]), (False, None, [None, None], ["Zimmetsiz demirbaş"], []))
        self.assertEqual(main.karsilastir([d("1")], [])["say"]["Bulunamadı"], 1)
        with self.assertRaises(ValueError):
            main.karsilastir([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "demirbaslar.csv", ORNEK / "sayim.csv", cikti)
            self.assertEqual((r["aktif"], r["deger"], r["say"], r["oran"], r["kayip_deger"]),
                             (10, D(208400), {"Yerinde": 6, "Farklı konumda": 2, "Bulunamadı": 1, "Mükerrer sayım": 1, "Kayıttan çıkmış, sayıldı": 1}, D("90.00"), D(32000)))
            self.assertEqual([(x["no"], x["sonuc"]) for x in r["demirbaslar"] if x["sonuc"] not in ("Yerinde", None)],
                             [("DMB-0002", "Farklı konumda"), ("DMB-0004", "Bulunamadı"), ("DMB-0007", "Mükerrer sayım"), ("DMB-0009", "Kayıttan çıkmış, sayıldı"), ("DMB-0011", "Farklı konumda")])
            self.assertEqual([k["no"] for k in r["kayit_disi"]], ["DMB-0099"])
            self.assertEqual([(g["ad"], g["adet"], g["deger"], g["Bulunamadı"]) for g in r["konumlar"]],
                             [("Kat 1 - Satış", 2, D(59500), 0), ("Kat 2 - Muhasebe", 4, D(46500), 0), ("Mutfak", 1, D(18000), 0), ("Toplantı Odası 1", 2, D(77000), 1), ("(konum yok)", 1, D(7400), 0)])
            self.assertEqual([(g["ad"], g["adet"], g["deger"]) for g in r["zimmet"]],
                             [("Ayşe Kaya", 2, D(12700)), ("Hakan Öz", 2, D(34400)), ("İdari İşler", 4, D(133000)), ("Mert Demir", 1, D(21500)), ("(zimmetsiz)", 1, D(6800))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Demirbaş bulunamadı", "DMB-0004 · Projeksiyon Cihazı"), ("Yüksek", "Mükerrer sayım", "DMB-0007 · Yazıcı"), ("Orta", "Farklı konumda bulundu", "DMB-0002 · Ofis Koltuğu"),
                              ("Orta", "Farklı konumda bulundu", "DMB-0011 · Evrak İmha Makinesi"), ("Orta", "Kayıt dışı demirbaş", "DMB-0099"), ("Orta", "Kayıttan çıkmış demirbaş sayıldı", "DMB-0009 · Buzdolabı"),
                              ("Orta", "Zimmetsiz demirbaş", "DMB-0003 · Çelik Dolap"), ("Bilgi", "Konum kaydı yok", "DMB-0011 · Evrak İmha Makinesi")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Demirbaşlar", "Sayım Farkları", "Konum Özeti", "Zimmet Özeti", "Uyarılar"])
            main.calistir(ORNEK / "demirbaslar.csv", None, cikti)
            self.assertEqual(load_workbook(cikti).sheetnames, ["Demirbaşlar", "Konum Özeti", "Zimmet Özeti", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, sy = Path(tmp) / "d.csv", Path(tmp) / "s.csv"
            y.write_text("Sicil No;Tanım;Grup;Yer;Kullanıcı;Durum;Tarih;Bedel\nA-1;Masa;;Oda;Ali;;01.02.2024;1.250,50\na 1;Kopya;;;;;;\nA-2;Dolap;M;;;HURDAYA AYRILDI;;x\n;;;;;;;\n", encoding="utf-8")
            sy.write_text("Barkod;Yer;Sayan\nA-1;Oda;S\n;Oda;S\n", encoding="utf-8")
            l, uy = main.oku_demirbas(y)
            self.assertEqual(([(x["no"], x["ad"], x["kategori"], x["konum"], x["zimmetli"], x["durum"], x["aktif"], x["alis"], x["deger"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("A-1", "Masa", "(kategori yok)", "Oda", "Ali", "Aktif", True, date(2024, 2, 1), D("1250.50")), ("A-2", "Dolap", "M", "", "", "Hurda", False, None, None)],
                              [("Mükerrer demirbaş numarası", "a 1"), ("Değer okunamadı", "A-2")]))
            self.assertEqual([(x["no"], x["konum"], x["sayan"]) for x in main.oku_sayim(sy)], [("A-1", "Oda", "S")])
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--demirbaslar", str(y), "--sayim", str(sy), "--cikti", c]), 0)
            self.assertEqual(main.main(["--demirbaslar", str(y), "--cikti", c]), 0)
            self.assertEqual(main.main(["--demirbaslar", str(y), "--sayim", str(Path(tmp) / "yok.csv"), "--cikti", c]), 1)
            self.assertEqual(main.main(["--demirbaslar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Ad;Konum\nA;1\n", "Demirbaş No;Konum\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(["--demirbaslar", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
