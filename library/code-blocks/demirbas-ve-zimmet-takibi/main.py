"""
Demirbaş ve Zimmet Takibi — Gerisi Bende kod bloğu
İdari İşler › İdari İşler Sorumlusu

Demirbaş kaydını sayım listesiyle karşılaştırır; konum, zimmet ve sayım farklarını raporlar.
  Yerinde                  = sayımda bulundu ve bulunduğu konum kayıttaki konumla aynı
  Farklı konumda           = sayımda bulundu ama konum kayıttakinden farklı (ya da kayıtta konum yok)
  Bulunamadı               = kayıtta aktif, sayım listesinde yok
  Mükerrer sayım           = aynı demirbaş numarası birden çok farklı konumda sayılmış
  Kayıt dışı               = sayımda var, demirbaş kaydında yok
  Kayıttan çıkmış, sayıldı = durumu hurda / satıldı / devredildi / kayıp olan demirbaş sayımda bulundu
  Sayım oranı %            = sayımda bulunan aktif demirbaş ÷ aktif demirbaş × 100
Sayım dosyası verilmezse yalnız envanter, zimmet ve konum özeti üretilir. Değer, kayıttaki tutardır; amortisman hesaplanmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
D_SUTUN = {"no": ("demirbas no", "demirbas kodu", "sicil no", "barkod", "no"), "ad": ("demirbas", "demirbas adi", "ad", "tanim", "aciklama"), "kategori": ("kategori", "grup", "tur"), "konum": ("konum", "yer", "lokasyon", "oda"),
           "zimmetli": ("zimmetli", "zimmetli kisi", "kullanici", "sorumlu"), "durum": ("durum",), "alis": ("alis tarihi", "edinim tarihi", "tarih"), "deger": ("deger", "alis bedeli", "tutar", "bedel")}
S_SUTUN = {"no": ("demirbas no", "demirbas kodu", "sicil no", "barkod", "no"), "konum": ("bulundugu konum", "sayim konumu", "konum", "yer", "lokasyon"), "sayan": ("sayan", "sayim yapan", "sorumlu"),
           "tarih": ("sayim tarihi", "tarih"), "nota": ("not", "aciklama")}
PASIF = {"hurda": "Hurda", "hurdaya ayrildi": "Hurda", "satildi": "Satıldı", "devredildi": "Devredildi", "kayip": "Kayıp", "pasif": "Pasif", "kullanim disi": "Pasif"}
SONUC_RENK = {"Yerinde": YESIL, "Farklı konumda": SARI, "Bulunamadı": KIRMIZI, "Mükerrer sayım": KIRMIZI, "Kayıttan çıkmış, sayıldı": SARI}
Z = Decimal(0)


def anahtar(no) -> str:
    return katla(no).replace(" ", "")


def oku_demirbas(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, D_SUTUN, ("no",)):
        no = metin(r["no"])
        if not anahtar(no):
            continue
        if anahtar(no) in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer demirbaş numarası", no, f"Kayıt dosyasında birden çok kez geçiyor (satır {r['_satir']}). İlk kayıt kullanıldı."))
            continue
        gorulen.add(anahtar(no))
        deger = para(r["deger"])
        if metin(r["deger"]) and (deger is None or deger < 0):
            uy.append(uyari("Orta", "Değer okunamadı", no, f"Değer '{metin(r['deger'])}' okunamadı; boş sayıldı."))
            deger = None
        durum = PASIF.get(katla(r["durum"]))
        l.append({"no": no, "ad": metin(r["ad"]), "kategori": metin(r["kategori"]) or "(kategori yok)", "konum": metin(r["konum"]), "zimmetli": metin(r["zimmetli"]), "durum": durum or "Aktif", "aktif": durum is None,
                  "alis": tarih(r["alis"]), "deger": deger})
    return l, uy


def oku_sayim(yol: Path) -> list[dict]:
    return [{"no": metin(r["no"]), "konum": metin(r["konum"]), "sayan": metin(r["sayan"]), "tarih": tarih(r["tarih"]), "nota": metin(r["nota"])} for r in kayitlar(yol, S_SUTUN, ("no",)) if anahtar(r["no"])]


def karsilastir(demirbaslar: list[dict], sayim: list[dict] | None = None) -> dict:
    if not demirbaslar:
        raise ValueError("Demirbaş kaydı boş.")
    uy, sayilan = [], defaultdict(list)
    for s in sayim or []:
        sayilan[anahtar(s["no"])].append(s)
    for d in demirbaslar:
        kim = f"{d['no']} · {d['ad']}".strip(" ·")
        bulunan = sayilan.get(anahtar(d["no"]), [])
        yerler = list(dict.fromkeys(s["konum"] for s in bulunan if s["konum"]))
        d["sayim_konum"], d["sonuc"] = ", ".join(yerler), None
        if d["aktif"] and not d["zimmetli"]:
            uy.append(uyari("Orta", "Zimmetsiz demirbaş", kim, f"Zimmetli kişi yazılmamış. Konum: {d['konum'] or 'belirtilmemiş'}."))
        if d["aktif"] and not d["konum"]:
            uy.append(uyari("Bilgi", "Konum kaydı yok", kim, "Kayıtta konum yazılmamış."))
        if sayim is None:
            continue
        if not d["aktif"]:
            if bulunan:
                d["sonuc"] = "Kayıttan çıkmış, sayıldı"
                uy.append(uyari("Orta", "Kayıttan çıkmış demirbaş sayıldı", kim, f"Kayıtta durumu '{d['durum']}' ama sayımda bulundu ({d['sayim_konum'] or 'konum yazılmamış'}). Kaydı veya sayımı düzeltin."))
            continue
        deger = f" Kayıtlı değeri {d['deger']:,.2f}." if d["deger"] else ""
        if not bulunan:
            d["sonuc"] = "Bulunamadı"
            uy.append(uyari("Yüksek", "Demirbaş bulunamadı", kim, f"Kayıttaki konum: {d['konum'] or 'belirtilmemiş'}; zimmetli: {d['zimmetli'] or 'yok'}.{deger} Sayımda yok."))
        elif len({katla(y) for y in yerler}) > 1:
            d["sonuc"] = "Mükerrer sayım"
            uy.append(uyari("Yüksek", "Mükerrer sayım", kim, f"Aynı numara birden çok konumda sayılmış: {d['sayim_konum']}. Etiket kopyası ya da okuma hatası olabilir."))
        elif d["konum"] and (not yerler or katla(yerler[0]) == katla(d["konum"])):
            d["sonuc"] = "Yerinde"
        else:
            d["sonuc"] = "Farklı konumda"
            uy.append(uyari("Orta", "Farklı konumda bulundu", kim, f"Kayıttaki konum: {d['konum'] or 'belirtilmemiş'}; sayımda bulunduğu yer: {d['sayim_konum'] or 'yazılmamış'}. Zimmetli: {d['zimmetli'] or 'yok'}."))
    kayitli = {anahtar(d["no"]) for d in demirbaslar}
    kayit_disi = []
    for k, liste in sayilan.items():
        if k not in kayitli:
            kayit_disi.append({"no": liste[0]["no"], "konum": ", ".join(dict.fromkeys(s["konum"] for s in liste if s["konum"])), "sayan": liste[0]["sayan"], "nota": liste[0]["nota"]})
            uy.append(uyari("Orta", "Kayıt dışı demirbaş", liste[0]["no"], f"Sayımda var ({kayit_disi[-1]['konum'] or 'konum yazılmamış'}), demirbaş kaydında yok."))
    aktif = [d for d in demirbaslar if d["aktif"]]

    def ozet(alan: str, bos: str) -> list[dict]:
        grup = {}
        for d in aktif:
            g = grup.setdefault(katla(d[alan]), {"ad": d[alan] or bos, "adet": 0, "deger": Z, "Yerinde": 0, "Farklı konumda": 0, "Bulunamadı": 0, "Mükerrer sayım": 0, "kayip_deger": Z})
            g["adet"] += 1
            g["deger"] += d["deger"] or Z
            if d["sonuc"]:
                g[d["sonuc"]] += 1
            if d["sonuc"] == "Bulunamadı":
                g["kayip_deger"] += d["deger"] or Z
        return sorted(grup.values(), key=lambda g: (g["ad"] == bos, katla(g["ad"])))

    say = {s: sum(1 for d in demirbaslar if d["sonuc"] == s) for s in SONUC_RENK}
    bulunan = say["Yerinde"] + say["Farklı konumda"] + say["Mükerrer sayım"]
    return {"demirbaslar": demirbaslar, "aktif": len(aktif), "deger": sum((d["deger"] or Z for d in aktif), Z), "say": say, "sayim_var": sayim is not None, "kayit_disi": sorted(kayit_disi, key=lambda k: anahtar(k["no"])),
            "oran": yuvarla(Decimal(bulunan * 100) / len(aktif)) if sayim is not None and aktif else None, "kayip_deger": sum((d["deger"] or Z for d in aktif if d["sonuc"] == "Bulunamadı"), Z),
            "konumlar": ozet("konum", "(konum yok)"), "zimmet": ozet("zimmetli", "(zimmetsiz)"), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Demirbaşlar"
    _baslik(ws, ["Demirbaş No", "Demirbaş", "Kategori", "Durum", "Kayıtlı Konum", "Zimmetli", "Alış Tarihi", "Değer", "Sayım Sonucu", "Sayımda Bulunduğu Yer"], (15, 30, 18, 11, 22, 22, 12, 14, 24, 24))
    for d in r["demirbaslar"]:
        _satir(ws, [d["no"], d["ad"], d["kategori"], d["durum"], d["konum"], d["zimmetli"], d["alis"], d["deger"], d["sonuc"], d["sayim_konum"]], para_s=(8,), tarih_s=(7,), dolgu={9: SONUC_RENK.get(d["sonuc"])})
    if r["sayim_var"]:
        ws = wb.create_sheet("Sayım Farkları")
        _baslik(ws, ["Fark", "Demirbaş No", "Demirbaş", "Kayıtlı Konum", "Sayımda Bulunduğu Yer", "Zimmetli", "Değer", "Açıklama (siz doldurun)"], (24, 15, 30, 22, 24, 22, 14, 36))
        for sonuc in ("Bulunamadı", "Mükerrer sayım", "Farklı konumda", "Kayıttan çıkmış, sayıldı"):
            for d in r["demirbaslar"]:
                if d["sonuc"] == sonuc:
                    _satir(ws, [sonuc, d["no"], d["ad"], d["konum"], d["sayim_konum"], d["zimmetli"], d["deger"], None], para_s=(7,), dolgu={1: SONUC_RENK[sonuc]})
        for k in r["kayit_disi"]:
            _satir(ws, ["Kayıt dışı", k["no"], k["nota"], None, k["konum"], None, None, None], dolgu={1: SARI})
    for baslik, alan, ad in (("Konum Özeti", "konumlar", "Kayıtlı Konum"), ("Zimmet Özeti", "zimmet", "Zimmetli")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Aktif Demirbaş", "Değer", "Yerinde", "Farklı Konumda", "Bulunamadı", "Mükerrer Sayım", "Bulunamayan Değer"], (28, 10, 15, 10, 10, 11, 10, 15))
        for g in r[alan]:
            _satir(ws, [g["ad"], g["adet"], g["deger"]] + ([g["Yerinde"], g["Farklı konumda"], g["Bulunamadı"], g["Mükerrer sayım"], g["kayip_deger"] or None] if r["sayim_var"] else [None] * 5), para_s=(3, 8),
                   dolgu={6: KIRMIZI if g["Bulunamadı"] else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, sayim_yol: Path | None, cikti: Path) -> dict:
    demirbaslar, uy = oku_demirbas(yol)
    r = karsilastir(demirbaslar, oku_sayim(sayim_yol) if sayim_yol else None)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Demirbaş kaydını sayım listesiyle karşılaştırır; konum, zimmet ve sayım farklarını raporlar.")
    p.add_argument("--demirbaslar", type=Path, help="Demirbaş No, Demirbaş, Kategori, Konum, Zimmetli, Durum, Alış Tarihi, Değer")
    p.add_argument("--sayim", type=Path, help="Sayım listesi: Demirbaş No, Bulunduğu Konum, Sayan, Sayım Tarihi, Not (isteğe bağlı)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "demirbas_zimmet_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.demirbaslar is None
    yol = ORNEK / "demirbaslar.csv" if ornek else a.demirbaslar
    sayim = ORNEK / "sayim.csv" if ornek and a.sayim is None else a.sayim
    for y in (yol, sayim):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    try:
        r = calistir(yol, sayim, a.cikti)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    s = r["say"]
    print(f"[OK] {r['aktif']} aktif demirbaş · kayıtlı değer {r['deger']:,.2f}" + (f" · sayım oranı %{r['oran']} · yerinde {s['Yerinde']}, farklı konumda {s['Farklı konumda']}, bulunamadı {s['Bulunamadı']}"
          f" ({r['kayip_deger']:,.2f}), kayıt dışı {len(r['kayit_disi'])}" if r["sayim_var"] else " · sayım dosyası verilmedi"))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
