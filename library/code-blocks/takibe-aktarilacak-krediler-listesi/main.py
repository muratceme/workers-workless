"""
Takibe Aktarılacak Krediler Listesi — Workers / Workless kod bloğu
Krediler İzleme ve Takip › Yasal Takip Uzmanı

Gecikmesi eşiği aşan veya kat edilen kredileri, gerekli evrak durumuyla birlikte yasal takibe aktarım listesine dönüştürür.
  Aday            = gecikme günü --gun (90) günü aşmış YA DA hesap kat tarihi dolu
  Evrak sütunları = bilinen sütunların dışındaki her sütun bir evrak sayılır; hücre: Var / Yok / Gerekmez
  Aktarıma hazır  = aday ve gereken evrakın tamamı "Var"
  Evrak eksik     = aday ve en az bir evrak "Yok" ya da boş
  Aday değil      = eşiği aşmamış ve kat edilmemiş
  Sıra            = Evrak eksik → Aktarıma hazır → Aday değil; içinde risk büyükten küçüğe
Hangi evrakın gerektiğini siz belirlersiniz (dosyanın sütunları). Eşik ve aktarım kararı bankanızın politikasına ve mevzuata göre verilir.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
BILINEN = {"no": ("musteri no", "musteri numarasi"), "ad": ("musteri", "musteri adi", "unvan"), "kredi": ("kredi no", "hesap no", "kredi numarasi"), "risk": ("risk", "risk bakiyesi", "bakiye", "toplam risk"),
           "gun": ("gecikme gun", "gecikme gunu", "gecikme"), "kat": ("hesap kat tarihi", "kat tarihi", "hesap kat"), "nota": ("not", "aciklama")}
VAR, GEREKMEZ = {"var", "evet", "tam", "x", "mevcut", "tamam"}, {"gerekmez", "gerekli degil", "yok gerekmez", "uygulanmaz", "n a"}
DURUMLAR = ("Evrak eksik", "Aktarıma hazır", "Aday değil")
DURUM_RENK = {"Evrak eksik": KIRMIZI, "Aktarıma hazır": YESIL, "Aday değil": None}
Z = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[str], list[dict]]:
    """Bilinen sütunları eşler; kalan her sütun bir evrak sütunudur. (krediler, evrak adları, uyarılar) döndürür."""
    satirlar = tablo_oku(yol)
    for bi, r in enumerate(satirlar[:15]):
        b = [katla(c) for c in r]
        es = {alan: next((b.index(a) for a in adlar if a in b), None) for alan, adlar in BILINEN.items()}
        if es["no"] is not None and es["risk"] is not None:
            break
    else:
        raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: Müşteri No, Risk")
    evrak_sutun = [(j, metin(c)) for j, c in enumerate(r) if metin(c) and j not in es.values()]
    al = lambda satir, j: satir[j] if j is not None and j < len(satir) else None   # noqa: E731
    uy, l = [], []
    for n, s in enumerate(satirlar[bi + 1:], bi + 2):
        no = metin(al(s, es["no"]))
        if not no:
            continue
        kim = f"{no} · {metin(al(s, es['ad']))} · {metin(al(s, es['kredi']))}".replace(" ·  ·", " ·").strip(" ·")
        risk, gun = para(al(s, es["risk"])), para(al(s, es["gun"])) if metin(al(s, es["gun"])) else Z
        if risk is None or risk < 0 or gun is None or gun < 0 or gun != gun.to_integral_value():
            uy.append(uyari("Yüksek", "Kredi satırı okunamadı", kim, f"Satır {n}: risk '{metin(al(s, es['risk']))}', gecikme gün '{metin(al(s, es['gun']))}'. Satır alınmadı."))
            continue
        kat_ham = al(s, es["kat"])
        if metin(kat_ham) and tarih(kat_ham) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", kim, f"Hesap kat tarihi '{metin(kat_ham)}' okunamadı; boş sayıldı."))
        evrak = {}
        for j, ad in evrak_sutun:
            k = katla(al(s, j))
            evrak[ad] = "Var" if k in VAR or metin(al(s, j)) in ("✓", "✔", "+") else "Gerekmez" if k in GEREKMEZ else "Yok"
        l.append({"no": no, "ad": metin(al(s, es["ad"])), "kredi": metin(al(s, es["kredi"])), "risk": risk, "gun": int(gun), "kat": tarih(kat_ham), "nota": metin(al(s, es["nota"])), "evrak": evrak})
    return l, [ad for _, ad in evrak_sutun], uy


def degerlendir(krediler: list[dict], evraklar: list[str], esik: int = 90) -> dict:
    if not krediler:
        raise ValueError("Değerlendirilecek kredi yok.")
    uy = []
    for k in krediler:
        kim = f"{k['no']} · {k['ad']} · {k['kredi']}".replace(" ·  ·", " ·").strip(" ·")
        k["aday"] = k["gun"] > esik or k["kat"] is not None
        k["eksikler"] = [e for e in evraklar if k["evrak"].get(e, "Yok") == "Yok"]
        k["neden"] = " ve ".join(([f"gecikme {k['gun']} gün"] if k["gun"] > esik else []) + ([f"hesap {k['kat']:%d.%m.%Y} tarihinde kat edildi"] if k["kat"] else []))
        k["durum"] = "Aday değil" if not k["aday"] else "Evrak eksik" if k["eksikler"] else "Aktarıma hazır"
        if k["durum"] == "Evrak eksik":
            uy.append(uyari("Yüksek", "Evrak eksik", kim, f"Aday ({k['neden']}); risk {k['risk']:,.2f}. Eksik: {', '.join(k['eksikler'])}."))
        if k["gun"] > esik and k["kat"] is None:
            uy.append(uyari("Orta", "Hesap kat edilmemiş", kim, f"Gecikme {k['gun']} gün (eşik {esik}) ama hesap kat tarihi yok. Risk {k['risk']:,.2f}."))
    krediler.sort(key=lambda k: (DURUMLAR.index(k["durum"]), -k["risk"], katla(k["no"])))
    ozet = [{"durum": d, "kredi": sum(1 for k in krediler if k["durum"] == d), "risk": sum((k["risk"] for k in krediler if k["durum"] == d), Z)} for d in DURUMLAR]
    adaylar = [k for k in krediler if k["aday"]]
    return {"krediler": krediler, "evraklar": evraklar, "ozet": ozet, "aday": len(adaylar), "aday_risk": sum((k["risk"] for k in adaylar), Z),
            "evrak_ozet": [{"evrak": e, "var": sum(1 for k in adaylar if k["evrak"].get(e) == "Var"), "eksik": sum(1 for k in adaylar if k["evrak"].get(e, "Yok") == "Yok"), "gerekmez": sum(1 for k in adaylar if k["evrak"].get(e) == "Gerekmez")}
                           for e in evraklar], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, esik: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aktarım Listesi"
    e = r["evraklar"]
    _baslik(ws, ["Durum", "Müşteri No", "Müşteri", "Kredi No", "Risk", "Gecikme (gün)", "Hesap Kat Tarihi", "Aday Olma Nedeni"] + e + ["Eksik Evrak", "Not"], (16, 12, 26, 12, 16, 10, 13, 40) + (13,) * len(e) + (40, 30))
    for k in r["krediler"]:
        _satir(ws, [k["durum"], k["no"], k["ad"], k["kredi"], k["risk"], k["gun"] or None, k["kat"], k["neden"]] + [k["evrak"].get(x, "Yok") for x in e] + [", ".join(k["eksikler"]) if k["aday"] else None, k["nota"]], para_s=(5,), tarih_s=(7,),
               dolgu={1: DURUM_RENK[k["durum"]]} | {9 + i: KIRMIZI for i, x in enumerate(e) if k["aday"] and k["evrak"].get(x, "Yok") == "Yok"}, sar=(8, 9 + len(e)))
    ws = wb.create_sheet("Özet")
    _baslik(ws, ["Durum", "Kredi", "Risk"], (24, 8, 18))
    for o in r["ozet"]:
        _satir(ws, [o["durum"], o["kredi"], o["risk"]], para_s=(3,), dolgu={1: DURUM_RENK[o["durum"]]})
    _satir(ws, ["Aday toplamı", r["aday"], r["aday_risk"]], para_s=(3,))
    ws.append([])
    _baslik(ws, ["Evrak (adaylarda)", "Var", "Eksik", "Gerekmez"], dondur="A2")
    for x in r["evrak_ozet"]:
        _satir(ws, [x["evrak"], x["var"], x["eksik"] or None, x["gerekmez"] or None], dolgu={3: KIRMIZI if x["eksik"] else None})
    ws.append([])
    ws.append([f"Aday: gecikme {esik} günü aşmış ya da hesap kat edilmiş krediler."])
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, esik: int = 90) -> dict:
    krediler, evraklar, uy = oku(yol)
    r = degerlendir(krediler, evraklar, esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Gecikmesi eşiği aşan veya kat edilen kredileri evrak durumuyla yasal takibe aktarım listesine dönüştürür.")
    p.add_argument("--krediler", type=Path, help="Müşteri No, Müşteri, Kredi No, Risk, Gecikme Gün, Hesap Kat Tarihi, Not ve evrak sütunları (her evrak için Var / Yok / Gerekmez)")
    p.add_argument("--gun", type=int, default=90, help="Gecikme bu günü aşarsa kredi aktarım adayı sayılır (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "takibe_aktarilacak_krediler.xlsx")
    a = p.parse_args(argv)
    yol = a.krediler or ORNEK / "krediler.csv"
    if a.gun < 0:
        print("[X] --gun sıfır veya pozitif olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, a.gun)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = {x["durum"]: x for x in r["ozet"]}
    print(f"[OK] {len(r['krediler'])} kredi · aday {r['aday']} ({r['aday_risk']:,.2f}) · aktarıma hazır {o['Aktarıma hazır']['kredi']} ({o['Aktarıma hazır']['risk']:,.2f}) · evrak eksik {o['Evrak eksik']['kredi']} ({o['Evrak eksik']['risk']:,.2f})"
          f" · {len(r['evraklar'])} evrak sütunu")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
