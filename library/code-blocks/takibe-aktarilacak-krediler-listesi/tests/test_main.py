import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
E, H, N = "Evrak eksik", "Aktarıma hazır", "Aday değil"
EV = ["Sözleşme", "İhtar"]


def k(no, gun=0, kat=None, risk=100, **evrak):
    return {"no": no, "ad": "Ad", "kredi": "K", "risk": D(str(risk)), "gun": gun, "kat": kat, "nota": "", "evrak": {"Sözleşme": evrak.get("s", "Var"), "İhtar": evrak.get("i", "Var")}}


class DegerlendirTesti(unittest.TestCase):
    def test_adaylik_ve_evrak(self):
        kat = date(2026, 9, 1)
        l = [k("90", 90), k("91", 91), k("kat", 10, kat), k("ikisi", 120, kat, i="Yok"), k("gerekmez", 91, kat, s="Gerekmez"), k("hepsi_yok", 95, s="Yok", i="Yok", risk=900), k("aday_değil_eksik", 5, s="Yok")]
        r = main.degerlendir(l, EV)
        x = {v["no"]: v for v in r["krediler"]}
        self.assertEqual({n: (v["durum"], v["eksikler"], v["neden"]) for n, v in x.items()},
                         {"90": (N, [], ""), "91": (H, [], "gecikme 91 gün"), "kat": (H, [], "hesap 01.09.2026 tarihinde kat edildi"), "ikisi": (E, ["İhtar"], "gecikme 120 gün ve hesap 01.09.2026 tarihinde kat edildi"),
                          "gerekmez": (H, [], "gecikme 91 gün ve hesap 01.09.2026 tarihinde kat edildi"), "hepsi_yok": (E, ["Sözleşme", "İhtar"], "gecikme 95 gün"), "aday_değil_eksik": (N, ["Sözleşme"], "")})
        self.assertEqual([v["no"] for v in r["krediler"]], ["hepsi_yok", "ikisi", "91", "gerekmez", "kat", "90", "aday_değil_eksik"])
        self.assertEqual((r["ozet"], r["aday"], r["aday_risk"]), ([{"durum": E, "kredi": 2, "risk": D(1000)}, {"durum": H, "kredi": 3, "risk": D(300)}, {"durum": N, "kredi": 2, "risk": D(200)}], 5, D(1300)))
        self.assertEqual(r["evrak_ozet"], [{"evrak": "Sözleşme", "var": 3, "eksik": 1, "gerekmez": 1}, {"evrak": "İhtar", "var": 3, "eksik": 2, "gerekmez": 0}])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]),
                         [("Orta", "Hesap kat edilmemiş", "91 · Ad · K"), ("Orta", "Hesap kat edilmemiş", "hepsi_yok · Ad · K"), ("Yüksek", "Evrak eksik", "hepsi_yok · Ad · K"), ("Yüksek", "Evrak eksik", "ikisi · Ad · K")])
        self.assertEqual(([v["durum"] for v in main.degerlendir([k("a", 91), k("b", 121)], EV, 120)["krediler"]], main.degerlendir([k("a", 1)], EV, 0)["krediler"][0]["durum"]), ([H, N], H))
        r = main.degerlendir([k("a", 91)], [])
        self.assertEqual((r["krediler"][0]["durum"], r["evrak_ozet"]), (H, []))
        with self.assertRaises(ValueError):
            main.degerlendir([], EV)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "krediler.csv", cikti)
            self.assertEqual(r["evraklar"], ["Kredi Sözleşmesi", "Hesap Kat İhtarı", "Tebliğ Şerhi", "Teminat Evrakı", "Kefil Bilgileri"])
            self.assertEqual([(x["no"], x["durum"], x["eksikler"]) for x in r["krediler"]],
                             [("5002", E, ["Tebliğ Şerhi"]), ("5005", E, ["Teminat Evrakı"]), ("5003", E, ["Hesap Kat İhtarı", "Tebliğ Şerhi"]), ("5001", H, []), ("5006", H, []), ("5004", N, ["Hesap Kat İhtarı", "Tebliğ Şerhi"])])
            self.assertEqual((r["ozet"], r["aday"], r["aday_risk"]), ([{"durum": E, "kredi": 3, "risk": D(4220000)}, {"durum": H, "kredi": 2, "risk": D(1200000)}, {"durum": N, "kredi": 1, "risk": D(600000)}], 5, D(5420000)))
            self.assertEqual([(x["evrak"], x["var"], x["eksik"], x["gerekmez"]) for x in r["evrak_ozet"]],
                             [("Kredi Sözleşmesi", 5, 0, 0), ("Hesap Kat İhtarı", 4, 1, 0), ("Tebliğ Şerhi", 3, 2, 0), ("Teminat Evrakı", 3, 1, 1), ("Kefil Bilgileri", 3, 0, 2)])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]],
                             [("Yüksek", "Evrak eksik", "5002 · Beta Yapı · K-12"), ("Yüksek", "Evrak eksik", "5003 · Gama Gıda · K-13"), ("Yüksek", "Evrak eksik", "5005 · Epsilon Otomotiv · K-15"), ("Orta", "Hesap kat edilmemiş", "5003 · Gama Gıda · K-13")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Aktarım Listesi", "Özet", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "k.csv"
            y.write_text("Rapor;;;;;;;\nİpotek Belgesi;Bakiye;Müşteri Numarası;Kat Tarihi;Sözleşme;Gecikme;Açıklama;Unvan\nEVET;1.250.000,50;7;x;;120;not;A Ltd.\nGerekli değil;5;8;01.09.2026;✓;;;\nvar;x;9;;;;;\nvar;5;10;;;1,5;;\nvar;5;;;;;;\n",
                         encoding="utf-8")
            l, evraklar, uy = main.oku(y)
            self.assertEqual((evraklar, l, [(u["tur"], u["kim"]) for u in uy]),
                             (["İpotek Belgesi", "Sözleşme"], [{"no": "7", "ad": "A Ltd.", "kredi": "", "risk": D("1250000.50"), "gun": 120, "kat": None, "nota": "not", "evrak": {"İpotek Belgesi": "Var", "Sözleşme": "Yok"}},
                                                               {"no": "8", "ad": "", "kredi": "", "risk": D(5), "gun": 0, "kat": date(2026, 9, 1), "nota": "", "evrak": {"İpotek Belgesi": "Gerekmez", "Sözleşme": "Var"}}],
                              [("Tarih okunamadı", "7 · A Ltd."), ("Kredi satırı okunamadı", "9"), ("Kredi satırı okunamadı", "10")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--krediler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--gun", "120"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--gun", "-1"]), 2)
            self.assertEqual(main.main(["--krediler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Müşteri No;Kredi No\nA;B\n", "Müşteri No;Risk\n;1\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            y.write_text("Müşteri No;Risk;Gecikme Gün\n7;100;95\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
