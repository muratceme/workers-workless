"""
MTBF ve MTTR Analizi — Gerisi Bende kod bloğu
Bakım Onarım › Bakım Mühendisi

Arıza kayıtlarından makine bazında arızalar arası ortalama süreyi (MTBF) ve ortalama onarım süresini (MTTR) hesaplar.
  Arıza süresi (saat)  = bitiş − başlangıç          (yoksa "Süre (dk)" ÷ 60)
  MTTR                 = toplam arıza süresi ÷ arıza sayısı
  Çalışma süresi       = planlı çalışma süresi − toplam arıza süresi
  MTBF                 = çalışma süresi ÷ arıza sayısı
  Kullanılabilirlik %  = çalışma süresi ÷ planlı çalışma süresi × 100        ( = MTBF ÷ (MTBF + MTTR) )
Planlı çalışma süresi makine dosyasından veya --planli-saat ile verilir; verilmezse yalnız MTTR hesaplanır.
"Arıza süresi" kayıttaki duruş süresidir: bekleme (parça, teknisyen) dahilse MTTR de onu içerir. Hedefler işletmenize özeldir.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"makine": ("makine", "ekipman", "makine kodu", "hat"), "bas": ("baslangic", "ariza baslangic", "baslangic zamani", "durus baslangic"), "bit": ("bitis", "ariza bitis", "bitis zamani", "durus bitis"),
         "sure": ("sure dk", "sure", "durus suresi dk", "ariza suresi dk"), "tip": ("ariza tipi", "tip", "ariza turu", "neden")}
MAKINE_SUTUN = {"makine": ("makine", "ekipman", "makine kodu", "hat"), "planli": ("planli calisma suresi saat", "planli calisma suresi", "planli sure saat", "planli saat", "calisma saati")}
Z = Decimal(0)


def zaman(x) -> datetime | None:
    if isinstance(x, datetime):
        return x
    if isinstance(x, date):
        return datetime(x.year, x.month, x.day)
    s = " ".join(str(x or "").replace("T", " ").split())
    for f in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def oku(yol: Path, makine_yol: Path | None = None) -> tuple[list[dict], dict, list[dict]]:
    uy, l, planli = [], [], {}
    for r in kayitlar(yol, SUTUN, ("makine",)):
        m = metin(r["makine"])
        if not m:
            continue
        kim = f"{m} (satır {r['_satir']})"
        bas, bit, sure = zaman(r["bas"]), zaman(r["bit"]), para(r["sure"])
        if metin(r["bas"]) and not bas:
            uy.append(uyari("Orta", "Başlangıç zamanı okunamadı", kim, f"'{metin(r['bas'])}'. GG.AA.YYYY SS:DD biçiminde yazın."))
        if bas and bit:
            if bit < bas:
                uy.append(uyari("Yüksek", "Bitiş başlangıçtan önce", kim, f"{bas:%d.%m.%Y %H:%M} – {bit:%d.%m.%Y %H:%M}. Kayıt alınmadı."))
                continue
            saat = Decimal((bit - bas).total_seconds()) / 3600
        elif sure is not None and sure >= 0:
            saat = sure / 60
        elif bas and not metin(r["bit"]):
            uy.append(uyari("Orta", "Arıza kapanmamış", kim, f"Başlangıç {bas:%d.%m.%Y %H:%M}, bitiş boş. Devam eden arıza hesaba alınmadı."))
            continue
        else:
            uy.append(uyari("Yüksek", "Arıza süresi okunamadı", kim, f"Başlangıç '{metin(r['bas'])}', bitiş '{metin(r['bit'])}', süre '{metin(r['sure'])}'. Kayıt alınmadı."))
            continue
        l.append({"makine": m, "anahtar": katla(m), "bas": bas, "saat": saat, "tip": metin(r["tip"]) or "Belirtilmemiş"})
    if makine_yol:
        for r in kayitlar(makine_yol, MAKINE_SUTUN, ("makine", "planli")):
            m, p = metin(r["makine"]), para(r["planli"])
            if not m:
                continue
            if p is None or p <= 0:
                uy.append(uyari("Yüksek", "Planlı çalışma süresi okunamadı", m, f"{makine_yol.name} satır {r['_satir']}: '{metin(r['planli'])}'. Makine için MTBF hesaplanmadı."))
                continue
            planli[katla(m)] = {"ad": m, "saat": p}
    return l, planli, uy


def analiz(arizalar: list[dict], planli: dict | None = None, planli_saat: Decimal | None = None, mtbf_hedef: Decimal | None = None, mttr_hedef: Decimal | None = None) -> dict:
    planli, uy, g = planli or {}, [], {}
    for a in arizalar:
        g.setdefault(a["anahtar"], {"ad": a["makine"], "l": []})["l"].append(a)
    for k, p in planli.items():
        g.setdefault(k, {"ad": p["ad"], "l": []})
    makineler = []
    for k, x in g.items():
        n, durus = len(x["l"]), sum((a["saat"] for a in x["l"]), Z)
        ps = planli[k]["saat"] if k in planli else planli_saat
        m = {"makine": x["ad"], "ariza": n, "durus": yuvarla(durus), "planli": ps, "mttr": yuvarla(durus / n) if n else None, "calisma": None, "mtbf": None, "kullanilabilirlik": None, "isaret": []}
        if ps is not None:
            if durus > ps:
                uy.append(uyari("Yüksek", "Arıza süresi planlı süreden büyük", x["ad"], f"Toplam arıza {yuvarla(durus)} saat > planlı çalışma {ps:f} saat. Dönem veya planlı süre hatalı olabilir; MTBF hesaplanmadı."))
            else:
                m["calisma"] = yuvarla(ps - durus)
                m["mtbf"] = yuvarla((ps - durus) / n) if n else None
                m["kullanilabilirlik"] = yuvarla((ps - durus) / ps * 100)
        elif n:
            uy.append(uyari("Orta", "Planlı çalışma süresi yok", x["ad"], "Makine dosyasında veya --planli-saat ile süre verilmedi; yalnız MTTR hesaplandı."))
        if mtbf_hedef is not None and m["mtbf"] is not None and m["mtbf"] < mtbf_hedef:
            m["isaret"].append("MTBF hedefin altında")
        if mttr_hedef is not None and m["mttr"] is not None and m["mttr"] > mttr_hedef:
            m["isaret"].append("MTTR hedefin üstünde")
        tipler = defaultdict(lambda: Z)
        for a in x["l"]:
            tipler[a["tip"]] += a["saat"]
        m["en_cok"] = max(tipler.items(), key=lambda t: (t[1], t[0]))[0] if tipler else ""
        makineler.append(m)
    makineler.sort(key=lambda m: (-m["durus"], m["makine"]))
    tip = {}
    for a in arizalar:
        t = tip.setdefault(a["tip"], {"tip": a["tip"], "ariza": 0, "durus": Z, "makineler": set()})
        t["ariza"] += 1
        t["durus"] += a["saat"]
        t["makineler"].add(a["anahtar"])
    toplam_durus = sum((a["saat"] for a in arizalar), Z)
    tipler, bir = [], Z
    for t in sorted(tip.values(), key=lambda t: (-t["durus"], t["tip"])):
        bir += t["durus"]
        tipler.append({"tip": t["tip"], "ariza": t["ariza"], "durus": yuvarla(t["durus"]), "mttr": yuvarla(t["durus"] / t["ariza"]), "makine": len(t["makineler"]),
                       "pay": yuvarla(t["durus"] / toplam_durus * 100) if toplam_durus else Z, "birikimli": yuvarla(bir / toplam_durus * 100) if toplam_durus else Z})
    tp = [m for m in makineler if m["calisma"] is not None]
    n, pl, cal = sum(m["ariza"] for m in tp), sum((m["planli"] for m in tp), Z), sum((m["calisma"] for m in tp), Z)
    genel = {"ariza": len(arizalar), "durus": yuvarla(toplam_durus), "mttr": yuvarla(toplam_durus / len(arizalar)) if arizalar else None, "mtbf": yuvarla(cal / n) if n else None,
             "kullanilabilirlik": yuvarla(cal / pl * 100) if pl else None, "makine": len(makineler)}
    return {"makineler": makineler, "tipler": tipler, "genel": genel, "uyarilar": uy, "hedef": (mtbf_hedef, mttr_hedef)}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, g = Workbook(), r["genel"]
    ws = wb.active
    ws.title = "Makine Bazında"
    _baslik(ws, ["Makine", "Arıza Sayısı", "Toplam Arıza Süresi (saat)", "Planlı Çalışma (saat)", "Çalışma Süresi (saat)", "MTBF (saat)", "MTTR (saat)", "Kullanılabilirlik %",
                 "En Çok Süre Kaybettiren Arıza Tipi", "İşaret", "Aksiyon (siz doldurun)"], (26, 9, 14, 13, 13, 12, 12, 13, 28, 34, 30))
    for m in r["makineler"]:
        _satir(ws, [m["makine"], m["ariza"], m["durus"], m["planli"], m["calisma"], m["mtbf"] if m["mtbf"] is not None else ("arıza yok" if m["calisma"] is not None else None), m["mttr"],
                    m["kullanilabilirlik"], m["en_cok"], ", ".join(m["isaret"]), ""], para_s=(3, 4, 5, 6, 7, 8), dolgu={10: KIRMIZI if m["isaret"] else None})
    _satir(ws, ["Genel", g["ariza"], g["durus"], None, None, g["mtbf"], g["mttr"], g["kullanilabilirlik"]], para_s=(3, 6, 7, 8))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append(["Genel MTBF ve kullanılabilirlik yalnız planlı çalışma süresi bilinen makinelerden hesaplanır. MTBF = çalışma süresi ÷ arıza sayısı; MTTR = arıza süresi ÷ arıza sayısı."])
    ws = wb.create_sheet("Arıza Tipi Bazında")
    _baslik(ws, ["Arıza Tipi", "Arıza Sayısı", "Toplam Süre (saat)", "MTTR (saat)", "Etkilenen Makine", "Süre Payı %", "Birikimli %"], (30, 9, 14, 12, 11, 11, 11))
    for t in r["tipler"]:
        _satir(ws, [t["tip"], t["ariza"], t["durus"], t["mttr"], t["makine"], t["pay"], t["birikimli"]], para_s=(3, 4, 6, 7))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, makine_yol: Path | None = None, **secenek) -> dict:
    arizalar, planli, uy = oku(yol, makine_yol)
    if not arizalar and not planli:
        raise ValueError(f"{yol.name}: hesaba alınabilecek arıza kaydı yok.")
    r = analiz(arizalar, planli, **secenek)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, r["uyarilar"])
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Arıza kayıtlarından makine bazında MTBF, MTTR ve kullanılabilirlik hesaplar.")
    p.add_argument("--arizalar", type=Path, default=ORNEK / "ariza_kayitlari.csv", help="Makine, Başlangıç, Bitiş (GG.AA.YYYY SS:DD) veya Süre (dk), Arıza Tipi")
    p.add_argument("--makineler", type=Path, help="Makine, Planlı Çalışma Süresi (saat): dönem için makine bazında")
    p.add_argument("--planli-saat", help="Makine dosyasında olmayan makineler için dönemin planlı çalışma süresi (saat)")
    p.add_argument("--mtbf-hedef", help="Bu değerin altındaki MTBF işaretlenir (saat)")
    p.add_argument("--mttr-hedef", help="Bu değerin üstündeki MTTR işaretlenir (saat)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "mtbf_mttr_analizi.xlsx")
    a = p.parse_args(argv)
    makine = a.makineler or (ORNEK / "makineler.csv" if a.arizalar == ORNEK / "ariza_kayitlari.csv" else None)
    for y in (a.arizalar, makine):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    ham = (a.planli_saat, a.mtbf_hedef, a.mttr_hedef)
    v = [para(x) if x else None for x in ham]
    if any(x and (d is None or d <= 0) for x, d in zip(ham, v)):
        print("[X] --planli-saat, --mtbf-hedef ve --mttr-hedef pozitif sayı olmalı")
        return 2
    try:
        r = calistir(a.arizalar, a.cikti, makine, planli_saat=v[0], mtbf_hedef=v[1], mttr_hedef=v[2])
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    g = r["genel"]
    print(f"[OK] {g['makine']} makine · {g['ariza']} arıza · {g['durus']} saat duruş · MTBF {g['mtbf'] if g['mtbf'] is not None else '—'} saat · MTTR {g['mttr'] if g['mttr'] is not None else '—'} saat")
    for m in r["makineler"][:5]:
        if m["ariza"]:
            print(f"     {m['makine']}: {m['ariza']} arıza · MTBF {m['mtbf'] if m['mtbf'] is not None else '—'} · MTTR {m['mttr']}" + (f" [{', '.join(m['isaret'])}]" if m["isaret"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
