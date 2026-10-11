import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def ar(makine, saat, tip="Mekanik"):
    return {"makine": makine, "anahtar": main.katla(makine), "bas": None, "saat": D(str(saat)), "tip": tip}


def pl(**kw):
    return {main.katla(k): {"ad": k, "saat": D(v)} for k, v in kw.items()}


class HesapTesti(unittest.TestCase):
    def test_mtbf_mttr(self):
        r = main.analiz([ar("M1", 2), ar("M1", 1.5, "Elektrik"), ar("M1", 4), ar("M1", 0.5, "Ayar")], pl(M1=480))
        m = r["makineler"][0]
        self.assertEqual((m["ariza"], m["durus"], m["calisma"], m["mttr"], m["mtbf"], m["kullanilabilirlik"], m["en_cok"]), (4, D("8.00"), D("472.00"), D("2.00"), D("118.00"), D("98.33"), "Mekanik"))
        # kullanılabilirlik = MTBF ÷ (MTBF + MTTR) = 118 ÷ 120
        self.assertEqual(main.yuvarla(m["mtbf"] / (m["mtbf"] + m["mttr"]) * 100), m["kullanilabilirlik"])

    def test_arizasiz_ve_planlisiz_makine(self):
        r = main.analiz([ar("M5", 3)], pl(M4=480))
        d = {m["makine"]: m for m in r["makineler"]}
        self.assertEqual((d["M4"]["ariza"], d["M4"]["mtbf"], d["M4"]["mttr"], d["M4"]["kullanilabilirlik"]), (0, None, None, D("100.00")))
        self.assertEqual((d["M5"]["mttr"], d["M5"]["mtbf"], d["M5"]["planli"], [u["tur"] for u in r["uyarilar"]]), (D("3.00"), None, None, ["Planlı çalışma süresi yok"]))
        r = main.analiz([ar("M5", 3)], pl(M4=480), planli_saat=D(300))
        m = next(m for m in r["makineler"] if m["makine"] == "M5")
        self.assertEqual((m["mtbf"], m["kullanilabilirlik"], r["uyarilar"]), (D("297.00"), D("99.00"), []))
        r = main.analiz([ar("M1", 500)], pl(M1=480))
        self.assertEqual((r["makineler"][0]["mtbf"], [u["tur"] for u in r["uyarilar"]]), (None, ["Arıza süresi planlı süreden büyük"]))

    def test_hedefler_tipler_ve_genel(self):
        l = [ar("M1", 2), ar("M1", 4), ar("M2", 12, "Hidrolik"), ar("M2", 12, "Hidrolik"), ar("M3", 1, "Ayar")]
        r = main.analiz(l, pl(M1=480, M2=480), mtbf_hedef=D(230), mttr_hedef=D(3))
        self.assertEqual([(m["makine"], m["mtbf"], m["mttr"], m["isaret"]) for m in r["makineler"]],
                         [("M2", D("228.00"), D("12.00"), ["MTBF hedefin altında", "MTTR hedefin üstünde"]), ("M1", D("237.00"), D("3.00"), []), ("M3", None, D("1.00"), [])])
        self.assertEqual([(t["tip"], t["ariza"], t["durus"], t["mttr"], t["makine"], t["pay"], t["birikimli"]) for t in r["tipler"]],
                         [("Hidrolik", 2, D("24.00"), D("12.00"), 1, D("77.42"), D("77.42")), ("Mekanik", 2, D("6.00"), D("3.00"), 1, D("19.35"), D("96.77")),
                          ("Ayar", 1, D("1.00"), D("1.00"), 1, D("3.23"), D("100.00"))])
        g = r["genel"]
        self.assertEqual((g["ariza"], g["durus"], g["mttr"], g["mtbf"], g["kullanilabilirlik"]), (5, D("31.00"), D("6.20"), D("232.50"), D("96.88")))    # (960 − 30) ÷ 4; M3 planlısız

    def test_zaman(self):
        self.assertEqual([main.zaman(x) for x in ("05.09.2026 10:00", "2026-09-05T10:00:00", "05.09.2026", "saat 10", None)],
                         [datetime(2026, 9, 5, 10), datetime(2026, 9, 5, 10), datetime(2026, 9, 5), None, None])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "m.xlsx"
            r = main.calistir(ORNEK / "ariza_kayitlari.csv", cikti, ORNEK / "makineler.csv")
            d = {m["makine"]: (m["ariza"], m["durus"], m["mtbf"], m["mttr"], m["kullanilabilirlik"]) for m in r["makineler"]}
            self.assertEqual(d, {"Enjeksiyon 2": (2, D("24.00"), D("228.00"), D("12.00"), D("95.00")), "Enjeksiyon 1": (4, D("8.00"), D("118.00"), D("2.00"), D("98.33")),
                                 "Pres 1": (6, D("4.00"), D("39.33"), D("0.67"), D("98.33")), "Konveyör": (1, D("3.00"), None, D("3.00"), None), "Robot 1": (0, D("0.00"), None, None, D("100.00"))})
            self.assertEqual({t["tip"]: (t["ariza"], t["durus"], t["mttr"]) for t in r["tipler"]},
                             {"Hidrolik": (2, D("24.00"), D("12.00")), "Mekanik": (3, D("9.00"), D("3.00")), "Elektrik": (4, D("3.25"), D("0.81")), "Ayar": (4, D("2.75"), D("0.69"))})
            g = r["genel"]
            self.assertEqual((g["ariza"], g["durus"], g["mttr"], g["mtbf"], g["kullanilabilirlik"]), (13, D("39.00"), D("3.00"), D("137.00"), D("97.86")))       # (1680 − 36) ÷ 12
            self.assertEqual([u["tur"] for u in r["uyarilar"]], ["Arıza kapanmamış", "Planlı çalışma süresi yok"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Makine Bazında", "Arıza Tipi Bazında", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, m = Path(tmp) / "a.csv", Path(tmp) / "m.csv"
            y.write_text("Makine;Başlangıç;Bitiş;Süre (dk);Arıza Tipi\nM1;05.09.2026 10:00;05.09.2026 11:30;;\nM1;;;90;Ayar\nM1;05.09.2026 12:00;05.09.2026 11:00;;\nM1;sabah;;;\nM2;06.09.2026 08:00;;;\n",
                         encoding="utf-8")
            m.write_text("Makine;Planlı Çalışma Süresi (saat)\nM1;200\nM2;yok\n", encoding="utf-8")
            l, p, uy = main.oku(y, m)
            self.assertEqual(([(x["saat"], x["tip"]) for x in l], p), ([(D("1.5"), "Belirtilmemiş"), (D("1.5"), "Ayar")], {"m1": {"ad": "M1", "saat": D(200)}}))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Arıza kapanmamış", "Arıza süresi okunamadı", "Başlangıç zamanı okunamadı", "Bitiş başlangıçtan önce", "Planlı çalışma süresi okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--arizalar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--makineler", str(m), "--mtbf-hedef", "100", "--mttr-hedef", "1"]), 0)
            self.assertEqual(main.main(arg + ["--planli-saat", "160"]), 0)
            self.assertEqual(main.main(arg + ["--planli-saat", "0"]), 2)
            self.assertEqual(main.main(["--arizalar", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Makine;Başlangıç;Bitiş\n", encoding="utf-8")
            self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
