"""
KKB ve Risk Merkezi Raporu Özeti — Workers / Workless kod bloğu
Krediler Tahsis › Kredi Tahsis Uzman Yardımcısı

Risk Merkezi / KKB raporundan tabloya aktarılmış limit-risk satırlarını özetler.
  Toplam limit, toplam risk (nakdi / gayrinakdi), çalışılan banka sayısı
  Limit kullanım oranı = risk ÷ limit                 (toplam ve banka bazında)
  Banka payı           = bankanın riski ÷ toplam risk (en büyük pay = yoğunlaşma)
  Dönemsel seyir       = dosyada birden çok dönem varsa toplam riskin ve banka sayısının değişimi
  Bulgular             = gecikme, limit aşımı, yüksek kullanım, hızlı risk artışı, olumsuz kayıtlar
Kod raporu yorumlamaz, kredi notu üretmez ve kredi kararı vermez; yalnız girilen satırları toplar.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"donem": ("donem", "rapor donemi", "ay"), "banka": ("banka", "kaynak", "kurulus", "finans kurulusu"), "tur": ("kredi turu", "tur", "urun"),
         "sinif": ("nakdi gayrinakdi", "sinif", "nakit gayrinakit"), "limit": ("limit", "toplam limit"), "risk": ("risk", "toplam risk", "bakiye"),
         "gun": ("gecikme gun", "gecikme gunu", "gecikme"), "geciken": ("geciken tutar", "gecikmedeki tutar", "gecikmis tutar")}
OLUMSUZ_SUTUN = {"tur": ("tur", "kayit turu"), "tarih": ("tarih",), "tutar": ("tutar",), "durum": ("durum",), "aciklama": ("aciklama", "not")}
GAYRINAKDI = ("gayri", "teminat mektubu", "akreditif", "kabul kredisi", "aval", "kefalet", "cek taahhut", "garanti")
KAPALI = ("kapandi", "odendi", "kapali", "tasfiye", "kaldirildi")
Z = Decimal(0)


def donem(x) -> tuple[int, int] | None:
    if isinstance(x, (date, datetime)):
        return x.year, x.month
    s = metin(x)
    m = re.fullmatch(r"(\d{4})[-./ ](\d{1,2})", s)
    if m:
        y, a = int(m[1]), int(m[2])
    else:
        m = re.fullmatch(r"(\d{1,2})[-./ ](\d{4})", s)
        if not m:
            t = tarih(s)
            return (t.year, t.month) if t else None
        y, a = int(m[2]), int(m[1])
    return (y, a) if 1 <= a <= 12 else None


def donem_ad(d) -> str:
    return f"{d[1]:02d}.{d[0]}" if d else "—"


def oran(a: Decimal, b: Decimal) -> Decimal | None:
    return yuvarla(a / b * 100) if b else None


def sinif_bul(sinif: str, tur: str) -> str:
    s = katla(sinif) or katla(tur)
    return "Gayrinakdi" if any(g in s for g in GAYRINAKDI) else "Nakdi"


def oku(yol: Path, olumsuz_yol: Path | None = None) -> tuple[list[dict], list[dict], list[dict]]:
    uy, satirlar = [], []
    for r in kayitlar(yol, SUTUN, ("banka", "risk")):
        b = metin(r["banka"])
        if not b:
            continue
        kim = f"{b} (satır {r['_satir']})"
        risk, limit = para(r["risk"]), para(r["limit"])
        if risk is None or risk < 0:
            uy.append(uyari("Yüksek", "Risk tutarı okunamadı", kim, f"'{metin(r['risk'])}'. Satır özete alınmadı."))
            continue
        if metin(r["limit"]) and (limit is None or limit < 0):
            uy.append(uyari("Orta", "Limit okunamadı", kim, f"'{metin(r['limit'])}'. Limit 0 sayıldı; kullanım oranı hesaplanamaz."))
            limit = None
        d = donem(r["donem"])
        if metin(r["donem"]) and not d:
            uy.append(uyari("Yüksek", "Dönem okunamadı", kim, f"'{metin(r['donem'])}'. AA.YYYY veya YYYY-AA yazın. Satır özete alınmadı."))
            continue
        gun = para(r["gun"])
        satirlar.append({"donem": d, "banka": b, "tur": metin(r["tur"]) or "Belirtilmemiş", "sinif": sinif_bul(metin(r["sinif"]), metin(r["tur"])), "limit": limit or Z, "risk": risk,
                         "gun": max(int(gun), 0) if gun is not None else 0, "geciken": para(r["geciken"]) or Z})
    if len({s["donem"] is None for s in satirlar}) > 1:
        uy.append(uyari("Orta", "Bazı satırlarda dönem boş", yol.name, "Dönemi boş satırlar ayrı bir dönem gibi değerlendirilir; dönem sütununu ya tümüyle doldurun ya boş bırakın."))
    olumsuz = []
    if olumsuz_yol:
        for r in kayitlar(olumsuz_yol, OLUMSUZ_SUTUN, ("tur",)):
            if not metin(r["tur"]):
                continue
            t = tarih(r["tarih"])
            if metin(r["tarih"]) and not t:
                uy.append(uyari("Orta", "Olumsuz kayıt tarihi okunamadı", metin(r["tur"]), f"{olumsuz_yol.name} satır {r['_satir']}: '{metin(r['tarih'])}'."))
            olumsuz.append({"tur": metin(r["tur"]), "tarih": t, "tutar": para(r["tutar"]), "durum": metin(r["durum"]), "aciklama": metin(r["aciklama"]),
                            "acik": not any(k in katla(r["durum"]) for k in KAPALI)})
    return satirlar, olumsuz, uy


def topla(satirlar: list[dict]) -> dict:
    t = {"limit": Z, "risk": Z, "nakdi": Z, "gayrinakdi": Z, "nakdi_limit": Z, "geciken": Z, "gun": 0}
    for s in satirlar:
        t["limit"] += s["limit"]
        t["risk"] += s["risk"]
        t["nakdi" if s["sinif"] == "Nakdi" else "gayrinakdi"] += s["risk"]
        t["geciken"] += s["geciken"]
        t["gun"] = max(t["gun"], s["gun"])
    t["kullanim"] = oran(t["risk"], t["limit"])
    t["banka"] = len({katla(s["banka"]) for s in satirlar if s["risk"] > 0 or s["limit"] > 0})
    t["riskli_banka"] = len({katla(s["banka"]) for s in satirlar if s["risk"] > 0})
    return t


def analiz(satirlar: list[dict], olumsuz: list[dict] | None = None, kullanim_esik: Decimal = Decimal(90), artis_esik: Decimal = Decimal(25), yogunlasma_esik: Decimal = Decimal(50)) -> dict:
    if not satirlar:
        raise ValueError("Özetlenecek satır yok.")
    donemler = sorted({s["donem"] for s in satirlar}, key=lambda d: d or (0, 0))
    son = [s for s in satirlar if s["donem"] == donemler[-1]]
    t, b = topla(son), []
    bankalar = defaultdict(list)
    for s in son:
        bankalar[katla(s["banka"])].append(s)
    for l in bankalar.values():
        x = topla(l) | {"banka": l[0]["banka"], "satir": len(l)}
        x["pay"] = oran(x["risk"], t["risk"])
        bankalar_asim = [s for s in l if s["limit"] > 0 and s["risk"] > s["limit"]]
        x["asim"] = sum((s["risk"] - s["limit"] for s in bankalar_asim), Z)
        b.append(x)
    b.sort(key=lambda x: (-x["risk"], x["banka"]))
    turler = defaultdict(list)
    for s in son:
        turler[(s["sinif"], s["tur"])].append(s)
    tur = sorted(({"sinif": k[0], "tur": k[1]} | topla(l) | {"pay": oran(sum(s["risk"] for s in l), t["risk"])} for k, l in turler.items()), key=lambda x: (x["sinif"] != "Nakdi", -x["risk"]))
    seyir, onceki = [], None
    for d in donemler:
        x = topla([s for s in satirlar if s["donem"] == d]) | {"donem": d}
        x["degisim"] = oran(x["risk"] - onceki["risk"], onceki["risk"]) if onceki else None
        seyir.append(x)
        onceki = x
    bulgu = []
    for s in son:
        if s["gun"] > 0 or s["geciken"] > 0:
            bulgu.append(("Yüksek", "Gecikmede kredi var", s["banka"], f"{s['tur']}: {s['gun']} gün gecikme, geciken tutar {tl(s['geciken'])}."))
        if s["limit"] > 0 and s["risk"] > s["limit"]:
            bulgu.append(("Yüksek", "Limit aşımı", s["banka"], f"{s['tur']}: risk {tl(s['risk'])} > limit {tl(s['limit'])} (aşım {tl(s['risk'] - s['limit'])})."))
        if s["limit"] == 0 and s["risk"] > 0:
            bulgu.append(("Bilgi", "Limitsiz risk", s["banka"], f"{s['tur']}: risk {tl(s['risk'])}, limit girilmemiş veya 0. Kullanım oranına bu satır yalnız risk tarafıyla girer."))
    for s in satirlar:
        if s["donem"] != donemler[-1] and (s["gun"] > 0 or s["geciken"] > 0):
            bulgu.append(("Orta", "Önceki dönemde gecikme", s["banka"], f"{donem_ad(s['donem'])} · {s['tur']}: {s['gun']} gün, geciken tutar {tl(s['geciken'])}."))
    if t["kullanim"] is not None and t["kullanim"] >= kullanim_esik:
        bulgu.append(("Orta", "Limit kullanımı yüksek", "Toplam", f"Toplam risk / toplam limit %{t['kullanim']} (eşik %{kullanim_esik}). Boş limit {tl(max(t['limit'] - t['risk'], Z))}."))
    if t["riskli_banka"] > 1 and b[0]["pay"] is not None and b[0]["pay"] >= yogunlasma_esik:
        bulgu.append(("Bilgi", "Risk tek bankada yoğunlaşmış", b[0]["banka"], f"Toplam riskin %{b[0]['pay']} kadarı bu bankada (eşik %{yogunlasma_esik})."))
    if len(seyir) > 1:
        s1, s0, ilk = seyir[-1], seyir[-2], seyir[0]
        if s1["degisim"] is not None and s1["degisim"] >= artis_esik:
            bulgu.append(("Orta", "Toplam risk hızlı arttı", donem_ad(s1["donem"]), f"{donem_ad(s0['donem'])} döneminde {tl(s0['risk'])} olan toplam risk {tl(s1['risk'])} oldu (%{s1['degisim']} artış; eşik %{artis_esik})."))
        if s1["riskli_banka"] - ilk["riskli_banka"] >= 2:
            bulgu.append(("Bilgi", "Riskli banka sayısı arttı", donem_ad(s1["donem"]), f"{donem_ad(ilk['donem'])} döneminde {ilk['riskli_banka']} olan riskli banka sayısı {s1['riskli_banka']} oldu."))
    for o in olumsuz or []:
        ek = (f"{o['tarih']:%d.%m.%Y} · " if o["tarih"] else "") + (f"{tl(o['tutar'])} · " if o["tutar"] is not None else "") + (o["durum"] or "durum belirtilmemiş")
        bulgu.append(("Yüksek" if o["acik"] else "Orta", "Olumsuz kayıt (açık)" if o["acik"] else "Olumsuz kayıt (kapanmış)", o["tur"], ek + (f" · {o['aciklama']}" if o["aciklama"] else "")))
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    bulgu.sort(key=lambda x: sira[x[0]])
    return {"donemler": donemler, "son": donemler[-1], "toplam": t, "bankalar": b, "turler": tur, "seyir": seyir, "bulgular": bulgu, "olumsuz": olumsuz or []}


def tl(x: Decimal) -> str:
    return f"{yuvarla(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, t = Workbook(), r["toplam"]
    oz = wb.active
    oz.title = "Özet"
    _baslik(oz, ["Gösterge", "Değer", "Açıklama"], (34, 20, 80))
    for ad, v, ac in (("Dönem", donem_ad(r["son"]), "Dosyadaki en son dönem" if len(r["donemler"]) > 1 else ""),
                      ("Çalışılan banka sayısı", t["banka"], "Limiti veya riski olan"), ("Riski olan banka sayısı", t["riskli_banka"], ""),
                      ("Toplam limit", t["limit"], ""), ("Toplam risk", t["risk"], ""), ("Nakdi risk", t["nakdi"], ""), ("Gayrinakdi risk", t["gayrinakdi"], ""),
                      ("Limit kullanım oranı %", t["kullanim"], "Toplam risk ÷ toplam limit"), ("Boş limit", max(t["limit"] - t["risk"], Z), "Aşım varsa 0 gösterilir"),
                      ("Geciken tutar", t["geciken"], ""), ("En uzun gecikme (gün)", t["gun"], ""),
                      ("Bulgu sayısı (yüksek / orta / bilgi)", " / ".join(str(sum(b[0] == d for b in r["bulgular"])) for d in ("Yüksek", "Orta", "Bilgi")), "Ayrıntı: Bulgular sayfası")):
        _satir(oz, [ad, v, ac], para_s=(2,) if isinstance(v, Decimal) else ())
    oz.append([])
    oz.append(["Bu özet, tabloya aktarılan satırların toplamıdır; Risk Merkezi / KKB raporunun yerine geçmez ve kredi kararı değildir."])
    ws = wb.create_sheet("Banka Bazında")
    _baslik(ws, ["Banka", "Satır", "Limit", "Risk", "Nakdi Risk", "Gayrinakdi Risk", "Kullanım %", "Toplam Riskteki Pay %", "Limit Aşımı", "Geciken Tutar", "Gecikme (gün)"],
            (26, 7, 18, 18, 18, 18, 11, 12, 16, 16, 11))
    for x in r["bankalar"]:
        _satir(ws, [x["banka"], x["satir"], x["limit"], x["risk"], x["nakdi"], x["gayrinakdi"], x["kullanim"], x["pay"], x["asim"] or None, x["geciken"] or None, x["gun"] or None],
               para_s=(3, 4, 5, 6, 7, 8, 9, 10), dolgu={9: KIRMIZI if x["asim"] else None, 10: KIRMIZI if x["geciken"] or x["gun"] else None, 11: KIRMIZI if x["gun"] else None})
    ws = wb.create_sheet("Tür Bazında")
    _baslik(ws, ["Sınıf", "Kredi Türü", "Limit", "Risk", "Kullanım %", "Toplam Riskteki Pay %", "Banka Sayısı"], (14, 30, 18, 18, 11, 12, 10))
    for x in r["turler"]:
        _satir(ws, [x["sinif"], x["tur"], x["limit"], x["risk"], x["kullanim"], x["pay"], x["banka"]], para_s=(3, 4, 5, 6))
    ws = wb.create_sheet("Dönemsel Seyir")
    _baslik(ws, ["Dönem", "Banka Sayısı", "Riski Olan Banka", "Toplam Limit", "Toplam Risk", "Nakdi Risk", "Gayrinakdi Risk", "Kullanım %", "Riskte Değişim %", "Geciken Tutar"],
            (12, 10, 10, 18, 18, 18, 18, 11, 12, 16))
    for x in r["seyir"]:
        _satir(ws, [donem_ad(x["donem"]), x["banka"], x["riskli_banka"], x["limit"], x["risk"], x["nakdi"], x["gayrinakdi"], x["kullanim"], x["degisim"], x["geciken"] or None],
               para_s=(4, 5, 6, 7, 8, 9, 10), dolgu={10: KIRMIZI if x["geciken"] else None})
    ws = wb.create_sheet("Bulgular")
    _baslik(ws, ["Düzey", "Bulgu", "Banka / Konu", "Açıklama", "Değerlendirme (siz doldurun)"], (9, 30, 26, 90, 40))
    for d, bas, kim, ac in r["bulgular"]:
        _satir(ws, [d, bas, kim, ac, ""], dolgu={1: RENK[d]}, sar=(4,))
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, olumsuz_yol: Path | None = None, **esik) -> dict:
    satirlar, olumsuz, uy = oku(yol, olumsuz_yol)
    r = analiz(satirlar, olumsuz, **esik)
    r["uyarilar"] = sirala(uy)
    rapor_yaz(cikti, r, uy)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Risk Merkezi / KKB raporundan aktarılan limit-risk satırlarını özetler.")
    p.add_argument("--rapor", type=Path, default=ORNEK / "risk_raporu.csv", help="Dönem, Banka, Kredi Türü, Nakdi/Gayrinakdi, Limit, Risk, Gecikme Gün, Geciken Tutar")
    p.add_argument("--olumsuz", type=Path, help="İsteğe bağlı olumsuz kayıtlar: Tür, Tarih, Tutar, Durum, Açıklama")
    p.add_argument("--kullanim", default="90", help="Limit kullanım uyarı eşiği %% (varsayılan 90)")
    p.add_argument("--artis", default="25", help="Dönemden döneme risk artışı uyarı eşiği %% (varsayılan 25)")
    p.add_argument("--yogunlasma", default="50", help="Tek banka payı bilgi eşiği %% (varsayılan 50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "risk_raporu_ozeti.xlsx")
    a = p.parse_args(argv)
    olumsuz = a.olumsuz or (ORNEK / "olumsuz_kayitlar.csv" if a.rapor == ORNEK / "risk_raporu.csv" else None)
    for y in (a.rapor, olumsuz):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    e = [para(x) for x in (a.kullanim, a.artis, a.yogunlasma)]
    if any(x is None or x <= 0 for x in e):
        print("[X] --kullanim, --artis ve --yogunlasma pozitif sayı olmalı")
        return 2
    try:
        r = calistir(a.rapor, a.cikti, olumsuz, kullanim_esik=e[0], artis_esik=e[1], yogunlasma_esik=e[2])
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] Dönem {donem_ad(r['son'])} · {t['banka']} banka · limit {tl(t['limit'])} · risk {tl(t['risk'])} · kullanım %{t['kullanim'] if t['kullanim'] is not None else '—'}")
    for d, bas, kim, ac in r["bulgular"]:
        if d == "Yüksek":
            print(f"[!] {bas} · {kim}: {ac}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
