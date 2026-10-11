import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def s(banka="B1", risk=100, limit=200, donem=(2026, 9), tur="Rotatif", sinif="Nakdi", gun=0, geciken=0):
    return {"donem": donem, "banka": banka, "tur": tur, "sinif": sinif, "limit": D(limit), "risk": D(risk), "gun": gun, "geciken": D(geciken)}


def basliklar(r, duzey=None):
    return [b[1] for b in r["bulgular"] if duzey in (None, b[0])]


class HesapTesti(unittest.TestCase):
    def test_toplam_ve_banka(self):
        r = main.analiz([s("B1", 300, 400), s("B1", 100, 100, tur="Teminat Mektubu", sinif="Gayrinakdi"), s("B2", 100, 500)])
        t = r["toplam"]
        self.assertEqual((t["limit"], t["risk"], t["nakdi"], t["gayrinakdi"], t["kullanim"], t["banka"]), (D(1000), D(500), D(400), D(100), D("50.00"), 2))
        b1, b2 = r["bankalar"]
        self.assertEqual((b1["banka"], b1["risk"], b1["kullanim"], b1["pay"], b1["satir"]), ("B1", D(400), D("80.00"), D("80.00"), 2))
        self.assertEqual((b2["kullanim"], b2["pay"]), (D("20.00"), D("20.00")))
        self.assertEqual(basliklar(r), ["Risk tek bankada yoğunlaşmış"])
        self.assertEqual([(x["sinif"], x["tur"], x["risk"]) for x in r["turler"]], [("Nakdi", "Rotatif", D(400)), ("Gayrinakdi", "Teminat Mektubu", D(100))])

    def test_gecikme_asim_limitsiz(self):
        r = main.analiz([s("B1", 120, 100), s("B2", 50, 0), s("B3", 10, 500, gun=12, geciken=4)])
        self.assertEqual(basliklar(r), ["Limit aşımı", "Gecikmede kredi var", "Limitsiz risk", "Risk tek bankada yoğunlaşmış"])
        self.assertEqual((r["bankalar"][0]["asim"], r["toplam"]["geciken"], r["toplam"]["gun"]), (D(20), D(4), 12))
        self.assertEqual(basliklar(main.analiz([s("B1", 100, 100)])), ["Limit kullanımı yüksek"])       # tam limit: aşım değil, kullanım %100
        self.assertEqual(basliklar(main.analiz([s("B1", 89, 100)])), [])
        self.assertEqual(basliklar(main.analiz([s("B1", 89, 100)], kullanim_esik=D(89))), ["Limit kullanımı yüksek"])

    def test_donemsel_seyir(self):
        l = [s("B1", 100, 1000, (2026, 7)), s("B1", 100, 1000, (2026, 8), gun=5), s("B1", 125, 1000), s("B2", 0, 1000), s("B3", 0, 1000, (2026, 7))]
        r = main.analiz(l)
        self.assertEqual([(x["donem"], x["risk"], x["degisim"], x["banka"], x["riskli_banka"]) for x in r["seyir"]],
                         [((2026, 7), D(100), None, 2, 1), ((2026, 8), D(100), D("0.00"), 1, 1), ((2026, 9), D(125), D("25.00"), 2, 1)])
        self.assertEqual(basliklar(r), ["Önceki dönemde gecikme", "Toplam risk hızlı arttı"])
        self.assertEqual(basliklar(main.analiz(l, artis_esik=D(26))), ["Önceki dönemde gecikme"])
        r = main.analiz([s("B1", 10, 1000, (2026, 7)), s("B1", 10, 1000), s("B2", 1, 1000), s("B3", 1, 1000)])
        self.assertEqual(basliklar(r), ["Risk tek bankada yoğunlaşmış", "Riskli banka sayısı arttı"])

    def test_olumsuz_ve_yardimcilar(self):
        o = [{"tur": "Karşılıksız çek", "tarih": date(2026, 9, 2), "tutar": D(75), "durum": "Açık", "aciklama": "", "acik": True},
             {"tur": "Protesto", "tarih": None, "tutar": None, "durum": "Ödendi", "aciklama": "x", "acik": False}]
        r = main.analiz([s()], o)
        self.assertEqual([(b[0], b[1]) for b in r["bulgular"]], [("Yüksek", "Olumsuz kayıt (açık)"), ("Orta", "Olumsuz kayıt (kapanmış)")])
        self.assertIn("02.09.2026 · 75,00 · Açık", r["bulgular"][0][3])
        self.assertEqual([main.donem(x) for x in ("09.2026", "2026-09", "30.09.2026", "13.2026", "eylül")], [(2026, 9), (2026, 9), (2026, 9), None, None])
        self.assertEqual([main.sinif_bul(a, b) for a, b in (("", "Teminat Mektubu"), ("Gayri Nakdi", "x"), ("", "Rotatif"), ("Nakdi", "Akreditif"))],
                         ["Gayrinakdi", "Gayrinakdi", "Nakdi", "Nakdi"])
        with self.assertRaises(ValueError):
            main.analiz([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            r = main.calistir(ORNEK / "risk_raporu.csv", cikti, ORNEK / "olumsuz_kayitlar.csv")
            t = r["toplam"]
            self.assertEqual((r["son"], t["banka"], t["limit"], t["risk"], t["nakdi"], t["gayrinakdi"], t["kullanim"]),
                             ((2026, 9), 5, D(16000000), D(16000000), D(14200000), D(1800000), D("100.00")))
            self.assertEqual([x["degisim"] for x in r["seyir"]], [None, D("23.66"), D("39.13")])       # 9,3 → 11,5 → 16,0 milyon
            self.assertEqual((r["bankalar"][0]["banka"], r["bankalar"][0]["pay"], r["bankalar"][0]["asim"]), ("Banka 1", D("44.38"), D(300000)))
            self.assertEqual(basliklar(r, "Yüksek"), ["Limit aşımı", "Gecikmede kredi var", "Olumsuz kayıt (açık)"])
            self.assertEqual(sorted(basliklar(r, "Orta")), ["Limit kullanımı yüksek", "Olumsuz kayıt (kapanmış)", "Toplam risk hızlı arttı"])
            self.assertEqual(sorted(basliklar(r, "Bilgi")), ["Limitsiz risk", "Riskli banka sayısı arttı"])
            self.assertEqual(r["uyarilar"], [])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Banka Bazında", "Tür Bazında", "Dönemsel Seyir", "Bulgular", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, o = Path(tmp) / "r.csv", Path(tmp) / "o.csv"
            y.write_text("Dönem;Banka;Kredi Türü;Limit;Risk\n09.2026;B1;Akreditif;1.000,00;400,00\n09.2026;B2;Spot;bin;100\n09.2026;B3;Spot;100;yok\neylül;B4;Spot;1;1\n;;;;\n",
                         encoding="utf-8")
            o.write_text("Tür;Tarih;Tutar;Durum\nProtesto;dün;10;Kapandı\n", encoding="utf-8")
            sat, ol, uy = main.oku(y, o)
            self.assertEqual([(x["banka"], x["sinif"], x["limit"], x["risk"]) for x in sat], [("B1", "Gayrinakdi", D("1000.00"), D("400.00")), ("B2", "Nakdi", D(0), D(100))])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Dönem okunamadı", "Limit okunamadı", "Olumsuz kayıt tarihi okunamadı", "Risk tutarı okunamadı"])
            self.assertEqual(ol[0]["acik"], False)
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--rapor", str(y), "--cikti", c]), 0)
            self.assertEqual(main.main(["--rapor", str(y), "--cikti", c, "--kullanim", "sıfır"]), 2)
            self.assertEqual(main.main(["--rapor", str(Path(tmp) / "yok.csv")]), 1)
            y.write_text("Banka;Risk\n", encoding="utf-8")
            self.assertEqual(main.main(["--rapor", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
