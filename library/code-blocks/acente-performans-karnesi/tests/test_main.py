import sys
import tempfile
import unittest
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def ac(ad="A", bolge="B1", uretim=600, onceki=500, hedef=800, kazanilmis=550, hasar=385, tahakkuk=600, tahsilat=510, yenilenecek=250, yenilenen=175):
    d = dict(uretim=uretim, onceki=onceki, hedef=hedef, kazanilmis=kazanilmis, hasar=hasar, tahakkuk=tahakkuk, tahsilat=tahsilat, yenilenecek=yenilenecek, yenilenen=yenilenen)
    return {"acente": ad, "bolge": bolge} | {k: (None if v is None else D(v)) for k, v in d.items()}


class KarneTesti(unittest.TestCase):
    def test_gostergeler_ve_puan(self):
        a = main.karne([ac()])["acenteler"][0]
        self.assertEqual(a["deger"], {"hedef": D("75.00"), "buyume": D("20.00"), "hp": D("70.00"), "tahsilat": D("85.00"), "yenileme": D("70.00")})
        self.assertEqual(a["puanlar"], {"hedef": D("50.00"), "buyume": D("40.00"), "hp": D("50.00"), "tahsilat": D("50.00"), "yenileme": D("50.00")})
        self.assertEqual((a["puan"], a["nota"], a["zayif"], a["pay"]), (D("48.50"), "D", "Büyüme %", D("100.00")))        # (25×50 + 15×40 + 25×50 + 20×50 + 15×50) ÷ 100

    def test_puan_sinirlari(self):
        o = main.varsayilan()
        p = lambda k, v: main.puanla(D(v), o[k])   # noqa: E731
        self.assertEqual([p("hedef", v) for v in (40, 50, 75, 100, 130)], [D(0), D(0), D(50), D(100), D(100)])
        self.assertEqual([p("hp", v) for v in (120, 100, 70, 40, 10)], [D(0), D(0), D(50), D(100), D(100)])              # düşük olan iyi
        self.assertEqual([p("buyume", v) for v in (-25, 25)], [D(0), D(50)])
        self.assertIsNone(main.puanla(None, o["hedef"]))

    def test_eksik_gosterge_ve_notlar(self):
        a = main.karne([ac(onceki=None, yenilenecek=0, yenilenen=0, hasar=137.5, tahsilat=540)])["acenteler"][0]
        self.assertEqual((a["deger"]["buyume"], a["deger"]["yenileme"], a["eksik"]), (None, None, ["Büyüme %", "Yenileme oranı %"]))
        self.assertEqual((a["puan"], a["nota"]), (D("72.62"), "B"))                                                       # (25×50 + 25×100 + 20×66,67) ÷ 70
        self.assertEqual(main.karne([ac()], notlar=(D(48), D(40), D(30)))["acenteler"][0]["nota"], "A")
        self.assertEqual(main.karne([ac(hedef=None, onceki=None, kazanilmis=None, tahakkuk=None, yenilenecek=None)])["acenteler"][0]["nota"], "—")
        o = main.varsayilan()
        o["buyume"]["agirlik"] = D(0)
        a = main.karne([ac()], o)["acenteler"][0]
        self.assertEqual((a["puan"], a["zayif"]), (D("50.00"), ""))

    def test_siralama_ve_bolge(self):
        r = main.karne([ac("X", "B1"), ac("Y", "B2", uretim=800, onceki=400, hasar=220), ac("Z", "B1", uretim=1200, hasar=220, tahsilat=600, yenilenen=250)])
        self.assertEqual([(a["acente"], a["sira"], a["bolge_sira"]) for a in r["acenteler"]], [("Z", 1, 1), ("Y", 2, 1), ("X", 3, 2)])
        b = {x["ad"]: x for x in r["bolge"]}
        self.assertEqual((b["B1"]["acente"], b["B1"]["uretim"], b["B1"]["deger"]["hedef"], b["B1"]["deger"]["hp"]), (2, D(1800), D("112.50"), D("55.00")))
        self.assertEqual((r["toplam"], r["acenteler"][0]["pay"]), (D(2600), D("46.15")))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "k.xlsx"
            r = main.calistir(ORNEK / "acente_verisi.csv", cikti, ORNEK / "olcutler.csv")
            self.assertEqual([(a["acente"][-2:], a["puan"], a["nota"]) for a in r["acenteler"]],
                             [("A1", D("94.79"), "A"), ("A5", D("86.67"), "A"), ("A6", D("72.62"), "B"), ("A3", D("56.67"), "C"), ("A2", D("48.50"), "D"), ("A4", D("30.00"), "D")])
            a4 = r["acenteler"][-1]
            self.assertEqual((a4["deger"], a4["zayif"]), ({"hedef": D("60.00"), "buyume": D("-25.00"), "hp": D("40.00"), "tahsilat": D("65.00"), "yenileme": D("40.00")}, "Büyüme %"))
            self.assertEqual((r["toplam"], r["acenteler"][0]["pay"], r["uyarilar"]), (D(39000000), D("30.77"), []))
            self.assertEqual(load_workbook(cikti).sheetnames, ["Karne", "Gösterge Puanları", "Bölge Özeti", "Ölçütler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y, o = Path(tmp) / "a.csv", Path(tmp) / "o.csv"
            y.write_text("Acente;Üretim;Hedef;Tahsilat\nA;1.000,00;çok;-5\na;5;;\nB;yok;;\nC;500;1.000;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual([(x["acente"], x["uretim"], x["hedef"], x["tahsilat"], x["bolge"]) for x in l], [("A", D("1000.00"), None, None, "Belirtilmemiş"), ("C", D(500), D(1000), None, "Belirtilmemiş")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Değer okunamadı"] * 3 + ["Mükerrer acente"])
            o.write_text("Gösterge;Ağırlık;Sıfır Puan;Tam Puan\nHedef gerçekleşme %;40;60;110\nH/P;30;90;50\nMemnuniyet;10;0;5\nBüyüme %;5;10;10\n", encoding="utf-8")
            ol, uy = main.olcut_oku(o)
            self.assertEqual((ol["hedef"]["agirlik"], ol["hedef"]["sifir"], ol["hp"]["tam"], ol["buyume"]["agirlik"]), (D(40), D(60), D(50), D(15)))
            self.assertEqual(sorted(u["tur"] for u in uy), ["Gösterge tanınmadı", "Ölçüt okunamadı"])
            c = str(Path(tmp) / "x.xlsx")
            self.assertEqual(main.main(["--acenteler", str(y), "--olcutler", str(o), "--cikti", c, "--notlar", "90,70,40"]), 0)
            self.assertEqual(main.main(["--acenteler", str(y), "--cikti", c, "--notlar", "50,60,70"]), 2)
            self.assertEqual(main.main(["--acenteler", str(Path(tmp) / "yok.csv")]), 1)
            o.write_text("Gösterge;Ağırlık;Sıfır Puan;Tam Puan\n" + "".join(f"{g};0;0;1\n" for g in ("Hedef", "Büyüme", "HP", "Tahsilat", "Yenileme")), encoding="utf-8")
            self.assertEqual(main.main(["--acenteler", str(y), "--olcutler", str(o), "--cikti", c]), 1)
            y.write_text("Acente;Üretim\n", encoding="utf-8")
            self.assertEqual(main.main(["--acenteler", str(y), "--cikti", c]), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
