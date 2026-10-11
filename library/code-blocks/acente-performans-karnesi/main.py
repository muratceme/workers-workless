"""
Acente Performans Karnesi — Workers / Workless kod bloğu
Satış ve Acente Yönetimi › Bölge Satış Yöneticisi

Acente bazında beş göstergeyi hesaplar, puana çevirir ve karşılaştırmalı karne üretir.
  Hedef gerçekleşme % = üretim ÷ hedef                       Büyüme %        = (üretim − önceki dönem üretim) ÷ önceki dönem üretim
  Hasar/prim %        = gerçekleşen hasar ÷ kazanılmış prim   Tahsilat oranı % = tahsilat ÷ tahakkuk (vadesi gelen prim)
  Yenileme oranı %    = yenilenen poliçe ÷ yenilenecek poliçe
  Gösterge puanı = (değer − sıfır puan değeri) ÷ (tam puan değeri − sıfır puan değeri) × 100, 0 ile 100 arasına sınırlanır
  Karne puanı    = Σ(gösterge puanı × ağırlık) ÷ Σ(hesaplanabilen göstergelerin ağırlığı)
  Not            = puan ≥ 80 A, ≥ 65 B, ≥ 50 C, altı D   (--notlar ile değişir)
Ağırlıklar ve puan sınırları örnek varsayılanlardır; ölçütler dosyasıyla kendi karnenizi tanımlarsınız.
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


SUTUN = {"acente": ("acente", "acente adi", "acente kodu"), "bolge": ("bolge", "bolge mudurlugu"), "uretim": ("uretim", "prim uretimi", "yazilan prim", "donem uretim"),
         "onceki": ("onceki donem uretim", "onceki uretim", "gecen yil uretim", "onceki donem"), "hedef": ("hedef", "uretim hedefi"),
         "kazanilmis": ("kazanilmis prim",), "hasar": ("gerceklesen hasar", "hasar"), "tahakkuk": ("tahakkuk", "vadesi gelen prim", "tahakkuk eden prim"),
         "tahsilat": ("tahsilat", "tahsil edilen"), "yenilenecek": ("yenilenecek police", "vadesi gelen police", "yenilenecek"), "yenilenen": ("yenilenen police", "yenilenen")}
# kod: (ad, ağırlık, sıfır puan değeri, tam puan değeri)
OLCUT = {"hedef": ("Hedef gerçekleşme %", 25, 50, 100), "buyume": ("Büyüme %", 15, 0, 50), "hp": ("Hasar/prim %", 25, 100, 40),
         "tahsilat": ("Tahsilat oranı %", 20, 70, 100), "yenileme": ("Yenileme oranı %", 15, 50, 90)}
OLCUT_ES = {"hedef gerceklesme": "hedef", "hedef": "hedef", "buyume": "buyume", "hasar prim": "hp", "h p": "hp", "hp": "hp", "tahsilat orani": "tahsilat", "tahsilat": "tahsilat",
            "yenileme orani": "yenileme", "yenileme": "yenileme"}
Z = Decimal(0)


def varsayilan() -> dict:
    return {k: {"ad": v[0], "agirlik": Decimal(v[1]), "sifir": Decimal(v[2]), "tam": Decimal(v[3])} for k, v in OLCUT.items()}


def olcut_oku(yol: Path) -> tuple[dict, list[dict]]:
    o, uy = varsayilan(), []
    for r in kayitlar(yol, {"ad": ("gosterge", "olcut", "kod"), "agirlik": ("agirlik", "puan"), "sifir": ("sifir puan", "sifir puan degeri", "alt sinir"),
                            "tam": ("tam puan", "tam puan degeri", "ust sinir")}, ("ad",)):
        if not metin(r["ad"]):
            continue
        kod = OLCUT_ES.get(katla(re.sub(r"%", "", metin(r["ad"]))))
        a, s, t = para(r["agirlik"]), para(r["sifir"]), para(r["tam"])
        if not kod:
            uy.append(uyari("Orta", "Gösterge tanınmadı", metin(r["ad"]), f"{yol.name} satır {r['_satir']}: geçerli göstergeler {', '.join(v[0] for v in OLCUT.values())}."))
        elif a is None or a < 0 or s is None or t is None or s == t:
            uy.append(uyari("Yüksek", "Ölçüt okunamadı", metin(r["ad"]), f"{yol.name} satır {r['_satir']}: ağırlık ≥ 0, sıfır ve tam puan değerleri farklı sayılar olmalı. Varsayılan kullanıldı."))
        else:
            o[kod] |= {"agirlik": a, "sifir": s, "tam": t}
    if not any(v["agirlik"] > 0 for v in o.values()):
        raise ValueError(f"{yol.name}: tüm ağırlıklar 0; en az bir gösterge ağırlıklı olmalı.")
    return o, uy


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l, gorulen = [], [], set()
    for r in kayitlar(yol, SUTUN, ("acente", "uretim")):
        ad = metin(r["acente"])
        if not ad:
            continue
        kim = f"{ad} (satır {r['_satir']})"
        if katla(ad) in gorulen:
            uy.append(uyari("Yüksek", "Mükerrer acente", kim, "Acente tekrar ediyor; ilk satır kullanıldı."))
            continue
        a = {"acente": ad, "bolge": metin(r["bolge"]) or "Belirtilmemiş"}
        for alan in ("uretim", "onceki", "hedef", "kazanilmis", "hasar", "tahakkuk", "tahsilat", "yenilenecek", "yenilenen"):
            a[alan] = para(r[alan])
            if a[alan] is None and metin(r[alan]) or (a[alan] is not None and a[alan] < 0):
                uy.append(uyari("Yüksek", "Değer okunamadı", kim, f"{SUTUN[alan][0]}: '{metin(r[alan])}'. Boş sayıldı; ilgili gösterge hesaplanmadı."))
                a[alan] = None
        if a["uretim"] is None:
            continue
        gorulen.add(katla(ad))
        l.append(a)
    return l, uy


def yuzde(a, b) -> Decimal | None:
    return None if a is None or b is None or b == 0 else a / b * 100


def puanla(deger: Decimal | None, o: dict) -> Decimal | None:
    if deger is None:
        return None
    return min(max((deger - o["sifir"]) / (o["tam"] - o["sifir"]), Z), Decimal(1)) * 100


def karne(acenteler: list[dict], olcut: dict | None = None, notlar: tuple = (Decimal(80), Decimal(65), Decimal(50))) -> dict:
    olcut = olcut or varsayilan()
    toplam = sum((a["uretim"] for a in acenteler), Z)
    for a in acenteler:
        d = {"hedef": yuzde(a["uretim"], a["hedef"]), "buyume": yuzde(a["uretim"] - a["onceki"], a["onceki"]) if a["onceki"] else None, "hp": yuzde(a["hasar"], a["kazanilmis"]),
             "tahsilat": yuzde(a["tahsilat"], a["tahakkuk"]), "yenileme": yuzde(a["yenilenen"], a["yenilenecek"])}
        p = {k: puanla(v, olcut[k]) for k, v in d.items()}
        agirlik = sum((olcut[k]["agirlik"] for k, v in p.items() if v is not None), Z)
        a["deger"] = {k: yuvarla(v) if v is not None else None for k, v in d.items()}
        a["puanlar"] = {k: yuvarla(v) if v is not None else None for k, v in p.items()}
        a["eksik"] = [olcut[k]["ad"] for k, v in p.items() if v is None and olcut[k]["agirlik"] > 0]
        a["puan"] = yuvarla(sum((v * olcut[k]["agirlik"] for k, v in p.items() if v is not None), Z) / agirlik) if agirlik else None
        a["nota"] = "—" if a["puan"] is None else "A" if a["puan"] >= notlar[0] else "B" if a["puan"] >= notlar[1] else "C" if a["puan"] >= notlar[2] else "D"
        a["pay"] = yuvarla(a["uretim"] / toplam * 100) if toplam else None
        zayif = [(v, k) for k, v in p.items() if v is not None and olcut[k]["agirlik"] > 0]
        a["zayif"] = olcut[min(zayif)[1]]["ad"] if zayif and min(zayif)[0] < 50 else ""
    acenteler.sort(key=lambda a: (a["puan"] is None, -(a["puan"] or Z), a["acente"]))
    for i, a in enumerate(acenteler, 1):
        a["sira"] = i
    for b in {a["bolge"] for a in acenteler}:
        for i, a in enumerate([a for a in acenteler if a["bolge"] == b], 1):
            a["bolge_sira"] = i
    bolge = {}
    for a in acenteler:
        x = bolge.setdefault(a["bolge"], {"ad": a["bolge"], "acente": 0} | {k: Z for k in ("uretim", "onceki", "hedef", "kazanilmis", "hasar", "tahakkuk", "tahsilat", "yenilenecek", "yenilenen")}
                             | {"puanlar": []})
        x["acente"] += 1
        for k in ("uretim", "onceki", "hedef", "kazanilmis", "hasar", "tahakkuk", "tahsilat", "yenilenecek", "yenilenen"):
            x[k] += a[k] or Z
        if a["puan"] is not None:
            x["puanlar"].append(a["puan"])
    for x in bolge.values():
        x["deger"] = {k: yuvarla(v) if v is not None else None for k, v in
                      {"hedef": yuzde(x["uretim"], x["hedef"]), "buyume": yuzde(x["uretim"] - x["onceki"], x["onceki"]), "hp": yuzde(x["hasar"], x["kazanilmis"]),
                       "tahsilat": yuzde(x["tahsilat"], x["tahakkuk"]), "yenileme": yuzde(x["yenilenen"], x["yenilenecek"])}.items()}
        x["ort_puan"] = yuvarla(sum(x["puanlar"], Z) / len(x["puanlar"])) if x["puanlar"] else None
    return {"acenteler": acenteler, "bolge": sorted(bolge.values(), key=lambda x: -x["uretim"]), "olcut": olcut, "toplam": toplam, "notlar": notlar}


def rapor_yaz(cikti: Path, r: dict, uy: list[dict]) -> None:
    wb, o = Workbook(), r["olcut"]
    renk = {"A": YESIL, "B": MAVI, "C": SARI, "D": KIRMIZI, "—": None}
    kodlar = list(OLCUT)
    ws = wb.active
    ws.title = "Karne"
    _baslik(ws, ["Sıra", "Acente", "Bölge", "Bölge Sırası", "Üretim", "Üretim Payı %", *[o[k]["ad"] for k in kodlar], "Karne Puanı", "Not", "En Zayıf Gösterge", "Hesaplanamayan",
                 "Aksiyon (siz doldurun)"], (6, 30, 16, 8, 18, 9, 12, 10, 10, 10, 10, 10, 7, 24, 30, 30))
    for a in r["acenteler"]:
        _satir(ws, [a["sira"], a["acente"], a["bolge"], a["bolge_sira"], a["uretim"], a["pay"], *[a["deger"][k] for k in kodlar], a["puan"], a["nota"], a["zayif"], ", ".join(a["eksik"]), ""],
               para_s=(5, 6, 7, 8, 9, 10, 11, 12), dolgu={13: renk[a["nota"]]}, sar=(15,))
    ws = wb.create_sheet("Gösterge Puanları")
    _baslik(ws, ["Acente", *[f"{o[k]['ad']} · puan" for k in kodlar], "Karne Puanı", "Not"], (30, 16, 14, 14, 14, 14, 11, 7))
    for a in r["acenteler"]:
        _satir(ws, [a["acente"], *[a["puanlar"][k] for k in kodlar], a["puan"], a["nota"]], para_s=(2, 3, 4, 5, 6, 7),
               dolgu={j: KIRMIZI for j, k in enumerate(kodlar, 2) if a["puanlar"][k] is not None and a["puanlar"][k] < 50} | {8: renk[a["nota"]]})
    ws = wb.create_sheet("Bölge Özeti")
    _baslik(ws, ["Bölge", "Acente", "Üretim", *[o[k]["ad"] for k in kodlar], "Ortalama Karne Puanı"], (20, 8, 18, 12, 10, 10, 10, 10, 12))
    for x in r["bolge"]:
        _satir(ws, [x["ad"], x["acente"], x["uretim"], *[x["deger"][k] for k in kodlar], x["ort_puan"]], para_s=(3, 4, 5, 6, 7, 8, 9))
    ws = wb.create_sheet("Ölçütler")
    _baslik(ws, ["Gösterge", "Ağırlık", "Sıfır Puan Değeri", "Tam Puan Değeri"], (26, 10, 16, 16))
    for k in kodlar:
        _satir(ws, [o[k]["ad"], o[k]["agirlik"], o[k]["sifir"], o[k]["tam"]])
    ws.append([])
    ws.append([f"Not sınırları: A ≥ {r['notlar'][0]}, B ≥ {r['notlar'][1]}, C ≥ {r['notlar'][2]}, altı D. Hesaplanamayan göstergenin ağırlığı diğerlerine dağıtılır."])
    uyari_sayfasi(wb, sirala(uy))
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, olcut_yol: Path | None = None, notlar: tuple = (Decimal(80), Decimal(65), Decimal(50))) -> dict:
    olcut, uy1 = olcut_oku(olcut_yol) if olcut_yol else (varsayilan(), [])
    acenteler, uy2 = oku(yol)
    if not acenteler:
        raise ValueError(f"{yol.name}: karnesi çıkarılacak acente yok.")
    r = karne(acenteler, olcut, notlar)
    r["uyarilar"] = sirala(uy1 + uy2)
    rapor_yaz(cikti, r, uy1 + uy2)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Acente bazında üretim, büyüme, hasar/prim, tahsilat ve yenileme göstergelerinden karne üretir.")
    p.add_argument("--acenteler", type=Path, default=ORNEK / "acente_verisi.csv",
                   help="Acente, Bölge, Üretim, Önceki Dönem Üretim, Hedef, Kazanılmış Prim, Gerçekleşen Hasar, Tahakkuk, Tahsilat, Yenilenecek Poliçe, Yenilenen Poliçe")
    p.add_argument("--olcutler", type=Path, help="İsteğe bağlı: Gösterge, Ağırlık, Sıfır Puan, Tam Puan")
    p.add_argument("--notlar", default="80,65,50", help="A, B, C için asgari puanlar (varsayılan 80,65,50)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "acente_karnesi.xlsx")
    a = p.parse_args(argv)
    for y in (a.acenteler, a.olcutler):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    n = [para(x) for x in a.notlar.split(",")]
    if len(n) != 3 or None in n or not 100 >= n[0] > n[1] > n[2] >= 0:
        print("[X] --notlar azalan üç sayı olmalı (ör. 80,65,50)")
        return 2
    try:
        r = calistir(a.acenteler, a.cikti, a.olcutler, tuple(n))
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['acenteler'])} acente · toplam üretim {r['toplam']:,.2f} · " + " · ".join(f"{x} {sum(a['nota'] == x for a in r['acenteler'])}" for x in "ABCD"))
    for x in r["acenteler"]:
        print(f"     {x['sira']:>2}. {x['acente']}: {x['puan'] if x['puan'] is not None else '—'} ({x['nota']})" + (f" · zayıf: {x['zayif']}" if x["zayif"] else ""))
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
