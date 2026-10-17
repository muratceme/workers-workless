"""
Müşteri Tahsilat Takibi — Gerisi Bende kod bloğu
Teknik ve Operasyon › Acente Yöneticisi

Müşterilerden tahsil edilecek prim taksitlerini vadeye göre, sigorta şirketine aktarılacak tutarlarla birlikte izler.
  Müşteriden kalan        = tutar − tahsil edilen;   gecikmiş = kalan > 0 ve vade < durum tarihi;   yaklaşan = vadeye --gun (7) gün veya az
  Şirkete kalan borç      = şirkete aktarılacak tutar − şirkete ödenen   (aktarılacak tutar boşsa taksit tutarı)
  Vadesi geçmiş şirket borcu = şirkete ödeme vadesi < durum tarihi olan kalan borç
  Tahsil edilip aktarılmayan = min(tahsil edilen, aktarılacak tutar) − şirkete ödenen   (sıfırdan büyükse)
  Acentenin finanse ettiği   = şirkete ödenen − tahsil edilen                            (sıfırdan büyükse)
  Yaşlandırma             = gecikmiş müşteri alacağı 1–30, 31–60, 61–90, 90+ gün
Aktarılacak tutarı (komisyon düşülmüş net tutar) siz yazarsınız; kod komisyon hesaplamaz, cari hesap mutabakatı yapmaz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"police": ("police no", "police numarasi", "police"), "musteri": ("musteri", "sigortali", "sigorta ettiren"), "sirket": ("sigorta sirketi", "sirket"), "vade": ("vade", "vade tarihi", "musteri vadesi"),
         "tutar": ("tutar", "taksit tutari", "prim"), "tahsil": ("tahsil edilen", "tahsilat", "odenen"), "s_vade": ("sirkete odeme vadesi", "sirket vadesi", "sirkete odeme tarihi"),
         "s_tutar": ("sirkete aktarilacak", "sirkete aktarilacak tutar", "sirkete odenecek"), "s_odenen": ("sirkete odenen", "sirkete aktarilan")}
YAS = ((30, "1–30"), (60, "31–60"), (90, "61–90"), (None, "90+"))
DURUM_RENK = {"Gecikmiş": KIRMIZI, "Yaklaşıyor": SARI, "Vadesi gelmedi": None, "Tahsil edildi": YESIL}
SIFIR = Decimal(0)


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("musteri", "vade", "tutar")):
        musteri = metin(r["musteri"])
        if not musteri:
            continue
        kim = f"{metin(r['police']) or 'Satır ' + str(r['_satir'])} · {musteri}"
        vade, d = tarih(r["vade"]), {a: para(r[a]) if metin(r[a]) else None for a in ("tutar", "tahsil", "s_tutar", "s_odenen")}
        if vade is None or d["tutar"] is None or d["tutar"] <= 0 or any(metin(r[a]) and (d[a] is None or d[a] < 0) for a in d):
            uy.append(uyari("Yüksek", "Taksit okunamadı", kim, f"Satır {r['_satir']}: vade ya da tutarlardan biri okunamadı. Satır alınmadı."))
            continue
        if metin(r["s_vade"]) and tarih(r["s_vade"]) is None:
            uy.append(uyari("Orta", "Tarih okunamadı", kim, f"Şirkete ödeme vadesi '{metin(r['s_vade'])}' okunamadı; boş sayıldı."))
        l.append({"police": metin(r["police"]), "musteri": musteri, "sirket": metin(r["sirket"]) or "(şirket yok)", "vade": vade, "tutar": d["tutar"], "tahsil": d["tahsil"] or SIFIR, "s_vade": tarih(r["s_vade"]),
                  "s_tutar": d["s_tutar"] if d["s_tutar"] is not None else d["tutar"], "s_odenen": d["s_odenen"] or SIFIR})
    return l, uy


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    return yuvarla(pay * 100 / payda) if payda else None


def degerlendir(taksitler: list[dict], bugun: date, gun: int = 7, gecikme: int = 30) -> dict:
    if not taksitler:
        raise ValueError("İzlenecek taksit yok.")
    uy = []
    for t in taksitler:
        kim = f"{t['police'] or '(poliçe yok)'} · {t['musteri']} · vade {t['vade']:%d.%m.%Y}"
        t["kalan"] = max(t["tutar"] - t["tahsil"], SIFIR)
        kalan_gun = (t["vade"] - bugun).days
        t["durum"] = "Tahsil edildi" if t["kalan"] == 0 else "Gecikmiş" if kalan_gun < 0 else "Yaklaşıyor" if kalan_gun <= gun else "Vadesi gelmedi"
        t["gecikme"] = -kalan_gun if t["durum"] == "Gecikmiş" else None
        t["yas"] = next(e for s, e in YAS if s is None or t["gecikme"] <= s) if t["gecikme"] else None
        t["s_kalan"] = max(t["s_tutar"] - t["s_odenen"], SIFIR)
        s_gun = (t["s_vade"] - bugun).days if t["s_vade"] else None
        t["s_gecmis"] = t["s_kalan"] if s_gun is not None and s_gun < 0 else SIFIR
        t["s_yakin"] = t["s_kalan"] if s_gun is not None and 0 <= s_gun <= gun else SIFIR
        t["emanet"] = max(min(t["tahsil"], t["s_tutar"]) - t["s_odenen"], SIFIR)
        t["finanse"] = max(t["s_odenen"] - t["tahsil"], SIFIR)
        if t["tahsil"] > t["tutar"]:
            uy.append(uyari("Orta", "Tahsilat taksitten büyük", kim, f"Taksit {t['tutar']:,.2f}, tahsil edilen {t['tahsil']:,.2f}. Fazla tahsilat başka taksite mahsup edilmedi."))
        if t["emanet"] > 0 and t["s_gecmis"] > 0:
            uy.append(uyari("Yüksek", "Tahsil edilen prim şirkete aktarılmamış", kim, f"Müşteriden {t['tahsil']:,.2f} tahsil edildi; {t['sirket']} için {t['emanet']:,.2f} aktarılmadı. Şirkete ödeme vadesi {t['s_vade']:%d.%m.%Y} geçti."))
        if t["s_gecmis"] - t["emanet"] > 0:
            uy.append(uyari("Orta", "Şirket vadesi geçti, müşteri ödemedi", kim, f"{t['sirket']} için ödeme vadesi {t['s_vade']:%d.%m.%Y} geçti; kalan borç {t['s_kalan']:,.2f}, müşteriden kalan {t['kalan']:,.2f}."))
        if t["finanse"] > 0:
            uy.append(uyari("Orta", "Tahsil edilmeden şirkete ödenmiş", kim, f"Şirkete {t['s_odenen']:,.2f} ödendi, müşteriden {t['tahsil']:,.2f} tahsil edildi; {t['finanse']:,.2f} acente kaynağından karşılandı."))

    def topla(l: list[dict], alan: str, kosul=lambda t: True) -> Decimal:
        return sum((t[alan] for t in l if kosul(t)), SIFIR)

    def m_ozet(l: list[dict]) -> dict:
        gec = [t for t in l if t["durum"] == "Gecikmiş"]
        return {"taksit": len(l), "tutar": topla(l, "tutar"), "tahsil": topla(l, "tahsil"), "kalan": topla(l, "kalan"), "gecikmis": topla(gec, "kalan"), "enb": max((t["gecikme"] for t in gec), default=None),
                "yaklasan": topla(l, "kalan", lambda t: t["durum"] == "Yaklaşıyor"), "oran": _oran(topla(l, "tahsil"), topla(l, "tutar"))}

    def s_ozet(l: list[dict]) -> dict:
        return {"taksit": len(l), "aktarilacak": topla(l, "s_tutar"), "odenen": topla(l, "s_odenen"), "kalan": topla(l, "s_kalan"), "gecmis": topla(l, "s_gecmis"), "yakin": topla(l, "s_yakin"), "emanet": topla(l, "emanet"),
                "finanse": topla(l, "finanse")}

    def kirilim(alan: str, ozet, sira) -> list[dict]:
        grup = {}
        for t in taksitler:
            grup.setdefault(katla(t[alan]), []).append(t)
        return sorted(({"ad": g[0][alan], **ozet(g)} for g in grup.values()), key=sira)

    musteriler = kirilim("musteri", m_ozet, lambda s: (-s["gecikmis"], -s["kalan"], katla(s["ad"])))
    for m in musteriler:
        if m["gecikmis"] > 0:
            uy.append(uyari("Yüksek" if m["enb"] >= gecikme else "Orta", "Gecikmiş tahsilat", m["ad"], f"Gecikmiş tutar {m['gecikmis']:,.2f}; en eski taksit {m['enb']} gün gecikmiş. Toplam kalan {m['kalan']:,.2f}."))
    oz = {**m_ozet(taksitler), **{f"s_{k}": v for k, v in s_ozet(taksitler).items() if k != "taksit"}}
    return {"taksitler": sorted(taksitler, key=lambda t: (list(DURUM_RENK).index(t["durum"]), t["vade"], katla(t["musteri"]))), "musteriler": musteriler,
            "sirketler": kirilim("sirket", s_ozet, lambda s: (-s["gecmis"], -s["kalan"], katla(s["ad"]))), "ozet": oz,
            "yaslar": [{"dilim": e, "taksit": sum(1 for t in taksitler if t["yas"] == e), "tutar": topla(taksitler, "kalan", lambda t: t["yas"] == e)} for _, e in YAS], "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date, gun: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Taksitler"
    _baslik(ws, ["Durum", "Müşteri", "Poliçe No", "Sigorta Şirketi", "Vade", "Tutar", "Tahsil Edilen", "Kalan", "Gecikme (gün)", "Şirkete Ödeme Vadesi", "Şirkete Aktarılacak", "Şirkete Ödenen", "Şirkete Kalan Borç",
                 "Tahsil Edilip Aktarılmayan", "Acentenin Finanse Ettiği", "Not (siz doldurun)"], (15, 28, 14, 22, 12, 14, 14, 14, 9, 12, 14, 14, 14, 14, 14, 30))
    for t in r["taksitler"]:
        _satir(ws, [t["durum"], t["musteri"], t["police"], t["sirket"], t["vade"], t["tutar"], t["tahsil"], t["kalan"] or None, t["gecikme"], t["s_vade"], t["s_tutar"], t["s_odenen"], t["s_kalan"] or None, t["emanet"] or None, t["finanse"] or None, None],
               para_s=(6, 7, 8, 11, 12, 13, 14, 15), tarih_s=(5, 10), dolgu={1: DURUM_RENK[t["durum"]], 13: KIRMIZI if t["s_gecmis"] > 0 else None, 14: SARI if t["emanet"] > 0 else None, 15: SARI if t["finanse"] > 0 else None})
    ws = wb.create_sheet("Özet")
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (48, 18))
    for etiket, deger, para_mi in (("Taksit", o["taksit"], False), ("Toplam prim", o["tutar"], True), ("Müşteriden tahsil edilen", o["tahsil"], True), ("Tahsilat oranı %", o["oran"], True), ("Müşteriden kalan", o["kalan"], True),
                                   ("Gecikmiş müşteri alacağı", o["gecikmis"], True), (f"{gun} gün içinde vadesi gelen müşteri alacağı", o["yaklasan"], True), ("Şirketlere aktarılacak toplam", o["s_aktarilacak"], True),
                                   ("Şirketlere ödenen", o["s_odenen"], True), ("Şirketlere kalan borç", o["s_kalan"], True), ("Vadesi geçmiş şirket borcu", o["s_gecmis"], True),
                                   (f"{gun} gün içinde ödenecek şirket borcu", o["s_yakin"], True), ("Tahsil edilip şirkete aktarılmayan", o["s_emanet"], True), ("Acentenin finanse ettiği", o["s_finanse"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else ())
    ws.append([])
    _baslik(ws, ["Gecikmiş müşteri alacağının yaşı (gün)", "Taksit", "Tutar"], dondur="A2")
    for y in r["yaslar"]:
        _satir(ws, [y["dilim"], y["taksit"], y["tutar"]], para_s=(3,))
    ws.column_dimensions["C"].width = 18
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y}"])
    ws = wb.create_sheet("Müşteriler")
    _baslik(ws, ["Müşteri", "Taksit", "Toplam Prim", "Tahsil Edilen", "Tahsilat %", "Kalan", "Gecikmiş Tutar", "En Uzun Gecikme (gün)", f"{gun} Gün İçinde Vadesi Gelen"], (30, 8, 16, 16, 11, 16, 16, 11, 16))
    for m in r["musteriler"]:
        _satir(ws, [m["ad"], m["taksit"], m["tutar"], m["tahsil"], m["oran"], m["kalan"] or None, m["gecikmis"] or None, m["enb"], m["yaklasan"] or None], para_s=(3, 4, 5, 6, 7, 9), dolgu={7: KIRMIZI if m["gecikmis"] > 0 else None})
    ws = wb.create_sheet("Şirketler")
    _baslik(ws, ["Sigorta Şirketi", "Taksit", "Aktarılacak Toplam", "Ödenen", "Kalan Borç", "Vadesi Geçmiş Borç", f"{gun} Gün İçinde Ödenecek", "Tahsil Edilip Aktarılmayan", "Acentenin Finanse Ettiği"], (28, 8, 16, 16, 16, 16, 16, 16, 16))
    for s in r["sirketler"]:
        _satir(ws, [s["ad"], s["taksit"], s["aktarilacak"], s["odenen"], s["kalan"] or None, s["gecmis"] or None, s["yakin"] or None, s["emanet"] or None, s["finanse"] or None], para_s=range(3, 10), dolgu={6: KIRMIZI if s["gecmis"] > 0 else None})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, gun: int = 7, gecikme: int = 30) -> dict:
    taksitler, uy = oku(yol)
    r = degerlendir(taksitler, bugun, gun, gecikme)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun, gun)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Müşterilerden tahsil edilecek primleri vadeye göre, sigorta şirketine aktarılacak tutarlarla birlikte izler.")
    p.add_argument("--taksitler", type=Path, help="Poliçe No, Müşteri, Sigorta Şirketi, Vade, Tutar, Tahsil Edilen, Şirkete Ödeme Vadesi, Şirkete Aktarılacak, Şirkete Ödenen")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--gun", type=int, default=7, help="Kaç gün içinde vadesi gelenler 'yaklaşıyor' sayılsın (varsayılan 7)")
    p.add_argument("--gecikme", type=int, default=30, help="Bu kadar gün ve üstü gecikme yüksek önemli sayılır (varsayılan 30)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "musteri_tahsilat_takibi.xlsx")
    a = p.parse_args(argv)
    ornek = a.taksitler is None
    yol = ORNEK / "taksitler.csv" if ornek else a.taksitler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    if bugun is None or a.gun < 0 or a.gecikme < 1:
        print("[X] --tarih GG.AA.YYYY; --gun eksi olamaz; --gecikme en az 1 olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.gun, a.gecikme)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['taksit']} taksit · müşteriden kalan {o['kalan']:,.2f} (gecikmiş {o['gecikmis']:,.2f}) · şirketlere kalan borç {o['s_kalan']:,.2f} (vadesi geçmiş {o['s_gecmis']:,.2f}) · aktarılmayan tahsilat {o['s_emanet']:,.2f}")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
