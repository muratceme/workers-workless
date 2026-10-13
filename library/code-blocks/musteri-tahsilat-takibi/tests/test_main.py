import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)


def t(police, musteri, sirket, vade, tutar, tahsil=0, s_vade=None, s_tutar=None, s_odenen=0):
    return {"police": police, "musteri": musteri, "sirket": sirket, "vade": vade, "tutar": D(tutar), "tahsil": D(tahsil), "s_vade": s_vade, "s_tutar": D(tutar if s_tutar is None else s_tutar), "s_odenen": D(s_odenen)}


def veri():
    return [t("1", "A", "X", date(2026, 9, 1), 1000, 0, date(2026, 9, 15), 850), t("2", "a", "X", date(2026, 10, 5), 1000, 1000, date(2026, 10, 10), 850), t("3", "B", "Y", date(2026, 10, 15), 500, 0, date(2026, 10, 20), None, 500),
            t("4", "B", "y", date(2026, 12, 1), 500, 200, date(2026, 12, 15), 400), t("5", "C", "X", date(2026, 10, 12), 300, 0, date(2026, 10, 19), 250)]


class DegerlendirTesti(unittest.TestCase):
    def test_taksitler(self):
        r = main.degerlendir(veri(), BUGUN)
        self.assertEqual({x["police"]: (x["durum"], x["kalan"], x["gecikme"], x["yas"], x["s_kalan"], x["s_gecmis"], x["s_yakin"], x["emanet"], x["finanse"]) for x in r["taksitler"]},
                         {"1": ("Gecikmiş", D(1000), 41, "31–60", D(850), D(850), D(0), D(0), D(0)), "2": ("Tahsil edildi", D(0), None, None, D(850), D(850), D(0), D(850), D(0)),
                          "3": ("Yaklaşıyor", D(500), None, None, D(0), D(0), D(0), D(0), D(500)), "4": ("Vadesi gelmedi", D(300), None, None, D(400), D(0), D(0), D(200), D(0)),
                          "5": ("Yaklaşıyor", D(300), None, None, D(250), D(0), D(250), D(0), D(0))})
        self.assertEqual([x["police"] for x in r["taksitler"]], ["1", "5", "3", "4", "2"])
        self.assertEqual([(m["ad"], m["taksit"], m["tutar"], m["tahsil"], m["oran"], m["kalan"], m["gecikmis"], m["enb"], m["yaklasan"]) for m in r["musteriler"]],
                         [("A", 2, D(2000), D(1000), D("50.00"), D(1000), D(1000), 41, D(0)), ("B", 2, D(1000), D(200), D("20.00"), D(800), D(0), None, D(500)), ("C", 1, D(300), D(0), D("0.00"), D(300), D(0), None, D(300))])
        self.assertEqual([(s["ad"], s["taksit"], s["aktarilacak"], s["odenen"], s["kalan"], s["gecmis"], s["yakin"], s["emanet"], s["finanse"]) for s in r["sirketler"]],
                         [("X", 3, D(1950), D(0), D(1950), D(1700), D(250), D(850), D(0)), ("Y", 2, D(900), D(500), D(400), D(0), D(0), D(200), D(500))])
        self.assertEqual(r["ozet"], {"taksit": 5, "tutar": D(3300), "tahsil": D(1200), "kalan": D(2100), "gecikmis": D(1000), "enb": 41, "yaklasan": D(800), "oran": D("36.36"), "s_aktarilacak": D(2850), "s_odenen": D(500), "s_kalan": D(2350),
                                     "s_gecmis": D(1700), "s_yakin": D(250), "s_emanet": D(1050), "s_finanse": D(500)})
        self.assertEqual([(y["dilim"], y["taksit"], y["tutar"]) for y in r["yaslar"]], [("1–30", 0, D(0)), ("31–60", 1, D(1000)), ("61–90", 0, D(0)), ("90+", 0, D(0))])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"][:1]) for u in r["uyarilar"]),
                         [("Orta", "Tahsil edilmeden şirkete ödenmiş", "3"), ("Orta", "Şirket vadesi geçti, müşteri ödemedi", "1"), ("Yüksek", "Gecikmiş tahsilat", "A"), ("Yüksek", "Tahsil edilen prim şirkete aktarılmamış", "2")])

    def test_sinirlar(self):
        d = lambda vade, **kw: main.degerlendir([t("1", "A", "X", vade, 100)], BUGUN, **kw)   # noqa: E731
        self.assertEqual([d(v)["taksitler"][0]["durum"] for v in (date(2026, 10, 11), date(2026, 10, 12), date(2026, 10, 19), date(2026, 10, 20))], ["Gecikmiş", "Yaklaşıyor", "Yaklaşıyor", "Vadesi gelmedi"])
        self.assertEqual([d(v)["taksitler"][0]["yas"] for v in (date(2026, 9, 12), date(2026, 9, 11), date(2026, 8, 13), date(2026, 7, 14), date(2026, 7, 13))], ["1–30", "31–60", "31–60", "61–90", "90+"])
        self.assertEqual(([u["onem"] for u in d(date(2026, 9, 13))["uyarilar"]], [u["onem"] for u in d(date(2026, 9, 12))["uyarilar"]], [u["onem"] for u in d(date(2026, 9, 13), gecikme=29)["uyarilar"]], d(date(2026, 10, 20), gun=8)["taksitler"][0]["durum"]),
                         (["Orta"], ["Yüksek"], ["Yüksek"], "Yaklaşıyor"))
        r = main.degerlendir([t("1", "A", "X", date(2026, 9, 1), 100, 150, date(2026, 9, 10), 80, 20), t("2", "A", "X", date(2026, 9, 1), 100, 30, None, 80, 0)], BUGUN)
        self.assertEqual(([(x["kalan"], x["emanet"], x["s_gecmis"], x["finanse"]) for x in sorted(r["taksitler"], key=lambda x: x["police"])], sorted(u["tur"] for u in r["uyarilar"])),
                         ([(D(0), D(60), D(60), D(0)), (D(70), D(30), D(0), D(0))], ["Gecikmiş tahsilat", "Tahsil edilen prim şirkete aktarılmamış", "Tahsilat taksitten büyük"]))
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "taksitler.csv", cikti, BUGUN)
            self.assertEqual([(m["ad"], m["tutar"], m["tahsil"], m["kalan"], m["gecikmis"], m["enb"], m["yaklasan"]) for m in r["musteriler"]],
                             [("Örnek Lojistik A.Ş.", D(24000), D(0), D(24000), D(24000), 41, D(0)), ("Kurgu İnşaat Ltd. Şti.", D(6000), D(2500), D(3500), D(3500), 48, D(0)), ("Numune Tekstil A.Ş.", D(24000), D(9000), D(15000), D(0), None, D(15000)),
                              ("Deneme Gıda Ltd. Şti.", D(13000), D(8000), D(5000), D(0), None, D(0))])
            self.assertEqual([(s["ad"], s["aktarilacak"], s["odenen"], s["kalan"], s["gecmis"], s["yakin"], s["emanet"], s["finanse"]) for s in r["sirketler"]],
                             [("Deneme Sigorta", D(11900), D(2500), D(9400), D(9400), D(0), D(6800), D(0)), ("Örnek Sigorta", D(37400), D(10200), D(27200), D(0), D(10200), D(0), D(10200)),
                              ("Numune Sigorta", D(7650), D(7650), D(0), D(0), D(0), D(0), D(0))])
            o = r["ozet"]
            self.assertEqual((o["taksit"], o["tutar"], o["tahsil"], o["kalan"], o["gecikmis"], o["yaklasan"], o["s_kalan"], o["s_gecmis"], o["s_emanet"], o["s_finanse"]),
                             (7, D(67000), D(19500), D(47500), D(27500), D(15000), D(36600), D(9400), D(6800), D(10200)))
            self.assertEqual([(y["dilim"], y["taksit"], y["tutar"]) for y in r["yaslar"]], [("1–30", 1, D(12000)), ("31–60", 2, D(15500)), ("61–90", 0, D(0)), ("90+", 0, D(0))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"][:5]) for u in r["uyarilar"]],
                             [("Yüksek", "Gecikmiş tahsilat", "Kurgu"), ("Yüksek", "Gecikmiş tahsilat", "Örnek"), ("Yüksek", "Tahsil edilen prim şirkete aktarılmamış", "P-202"), ("Orta", "Tahsil edilmeden şirkete ödenmiş", "P-201"),
                              ("Orta", "Şirket vadesi geçti, müşteri ödemedi", "P-204")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Taksitler", "Özet", "Müşteriler", "Şirketler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Poliçe;Sigortalı;Şirket;Vade Tarihi;Taksit Tutarı;Tahsilat;Şirket Vadesi;Şirkete Ödenecek;Şirkete Aktarılan\nP1;Ali;X;01.09.2026;1.250,50;250,50;x;1.000;100\n;Veli;;02.09.2026;500;;;;\nP3;Ali;;x;500;;;;\n"
                         "P4;Ali;;02.09.2026;0;;;;\nP5;Ali;;02.09.2026;10;-1;;;\nP6;;;02.09.2026;10;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["police"], x["musteri"], x["sirket"], x["vade"], x["tutar"], x["tahsil"], x["s_vade"], x["s_tutar"], x["s_odenen"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("P1", "Ali", "X", date(2026, 9, 1), D("1250.50"), D("250.50"), None, D(1000), D(100)), ("", "Veli", "(şirket yok)", date(2026, 9, 2), D(500), D(0), None, D(500), D(0))],
                              [("Tarih okunamadı", "P1 · Ali"), ("Taksit okunamadı", "P3 · Ali"), ("Taksit okunamadı", "P4 · Ali"), ("Taksit okunamadı", "P5 · Ali")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--taksitler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--gun", "15", "--gecikme", "45"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--gun", "-1"], ["--gecikme", "0"])], [2] * 3)
            self.assertEqual(main.main(["--taksitler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Müşteri;Vade\nA;01.09.2026\n", "Müşteri;Vade;Tutar\n;01.09.2026;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
