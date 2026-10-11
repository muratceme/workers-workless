import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def iade(aciklama, sube="Merkez", kanal="Şube", musteri="", iban="", tutar=100, gun=1, sure=0, kullanici=""):
    t = main.siniflandir(aciklama, main.HATA_TURLERI)
    return {"ref": "R", "gonderim": date(2026, 9, gun), "iade": date(2026, 9, gun + sure), "sure": sure, "sube": sube, "kanal": kanal, "sistem": "FAST", "musteri": musteri,
            "iban": iban, "alici": "", "tutar": D(tutar), "kod": "", "aciklama": aciklama, "kullanici": kullanici, "tur": t[0], "kaynak": t[2], "oneri": t[3]}


class SiniflandirmaTesti(unittest.TestCase):
    def test_hata_turleri(self):
        def s(a):
            return main.siniflandir(a, main.HATA_TURLERI)[0]
        self.assertEqual(s("ALICI ADI İLE IBAN UYUŞMUYOR"), "Alıcı adı uyuşmuyor")
        self.assertEqual(s("Hesap kapalı"), "Hesap kapalı")
        self.assertEqual(s("Geçersiz IBAN"), "IBAN / hesap numarası hatalı")
        self.assertEqual(s("Hesap bulunamadı"), "IBAN / hesap numarası hatalı")
        self.assertEqual(s("Hesap üzerinde bloke var"), "Blokeli / kısıtlı hesap")
        self.assertEqual(s("Mükerrer gönderim"), "Mükerrer gönderim")
        self.assertEqual(s("Alıcının talebi üzerine iade"), "Alıcı talebiyle iade")
        self.assertEqual(s("Limit aşımı"), "Tutar / limit")
        self.assertEqual(s("Bağlantı zaman aşımı"), "Teknik / sistemsel")
        self.assertEqual(s("Kurum kodu tanımsız"), "Diğer")

    def test_kullanici_turleri_once(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "t.csv"
            y.write_text("Hata Türü;Anahtar İfadeler;Kaynak;Öneri\nKurum / abone hatası;kurum kodu, abone;Şube girişi;Kurum listesini güncelleyin\n;boş;;\n", encoding="utf-8")
            t = main.tur_dosyasi(y)
            self.assertEqual(len(t), 1)
            self.assertEqual(main.siniflandir("Abone numarası bulunamadı", t + main.HATA_TURLERI)[:3:2], ("Kurum / abone hatası", "Şube girişi"))


class AnalizTesti(unittest.TestCase):
    def test_pareto_ve_paylar(self):
        i = [iade("Hesap kapalı")] * 5 + [iade("IBAN hatalı")] * 3 + [iade("Limit aşımı")] + [iade("Hesap blokeli")]
        r = main.analiz(i)
        self.assertEqual([(t["ad"], t["adet"], t["pay"], t["kumulatif"], t["pareto"]) for t in r["turler"]],
                         [("Hesap kapalı", 5, D("50.0"), D("50.0"), "A"), ("IBAN / hesap numarası hatalı", 3, D("30.0"), D("80.0"), "A"),
                          ("Blokeli / kısıtlı hesap", 1, D("10.0"), D("90.0"), "B"), ("Tutar / limit", 1, D("10.0"), D("100.0"), "B")])
        self.assertEqual({k["ad"]: k["adet"] for k in r["kaynak"]}, {"Gönderen müşteri": 5, "Şube girişi": 4, "Alıcı taraf": 1})
        self.assertEqual(r["toplam_tutar"], D(1000))

    def test_iade_orani(self):
        i = [iade("IBAN hatalı", "A")] * 6 + [iade("Hesap kapalı", "B")] * 2 + [iade("Hesap kapalı", "C", "Mobil")] * 2
        g = {("A", "Şube"): 100, ("B", "Şube"): 400, ("C", "Mobil"): 500}
        r = main.analiz(i, g)
        self.assertEqual({s["ad"]: (s["gonderilen"], s["oran"], s["onlenebilir"]) for s in r["sube"]}, {"A": (100, D("6.00"), 6), "B": (400, D("0.50"), 0), "C": (500, D("0.40"), 0)})
        self.assertEqual({s["ad"]: s["oran"] for s in r["kanal"]}, {"Şube": D("1.60"), "Mobil": D("0.40")})
        self.assertEqual(r["genel_oran"], D(1))
        self.assertEqual([(u["tur"], u["kim"]) for u in r["uyarilar"]], [("İade oranı yüksek şube", "A")])

    def test_tekrarlar(self):
        i = [iade("Hesap kapalı", musteri="M1", iban="TR1", gun=g) for g in (1, 8, 15)] + [iade("IBAN hatalı", musteri="M2", iban="TR2"), iade("Limit aşımı", musteri="M2", iban="TR3")] \
            + [iade("Hesap kapalı", musteri="M3", iban="TR9")]
        r = main.analiz(i)
        self.assertEqual([(t["tip"], t["kim"], t["adet"], t["ayni_neden"]) for t in r["tekrarlar"]],
                         [("Gönderen müşteri", "M1", 3, True), ("Alıcı IBAN", "TR1", 3, True), ("Gönderen müşteri", "M2", 2, False)])
        self.assertEqual((r["tekrarlar"][0]["ilk"], r["tekrarlar"][0]["son"]), (date(2026, 9, 1), date(2026, 9, 15)))
        self.assertEqual(sorted(u["kim"] for u in r["uyarilar"]), ["Alıcı IBAN: TR1", "Gönderen müşteri: M1"])
        self.assertEqual(len(main.analiz(i, tekrar=3)["tekrarlar"]), 2)

    def test_diger_orani_ve_kullanici(self):
        i = [iade("Kurum kodu tanımsız")] * 3 + [iade("Abone bulunamadı")] + [iade("IBAN hatalı", kullanici="K1")] * 3 + [iade("Hesap kapalı", kullanici="K1")] * 9
        r = main.analiz(i)
        u = [x for x in r["uyarilar"] if x["tur"] == "Sınıflandırılamayan iadeler"]
        self.assertEqual(len(u), 1)                                                   # 4 / 16 = %25
        self.assertTrue(u[0]["aciklama"].endswith("Kurum kodu tanımsız; Abone bulunamadı"))
        self.assertEqual([(k["ad"], k["adet"]) for k in r["kullanici"]], [("K1", 3)])  # yalnız giriş kaynaklı iadeler sayılır


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "i.xlsx"
            r = main.calistir(ORNEK / "iadeler.csv", cikti, ORNEK / "gonderilen.csv")
            self.assertEqual((len(r["iadeler"]), r["turler"][0]["ad"], r["turler"][0]["adet"]), (105, "Alıcı adı uyuşmuyor", 33))
            self.assertEqual(sum(t["adet"] for t in r["turler"]), 105)
            self.assertEqual(r["turler"][-1]["kumulatif"], D("100.0"))
            self.assertIn(("Gönderen müşteri", "M30200", 6, True), [(t["tip"], t["kim"], t["adet"], t["ayni_neden"]) for t in r["tekrarlar"]])
            self.assertEqual([u["kim"] for u in r["uyarilar"] if u["onem"] == "Yüksek"], ["Sanayi"])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Hata Türleri", "Şube Bazında", "Kanal Bazında", "Tekrarlayanlar", "Günlük", "İadeler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "i.csv"
            y.write_text("Referans;Gönderim Tarihi;İade Tarihi;Şube;Tutar;İade Açıklaması\nA1;01.09.2026;03.09.2026;Merkez;1.250,50;Hesap kapalı\n"
                         "A2;05.09.2026;04.09.2026;Merkez;bin;IBAN hatalı\nA3;;;;;Alıcı ismi hatalı\n", encoding="utf-8")
            i, uy = main.oku(y, main.HATA_TURLERI)
            self.assertEqual([(x["ref"], x["sure"], x["tutar"], x["sube"], x["tur"]) for x in i],
                             [("A1", 2, D("1250.50"), "Merkez", "Hesap kapalı"), ("A2", None, D(0), "Merkez", "IBAN / hesap numarası hatalı"), ("A3", None, D(0), "Belirtilmemiş", "Alıcı adı uyuşmuyor")])
            self.assertEqual(sorted(u["tur"] for u in uy), ["Tutar okunamadı", "İade tarihi gönderimden önce"])
            arg = ["--iadeler", str(y), "--cikti", str(Path(tmp) / "x.xlsx")]
            self.assertEqual(main.main(arg), 0)
            self.assertEqual(main.main(arg + ["--tekrar", "1"]), 2)
            self.assertEqual(main.main(["--iadeler", str(Path(tmp) / "yok.csv")]), 1)
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "o.xlsx")]), 0)


if __name__ == "__main__":
    unittest.main()
