"""
İade Gelen EFT/FAST Analizi — Gerisi Bende kod bloğu
Bankacılık Operasyonları › Ödeme Sistemleri Operasyon Uzmanı

İade gelen ödemeleri iade nedeni, şube, kanal ve hata türüne göre sınıflandırır; tekrarlayan hata nedenlerini raporlar.
  - Hata türü: iade açıklamasındaki anahtar ifadelerden bulunur (IBAN / hesap no hatalı, hesap kapalı, alıcı adı uyuşmuyor,
    alıcı talebiyle iade, mükerrer gönderim, tutar / limit, blokeli hesap, teknik...). Eşleşmeyenler "Diğer" olur.
    Kendi anahtar ifadelerinizi --hata-turleri dosyasıyla verebilirsiniz (Hata Türü; Anahtar İfadeler; Kaynak; Öneri).
  - Kaynak: hatanın nerede önlenebileceği (Gönderen müşteri / Şube girişi / Alıcı taraf / Sistem).
  - Tekrar: aynı gönderen müşteriden veya aynı alıcı IBAN'ına --tekrar (varsayılan 2) ve üzeri iade.
  - İade oranı: gönderilen işlem adedi dosyası verilirse şube / kanal bazında iade ÷ gönderilen.
  - İade süresi: iade tarihi − gönderim tarihi (takvim günü).
  - Pareto: iadelerin %80'ini oluşturan hata türleri "A" olarak işaretlenir.
İade kodlarının anlamı sistemden sisteme değişir; kod sabit bir kod tablosu içermez, açıklama metnini kullanır.
İnternete bağlanmaz.

Kullanım:
    python main.py                                                 # örnek: 1 aylık 96 iade
    python main.py --iadeler iadeler.xlsx --gonderilen gonderilen.xlsx --tekrar 3
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


IADE_SUTUNLARI = {"ref": ("referans", "islem no", "sorgu no", "referans no"), "gonderim": ("gonderim tarihi", "islem tarihi", "tarih"), "iade": ("iade tarihi",),
                  "sube": ("sube",), "kanal": ("kanal",), "sistem": ("sistem", "odeme sistemi"), "musteri": ("gonderen musteri", "gonderen", "musteri", "musteri no"),
                  "iban": ("alici iban", "iban"), "alici": ("alici", "alici adi"), "tutar": ("tutar",), "kod": ("iade kodu", "kod"),
                  "aciklama": ("iade aciklamasi", "iade nedeni", "aciklama"), "kullanici": ("kullanici", "giren kullanici", "personel")}
GONDERILEN_SUTUNLARI = {"sube": ("sube",), "kanal": ("kanal",), "adet": ("gonderilen adet", "adet", "islem adedi")}
TUR_SUTUNLARI = {"tur": ("hata turu", "tur"), "ifadeler": ("anahtar ifadeler", "ifadeler"), "kaynak": ("kaynak",), "oneri": ("oneri", "onlem")}
# (hata türü, anahtar ifadeler — katlanmış, kaynak, öneri). Sıra önemlidir: ilk eşleşen alınır.
HATA_TURLERI = [
    ("Alıcı adı uyuşmuyor", ("alici adi", "isim uyus", "ad soyad uyus", "unvan uyus", "isim hatali", "alici ismi"), "Gönderen müşteri",
     "Alıcı adını hesap sahibinin resmî adıyla aynı girdirin; giriş ekranında ad-IBAN uyum sorgusu varsa kullanın."),
    ("Hesap kapalı", ("hesap kapali", "kapali hesap", "kapanmis hesap", "hesap kapanmis"), "Gönderen müşteri", "Müşterinin kayıtlı alıcı listesini güncelletin; düzenli ödemelerde alıcı bilgisini teyit edin."),
    ("IBAN / hesap numarası hatalı", ("iban hatali", "hatali iban", "gecersiz iban", "iban gecersiz", "hesap no hatali", "hesap bulunamadi", "hesap numarasi hatali", "boyle bir hesap"),
     "Şube girişi", "Girişte IBAN kontrol basamağı doğrulaması yapın; elle yazım yerine kopyalama / kayıtlı alıcı kullanın."),
    ("Blokeli / kısıtlı hesap", ("bloke", "kisitli", "haciz", "tedbir"), "Alıcı taraf", "Müşteriyi bilgilendirin; önlenebilir bir giriş hatası değildir."),
    ("Mükerrer gönderim", ("mukerrer", "cift gonderim", "iki kez", "tekrar gonderilmis"), "Şube girişi", "Aynı gün aynı alıcı ve tutar için uyarı / ikinci onay uygulayın."),
    ("Alıcı talebiyle iade", ("alici talebi", "alicinin talebi", "musteri talebi", "kabul edilmedi", "alici kabul etmedi", "iade talebi"), "Alıcı taraf",
     "Gönderen müşteriyi bilgilendirin; tekrarlıyorsa ticari anlaşmazlık olabilir."),
    ("Tutar / limit", ("limit", "tutar asimi", "tutar hatali", "azami tutar"), "Şube girişi", "Sistem ve işlem limitlerini giriş ekranında kontrol edin; doğru ödeme sistemini seçin."),
    ("Açıklama / zorunlu alan eksik", ("aciklama eksik", "zorunlu alan", "eksik bilgi", "kimlik no", "vergi no", "tckn"), "Şube girişi", "Zorunlu alanları giriş ekranında denetleyin."),
    ("Para birimi / hesap türü uyumsuz", ("para birimi", "doviz hesabi", "hesap turu", "yabanci para"), "Gönderen müşteri", "Alıcı hesabın para birimini ve türünü teyit ettirin."),
    ("Teknik / sistemsel", ("teknik", "sistem", "zaman asimi", "baglanti", "timeout"), "Sistem", "BT / ödeme sistemleri ekibine iletin; aynı saat aralığında yığılma var mı bakın."),
]
DIGER = ("Diğer", (), "Belirsiz", "İade açıklamasını elle inceleyin; sık geliyorsa --hata-turleri dosyasına anahtar ifade ekleyin.")


def tur_dosyasi(yol: Path) -> list[tuple]:
    turler = []
    for r in kayitlar(yol, TUR_SUTUNLARI, ("tur", "ifadeler")):
        ifadeler = tuple(katla(x) for x in re.split(r"[,|/]", metin(r["ifadeler"])) if katla(x))
        if metin(r["tur"]) and ifadeler:
            turler.append((metin(r["tur"]), ifadeler, metin(r["kaynak"]) or "Belirsiz", metin(r["oneri"])))
    return turler


def siniflandir(aciklama: str, turler: list[tuple]) -> tuple:
    k = katla(aciklama)
    return next((t for t in turler if any(i in k for i in t[1])), DIGER)


def oku(yol: Path, turler: list[tuple]) -> tuple[list[dict], list[dict]]:
    uy, iadeler = [], []
    for r in kayitlar(yol, IADE_SUTUNLARI, ("aciklama",)):
        if not (metin(r["aciklama"]) or metin(r["kod"]) or metin(r["ref"])):
            continue
        t = para(r["tutar"])
        g, i = tarih(r["gonderim"]), tarih(r["iade"])
        if metin(r["tutar"]) and (t is None or t < 0):
            uy.append(uyari("Orta", "Tutar okunamadı", r["ref"] or f"Satır {r['_satir']}", f"'{metin(r['tutar'])}'; tutar 0 alındı."))
        sure = (i - g).days if g and i else None
        if sure is not None and sure < 0:
            uy.append(uyari("Orta", "İade tarihi gönderimden önce", r["ref"] or f"Satır {r['_satir']}", f"Gönderim {g:%d.%m.%Y}, iade {i:%d.%m.%Y}. Süre hesaba katılmadı."))
            sure = None
        tur = siniflandir(metin(r["aciklama"]), turler)
        iadeler.append({"ref": metin(r["ref"]), "gonderim": g, "iade": i, "sure": sure, "sube": metin(r["sube"]) or "Belirtilmemiş", "kanal": metin(r["kanal"]) or "Belirtilmemiş",
                        "sistem": metin(r["sistem"]).upper() or "—", "musteri": metin(r["musteri"]), "iban": re.sub(r"\s", "", metin(r["iban"])).upper(), "alici": metin(r["alici"]),
                        "tutar": t if t is not None and t >= 0 else Decimal(0), "kod": metin(r["kod"]), "aciklama": metin(r["aciklama"]), "kullanici": metin(r["kullanici"]),
                        "tur": tur[0], "kaynak": tur[2], "oneri": tur[3]})
    return iadeler, uy


def gonderilen_oku(yol: Path) -> dict:
    g = defaultdict(int)
    for r in kayitlar(yol, GONDERILEN_SUTUNLARI, ("adet",)):
        a = para(r["adet"])
        if a is not None and a >= 0:
            g[(metin(r["sube"]) or "Belirtilmemiş", metin(r["kanal"]) or "Belirtilmemiş")] += int(a)
    return dict(g)


def grupla(iadeler: list[dict], anahtar) -> list[dict]:
    g = defaultdict(list)
    for i in iadeler:
        g[anahtar(i)].append(i)
    toplam = len(iadeler)
    return sorted(({"ad": k, "adet": len(l), "tutar": sum((x["tutar"] for x in l), Decimal(0)), "pay": yuvarla(Decimal(len(l)) / toplam * 100, "0.1"), "liste": l}
                   for k, l in g.items()), key=lambda x: (-x["adet"], str(x["ad"])))


def analiz(iadeler: list[dict], gonderilen: dict | None = None, tekrar: int = 2) -> dict:
    uy = []
    turler = grupla(iadeler, lambda i: i["tur"])
    kum = 0
    for t in turler:
        t["pareto"] = "A" if Decimal(kum) / len(iadeler) * 100 < 80 else "B"
        kum += t["adet"]
        t["kumulatif"] = yuvarla(Decimal(kum) / len(iadeler) * 100, "0.1")
        t["kaynak"], t["oneri"] = t["liste"][0]["kaynak"], t["liste"][0]["oneri"]
        t["kodlar"] = ", ".join(sorted({x["kod"] for x in t["liste"] if x["kod"]}))
        sureler = [x["sure"] for x in t["liste"] if x["sure"] is not None]
        t["sure"] = yuvarla(Decimal(sum(sureler)) / len(sureler), "0.1") if sureler else None
    diger = next((t for t in turler if t["ad"] == "Diğer"), None)
    if diger and diger["pay"] > 15:
        sayim = defaultdict(int)
        for x in diger["liste"]:
            sayim[x["aciklama"]] += 1
        uy.append(uyari("Orta", "Sınıflandırılamayan iadeler", f"{diger['adet']} iade", f"İadelerin %{diger['pay']} kadarı 'Diğer'. En sık açıklamalar: " +
                        "; ".join(a for a, _ in sorted(sayim.items(), key=lambda z: (-z[1], z[0]))[:5])))
    sube = grupla(iadeler, lambda i: i["sube"])
    kanal = grupla(iadeler, lambda i: i["kanal"])
    for liste, idx in ((sube, 0), (kanal, 1)):
        for s in liste:
            enler = grupla(s["liste"], lambda i: i["tur"])
            s["en_sik"] = f"{enler[0]['ad']} ({enler[0]['adet']})"
            s["onlenebilir"] = sum(x["kaynak"] == "Şube girişi" for x in s["liste"])
            if gonderilen:
                gon = sum(v for k, v in gonderilen.items() if k[idx] == s["ad"])
                s["gonderilen"], s["oran"] = gon, (yuvarla(Decimal(s["adet"]) / gon * 100, "0.01") if gon else None)
    if gonderilen:
        tg = sum(gonderilen.values())
        genel_oran = Decimal(len(iadeler)) / tg * 100 if tg else None
        for s in sube:
            if s.get("oran") is not None and genel_oran and s["oran"] > genel_oran * 2 and s["adet"] >= 5:
                uy.append(uyari("Yüksek", "İade oranı yüksek şube", s["ad"], f"İade oranı %{s['oran']}; genel ortalama %{yuvarla(genel_oran, '0.01')}. En sık neden: {s['en_sik']}."))
    else:
        genel_oran = None
    tekrarlar = []
    for ad, anahtar in (("Gönderen müşteri", "musteri"), ("Alıcı IBAN", "iban")):
        for g in grupla([i for i in iadeler if i[anahtar]], lambda i: i[anahtar]):
            if g["adet"] >= tekrar:
                nedenler = grupla(g["liste"], lambda i: i["tur"])
                tekrarlar.append({"tip": ad, "kim": g["ad"], "adet": g["adet"], "tutar": g["tutar"], "nedenler": ", ".join(f"{n['ad']} ({n['adet']})" for n in nedenler),
                                  "ayni_neden": len(nedenler) == 1, "en_cok": nedenler[0]["adet"], "ilk": min((x["gonderim"] for x in g["liste"] if x["gonderim"]), default=None),
                                  "son": max((x["gonderim"] for x in g["liste"] if x["gonderim"]), default=None), "subeler": ", ".join(sorted({x["sube"] for x in g["liste"]}))})
    tekrarlar.sort(key=lambda t: (-t["adet"], t["kim"]))
    for t in tekrarlar:
        if t["en_cok"] >= max(3, tekrar + 1):                         # aynı neden en az 3 kez
            uy.append(uyari("Orta", "Aynı nedenle tekrarlayan iade", f"{t['tip']}: {t['kim']}", f"{t['adet']} iade ({t['nedenler']}). Kök neden giderilmeden gönderim sürüyor."))
    kullanici = [k for k in grupla([i for i in iadeler if i["kullanici"] and i["kaynak"] == "Şube girişi"], lambda i: i["kullanici"]) if k["adet"] >= 3]
    gunluk = grupla([i for i in iadeler if i["iade"] or i["gonderim"]], lambda i: i["iade"] or i["gonderim"])
    gunluk.sort(key=lambda g: g["ad"])
    kaynak = grupla(iadeler, lambda i: i["kaynak"])
    return {"iadeler": iadeler, "turler": turler, "sube": sube, "kanal": kanal, "tekrarlar": tekrarlar, "kullanici": kullanici, "gunluk": gunluk, "kaynak": kaynak,
            "uyarilar": uy, "toplam_tutar": sum((i["tutar"] for i in iadeler), Decimal(0)), "genel_oran": genel_oran, "oran_var": bool(gonderilen)}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

def rapor_yaz(cikti: Path, r: dict) -> None:
    wb = Workbook()
    ht = wb.active
    ht.title = "Hata Türleri"
    ht.append([f"{len(r['iadeler'])} iade · toplam {r['toplam_tutar']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") +
               (f" · genel iade oranı %{yuvarla(r['genel_oran'], '0.01')}" if r["genel_oran"] is not None else "")])
    ht["A1"].font = Font(bold=True)
    _baslik(ht, ["Hata Türü", "Adet", "Pay %", "Kümülatif %", "Pareto", "Tutar", "Ortalama İade Süresi (gün)", "Nerede Önlenir", "İade Kodları", "Öneri"],
            (32, 7, 8, 10, 8, 18, 12, 18, 18, 90), "B3")
    for t in r["turler"]:
        _satir(ht, [t["ad"], t["adet"], t["pay"], t["kumulatif"], t["pareto"], t["tutar"], t["sure"], t["kaynak"], t["kodlar"], t["oneri"]], para_s=(6,),
               dolgu={5: KIRMIZI if t["pareto"] == "A" else None}, sar=(10,))
    ht.append([])
    _baslik(ht, ["Nerede Önlenir", "Adet", "Pay %"], dondur=None)
    for k in r["kaynak"]:
        _satir(ht, [k["ad"], k["adet"], k["pay"]])

    for ad, veri in (("Şube Bazında", r["sube"]), ("Kanal Bazında", r["kanal"])):
        ws = wb.create_sheet(ad)
        _baslik(ws, [ad.split()[0], "İade Adedi", "Pay %", "Tutar", "Gönderilen Adet", "İade Oranı %", "Giriş Kaynaklı İade", "En Sık Neden"], (22, 9, 8, 18, 11, 10, 11, 40))
        for s in (sorted(veri, key=lambda s: -(s.get("oran") or 0)) if r["oran_var"] else veri):
            _satir(ws, [s["ad"], s["adet"], s["pay"], s["tutar"], s.get("gonderilen"), s.get("oran"), s["onlenebilir"], s["en_sik"]], para_s=(4,),
                   dolgu={6: KIRMIZI if r["genel_oran"] and s.get("oran") and s["oran"] > r["genel_oran"] * 2 else None})

    tk = wb.create_sheet("Tekrarlayanlar")
    _baslik(tk, ["Tür", "Kim", "İade Adedi", "Tutar", "Nedenler", "Hep Aynı Neden", "İlk Gönderim", "Son Gönderim", "Şubeler", "Aksiyon"], (18, 30, 9, 18, 50, 10, 12, 12, 24, 30))
    for t in r["tekrarlar"]:
        _satir(tk, [t["tip"], t["kim"], t["adet"], t["tutar"], t["nedenler"], "Evet" if t["ayni_neden"] else "Hayır", t["ilk"], t["son"], t["subeler"], ""], para_s=(4,), tarih_s=(7, 8),
               dolgu={6: KIRMIZI if t["ayni_neden"] else None, 10: SARI}, sar=(5,))
    if r["kullanici"]:
        tk.append([])
        _baslik(tk, ["Giriş kaynaklı iade", "Kullanıcı", "İade Adedi", "Tutar"], dondur=None)
        for k in r["kullanici"]:
            _satir(tk, ["", k["ad"], k["adet"], k["tutar"]], para_s=(4,))
        tk.append(["Kullanıcı bazlı sayım eğitim ihtiyacını göstermek içindir; işlem hacmine oranlanmadan performans ölçüsü olarak kullanılmamalıdır."])

    gn = wb.create_sheet("Günlük")
    _baslik(gn, ["Tarih", "İade Adedi", "Tutar"], (12, 10, 18))
    ort = Decimal(len(r["iadeler"])) / len(r["gunluk"]) if r["gunluk"] else Decimal(0)
    for g in r["gunluk"]:
        _satir(gn, [g["ad"], g["adet"], g["tutar"]], para_s=(3,), tarih_s=(1,), dolgu={2: KIRMIZI if g["adet"] > ort * 2 and g["adet"] >= 5 else None})

    dt = wb.create_sheet("İadeler")
    _baslik(dt, ["Referans", "Gönderim", "İade", "Süre (gün)", "Şube", "Kanal", "Sistem", "Gönderen", "Alıcı", "Alıcı IBAN", "Tutar", "İade Kodu", "İade Açıklaması", "Hata Türü",
                 "Nerede Önlenir"], (14, 11, 11, 7, 14, 14, 8, 20, 24, 30, 16, 9, 40, 28, 16))
    for i in r["iadeler"]:
        _satir(dt, [i["ref"], i["gonderim"], i["iade"], i["sure"], i["sube"], i["kanal"], i["sistem"], i["musteri"], i["alici"], i["iban"], i["tutar"], i["kod"], i["aciklama"],
                    i["tur"], i["kaynak"]], para_s=(11,), tarih_s=(2, 3), dolgu={14: SARI if i["tur"] == "Diğer" else None})
    dt.auto_filter.ref = f"A1:O{dt.max_row}"
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, g_yol: Path | None = None, tur_yol: Path | None = None, tekrar: int = 2) -> dict:
    turler = (tur_dosyasi(tur_yol) if tur_yol else []) + HATA_TURLERI                # kullanıcının türleri önce denenir
    iadeler, uy = oku(yol, turler)
    if not iadeler:
        raise ValueError(f"{yol.name}: iade satırı bulunamadı")
    r = analiz(iadeler, gonderilen_oku(g_yol) if g_yol else None, tekrar)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="İade gelen EFT / FAST ödemelerini hata türü, şube ve kanala göre sınıflandırır; tekrarlayan nedenleri raporlar.")
    p.add_argument("--iadeler", type=Path, default=ORNEK / "iadeler.csv",
                   help="Referans, Gönderim Tarihi, İade Tarihi, Şube, Kanal, Sistem, Gönderen Müşteri, Alıcı, Alıcı IBAN, Tutar, İade Kodu, İade Açıklaması, Kullanıcı")
    p.add_argument("--gonderilen", type=Path, help="Şube, Kanal, Gönderilen Adet (iade oranı için)")
    p.add_argument("--hata-turleri", type=Path, help="Hata Türü, Anahtar İfadeler (virgülle), Kaynak, Öneri")
    p.add_argument("--tekrar", type=int, default=2, help="Aynı gönderen / alıcı için tekrar eşiği (varsayılan 2)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "iade_analizi.xlsx")
    a = p.parse_args(argv)
    gon = a.gonderilen or (ORNEK / "gonderilen.csv" if a.iadeler == ORNEK / "iadeler.csv" else None)
    for y in (a.iadeler, gon, a.hata_turleri):
        if y and not y.exists():
            print(f"[X] Dosya bulunamadı: {y}")
            return 1
    if a.tekrar < 2:
        print("[X] --tekrar en az 2 olmalı")
        return 2
    try:
        r = calistir(a.iadeler, a.cikti, gon, a.hata_turleri, a.tekrar)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(r['iadeler'])} iade · {len(r['turler'])} hata türü · {len(r['tekrarlar'])} tekrarlayan gönderen / alıcı")
    for t in r["turler"][:5]:
        print(f"     {t['ad']}: {t['adet']} (%{t['pay']})")
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
