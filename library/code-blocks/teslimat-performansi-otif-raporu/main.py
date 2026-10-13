"""
Teslimat Performansı (OTIF) Raporu — Workers / Workless kod bloğu
Lojistik ve Depo › Lojistik Uzmanı

Siparişlerin zamanında ve eksiksiz teslim oranını müşteri ve taşıyıcı bazında raporlar.
  Zamanında (satır)    = teslim tarihi ≤ söz verilen tarih + tolerans günü
  Eksiksiz (satır)     = teslim miktarı ≥ sipariş miktarı × (1 − eksik toleransı ÷ 100)
  Sipariş zamanında    = bütün satırları zamanında;   sipariş eksiksiz = bütün satırları eksiksiz
  OTIF                 = sipariş hem zamanında hem eksiksiz
  Zamanında %, Eksiksiz %, OTIF % = ilgili sipariş sayısı ÷ değerlendirilen sipariş sayısı × 100
  Miktar karşılama %   = Σ min(teslim, sipariş) ÷ Σ sipariş miktarı × 100
  Açık                 = teslim edilmemiş ve söz tarihi henüz geçmemiş satırı olan sipariş; oranlara girmez
  Teslim edilmemiş ve söz tarihi geçmiş satır hem geç hem eksik sayılır; gecikme durum tarihine kadar hesaplanır.
Erken teslim zamanında sayılır. Toleranslar komut satırından verilir (varsayılan sıfır).
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
SUTUN = {"siparis": ("siparis no", "siparis numarasi", "no"), "musteri": ("musteri", "cari", "firma"), "tasiyici": ("tasiyici", "nakliyeci", "kargo firmasi"),
         "soz": ("soz verilen tarih", "termin tarihi", "termin", "istenen teslim tarihi", "soz tarihi"), "teslim": ("teslim tarihi", "gerceklesen teslim tarihi", "teslim edilen tarih"),
         "sip_miktar": ("siparis miktari", "istenen miktar", "miktar"), "tes_miktar": ("teslim miktari", "teslim edilen miktar", "sevk miktari")}
SONUCLAR = ("OTIF", "Geç", "Eksik", "Geç ve eksik", "Açık")
SONUC_RENK = {"OTIF": YESIL, "Geç": SARI, "Eksik": SARI, "Geç ve eksik": KIRMIZI, "Açık": MAVI}
Z, YUZ = Decimal(0), Decimal(100)


def olcu(x) -> Decimal | None:
    """Miktar okur: virgül ya da tek nokta ondalık ayırıcıdır; '1.250,5' binlik ayırıcılıdır."""
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x or ""))
    try:
        return Decimal(s.replace(".", "").replace(",", ".") if "," in s else s) if s else None
    except InvalidOperation:
        return None


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("soz", "sip_miktar")):
        if not metin(r["soz"]) and not metin(r["sip_miktar"]) and not metin(r["siparis"]):
            continue
        kim = metin(r["siparis"]) or f"Satır {r['_satir']}"
        soz, teslim, sm, tm = tarih(r["soz"]), tarih(r["teslim"]), olcu(r["sip_miktar"]), olcu(r["tes_miktar"])
        if soz is None or sm is None or sm <= 0 or (metin(r["teslim"]) and teslim is None) or (tm is not None and tm < 0) or (metin(r["tes_miktar"]) and tm is None):
            uy.append(uyari("Yüksek", "Satır okunamadı", kim, f"Satır {r['_satir']}: söz tarihi '{metin(r['soz'])}', teslim tarihi '{metin(r['teslim'])}', sipariş miktarı '{metin(r['sip_miktar'])}', teslim miktarı '{metin(r['tes_miktar'])}'. "
                            "Satır alınmadı."))
            continue
        l.append({"siparis": metin(r["siparis"]) or f"(satır {r['_satir']})", "musteri": metin(r["musteri"]) or "(müşteri yok)", "tasiyici": metin(r["tasiyici"]) or "(taşıyıcı yok)", "soz": soz, "teslim": teslim, "sip_miktar": sm,
                  "tes_miktar": tm if tm is not None else (sm if teslim else Z), "miktar_yok": tm is None and teslim is not None})
    return l, uy


def _oran(pay, payda) -> Decimal | None:
    return yuvarla(Decimal(pay) * YUZ / Decimal(payda)) if payda else None


def _ozet(siparisler: list[dict]) -> dict:
    d = [s for s in siparisler if s["sonuc"] != "Açık"]
    gec = [s["gecikme"] for s in d if s["gecikme"] > 0]
    sm, tm = sum((s["sip_miktar"] for s in d), Z), sum((s["karsilanan"] for s in d), Z)
    return {"siparis": len(siparisler), "acik": len(siparisler) - len(d), "degerlendirilen": len(d), "zamaninda": sum(1 for s in d if s["zamaninda"]), "eksiksiz": sum(1 for s in d if s["eksiksiz"]), "otif": sum(1 for s in d if s["sonuc"] == "OTIF"),
            "zamaninda_oran": _oran(sum(1 for s in d if s["zamaninda"]), len(d)), "eksiksiz_oran": _oran(sum(1 for s in d if s["eksiksiz"]), len(d)), "otif_oran": _oran(sum(1 for s in d if s["sonuc"] == "OTIF"), len(d)),
            "karsilama": _oran(tm, sm), "ort_gecikme": yuvarla(Decimal(sum(gec)) / len(gec)) if gec else None}


def degerlendir(satirlar: list[dict], bugun: date, tolerans: int = 0, eksik_tolerans: Decimal = Z) -> dict:
    if not satirlar:
        raise ValueError("Değerlendirilecek sipariş satırı yok.")
    uy, grup = [], {}
    for x in satirlar:
        x["acik"] = x["teslim"] is None and x["soz"] + timedelta(days=tolerans) >= bugun
        x["gecikme"] = max(((x["teslim"] or bugun) - x["soz"]).days - tolerans, 0) if not x["acik"] else 0
        x["zamaninda"] = not x["acik"] and x["teslim"] is not None and x["gecikme"] == 0
        x["eksiksiz"] = not x["acik"] and x["teslim"] is not None and x["tes_miktar"] >= x["sip_miktar"] * (1 - eksik_tolerans / YUZ)
        grup.setdefault(katla(x["siparis"]), []).append(x)
        if x["miktar_yok"]:
            uy.append(uyari("Bilgi", "Teslim miktarı yazılmamış", x["siparis"], "Teslim tarihi var ama teslim miktarı boş; sipariş miktarı kadar teslim edildiği varsayıldı."))
        if x["tes_miktar"] > x["sip_miktar"]:
            uy.append(uyari("Bilgi", "Fazla teslim", x["siparis"], f"Sipariş {x['sip_miktar']:f}, teslim {x['tes_miktar']:f}. Eksiksiz sayıldı; fazla miktar karşılama oranına katılmadı."))
    siparisler = []
    for g in grup.values():
        s = {"siparis": g[0]["siparis"], "musteri": g[0]["musteri"], "tasiyici": g[0]["tasiyici"], "satir": len(g), "soz": min(x["soz"] for x in g), "teslim": max((x["teslim"] for x in g if x["teslim"]), default=None),
             "sip_miktar": sum((x["sip_miktar"] for x in g), Z), "tes_miktar": sum((x["tes_miktar"] for x in g), Z), "karsilanan": sum((min(x["tes_miktar"], x["sip_miktar"]) for x in g), Z), "gecikme": max(x["gecikme"] for x in g)}
        s["ay"] = f"{s['soz']:%Y-%m}"
        if any(x["acik"] for x in g):
            s["zamaninda"] = s["eksiksiz"] = None
            s["sonuc"] = "Açık"
        else:
            s["zamaninda"], s["eksiksiz"] = all(x["zamaninda"] for x in g), all(x["eksiksiz"] for x in g)
            s["sonuc"] = "OTIF" if s["zamaninda"] and s["eksiksiz"] else "Geç" if s["eksiksiz"] else "Eksik" if s["zamaninda"] else "Geç ve eksik"
            bekleyen = [x for x in g if x["teslim"] is None]
            if bekleyen:
                uy.append(uyari("Yüksek", "Teslim edilmemiş gecikmiş sipariş", s["siparis"], f"{s['musteri']} · {s['tasiyici']}: söz verilen tarih {min(x['soz'] for x in bekleyen):%d.%m.%Y}, "
                                f"{max(x['gecikme'] for x in bekleyen)} gündür teslim edilmedi ({sum((x['sip_miktar'] for x in bekleyen), Z):f} birim)."))
        siparisler.append(s)
    siparisler.sort(key=lambda s: (s["soz"], katla(s["siparis"])))

    def kirilim(alan: str) -> list[dict]:
        g = {}
        for s in siparisler:
            g.setdefault(s[alan] if alan == "ay" else katla(s[alan]), []).append(s)
        return [{"ad": liste[0][alan], **_ozet(liste)} for _, liste in sorted(g.items())]

    return {"siparisler": siparisler, "ozet": _ozet(siparisler), "say": {k: sum(1 for s in siparisler if s["sonuc"] == k) for k in SONUCLAR}, "musteriler": kirilim("musteri"), "tasiyicilar": kirilim("tasiyici"), "aylar": kirilim("ay"), "uyarilar": uy}


def rapor_yaz(cikti: Path, r: dict, bugun: date, hedef: Decimal) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    o = r["ozet"]
    _baslik(ws, ["Gösterge", "Değer"], (44, 14))
    for etiket, deger, para_mi in (("Sipariş sayısı", o["siparis"], False), ("Açık (söz tarihi gelmemiş)", o["acik"], False), ("Değerlendirilen sipariş", o["degerlendirilen"], False), ("Zamanında teslim %", o["zamaninda_oran"], True),
                                   ("Eksiksiz teslim %", o["eksiksiz_oran"], True), ("OTIF %", o["otif_oran"], True), ("Miktar karşılama %", o["karsilama"], True), ("Geç siparişlerde ortalama gecikme (gün)", o["ort_gecikme"], True)):
        _satir(ws, [etiket, deger], para_s=(2,) if para_mi else (), dolgu={2: (YESIL if deger is not None and deger >= hedef else KIRMIZI) if etiket == "OTIF %" else None})
    ws.append([])
    _baslik(ws, ["Sonuç", "Sipariş"], dondur="A2")
    for k, n in r["say"].items():
        _satir(ws, [k, n], dolgu={1: SONUC_RENK[k]})
    ws.append([])
    ws.append([f"Durum tarihi: {bugun:%d.%m.%Y} · OTIF hedefi %{hedef:f}"])
    for baslik, alan, ad in (("Müşteriler", "musteriler", "Müşteri"), ("Taşıyıcılar", "tasiyicilar", "Taşıyıcı"), ("Aylar", "aylar", "Ay (söz tarihi)")):
        ws = wb.create_sheet(baslik)
        _baslik(ws, [ad, "Sipariş", "Açık", "Değerlendirilen", "Zamanında", "Eksiksiz", "OTIF", "Zamanında %", "Eksiksiz %", "OTIF %", "Miktar Karşılama %", "Ort. Gecikme (gün)"], (26, 9, 8, 13, 11, 10, 8, 11, 10, 9, 12, 11))
        for g in r[alan]:
            _satir(ws, [g["ad"], g["siparis"], g["acik"], g["degerlendirilen"], g["zamaninda"], g["eksiksiz"], g["otif"], g["zamaninda_oran"], g["eksiksiz_oran"], g["otif_oran"], g["karsilama"], g["ort_gecikme"]], para_s=(8, 9, 10, 11, 12),
                   dolgu={10: None if g["otif_oran"] is None else YESIL if g["otif_oran"] >= hedef else KIRMIZI})
    ws = wb.create_sheet("Siparişler")
    _baslik(ws, ["Sonuç", "Sipariş No", "Müşteri", "Taşıyıcı", "Satır", "Söz Verilen Tarih", "Son Teslim Tarihi", "Gecikme (gün)", "Sipariş Miktarı", "Teslim Miktarı"], (14, 14, 26, 22, 7, 13, 13, 10, 14, 14))
    for s in r["siparisler"]:
        _satir(ws, [s["sonuc"], s["siparis"], s["musteri"], s["tasiyici"], s["satir"], s["soz"], s["teslim"], s["gecikme"] or None, s["sip_miktar"], s["tes_miktar"]], para_s=(9, 10), tarih_s=(6, 7), dolgu={1: SONUC_RENK[s["sonuc"]]})
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, bugun: date, tolerans: int = 0, eksik_tolerans: Decimal = Z, hedef: Decimal = Decimal(95)) -> dict:
    satirlar, uy = oku(yol)
    r = degerlendir(satirlar, bugun, tolerans, eksik_tolerans)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, bugun, hedef)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Siparişlerin zamanında ve eksiksiz teslim oranını (OTIF) müşteri ve taşıyıcı bazında raporlar.")
    p.add_argument("--siparisler", type=Path, help="Sipariş No, Müşteri, Taşıyıcı, Söz Verilen Tarih, Teslim Tarihi, Sipariş Miktarı, Teslim Miktarı (her satır bir sipariş kalemi)")
    p.add_argument("--tarih", help="Durum tarihi (GG.AA.YYYY); verilmezse bugün")
    p.add_argument("--tolerans", type=int, default=0, help="Söz tarihinden sonra kaç gün zamanında sayılsın (varsayılan 0)")
    p.add_argument("--eksik-tolerans", default="0", help="Sipariş miktarının yüzde kaçı eksik olsa da eksiksiz sayılsın (varsayılan 0)")
    p.add_argument("--hedef", default="95", help="OTIF hedefi, yüzde; yalnız renklendirme için (varsayılan 95)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "otif_raporu.xlsx")
    a = p.parse_args(argv)
    ornek = a.siparisler is None
    yol = ORNEK / "siparisler.csv" if ornek else a.siparisler
    bugun = tarih(a.tarih) if a.tarih else date(2026, 10, 12) if ornek else date.today()
    et, hedef = para(a.eksik_tolerans), para(a.hedef)
    if bugun is None or a.tolerans < 0 or et is None or not 0 <= et < 100 or hedef is None or not 0 < hedef <= 100:
        print("[X] --tarih GG.AA.YYYY; --tolerans sıfır veya pozitif; --eksik-tolerans 0–100 arası (100 hariç); --hedef 0–100 arasında olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, bugun, a.tolerans, et, hedef)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    o = r["ozet"]
    print(f"[OK] {o['siparis']} sipariş ({o['acik']} açık)" + (f" · zamanında %{o['zamaninda_oran']} · eksiksiz %{o['eksiksiz_oran']} · OTIF %{o['otif_oran']} · miktar karşılama %{o['karsilama']}" if o["degerlendirilen"] else " · değerlendirilecek sipariş yok"))
    for u in r["uyarilar"]:
        if u["onem"] == "Yüksek":
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
