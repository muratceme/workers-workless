import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK
BUGUN = date(2026, 10, 12)
SOZ = date(2026, 9, 10)


def x(siparis, teslim=0, sm=10, tm=None, soz=SOZ, musteri="M", tasiyici="T"):
    """teslim: söz tarihinden kaç gün sonra teslim edildi (None = teslim edilmedi)."""
    t = None if teslim is None else soz + timedelta(days=teslim)
    return {"siparis": siparis, "musteri": musteri, "tasiyici": tasiyici, "soz": soz, "teslim": t, "sip_miktar": D(str(sm)), "tes_miktar": D(str(tm)) if tm is not None else (D(str(sm)) if t else D(0)), "miktar_yok": False}


def sonuc(satirlar, **kw):
    return {s["siparis"]: (s["sonuc"], s["gecikme"]) for s in main.degerlendir(satirlar, BUGUN, **kw)["siparisler"]}


class DegerlendirTesti(unittest.TestCase):
    def test_satir_kurallari(self):
        l = [x("aynı gün"), x("erken", -3), x("bir gün geç", 1), x("eksik", 0, 10, "9.99"), x("fazla", 0, 10, 12), x("geç eksik", 4, 10, 5), x("teslimsiz", None), x("açık", None, soz=BUGUN), x("açık2", None, soz=BUGUN + timedelta(days=9)),
             x("dün", None, soz=BUGUN - timedelta(days=1))]
        self.assertEqual(sonuc(l), {"aynı gün": ("OTIF", 0), "erken": ("OTIF", 0), "bir gün geç": ("Geç", 1), "eksik": ("Eksik", 0), "fazla": ("OTIF", 0), "geç eksik": ("Geç ve eksik", 4), "teslimsiz": ("Geç ve eksik", 32),
                                    "açık": ("Açık", 0), "açık2": ("Açık", 0), "dün": ("Geç ve eksik", 1)})
        s = sonuc(l, tolerans=1, eksik_tolerans=D("0.1"))
        self.assertEqual((s["bir gün geç"], s["eksik"], s["geç eksik"], s["dün"], s["teslimsiz"]), (("OTIF", 0), ("OTIF", 0), ("Geç ve eksik", 3), ("Açık", 0), ("Geç ve eksik", 31)))
        self.assertEqual((sonuc([x("a", 2)], tolerans=2)["a"], sonuc([x("a", 3)], tolerans=2)["a"], sonuc([x("a", 0, 100, 95)], eksik_tolerans=D(5))["a"], sonuc([x("a", 0, 100, "94.99")], eksik_tolerans=D(5))["a"]),
                         (("OTIF", 0), ("Geç", 1), ("OTIF", 0), ("Eksik", 0)))
        r = main.degerlendir(l, BUGUN)
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Bilgi", "Fazla teslim", "fazla"), ("Yüksek", "Teslim edilmemiş gecikmiş sipariş", "dün"), ("Yüksek", "Teslim edilmemiş gecikmiş sipariş", "teslimsiz")])
        self.assertEqual(r["say"], {"OTIF": 3, "Geç": 1, "Eksik": 1, "Geç ve eksik": 3, "Açık": 2})
        self.assertEqual(r["ozet"], {"siparis": 10, "acik": 2, "degerlendirilen": 8, "zamaninda": 4, "eksiksiz": 4, "otif": 3, "zamaninda_oran": D("50.00"), "eksiksiz_oran": D("50.00"), "otif_oran": D("37.50"), "karsilama": D("68.74"),
                                     "ort_gecikme": D("9.50")})
        with self.assertRaises(ValueError):
            main.degerlendir([], BUGUN)

    def test_cok_satirli_ve_kirilim(self):
        l = [x("S1", 0, musteri="A", tasiyici="X"), x("s1", 2, 20, musteri="B", tasiyici="Y"), x("S2", 0, 10, 4, musteri="A", tasiyici="X"), x("S2", 0, 10, 10), x("S3", 0, musteri="a", tasiyici="Y", soz=date(2026, 10, 1)),
             x("S3", None, soz=BUGUN + timedelta(days=3)), x("S4", 1, musteri="B", tasiyici="Y", soz=date(2026, 10, 2)), x("S4", None, soz=date(2026, 10, 5)), x("S5", -1, musteri="B", tasiyici="x", soz=date(2026, 10, 3))]
        r = main.degerlendir(l, BUGUN)
        self.assertEqual([(s["siparis"], s["musteri"], s["tasiyici"], s["satir"], s["soz"], s["teslim"], s["sip_miktar"], s["tes_miktar"], s["karsilanan"], s["gecikme"], s["sonuc"]) for s in r["siparisler"]],
                         [("S1", "A", "X", 2, SOZ, date(2026, 9, 12), D(30), D(30), D(30), 2, "Geç"), ("S2", "A", "X", 2, SOZ, SOZ, D(20), D(14), D(14), 0, "Eksik"), ("S3", "a", "Y", 2, date(2026, 10, 1), date(2026, 10, 1), D(20), D(10), D(10), 0, "Açık"),
                          ("S4", "B", "Y", 2, date(2026, 10, 2), date(2026, 10, 3), D(20), D(10), D(10), 7, "Geç ve eksik"), ("S5", "B", "x", 1, date(2026, 10, 3), date(2026, 10, 2), D(10), D(10), D(10), 0, "OTIF")])
        k = lambda alan: [(g["ad"], g["siparis"], g["acik"], g["degerlendirilen"], g["zamaninda"], g["eksiksiz"], g["otif"], g["zamaninda_oran"], g["eksiksiz_oran"], g["otif_oran"], g["karsilama"], g["ort_gecikme"]) for g in r[alan]]   # noqa: E731
        self.assertEqual(k("musteriler"), [("A", 3, 1, 2, 1, 1, 0, D("50.00"), D("50.00"), D("0.00"), D("88.00"), D("2.00")), ("B", 2, 0, 2, 1, 1, 1, D("50.00"), D("50.00"), D("50.00"), D("66.67"), D("7.00"))])
        self.assertEqual(k("tasiyicilar"), [("X", 3, 0, 3, 2, 2, 1, D("66.67"), D("66.67"), D("33.33"), D("90.00"), D("2.00")), ("Y", 2, 1, 1, 0, 0, 0, D("0.00"), D("0.00"), D("0.00"), D("50.00"), D("7.00"))])
        self.assertEqual(k("aylar"), [("2026-09", 2, 0, 2, 1, 1, 0, D("50.00"), D("50.00"), D("0.00"), D("88.00"), D("2.00")), ("2026-10", 3, 1, 2, 1, 1, 1, D("50.00"), D("50.00"), D("50.00"), D("66.67"), D("7.00"))])
        o = main.degerlendir([x("a", None, soz=BUGUN)], BUGUN)["ozet"]
        self.assertEqual((o["degerlendirilen"], o["otif_oran"], o["karsilama"], o["ort_gecikme"]), (0, None, None, None))


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "siparisler.csv", cikti, BUGUN)
            self.assertEqual(r["ozet"], {"siparis": 12, "acik": 1, "degerlendirilen": 11, "zamaninda": 7, "eksiksiz": 8, "otif": 6, "zamaninda_oran": D("63.64"), "eksiksiz_oran": D("72.73"), "otif_oran": D("54.55"), "karsilama": D("92.31"),
                                         "ort_gecikme": D("4.25")})
            self.assertEqual([(s["siparis"], s["sonuc"], s["gecikme"]) for s in r["siparisler"]],
                             [("SP-01", "OTIF", 0), ("SP-02", "OTIF", 0), ("SP-03", "Geç", 2), ("SP-04", "Eksik", 0), ("SP-05", "Geç ve eksik", 3), ("SP-06", "OTIF", 0), ("SP-07", "Geç", 2), ("SP-08", "OTIF", 0), ("SP-09", "Geç ve eksik", 10),
                              ("SP-10", "OTIF", 0), ("SP-12", "OTIF", 0), ("SP-11", "Açık", 0)])
            self.assertEqual(r["say"], {"OTIF": 6, "Geç": 2, "Eksik": 1, "Geç ve eksik": 2, "Açık": 1})
            k = lambda alan: [(g["ad"], g["degerlendirilen"], g["acik"], g["zamaninda_oran"], g["eksiksiz_oran"], g["otif_oran"], g["karsilama"], g["ort_gecikme"]) for g in r[alan]]   # noqa: E731
            self.assertEqual(k("musteriler"), [("Alfa Market", 4, 0, D("75.00"), D("100.00"), D("75.00"), D("100.00"), D("2.00")), ("Beta Yapı", 3, 0, D("66.67"), D("33.33"), D("33.33"), D("82.14"), D("3.00")),
                                               ("Delta Tekstil", 1, 1, D("100.00"), D("100.00"), D("100.00"), D("100.00"), None), ("Gama Gıda", 3, 0, D("33.33"), D("66.67"), D("33.33"), D("83.64"), D("6.00"))])
            self.assertEqual([x_[:6] for x_ in k("tasiyicilar")], [("Ege Taşıma", 3, 0, D("33.33"), D("66.67"), D("33.33")), ("Hızlı Nakliyat", 4, 1, D("100.00"), D("100.00"), D("100.00")), ("Yıldız Lojistik", 4, 0, D("50.00"), D("50.00"), D("25.00"))])
            self.assertEqual([x_[:6] for x_ in k("aylar")], [("2026-09", 8, 0, D("62.50"), D("75.00"), D("50.00")), ("2026-10", 3, 1, D("66.67"), D("66.67"), D("66.67"))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Yüksek", "Teslim edilmemiş gecikmiş sipariş", "SP-09"), ("Bilgi", "Fazla teslim", "SP-08")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Özet", "Müşteriler", "Taşıyıcılar", "Aylar", "Siparişler", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("No;Cari;Nakliyeci;Termin;Teslim Edilen Tarih;Miktar;Sevk Miktarı\nS1;;;05.09.2026;06.09.2026;1.250,5;\n;A;B;05.09.2026;;12.5;\nS3;A;B;x;;5;\nS4;A;B;05.09.2026;x;5;\nS5;A;B;05.09.2026;;0;\nS6;A;B;05.09.2026;;5;-1\n"
                         ";;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(v["siparis"], v["musteri"], v["tasiyici"], v["soz"], v["teslim"], v["sip_miktar"], v["tes_miktar"], v["miktar_yok"]) for v in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([("S1", "(müşteri yok)", "(taşıyıcı yok)", date(2026, 9, 5), date(2026, 9, 6), D("1250.5"), D("1250.5"), True), ("(satır 3)", "A", "B", date(2026, 9, 5), None, D("12.5"), D(0), False)],
                              [("Satır okunamadı", f"S{n}") for n in (3, 4, 5, 6)]))
            self.assertEqual([u["tur"] for u in main.degerlendir(l, BUGUN)["uyarilar"]], ["Teslim miktarı yazılmamış", "Teslim edilmemiş gecikmiş sipariş"])
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--siparisler", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--tarih", "12.10.2026", "--tolerans", "1", "--eksik-tolerans", "2,5", "--hedef", "90"]), 0)
            self.assertEqual(main.main(arg), 0)
            self.assertEqual([main.main(arg + v) for v in (["--tarih", "x"], ["--tolerans", "-1"], ["--eksik-tolerans", "100"], ["--eksik-tolerans", "x"], ["--hedef", "0"], ["--hedef", "101"])], [2] * 6)
            self.assertEqual(main.main(["--siparisler", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Sipariş No;Teslim Tarihi\nA;05.09.2026\n", "Söz Verilen Tarih;Sipariş Miktarı\n;\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
