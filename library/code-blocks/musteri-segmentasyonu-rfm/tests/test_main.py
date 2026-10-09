import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


class PuanTesti(unittest.TestCase):
    def test_puan_esitlik_ve_yon(self):
        tum = [1, 1, 1, 1, 1, 2, 3, 4, 5, 9]
        self.assertEqual([main.puan(v, tum) for v in (1, 2, 3, 4, 5, 9)], [1, 3, 4, 4, 5, 5])   # 1 + ⌊5 × kötü / 10⌋
        gun = [5, 10, 10, 40, 90, 200, 300, 300, 310, 365]
        self.assertEqual(main.puan(5, gun, buyuk_iyi=False), 5)        # 9 müşteri daha kötü → 1 + 4 = 5
        self.assertEqual(main.puan(10, gun, buyuk_iyi=False), 4)       # 7 daha kötü → 1 + 3
        self.assertEqual(main.puan(365, gun, buyuk_iyi=False), 1)
        self.assertEqual(main.puan(300, gun, buyuk_iyi=False), 2)      # 2 daha kötü → 1 + 1

    def test_segment_haritasi_tam(self):
        for r in range(1, 6):
            for f in range(1, 6):
                self.assertNotEqual(main.segment_bul(r, f)[0], "Diğer", (r, f))
        self.assertEqual(main.segment_bul(5, 5)[0], "Şampiyonlar")
        self.assertEqual(main.segment_bul(1, 5)[0], "Kaybedilmemesi Gerekenler")
        self.assertEqual(main.segment_bul(5, 1)[0], "Yeni Müşteriler")

    def test_iade_belge_pencere(self):
        with tempfile.TemporaryDirectory() as tmp:
            t = Path(tmp) / "s.csv"
            t.write_text("Müşteri No;Fatura Tarihi;Fatura No;Tutar\n"
                         "A;01.09.2026;F1;1.000,00\nA;01.09.2026;F1;500,00\nA;15.09.2026;F2;2.000,00\nA;20.09.2026;;-300,00\n"
                         "B;01.03.2026;F3;800,00\nB;01.09.2025;F4;9.000,00\nB;01.11.2026;F5;100,00\n", encoding="utf-8")
            m, uy = main.oku(t, date(2026, 10, 9), 12)
            s = main.analiz_et(m, date(2026, 10, 9))
            a = next(x for x in s["musteriler"] if x.no == "A")
            self.assertEqual((a.f, a.m, a.r, a.iade), (2, D("3200.00"), 24, D("300.00")))     # F1 iki satır = 1 alım; iade düşülür
            b = next(x for x in s["musteriler"] if x.no == "B")
            self.assertEqual((b.f, b.m), (1, D("800.00")))                                       # pencere dışı ve gelecek tarih alınmadı
            self.assertEqual({u["tur"] for u in uy}, {"Pencere dışı", "Gelecek tarihli satır"})


class OrnekVeriTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "r.xlsx"
            s = main.calistir(ORNEK / "satislar.csv", cikti, date(2026, 10, 9), segment_dosyalari=True)
            self.assertEqual(len(s["musteriler"]), 120)
            self.assertEqual(sum(o["sayi"] for o in s["ozet"]), 120)
            self.assertAlmostEqual(sum(o["ciro_pay"] for o in s["ozet"]), 1.0, places=6)
            self.assertEqual(sum(s["matris"].values()), 120)
            self.assertIn(("Net tutar sıfır / eksi", "C0121"), {(u["tur"], u["kim"]) for u in s["uyarilar"]})
            c3 = next(m for m in s["musteriler"] if m.no == "C0003")
            self.assertEqual(c3.iade, D("4500"))
            for m in s["kayiplar"]:
                self.assertTrue(m.f >= 3 and m.r > 2 * m.aralik)
            wb = load_workbook(cikti)
            self.assertEqual(wb.sheetnames, ["Segment Özeti", "Müşteriler", "R × F Matrisi", "Kayıp Sinyali", "Uyarılar"])
            self.assertEqual(len(s["csv"]), sum(1 for o in s["ozet"] if o["sayi"]))
            self.assertTrue(s["csv"][0].read_text(encoding="utf-8-sig").startswith("Müşteri No;"))

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main.main(["--cikti", str(Path(tmp) / "x.xlsx")]), 0)
            self.assertEqual(main.main(["--bugun", "x"]), 2)
            self.assertEqual(main.main(["--satislar", str(Path(tmp) / "yok.csv")]), 1)


if __name__ == "__main__":
    unittest.main()
