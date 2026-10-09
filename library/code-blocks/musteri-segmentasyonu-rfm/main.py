"""
Müşteri Segmentasyonu (RFM) — Workers / Workless kod bloğu
Pazarlama › Pazarlama Uzmanı

Satış geçmişinden müşterileri yenilik (R), sıklık (F) ve tutara (M) göre segmentlere ayırır:
  - Pencere: analiz tarihinden geriye --pencere-ay (12) ay.
  - R = analiz tarihi − son alım (gün); F = farklı fatura / sipariş sayısı (iade satırları sayılmaz); M = net tutar
    (iadeler düşülür). Net tutarı 0 veya altı olan müşteri puanlanmaz.
  - Puan 1–5: değeri kendisinden kötü olan müşteri oranına göre (1 + ⌊5 × oran⌋, en çok 5). Eşit değerler aynı puanı
    alır. R'de az gün iyidir.
  - Segment R ve F puanından (yaygın RFM segment haritası; README'de tablo); her segment için aksiyon önerisi.
  - Kayıp sinyali: en az 3 alımı olan ve son alımdan bu yana geçen süre ortalama alım aralığının --kayip-kat (2) katını
    aşan müşteri.
Rapor: segment özeti, müşteriler, R × F matrisi, kayıp sinyali, uyarılar; isteğe bağlı segment başına CSV (CRM'e
aktarım için). İnternete bağlanmaz.

Kullanım:
    python main.py                                                  # örnek: 120 müşteri, Ekim 2025 – Eylül 2026
    python main.py --satislar satislar.xlsx --bugun 01.10.2026 --pencere-ay 12 --segment-dosyalari
"""
from __future__ import annotations

import argparse
import calendar
import csv
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
K2 = Decimal("0.01")
SIFIR = Decimal(0)

# (segment, R puanları, F puanları, aksiyon) — ilk uyan
SEGMENTLER = [
    ("Şampiyonlar", {5}, {4, 5}, "Ödüllendirin; yeni ürünleri önce bu gruba duyurun; referans ve değerlendirme isteyin."),
    ("Sadık Müşteriler", {3, 4}, {4, 5}, "Sadakat programı ve çapraz / üst satış önerileri; ihtiyaçlarını sorun."),
    ("Potansiyel Sadıklar", {4, 5}, {2, 3}, "Sadakat programına davet; geçmiş alımlara göre kişiselleştirilmiş öneri."),
    ("Yeni Müşteriler", {5}, {1}, "Hoş geldin iletişimi ve kullanım desteği; ikinci alımı teşvik edin."),
    ("Umut Vaat Edenler", {4}, {1}, "Marka bilinirliği iletişimi; ikinci alım için küçük ve süreli teşvik."),
    ("İlgi Bekleyenler", {3}, {3}, "Sınırlı süreli teklif; geçmiş alımlarına göre ürün önerisi."),
    ("Uyumak Üzere", {3}, {1, 2}, "Popüler ürün önerisi ve yeniden bağlantı kampanyası; ilgiyi canlandırın."),
    ("Risk Altında", {1, 2}, {3, 4}, "Kişisel iletişim ve geri kazanma kampanyası; ayrılma nedenini öğrenin."),
    ("Kaybedilmemesi Gerekenler", {1, 2}, {5}, "Satış temsilcisi araması ve özel teklif; şikâyet veya sorun olup olmadığını kontrol edin."),
    ("Uykudakiler", {1, 2}, {1, 2}, "Düşük maliyetli yeniden aktivasyon (e-posta); pahalı kampanya yapmayın."),
]
SEGMENT_SIRASI = [s[0] for s in SEGMENTLER]
RENKLER = {"Şampiyonlar": "E3F4E1", "Sadık Müşteriler": "E3F4E1", "Potansiyel Sadıklar": "E8F0FE", "Yeni Müşteriler": "E8F0FE", "Umut Vaat Edenler": "E8F0FE",
           "İlgi Bekleyenler": "FFF4CE", "Uyumak Üzere": "FFF4CE", "Risk Altında": "FDE2E1", "Kaybedilmemesi Gerekenler": "FDE2E1", "Uykudakiler": "EEEEEE"}

SUTUNLAR = {"musteri": ("musteri no", "musteri kodu", "cari kod", "musteri id"), "ad": ("musteri adi", "musteri", "unvan", "cari adi"),
            "tarih": ("fatura tarihi", "siparis tarihi", "tarih"), "belge": ("fatura no", "siparis no", "belge no"), "tutar": ("tutar", "net tutar", "satis tutari"),
            "bolge": ("bolge", "sehir", "il"), "kanal": ("kanal", "satis kanali")}


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return Decimal(str(x))
    s = str(x).strip().replace("TL", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def tl(x: Decimal | None) -> str:
    return "—" if x is None else f"{x.quantize(K2, ROUND_HALF_UP):,.0f}".replace(",", ".")


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n} for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(zorunlu)}")


def ay_geri(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 - n, 12)
    yil, ay = d.year + y, m + 1
    return date(yil, ay, min(d.day, calendar.monthrange(yil, ay)[1]))


# ----------------------------------------------------------------------------
# Veri ve puanlama
# ----------------------------------------------------------------------------

@dataclass
class Musteri:
    no: str
    ad: str = ""
    bolge: str = ""
    kanal: str = ""
    alimlar: dict = field(default_factory=dict)     # belge → (tarih, tutar)
    iade: Decimal = SIFIR
    r: int = 0
    f: int = 0
    m: Decimal = SIFIR
    rp: int = 0
    fp: int = 0
    mp: int = 0
    segment: str = ""
    ilk: date | None = None
    son: date | None = None
    aralik: float | None = None
    kayip: bool = False


def oku(yol: Path, bugun: date, pencere_ay: int) -> tuple[dict[str, Musteri], list[dict]]:
    bas = ay_geri(bugun, pencere_ay)
    musteriler, uy, disarida, gelecek = {}, [], 0, 0
    for r in kayitlar(yol, SUTUNLAR, ("musteri", "tarih", "tutar")):
        no = metin(r.get("musteri"))
        t, tutar = tarih(r.get("tarih")), para(r.get("tutar"))
        if not no:
            continue
        if not t or tutar is None:
            uy.append({"onem": "Orta", "tur": "Okunamayan satır", "kim": no, "aciklama": f"Satır {r['_satir']}: tarih veya tutar okunamadı"})
            continue
        if t > bugun:
            gelecek += 1
            continue
        if t <= bas:
            disarida += 1
            continue
        m = musteriler.setdefault(no, Musteri(no, metin(r.get("ad")), metin(r.get("bolge")), metin(r.get("kanal"))))
        if tutar < 0:
            m.iade += -tutar
            continue
        belge = metin(r.get("belge")) or f"{t.isoformat()}"
        eski = m.alimlar.get(belge)
        m.alimlar[belge] = (min(eski[0], t), eski[1] + tutar) if eski else (t, tutar)
    if disarida:
        uy.append({"onem": "Bilgi", "tur": "Pencere dışı", "kim": "", "aciklama": f"{disarida} satır {bas:%d.%m.%Y} ve öncesinde; analize alınmadı (--pencere-ay)"})
    if gelecek:
        uy.append({"onem": "Orta", "tur": "Gelecek tarihli satır", "kim": "", "aciklama": f"{gelecek} satır analiz tarihinden sonra; alınmadı"})
    return musteriler, uy


def puan(deger, tum: list, buyuk_iyi: bool = True) -> int:
    """Eşit değerler aynı puanı alır: 1 + ⌊5 × (kendisinden kötü olanların oranı)⌋."""
    n = len(tum)
    kotu = sum(1 for v in tum if (v < deger if buyuk_iyi else v > deger))
    return min(5, 1 + (5 * kotu) // n)


def segment_bul(rp: int, fp: int) -> tuple[str, str]:
    for ad, rs, fs, aksiyon in SEGMENTLER:
        if rp in rs and fp in fs:
            return ad, aksiyon
    return "Diğer", ""


def analiz_et(musteriler: dict[str, Musteri], bugun: date, kayip_kat: float = 2.0) -> dict:
    uy = []
    gecerli = []
    for m in musteriler.values():
        if not m.alimlar:
            uy.append({"onem": "Bilgi", "tur": "Yalnız iade", "kim": m.no, "aciklama": "Pencerede satış yok, yalnız iade var; puanlanmadı"})
            continue
        tarihler = sorted(t for t, _ in m.alimlar.values())
        m.ilk, m.son = tarihler[0], tarihler[-1]
        m.r = (bugun - m.son).days
        m.f = len(m.alimlar)
        m.m = sum((v for _, v in m.alimlar.values()), SIFIR) - m.iade
        if m.m <= 0:
            uy.append({"onem": "Bilgi", "tur": "Net tutar sıfır / eksi", "kim": m.no, "aciklama": f"Satış {tl(m.m + m.iade)}, iade {tl(m.iade)}; puanlanmadı"})
            continue
        if len(tarihler) >= 2:
            m.aralik = (tarihler[-1] - tarihler[0]).days / (len(tarihler) - 1)
        gecerli.append(m)
    rler, fler, mler = [m.r for m in gecerli], [m.f for m in gecerli], [m.m for m in gecerli]
    for m in gecerli:
        m.rp, m.fp, m.mp = puan(m.r, rler, buyuk_iyi=False), puan(m.f, fler), puan(m.m, mler)
        m.segment, _ = segment_bul(m.rp, m.fp)
        if m.f >= 3 and m.aralik and m.r > kayip_kat * max(m.aralik, 1):
            m.kayip = True
    toplam_m = sum(mler, SIFIR)
    ozet = []
    for ad, rs, fs, aksiyon in SEGMENTLER:
        lst = [m for m in gecerli if m.segment == ad]
        if not lst:
            ozet.append({"segment": ad, "sayi": 0, "pay": 0.0, "ciro": SIFIR, "ciro_pay": 0.0, "r": None, "f": None, "m": None, "aksiyon": aksiyon, "kural": (rs, fs)})
            continue
        ciro = sum((m.m for m in lst), SIFIR)
        ozet.append({"segment": ad, "sayi": len(lst), "pay": len(lst) / len(gecerli), "ciro": ciro, "ciro_pay": float(ciro / toplam_m) if toplam_m else 0.0,
                     "r": statistics.median(m.r for m in lst), "f": statistics.median(m.f for m in lst), "m": ciro / len(lst), "aksiyon": aksiyon, "kural": (rs, fs)})
    for o in ozet:
        if o["segment"] in ("Kaybedilmemesi Gerekenler", "Risk Altında") and o["sayi"]:
            uy.append({"onem": "Orta", "tur": "Riskli segment", "kim": o["segment"],
                       "aciklama": f"{o['sayi']} müşteri, ciro payı %{o['ciro_pay'] * 100:.0f}; geri kazanma çalışması planlayın"})
    kayiplar = sorted((m for m in gecerli if m.kayip), key=lambda m: -m.m)
    if kayiplar:
        uy.append({"onem": "Orta", "tur": "Kayıp sinyali", "kim": f"{len(kayiplar)} müşteri",
                   "aciklama": f"Alım aralığının {kayip_kat:g} katından uzun süredir alım yok; toplam ciro {tl(sum((m.m for m in kayiplar), SIFIR))} TL"})
    matris = defaultdict(int)
    for m in gecerli:
        matris[(m.rp, m.fp)] += 1
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    uy.sort(key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    return {"musteriler": gecerli, "ozet": ozet, "matris": matris, "kayiplar": kayiplar, "uyarilar": uy, "bugun": bugun, "toplam": toplam_m}


# ----------------------------------------------------------------------------
# Rapor
# ----------------------------------------------------------------------------

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
UST = Alignment(vertical="top", wrap_text=True)
PF = "#,##0"


def _baslik(ws, basliklar, genislik=()):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


def _kural(k) -> str:
    rs, fs = k
    return f"R {min(rs)}–{max(rs)}" if len(rs) > 1 else f"R {min(rs)}", f"F {min(fs)}–{max(fs)}" if len(fs) > 1 else f"F {min(fs)}"


def rapor_yaz(cikti: Path, s: dict, pencere: str) -> None:
    wb = Workbook()
    so = wb.active
    so.title = "Segment Özeti"
    _baslik(so, ["Segment", "Kural", "Müşteri", "Müşteri Payı", "Ciro (net)", "Ciro Payı", "Medyan R (gün)", "Medyan F", "Ort. M", "Önerilen Aksiyon"],
            (26, 14, 9, 9, 14, 9, 10, 9, 12, 80))
    for o in s["ozet"]:
        r, f = _kural(o["kural"])
        so.append([o["segment"], f"{r}, {f}", o["sayi"], o["pay"], float(o["ciro"]), o["ciro_pay"], o["r"], o["f"], float(o["m"]) if o["m"] is not None else None,
                   o["aksiyon"]])
        so.cell(so.max_row, 1).fill = PatternFill("solid", fgColor=RENKLER[o["segment"]])
        so.cell(so.max_row, 4).number_format = so.cell(so.max_row, 6).number_format = "0.0%"
        so.cell(so.max_row, 5).number_format = so.cell(so.max_row, 9).number_format = PF
        so.cell(so.max_row, 10).alignment = UST
    so.append([])
    so.append([f"Analiz tarihi {s['bugun']:%d.%m.%Y}; pencere {pencere}. Puanlanan müşteri {len(s['musteriler'])}, net ciro {tl(s['toplam'])} TL."])
    so.append(["R: son alımdan bu yana gün (az iyi); F: alım (fatura) sayısı; M: net tutar. Puan 1–5, eşit değerler aynı puan."])

    mu = wb.create_sheet("Müşteriler")
    _baslik(mu, ["Müşteri No", "Müşteri Adı", "Bölge", "Kanal", "İlk Alım", "Son Alım", "R (gün)", "F", "M (net)", "İade", "R Puanı", "F Puanı", "M Puanı",
                 "RFM", "Segment", "Ort. Alım Aralığı (gün)", "Kayıp Sinyali"], (11, 26, 12, 10, 11, 11, 7, 5, 12, 10, 6, 6, 6, 6, 24, 10, 9))
    for m in sorted(s["musteriler"], key=lambda m: (SEGMENT_SIRASI.index(m.segment) if m.segment in SEGMENT_SIRASI else 99, -m.m)):
        mu.append([m.no, m.ad, m.bolge, m.kanal, m.ilk, m.son, m.r, m.f, float(m.m), float(m.iade), m.rp, m.fp, m.mp, f"{m.rp}{m.fp}{m.mp}", m.segment,
                   round(m.aralik, 1) if m.aralik else None, "Evet" if m.kayip else ""])
        mu.cell(mu.max_row, 5).number_format = mu.cell(mu.max_row, 6).number_format = "DD.MM.YYYY"
        mu.cell(mu.max_row, 9).number_format = mu.cell(mu.max_row, 10).number_format = PF
        mu.cell(mu.max_row, 15).fill = PatternFill("solid", fgColor=RENKLER.get(m.segment, "FFFFFF"))
    mu.auto_filter.ref = f"A1:Q{mu.max_row}"

    rf = wb.create_sheet("R × F Matrisi")
    rf.append(["Müşteri sayısı", "F 1", "F 2", "F 3", "F 4", "F 5"])
    for h in rf[1]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
    for r in range(5, 0, -1):
        rf.append([f"R {r}"] + [s["matris"].get((r, f), 0) for f in range(1, 6)])
        for f in range(1, 6):
            seg, _ = segment_bul(r, f)
            rf.cell(rf.max_row, f + 1).fill = PatternFill("solid", fgColor=RENKLER.get(seg, "FFFFFF"))
    rf.column_dimensions["A"].width = 14
    rf.append([])
    rf.append(["Renkler segmentleri gösterir (Segment Özeti sayfasındaki kurallar)."])

    ka = wb.create_sheet("Kayıp Sinyali")
    _baslik(ka, ["Müşteri No", "Müşteri Adı", "Segment", "Son Alım", "R (gün)", "Ort. Alım Aralığı", "Kat", "F", "M (net)", "Aksiyon / Sorumlu"],
            (11, 26, 24, 11, 7, 10, 6, 5, 12, 26))
    for m in s["kayiplar"]:
        ka.append([m.no, m.ad, m.segment, m.son, m.r, round(m.aralik, 1), round(m.r / m.aralik, 1), m.f, float(m.m), ""])
        ka.cell(ka.max_row, 4).number_format = "DD.MM.YYYY"
        ka.cell(ka.max_row, 9).number_format = PF
        ka.cell(ka.max_row, 10).fill = PatternFill("solid", fgColor="FFF4CE")

    uy = wb.create_sheet("Uyarılar")
    _baslik(uy, ["Önem", "Tür", "Kim", "Açıklama"], (9, 24, 24, 100))
    for u in s["uyarilar"]:
        uy.append([u["onem"], u["tur"], u["kim"], u["aciklama"]])
        uy.cell(uy.max_row, 1).fill = PatternFill("solid", fgColor=RENK[u["onem"]])
        uy.cell(uy.max_row, 4).alignment = UST
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def segment_csv_yaz(klasor: Path, s: dict) -> list[Path]:
    klasor.mkdir(parents=True, exist_ok=True)
    yollar = []
    for seg in SEGMENT_SIRASI:
        lst = [m for m in s["musteriler"] if m.segment == seg]
        if not lst:
            continue
        yol = klasor / (re.sub(r"[^a-z0-9]+", "_", katla(seg)).strip("_") + ".csv")
        with open(yol, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Müşteri No", "Müşteri Adı", "Segment", "R", "F", "M", "RFM"])
            for m in sorted(lst, key=lambda m: -m.m):
                w.writerow([m.no, m.ad, m.segment, m.r, m.f, f"{m.m:.2f}".replace(".", ","), f"{m.rp}{m.fp}{m.mp}"])
        yollar.append(yol)
    return yollar


def calistir(yol: Path, cikti: Path, bugun: date, pencere_ay: int = 12, kayip_kat: float = 2.0, segment_dosyalari: bool = False) -> dict:
    musteriler, uy = oku(yol, bugun, pencere_ay)
    if not musteriler:
        raise ValueError(f"{yol.name}: pencere içinde satış bulunamadı")
    s = analiz_et(musteriler, bugun, kayip_kat)
    if len(s["musteriler"]) < 20:
        uy.append({"onem": "Orta", "tur": "Az müşteri", "kim": "", "aciklama": f"{len(s['musteriler'])} müşteri; 5'li puanlama küçük gruplarda anlamlı değildir"})
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    s["uyarilar"] = sorted(uy + s["uyarilar"], key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))
    rapor_yaz(cikti, s, f"{ay_geri(bugun, pencere_ay):%d.%m.%Y} sonrası – {bugun:%d.%m.%Y}")
    s["csv"] = segment_csv_yaz(cikti.parent / "segmentler", s) if segment_dosyalari else []
    return s


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Satış geçmişinden müşterileri RFM (yenilik, sıklık, tutar) segmentlerine ayırır.")
    p.add_argument("--satislar", type=Path, default=ORNEK / "satislar.csv",
                   help="Müşteri No, Müşteri Adı, Fatura Tarihi, Fatura No, Tutar (iade eksi), Bölge, Kanal")
    p.add_argument("--bugun", help="Analiz tarihi GG.AA.YYYY (örnek veride 09.10.2026)")
    p.add_argument("--pencere-ay", type=int, default=12, help="Kaç aylık satış analiz edilsin (varsayılan 12)")
    p.add_argument("--kayip-kat", type=float, default=2.0, help="Son alımdan geçen süre ortalama aralığın kaç katını aşarsa kayıp sinyali (varsayılan 2)")
    p.add_argument("--segment-dosyalari", action="store_true", help="Her segment için segmentler/ klasörüne CSV yaz (CRM'e aktarım)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "rfm_segmentasyonu.xlsx")
    a = p.parse_args(argv)
    if not a.satislar.exists():
        print(f"[X] Dosya bulunamadı: {a.satislar}")
        return 1
    ornek = a.satislar == ORNEK / "satislar.csv"
    bugun = tarih(a.bugun) if a.bugun else (date(2026, 10, 9) if ornek else date.today())
    if not bugun:
        print("[X] --bugun GG.AA.YYYY biçiminde olmalı")
        return 2
    try:
        s = calistir(a.satislar, a.cikti, bugun, a.pencere_ay, a.kayip_kat, a.segment_dosyalari)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    print(f"[OK] {len(s['musteriler'])} müşteri puanlandı · net ciro {tl(s['toplam'])} TL")
    for o in s["ozet"]:
        if o["sayi"]:
            print(f"     {o['segment']:<26} {o['sayi']:>4} müşteri  ciro payı %{o['ciro_pay'] * 100:.1f}".replace(".", ","))
    for u in s["uyarilar"]:
        if u["onem"] != "Bilgi":
            print(f"[!] {u['kim']} · {u['tur']}: {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    if s["csv"]:
        print(f"[OK] {len(s['csv'])} segment dosyası: {s['csv'][0].parent.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
