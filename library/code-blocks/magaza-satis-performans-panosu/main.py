"""
Mağaza Satış Performans Panosu — Gerisi Bende kod bloğu
Mağaza Operasyonları › Mağaza Müdürü

Ciro, fiş adedi, sepet ortalaması, dönüşüm oranı ve geçen yıl karşılaştırmasını (LFL) mağaza ve gün bazında panoda toplar.
  Sepet ortalaması   = ciro ÷ fiş adedi
  Dönüşüm oranı %    = fiş adedi ÷ ziyaretçi sayısı × 100              (ziyaretçi sayısı yazılmış günler üzerinden)
  Fiş başına adet    = satılan adet ÷ fiş adedi                         (satılan adet yazılmış günler üzerinden)
  LFL büyüme %       = (ciro − geçen yıl ciro) ÷ geçen yıl ciro × 100   (yalnız geçen yıl cirosu olan günler karşılaştırılır)
  Hedef gerçekleşme % = ciro ÷ hedef ciro × 100                         (hedefi yazılmış günler üzerinden)
  Uyarı              = mağazanın LFL büyümesi --lfl (−%5) altında · hedef gerçekleşmesi --hedef (%90) altında
Oranlar günlük oranların ortalaması değil, toplamların oranıdır. Geçen yılın hangi günle eşleştiğini siz belirlersiniz.
"""
from __future__ import annotations
from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
ORNEK = BURASI / "ornek_veri"
_TR = str.maketrans("çğıöşüâîûÇĞİÖŞÜÂÎÛ", "cgiosuaiuCGIOSUAIU")
BASLIK_DOLGU, BASLIK_YAZI = PatternFill("solid", fgColor="1D1D1F"), Font(bold=True, color="FFFFFF")
RENK = {"Yüksek": "FDE2E1", "Orta": "FFF4CE", "Bilgi": "E8F0FE"}
KIRMIZI, SARI, MAVI, YESIL = "FDE2E1", "FFF4CE", "E8F0FE", "E3F4E1"
UST = Alignment(vertical="top", wrap_text=True)
PF, TF = "#,##0.00", "DD.MM.YYYY"
KURUS = Decimal("0.01")


def katla(s) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(s or "").translate(_TR).lower()).split())


def tarih(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x or "").strip()[:10]
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def para(x) -> Decimal | None:
    """1.250,50 · 1250.50 · 750.000 · %12,5 yazımlarını Decimal'e çevirir."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float, Decimal)):
        return Decimal(str(x))
    s = re.sub(r"[^\d,.\-]", "", str(x))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def yuvarla(x: Decimal, hane: str = "0.01") -> Decimal:
    return x.quantize(Decimal(hane), rounding=ROUND_HALF_UP)


def metin(x) -> str:
    return str(x if x is not None else "").strip()


def tablo_oku(yol: Path) -> list[list]:
    if yol.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(yol, data_only=True, read_only=True)
        try:
            satirlar = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        icerik = None
        for kod in ("utf-8-sig", "cp1254"):
            try:
                icerik = yol.read_text(encoding=kod)
                break
            except UnicodeDecodeError:
                continue
        ilk = "\n".join(icerik.splitlines()[:10])
        satirlar = list(csv.reader(icerik.splitlines(), delimiter=max(";\t", key=ilk.count)))
    return [r for r in satirlar if any(c not in (None, "") for c in r)]


def kayitlar(yol: Path, sozluk: dict, zorunlu: tuple[str, ...]) -> list[dict]:
    """Başlık satırını bulur (ilk 15 satırda); sütun adlarını Türkçe karakter ve büyük-küçük harf ayırmadan eşleştirir."""
    s = tablo_oku(yol)
    for bi, r in enumerate(s[:15]):
        b = [katla(c) for c in r]
        es = {}
        for alan, adlar in sozluk.items():
            j = next((b.index(a) for a in adlar if a in b and b.index(a) not in es.values()), None)
            if j is not None:
                es[alan] = j
        if all(z in es for z in zorunlu):
            return [{a: None for a in sozluk} | {a: (r[j] if j < len(r) else None) for a, j in es.items()} | {"_satir": n}
                    for n, r in enumerate(s[bi + 1:], bi + 2)]
    raise ValueError(f"{yol.name}: başlık satırı bulunamadı; gerekli sütunlar: {', '.join(sozluk[z][0] for z in zorunlu)}")


def uyari(onem: str, tur: str, kim, aciklama: str) -> dict:
    return {"onem": onem, "tur": tur, "kim": metin(kim), "aciklama": aciklama}


def sirala(uyarilar: list[dict]) -> list[dict]:
    sira = {"Yüksek": 0, "Orta": 1, "Bilgi": 2}
    return sorted(uyarilar, key=lambda u: (sira[u["onem"]], u["tur"], u["kim"]))


def _baslik(ws, basliklar, genislik=(), dondur="A2"):
    ws.append(basliklar)
    for h in ws[ws.max_row]:
        h.fill, h.font = BASLIK_DOLGU, BASLIK_YAZI
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(genislik, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = dondur


def _satir(ws, degerler, para_s=(), tarih_s=(), dolgu=None, sar=()):
    """Satır ekler; para_s / tarih_s sütunlarına biçim, dolgu = {sütun: renk}, sar = metni kaydırılacak sütunlar."""
    ws.append([float(d) if isinstance(d, Decimal) else d for d in degerler])
    n = ws.max_row
    for j in para_s:
        ws.cell(n, j).number_format = PF
    for j in tarih_s:
        ws.cell(n, j).number_format = TF
    for j, renk in (dolgu or {}).items():
        if renk:
            ws.cell(n, j).fill = PatternFill("solid", fgColor=renk)
    for j in sar:
        ws.cell(n, j).alignment = UST
    return n


def uyari_sayfasi(wb, uyarilar: list[dict]) -> None:
    ws = wb.create_sheet("Uyarılar")
    _baslik(ws, ["Önem", "Tür", "Kim / Ne", "Açıklama"], (9, 32, 28, 100))
    for u in uyarilar:
        _satir(ws, [u["onem"], u["tur"], u["kim"], u["aciklama"]], dolgu={1: RENK[u["onem"]]}, sar=(4,))


def kaydet(wb, cikti: Path) -> None:
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)

SUTUN = {"tarih": ("tarih", "gun"), "magaza": ("magaza", "magaza adi", "sube"), "ciro": ("ciro", "satis", "net satis", "satis tutari"), "fis": ("fis adedi", "fis sayisi", "islem adedi"),
         "ziyaretci": ("ziyaretci", "ziyaretci sayisi", "kisi sayisi", "trafik"), "adet": ("satilan adet", "adet", "urun adedi"), "gy": ("gecen yil ciro", "gecen yil", "onceki yil ciro"), "hedef": ("hedef ciro", "hedef", "butce")}
SAYI = ("ciro", "fis", "ziyaretci", "adet", "gy", "hedef")
SIFIR = Decimal(0)
GUN_AD = ("Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz")


def oku(yol: Path) -> tuple[list[dict], list[dict]]:
    uy, l = [], []
    for r in kayitlar(yol, SUTUN, ("tarih", "magaza", "ciro")):
        magaza = metin(r["magaza"])
        if not magaza:
            continue
        t, d = tarih(r["tarih"]), {a: para(r[a]) if metin(r[a]) else None for a in SAYI}
        if t is None or d["ciro"] is None or any(metin(r[a]) and (d[a] is None or (a != "ciro" and d[a] < 0)) for a in SAYI):
            uy.append(uyari("Yüksek", "Satır okunamadı", f"{magaza} · {metin(r['tarih'])}", f"Satır {r['_satir']}: tarih ya da sayılardan biri okunamadı. Satır alınmadı."))
            continue
        if d["ciro"] > 0 and not d["fis"]:
            uy.append(uyari("Orta", "Ciro var, fiş adedi yok", f"{magaza} · {t:%d.%m.%Y}", "Sepet ortalaması ve dönüşüm bu gün için hesaplanamadı."))
        l.append({"tarih": t, "magaza": magaza, **d})
    return l, uy


def _oran(pay, payda, hane="0.01", kat=100) -> Decimal | None:
    return yuvarla(Decimal(pay) * kat / Decimal(payda), hane) if payda else None


def gosterge(l: list[dict]) -> dict:
    top = lambda alan, kosul=lambda g: True: sum((g[alan] or SIFIR for g in l if kosul(g)), SIFIR)   # noqa: E731
    ciro, fis = top("ciro"), top("fis")
    z, a, gy, h = (lambda g: bool(g["ziyaretci"])), (lambda g: g["adet"] is not None and bool(g["fis"])), (lambda g: bool(g["gy"])), (lambda g: bool(g["hedef"]))
    return {"gun": len(l), "ciro": ciro, "fis": fis, "sepet": _oran(ciro, fis, kat=1), "ziyaretci": top("ziyaretci"), "donusum": _oran(top("fis", z), top("ziyaretci", z)), "adet": top("adet"), "upt": _oran(top("adet", a), top("fis", a), kat=1),
            "gy": top("gy"), "lfl_ciro": top("ciro", gy), "lfl": _oran(top("ciro", gy) - top("gy"), top("gy")), "hedef": top("hedef"), "hedef_oran": _oran(top("ciro", h), top("hedef")), "lfl_gun": sum(1 for g in l if gy(g))}


def pano(gunler: list[dict], lfl_esik: Decimal = Decimal(-5), hedef_esik: Decimal = Decimal(90)) -> dict:
    if not gunler:
        raise ValueError("Panoya girecek satış kaydı yok.")
    uy, mg, tg = [], {}, {}
    for g in gunler:
        mg.setdefault(katla(g["magaza"]), []).append(g)
        tg.setdefault(g["tarih"], []).append(g)
    toplam = gosterge(gunler)
    magazalar = sorted(({"magaza": l[0]["magaza"], **gosterge(l)} for l in mg.values()), key=lambda m: (-m["ciro"], katla(m["magaza"])))
    for m in magazalar:
        m["pay"] = _oran(m["ciro"], toplam["ciro"])
        if m["lfl"] is not None and m["lfl"] < lfl_esik:
            uy.append(uyari("Orta", "Geçen yılın gerisinde", m["magaza"], f"Karşılaştırılabilir {m['lfl_gun']} günde ciro {m['lfl_ciro']:,.2f}, geçen yıl {m['gy']:,.2f} (LFL %{m['lfl']}; eşik %{lfl_esik})."))
        elif m["lfl"] is None:
            uy.append(uyari("Bilgi", "Geçen yıl verisi yok", m["magaza"], "Karşılaştırılabilir gün olmadığı için LFL hesaplanmadı (yeni mağaza ya da geçen yıl kapalı)."))
        if m["hedef_oran"] is not None and m["hedef_oran"] < hedef_esik:
            uy.append(uyari("Orta", "Hedefin gerisinde", m["magaza"], f"Ciro {m['ciro']:,.2f}, hedef {m['hedef']:,.2f} (%{m['hedef_oran']}; eşik %{hedef_esik})."))
    detay = sorted(({"tarih": g["tarih"], "magaza": g["magaza"], **gosterge([g])} for g in gunler), key=lambda g: (g["tarih"], katla(g["magaza"])))
    return {"detay": detay, "magazalar": magazalar, "gunler": [{"tarih": t, **gosterge(l)} for t, l in sorted(tg.items())], "toplam": toplam, "uyarilar": uy}


BASLIK = ["Ciro", "Fiş Adedi", "Sepet Ortalaması", "Ziyaretçi", "Dönüşüm %", "Satılan Adet", "Fiş Başına Adet", "Geçen Yıl Ciro", "LFL Büyüme %", "Hedef Ciro", "Hedef Gerçekleşme %"]


def _deger(g: dict) -> list:
    return [g["ciro"], g["fis"], g["sepet"], g["ziyaretci"] or None, g["donusum"], g["adet"] or None, g["upt"], g["gy"] or None, g["lfl"], g["hedef"] or None, g["hedef_oran"]]


def rapor_yaz(cikti: Path, r: dict, lfl_esik: Decimal, hedef_esik: Decimal) -> None:
    wb = Workbook()

    def satir(ws, on: list, g: dict, kalin: bool = False) -> None:
        k = _satir(ws, on + _deger(g), para_s=[len(on) + i for i in (1, 3, 5, 7, 8, 9, 10, 11)], tarih_s=(1,) if isinstance(on[0], date) else (),
                               dolgu={len(on) + 9: None if g["lfl"] is None else KIRMIZI if g["lfl"] < lfl_esik else YESIL if g["lfl"] >= 0 else None,
                                      len(on) + 11: None if g["hedef_oran"] is None else KIRMIZI if g["hedef_oran"] < hedef_esik else YESIL if g["hedef_oran"] >= 100 else None})
        if kalin:
            for h in ws[k]:
                h.font = Font(bold=True)

    ws = wb.active
    ws.title = "Mağazalar"
    _baslik(ws, ["Mağaza", "Gün", "Ciro Payı %"] + BASLIK, (24, 6, 10, 16, 10, 12, 11, 10, 11, 10, 16, 11, 16, 12), "B2")
    for m in r["magazalar"]:
        satir(ws, [m["magaza"], m["gun"], m["pay"]], m)
    satir(ws, ["TOPLAM", r["toplam"]["gun"], None], r["toplam"], True)
    ws.append([])
    ws.append([f"LFL: yalnız geçen yıl cirosu yazılmış günler karşılaştırılır. Eşikler: LFL %{lfl_esik}, hedef %{hedef_esik}."])
    ws = wb.create_sheet("Günler")
    _baslik(ws, ["Tarih", "Gün", "Mağaza Sayısı"] + BASLIK, (12, 6, 9, 16, 10, 12, 11, 10, 11, 10, 16, 11, 16, 12), "B2")
    for g in r["gunler"]:
        satir(ws, [g["tarih"], GUN_AD[g["tarih"].weekday()], g["gun"]], g)
    ws = wb.create_sheet("Günlük Pano")
    _baslik(ws, ["Tarih", "Gün", "Mağaza"] + BASLIK, (12, 6, 24, 16, 10, 12, 11, 10, 11, 10, 16, 11, 16, 12), "D2")
    for g in r["detay"]:
        satir(ws, [g["tarih"], GUN_AD[g["tarih"].weekday()], g["magaza"]], g)
    uyari_sayfasi(wb, r["uyarilar"])
    kaydet(wb, cikti)


def calistir(yol: Path, cikti: Path, lfl_esik: Decimal = Decimal(-5), hedef_esik: Decimal = Decimal(90)) -> dict:
    gunler, uy = oku(yol)
    r = pano(gunler, lfl_esik, hedef_esik)
    r["uyarilar"] = sirala(uy + r["uyarilar"])
    rapor_yaz(cikti, r, lfl_esik, hedef_esik)
    return r


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Ciro, fiş adedi, sepet ortalaması, dönüşüm oranı ve geçen yıl karşılaştırmasını (LFL) günlük panoda toplar.")
    p.add_argument("--satislar", type=Path, help="Tarih, Mağaza, Ciro, Fiş Adedi, Ziyaretçi, Satılan Adet, Geçen Yıl Ciro, Hedef Ciro")
    p.add_argument("--lfl", default="-5", help="LFL büyümesi bu yüzdenin altındaysa uyarılır (varsayılan -5)")
    p.add_argument("--hedef", default="90", help="Hedef gerçekleşmesi bu yüzdenin altındaysa uyarılır (varsayılan 90)")
    p.add_argument("--cikti", type=Path, default=Path("cikti") / "magaza_satis_panosu.xlsx")
    a = p.parse_args(argv)
    yol = a.satislar or ORNEK / "gunluk_satislar.csv"
    lfl, hedef = para(a.lfl), para(a.hedef)
    if lfl is None or hedef is None or hedef < 0:
        print("[X] --lfl bir yüzde (ör. -5), --hedef sıfır ya da üstü bir yüzde olmalı")
        return 2
    if not yol.exists():
        print(f"[X] Dosya bulunamadı: {yol}")
        return 1
    try:
        r = calistir(yol, a.cikti, lfl, hedef)
    except ValueError as h:
        print(f"[X] {h}")
        return 1
    t = r["toplam"]
    print(f"[OK] {len(r['magazalar'])} mağaza · {len(r['gunler'])} gün · ciro {t['ciro']:,.2f} · fiş {t['fis']:,.0f}" + (f" · sepet {t['sepet']:,.2f}" if t["sepet"] is not None else "") + (f" · dönüşüm %{t['donusum']}" if t["donusum"] is not None else "")
          + (f" · LFL %{t['lfl']}" if t["lfl"] is not None else "") + (f" · hedef %{t['hedef_oran']}" if t["hedef_oran"] is not None else ""))
    for u in r["uyarilar"]:
        if u["onem"] in ("Yüksek", "Orta"):
            print(f"[!] {u['tur']}: {u['kim']} · {u['aciklama']}")
    print(f"[OK] Rapor: {a.cikti.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
