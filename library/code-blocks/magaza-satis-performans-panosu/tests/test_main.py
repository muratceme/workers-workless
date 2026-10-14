import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ORNEK = main.ORNEK


def g(magaza, gun, ciro, fis, ziyaretci=None, adet=None, gy=None, hedef=None):
    c = lambda x: None if x is None else D(str(x))   # noqa: E731
    return {"tarih": date(2026, 10, gun), "magaza": magaza, "ciro": D(ciro), "fis": c(fis), "ziyaretci": c(ziyaretci), "adet": c(adet), "gy": c(gy), "hedef": c(hedef)}


def veri():
    return [g("M1", 5, 10000, 100, 400, 250, 8000, 12000), g("m1", 6, 6000, 50, 300, 100, 0, 6000), g("M2", 5, 5000, 80, gy=6000)]


def k(x):
    return (x["gun"], x["ciro"], x["fis"], x["sepet"], x["donusum"], x["upt"], x["lfl"], x["hedef_oran"])


class PanoTesti(unittest.TestCase):
    def test_gostergeler(self):
        r = main.pano(veri())
        self.assertEqual([(m["magaza"], m["pay"]) + k(m) for m in r["magazalar"]],
                         [("M1", D("76.19"), 2, D(16000), D(150), D("106.67"), D("21.43"), D("2.33"), D("25.00"), D("88.89")), ("M2", D("23.81"), 1, D(5000), D(80), D("62.50"), None, None, D("-16.67"), None)])
        self.assertEqual(k(r["toplam"]), (3, D(21000), D(230), D("91.30"), D("21.43"), D("2.33"), D("7.14"), D("88.89")))
        self.assertEqual((r["toplam"]["lfl_ciro"], r["toplam"]["gy"], r["toplam"]["lfl_gun"], r["toplam"]["hedef"], r["toplam"]["ziyaretci"]), (D(15000), D(14000), 2, D(18000), D(700)))
        self.assertEqual([(x["tarih"].day,) + k(x) for x in r["gunler"]], [(5, 2, D(15000), D(180), D("83.33"), D("25.00"), D("2.50"), D("7.14"), D("83.33")), (6, 1, D(6000), D(50), D("120.00"), D("16.67"), D("2.00"), None, D("100.00"))])
        self.assertEqual([(x["tarih"].day, x["magaza"], x["sepet"], x["lfl"]) for x in r["detay"]], [(5, "M1", D("100.00"), D("25.00")), (5, "M2", D("62.50"), D("-16.67")), (6, "m1", D("120.00"), None)])
        self.assertEqual(sorted((u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]), [("Orta", "Geçen yılın gerisinde", "M2"), ("Orta", "Hedefin gerisinde", "M1")])

    def test_esikler_ve_eksik_veri(self):
        u = lambda **kw: sorted((x["tur"], x["kim"]) for x in main.pano(veri(), **kw)["uyarilar"])   # noqa: E731
        self.assertEqual((u(lfl_esik=D("-16.67"), hedef_esik=D("88.89")), u(lfl_esik=D(30), hedef_esik=D(0))), ([], [("Geçen yılın gerisinde", "M1"), ("Geçen yılın gerisinde", "M2")]))
        r = main.pano([g("M", 5, 1000, None), g("M", 6, 0, 0, 100)])
        t = r["toplam"]
        self.assertEqual((t["sepet"], t["donusum"], t["upt"], t["lfl"], t["hedef_oran"], [x["tur"] for x in r["uyarilar"]]), (None, D("0.00"), None, None, None, ["Geçen yıl verisi yok"]))
        with self.assertRaises(ValueError):
            main.pano([])


class DosyaTesti(unittest.TestCase):
    def test_ornek(self):
        with tempfile.TemporaryDirectory() as tmp:
            cikti = Path(tmp) / "t.xlsx"
            r = main.calistir(ORNEK / "gunluk_satislar.csv", cikti)
            self.assertEqual([(m["magaza"], m["pay"]) + k(m) for m in r["magazalar"]],
                             [("Merkez", D("42.03"), 7, D(620000), D(3100), D("200.00"), D("23.85"), D("2.52"), D("12.73"), D("97.64")), ("AVM", D("33.22"), 7, D(490000), D(2790), D("175.63"), D("16.61"), D("2.33"), None, D("89.09")),
                              ("Sahil", D("24.75"), 7, D(365000), D(2400), D("152.08"), D("18.90"), D("2.04"), D("-7.36"), D("91.25"))])
            self.assertEqual(k(r["toplam"]), (21, D(1475000), D(8290), D("177.93"), D("19.51"), D("2.32"), D("4.34"), D("93.06")))
            self.assertEqual([(x["tarih"].day, x["gun"], x["ciro"]) for x in r["gunler"]], [(d, 3, D(185000)) for d in range(5, 10)] + [(10, 3, D(275000)), (11, 3, D(275000))])
            self.assertEqual([(u["onem"], u["tur"], u["kim"]) for u in r["uyarilar"]], [("Orta", "Geçen yılın gerisinde", "Sahil"), ("Orta", "Hedefin gerisinde", "AVM"), ("Bilgi", "Geçen yıl verisi yok", "AVM")])
            self.assertEqual(load_workbook(cikti).sheetnames, ["Mağazalar", "Günler", "Günlük Pano", "Uyarılar"])

    def test_okuma_ve_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            y = Path(tmp) / "s.csv"
            y.write_text("Gün;Şube;Satış;Fiş Sayısı;Trafik;Adet;Geçen Yıl;Bütçe\n05.10.2026;A;12.500,50;1.200;4.000;2.500;10.000;13.000\n06.10.2026;A;100;;;;;\n07.10.2026;A;-50;1;;;;\nx;A;5;;;;;\n08.10.2026;A;y;;;;;\n08.10.2026;A;5;-1;;;;\n"
                         "08.10.2026;;5;;;;;\n", encoding="utf-8")
            l, uy = main.oku(y)
            self.assertEqual(([(x["tarih"].day, x["magaza"], x["ciro"], x["fis"], x["ziyaretci"], x["adet"], x["gy"], x["hedef"]) for x in l], [(u["tur"], u["kim"]) for u in uy]),
                             ([(5, "A", D("12500.50"), D(1200), D(4000), D(2500), D(10000), D(13000)), (6, "A", D(100), None, None, None, None, None), (7, "A", D(-50), D(1), None, None, None, None)],
                              [("Ciro var, fiş adedi yok", "A · 06.10.2026"), ("Satır okunamadı", "A · x"), ("Satır okunamadı", "A · 08.10.2026"), ("Satır okunamadı", "A · 08.10.2026")]))
            c = str(Path(tmp) / "x.xlsx")
            arg = ["--satislar", str(y), "--cikti", c]
            self.assertEqual(main.main(arg + ["--lfl", "0", "--hedef", "95"]), 0)
            self.assertEqual([main.main(arg + v) for v in (["--lfl", "x"], ["--hedef", "-1"])], [2] * 2)
            self.assertEqual(main.main(["--satislar", str(Path(tmp) / "yok.csv")]), 1)
            for icerik in ("Tarih;Mağaza\n05.10.2026;A\n", "Tarih;Mağaza;Ciro\n05.10.2026;;5\n"):
                y.write_text(icerik, encoding="utf-8")
                self.assertEqual(main.main(arg), 1)
            self.assertEqual(main.main(["--cikti", c]), 0)


if __name__ == "__main__":
    unittest.main()
