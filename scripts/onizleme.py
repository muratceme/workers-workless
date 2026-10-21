"""Her kod bloğunu örnek veriyle çalıştırıp çıktı Excel'inin ilk satırlarını data/onizlemeler.json'a yazar.

Sitedeki "Örnek çıktı" önizlemesi bu dosyadan gelir. Örnek veri ya da kod değiştiğinde yeniden çalıştırın:
    .venv/Scripts/python.exe scripts/onizleme.py            # hepsi
    .venv/Scripts/python.exe scripts/onizleme.py abc-def    # yalnız seçilenler
Sizden Gelenler kod blokları (library/istekler) "istekler/<id>" anahtarıyla yazılır.
"""
import concurrent.futures as cf, datetime as dt, hashlib, json, os, shutil, subprocess, sys, tempfile
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
LIB = KOK / "library" / "code-blocks"
ISTEK = KOK / "library" / "istekler"
CIKTI = KOK / "data" / "onizlemeler.json"
SATIR, SUTUN, SAYFA = 7, 7, 3


def klasor(id_: str) -> Path:
    """"abc" → library/code-blocks/abc · "istekler/abc" → library/istekler/abc"""
    return ISTEK / id_.split("/", 1)[1] if id_.startswith("istekler/") else LIB / id_


def anahtarlar() -> list[str]:
    """Önizlemesi olan paketler: tüm kod blokları ve Sizden Gelenler'deki kod blokları."""
    ids = [p.name for p in sorted(LIB.iterdir()) if (p / "task.json").exists()]
    if ISTEK.is_dir():
        ids += [f"istekler/{p.name}" for p in sorted(ISTEK.iterdir())
                if (p / "task.json").exists() and json.loads((p / "task.json").read_text(encoding="utf-8")).get("tur") == "kod"]
    return ids


def imza(id_: str) -> str:
    """Paketin önizlemeyi etkileyen içeriğinin özeti: ana dosya, requirements, ornek_veri ve ortak dosyalar.
    kontrol.py bunu data/onizlemeler.json'daki değerle karşılaştırır; farklıysa önizleme eskimiştir."""
    k = klasor(id_)
    m = json.loads((k / "task.json").read_text(encoding="utf-8"))
    posix = lambda p: p.relative_to(KOK).as_posix()  # sıralama her işletim sisteminde aynı olsun (Windows Path karşılaştırması büyük/küçük harfe duyarsız)
    yollar = [k / m["ana_dosya"], k / "requirements.txt"] + sorted((p for p in (k / "ornek_veri").rglob("*") if p.is_file()), key=posix)
    for rel in m.get("ortak_dosyalar", {}).values():
        src = KOK / "library" / rel
        yollar += sorted((p for p in src.rglob("*") if p.is_file()), key=posix) if src.is_dir() else [src]
    h = hashlib.sha256()
    for p in yollar:
        if p.exists():
            veri = p.read_bytes()
            if b"\x00" not in veri[:4096]:  # metin dosyası: satır sonu farkı (CRLF / LF) imzayı değiştirmesin (git eol=lf normalizasyonu)
                veri = veri.replace(b"\r\n", b"\n")
            h.update(str(p.relative_to(KOK)).replace("\\", "/").encode()); h.update(veri)
    return h.hexdigest()[:16]


def yazi(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(v, int) and not isinstance(v, bool):
        return f"{v:,}".replace(",", ".")
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y")
    s = str(v).strip()
    return s if len(s) <= 60 else s[:57] + "…"


def oku(xlsx: Path):
    from openpyxl import load_workbook  # kontrol.py bu modülü içe aktarır; openpyxl yalnız burada gerekir
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    sayfalar = []
    for ws in wb.worksheets[:SAYFA]:
        satirlar = []
        for row in ws.iter_rows(max_row=40, max_col=SUTUN, values_only=True):
            r = [yazi(v) for v in row]
            if any(r):
                satirlar.append(r)
            if len(satirlar) >= SATIR:
                break
        if satirlar:
            gen = max(len(r) for r in satirlar)
            while gen and not any(len(r) >= gen and r[gen - 1] for r in satirlar):
                gen -= 1
            sayfalar.append({"ad": ws.title, "satirlar": [r[:gen] + [""] * (gen - len(r[:gen])) for r in satirlar],
                             "toplam_satir": ws.max_row or 0})
    wb.close()
    return sayfalar


def bir(id_):
    try:
        return _bir(id_)
    except Exception as e:  # bir paketin hatası diğerlerini durdurmasın
        return id_, {"hata": f"{type(e).__name__}: {e}"[:300]}


def _bir(id_):
    kaynak = klasor(id_)
    m = json.loads((kaynak / "task.json").read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory() as td:
        d = Path(td) / id_
        shutil.copytree(kaynak, d, ignore=shutil.ignore_patterns("cikti", "__pycache__", "tests"))
        for hedef, rel in m.get("ortak_dosyalar", {}).items():
            src, dst = KOK / "library" / rel, d / hedef
            if src.is_dir():
                shutil.copytree(src, dst, dirs_exist_ok=True)
            else:
                dst.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy(src, dst)
        ortam = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
        try:
            r = subprocess.run([sys.executable, m["ana_dosya"]], cwd=d, env=ortam, capture_output=True, text=True,
                               encoding="utf-8", errors="replace", timeout=180, stdin=subprocess.DEVNULL)
        except subprocess.TimeoutExpired:
            return id_, {"hata": "zaman aşımı"}
        if r.returncode != 0:
            return id_, {"hata": (r.stderr or r.stdout)[-300:]}
        xs = sorted((p for p in d.rglob("*.xlsx") if "ornek_veri" not in p.parts), key=lambda p: (("cikti" not in p.parts), p.name))
        if not xs:
            return id_, {"hata": "xlsx çıktı yok"}
        return id_, {"imza": imza(id_), "dosya": xs[0].name, "sayfalar": oku(xs[0])}


def main():
    secili = [a for a in sys.argv[1:] if not a.startswith("--")]
    if "--imza-guncelle" in sys.argv:  # yalnız imzaları yaz (içerik değişmediyse yeniden çalıştırmaya gerek yok)
        d = json.loads(CIKTI.read_text(encoding="utf-8"))
        for id_, v in d.items():
            if "hata" not in v and klasor(id_).exists():
                v["imza"] = imza(id_)
        CIKTI.write_text(json.dumps(dict(sorted(d.items())), ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        print("imzalar güncellendi:", sum("imza" in v for v in d.values())); return
    ids = [a for a in anahtarlar() if not secili or a in secili or a.split("/")[-1] in secili]
    eski = json.loads(CIKTI.read_text(encoding="utf-8")) if CIKTI.exists() else {}
    with cf.ThreadPoolExecutor(max(2, (os.cpu_count() or 4) // 2)) as ex:
        for i, (id_, s) in enumerate(ex.map(bir, ids), 1):
            eski[id_] = s
            print(f"{i}/{len(ids)} {'HATA ' + s['hata'][:80] if 'hata' in s else 'ok'} · {id_}", flush=True)
    CIKTI.parent.mkdir(exist_ok=True)
    CIKTI.write_text(json.dumps(dict(sorted(eski.items())), ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    ok = sum("hata" not in v for v in eski.values())
    print(f"önizleme: {ok}/{len(eski)} hazır → {CIKTI.relative_to(KOK)}")


if __name__ == "__main__":
    main()
